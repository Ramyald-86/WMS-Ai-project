# FC_R145F17_PVZ_RECEIPT_EXCEL_SPLIT: full receiving-session Excel uses shared Excel-safe serialization.
from __future__ import annotations

import io
import uuid
from collections import defaultdict
from datetime import datetime, timezone, timedelta
from decimal import Decimal

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from openpyxl import Workbook
from pydantic import BaseModel, Field
from sqlalchemy import func, select, text
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.models import Container
from app.routers.pvz import PvzReceive, _product
from app.routers.receiving import _ensure_container, _ensure_location
from app.routers.warehouse_ui import scope
from app.schemas.inventory import OperationCreate, StockLine
from app.services.inventory import execute_operation
from app.services.stock_core_v2 import reconcile_stock_bucket
from app.services import pvz as pvz_service
from app.services.excel_safety import append_safe, workbook_bytes

router = APIRouter(prefix="/pvz/receiving-sessions", tags=["ПВЗ · единая приёмка"])
_SCHEMA_READY = False


class SessionCreate(BaseModel):
    client_id: uuid.UUID
    warehouse_id: uuid.UUID
    reference: str = Field(default="", max_length=300)
    performed_by: str = Field(default="Администратор", max_length=160)


class SessionScan(BaseModel):
    client_id: uuid.UUID
    warehouse_id: uuid.UUID
    barcode: str = Field(min_length=1, max_length=120)
    article: str = Field(default="", max_length=180)
    size: str = Field(default="", max_length=100)
    name: str = Field(default="", max_length=240)
    quantity: int = Field(default=1, ge=1, le=100)
    location_code: str = Field(default="ПРИЕМКА", max_length=120)
    container_code: str = Field(default="", max_length=160)
    comment: str = Field(default="", max_length=1000)
    performed_by: str = Field(default="Администратор", max_length=160)
    source: str = Field(default="SCAN", max_length=24)


class SessionAction(BaseModel):
    client_id: uuid.UUID
    warehouse_id: uuid.UUID
    performed_by: str = Field(default="Администратор", max_length=160)


class SessionMetaUpdate(BaseModel):
    client_id: uuid.UUID
    warehouse_id: uuid.UUID
    reference: str = Field(min_length=1, max_length=300)
    performed_by: str = Field(min_length=1, max_length=160)


def _required(value: str, message: str) -> str:
    clean = str(value or "").strip()
    if not clean:
        raise HTTPException(422, detail=message)
    return clean


def _assert_worker(session: dict, worker: str) -> str:
    clean = _required(worker, "Укажите сотрудника, который выполняет приёмку")
    assigned = str(session.get("performed_by") or "").strip()
    if assigned and clean.casefold() != assigned.casefold():
        raise HTTPException(409, detail=f"Эта приёмка закреплена за сотрудником: {assigned}")
    return assigned or clean


def install_schema(db: Session) -> None:
    global _SCHEMA_READY
    pvz_service.install_schema(db)
    db.execute(text("ALTER TABLE fc_pvz_batches ADD COLUMN IF NOT EXISTS receiving_session_id UUID NULL"))
    db.execute(text("CREATE INDEX IF NOT EXISTS ix_fc_pvz_batches_receiving_session ON fc_pvz_batches(receiving_session_id)"))
    db.execute(text('''
        CREATE TABLE IF NOT EXISTS fc_pvz_receiving_sessions (
          id UUID PRIMARY KEY,
          client_id UUID NOT NULL,
          warehouse_id UUID NOT NULL,
          status VARCHAR(16) NOT NULL DEFAULT 'DRAFT' CHECK(status IN ('DRAFT','COMPLETED','CANCELLED')),
          reference TEXT NOT NULL DEFAULT '',
          performed_by TEXT NOT NULL DEFAULT '',
          stock_operation_id UUID NULL,
          created_at TIMESTAMPTZ NOT NULL DEFAULT now(),
          completed_at TIMESTAMPTZ NULL,
          cancelled_at TIMESTAMPTZ NULL
        )
    '''))
    db.execute(text("CREATE INDEX IF NOT EXISTS ix_fc_pvz_receiving_sessions_scope ON fc_pvz_receiving_sessions(client_id,warehouse_id,status,created_at DESC)"))
    db.execute(text('''
        CREATE TABLE IF NOT EXISTS fc_pvz_receiving_boxes (
          id UUID PRIMARY KEY,
          session_id UUID NOT NULL REFERENCES fc_pvz_receiving_sessions(id) ON DELETE CASCADE,
          sequence INTEGER NOT NULL,
          container_code TEXT NOT NULL,
          container_name TEXT NOT NULL DEFAULT '',
          location_code TEXT NOT NULL DEFAULT 'ПРИЕМКА',
          created_at TIMESTAMPTZ NOT NULL DEFAULT now(),
          UNIQUE(session_id, sequence),
          UNIQUE(session_id, container_code)
        )
    '''))
    db.execute(text("CREATE INDEX IF NOT EXISTS ix_fc_pvz_receiving_boxes_session ON fc_pvz_receiving_boxes(session_id,sequence)"))
    db.execute(text('''
        CREATE TABLE IF NOT EXISTS fc_pvz_receiving_scans (
          id UUID PRIMARY KEY,
          session_id UUID NOT NULL REFERENCES fc_pvz_receiving_sessions(id) ON DELETE CASCADE,
          box_id UUID NOT NULL REFERENCES fc_pvz_receiving_boxes(id) ON DELETE CASCADE,
          sku_id UUID NOT NULL,
          barcode VARCHAR(120) NOT NULL,
          article VARCHAR(180) NOT NULL DEFAULT '',
          size VARCHAR(100) NOT NULL DEFAULT '',
          name TEXT NOT NULL DEFAULT '',
          location_code TEXT NOT NULL DEFAULT 'ПРИЕМКА',
          container_code TEXT NOT NULL,
          comment TEXT NOT NULL DEFAULT '',
          source VARCHAR(24) NOT NULL DEFAULT 'SCAN',
          performed_by TEXT NOT NULL DEFAULT '',
          scanned_at TIMESTAMPTZ NOT NULL DEFAULT now()
        )
    '''))
    db.execute(text("CREATE INDEX IF NOT EXISTS ix_fc_pvz_receiving_scans_session ON fc_pvz_receiving_scans(session_id,scanned_at,id)"))
    db.execute(text("ALTER TABLE fc_pvz_receiving_scans ALTER COLUMN box_id DROP NOT NULL"))
    db.execute(text("CREATE INDEX IF NOT EXISTS ix_fc_pvz_receiving_scans_box ON fc_pvz_receiving_scans(box_id,scanned_at,id)"))
    db.flush()
    _SCHEMA_READY = True


def _ensure_schema(db: Session) -> None:
    if not _SCHEMA_READY:
        install_schema(db)


def _session(db: Session, session_id: uuid.UUID, client_id: uuid.UUID, warehouse_id: uuid.UUID, *, lock=False):
    sql = "SELECT * FROM fc_pvz_receiving_sessions WHERE id=:id AND client_id=:c AND warehouse_id=:w"
    if lock:
        sql += " FOR UPDATE"
    row = db.execute(text(sql), {"id": str(session_id), "c": str(client_id), "w": str(warehouse_id)}).mappings().first()
    if not row:
        raise HTTPException(404, detail="Приёмка ПВЗ не найдена")
    return dict(row)


def _box_display(db: Session, client_id: uuid.UUID, warehouse_id: uuid.UUID, code: str) -> str:
    clean = str(code or "").strip()
    if not clean:
        return ""
    row = db.execute(select(Container).where(
        Container.client_id == client_id,
        Container.warehouse_id == warehouse_id,
        ((Container.barcode == clean) | (Container.name == clean)),
    )).scalars().first()
    return str((row.name or row.barcode) if row else clean)


def _ensure_live_storage_box(db: Session, client_id: uuid.UUID, warehouse_id: uuid.UUID, code: str, location_code: str):
    clean = str(code or "").strip()
    if not clean:
        return None, False
    existing = db.execute(select(Container).where(
        Container.client_id == client_id,
        Container.warehouse_id == warehouse_id,
        ((func.lower(Container.barcode) == clean.casefold()) | (func.lower(Container.name) == clean.casefold())),
    )).scalars().first()
    if existing:
        if str(existing.container_type or "").upper() == "DEFECT_BOX":
            raise HTTPException(409, detail="Это короб брака. Для обычной приёмки используйте другое имя")
        if str(existing.status or "ACTIVE").upper() != "ACTIVE":
            raise HTTPException(409, detail="Короб с таким именем уже есть, но он неактивен")
        return existing, False
    loc = _ensure_location(db, warehouse_id, str(location_code or "ПРИЕМКА").strip() or "ПРИЕМКА")
    box = _ensure_container(db, client_id, warehouse_id, loc.id, clean)
    if str(box.name or "").strip() != clean:
        box.name = clean
        db.flush()
    return box, True


def _get_or_create_session_box(db: Session, session, code: str, location_code: str):
    clean = str(code or "").strip()
    if not clean:
        raise HTTPException(422, detail="Сначала отсканируйте или укажите короб")
    row = db.execute(text('''SELECT * FROM fc_pvz_receiving_boxes
        WHERE session_id=:s AND lower(container_code)=lower(:code) LIMIT 1'''), {"s": str(session["id"]), "code": clean}).mappings().first()
    if row:
        return dict(row)
    seq = int(db.execute(text("SELECT COALESCE(MAX(sequence),0)+1 FROM fc_pvz_receiving_boxes WHERE session_id=:s"), {"s": str(session["id"])}).scalar() or 1)
    box_id = uuid.uuid4()
    display = _box_display(db, uuid.UUID(str(session["client_id"])), uuid.UUID(str(session["warehouse_id"])), clean)
    db.execute(text('''INSERT INTO fc_pvz_receiving_boxes(id,session_id,sequence,container_code,container_name,location_code)
        VALUES(:id,:s,:seq,:code,:name,:loc)'''), {
        "id": str(box_id), "s": str(session["id"]), "seq": seq, "code": clean,
        "name": display, "loc": str(location_code or "ПРИЕМКА").strip() or "ПРИЕМКА",
    })
    return {"id": box_id, "session_id": session["id"], "sequence": seq, "container_code": clean, "container_name": display, "location_code": location_code or "ПРИЕМКА"}


def _public(db: Session, session: dict):
    boxes = [dict(r) for r in db.execute(text('''SELECT b.*,
        COUNT(s.id)::int units, MIN(s.scanned_at) first_scan_at, MAX(s.scanned_at) last_scan_at
        FROM fc_pvz_receiving_boxes b LEFT JOIN fc_pvz_receiving_scans s ON s.box_id=b.id
        WHERE b.session_id=:s GROUP BY b.id ORDER BY b.sequence'''), {"s": str(session["id"])}).mappings().all()]
    scans = [dict(r) for r in db.execute(text('''SELECT s.*,b.sequence box_sequence,b.container_name
        FROM fc_pvz_receiving_scans s LEFT JOIN fc_pvz_receiving_boxes b ON b.id=s.box_id
        WHERE s.session_id=:s ORDER BY s.scanned_at,s.id'''), {"s": str(session["id"])}).mappings().all()]
    return {
        "id": str(session["id"]), "status": session["status"], "reference": session.get("reference") or "",
        "performed_by": session.get("performed_by") or "", "created_at": session.get("created_at"),
        "completed_at": session.get("completed_at"), "stock_operation_id": str(session["stock_operation_id"]) if session.get("stock_operation_id") else None,
        "boxes": [{**b, "id": str(b["id"]), "session_id": str(b["session_id"])} for b in boxes],
        "scans": [{**r, "id": str(r["id"]), "session_id": str(r["session_id"]), "box_id": str(r["box_id"]) if r.get("box_id") else None, "sku_id": str(r["sku_id"])} for r in scans],
        "units": len(scans), "box_count": len(boxes),
        "cell_units": sum(1 for r in scans if not r.get("box_id")),
    }


@router.post("")
def create_session(payload: SessionCreate, db: Session = Depends(get_db)):
    scope(db, payload.client_id, payload.warehouse_id); _ensure_schema(db)
    reference = _required(payload.reference, "Укажите название приёмки")
    worker = _required(payload.performed_by, "Укажите сотрудника, который выполняет приёмку")
    existing = db.execute(text('''SELECT * FROM fc_pvz_receiving_sessions
        WHERE client_id=:c AND warehouse_id=:w AND status='DRAFT' AND lower(performed_by)=lower(:p)
        ORDER BY created_at DESC LIMIT 1'''), {"c": str(payload.client_id), "w": str(payload.warehouse_id), "p": worker}).mappings().first()
    if existing:
        current = dict(existing)
        if str(current.get("reference") or "").strip().casefold() != reference.casefold():
            raise HTTPException(409, detail=f"У сотрудника {worker} уже открыта приёмка: {current.get('reference') or str(current['id'])[:8]}")
        return _public(db, current)
    sid = uuid.uuid4()
    db.execute(text('''INSERT INTO fc_pvz_receiving_sessions(id,client_id,warehouse_id,status,reference,performed_by)
        VALUES(:id,:c,:w,'DRAFT',:r,:p)'''), {"id": str(sid), "c": str(payload.client_id), "w": str(payload.warehouse_id), "r": reference, "p": worker})
    db.commit()
    return _public(db, _session(db, sid, payload.client_id, payload.warehouse_id))


@router.patch("/{session_id}/meta")
def update_session_meta(session_id: uuid.UUID, payload: SessionMetaUpdate, db: Session = Depends(get_db)):
    scope(db, payload.client_id, payload.warehouse_id); _ensure_schema(db)
    session = _session(db, session_id, payload.client_id, payload.warehouse_id, lock=True)
    if session["status"] != "DRAFT":
        raise HTTPException(409, detail="Название и сотрудника можно менять только у незавершённой приёмки")
    reference = _required(payload.reference, "Укажите название приёмки")
    worker = _required(payload.performed_by, "Укажите сотрудника, который выполняет приёмку")
    conflict = db.execute(text('''SELECT id,reference FROM fc_pvz_receiving_sessions
        WHERE client_id=:c AND warehouse_id=:w AND status='DRAFT' AND id<>:id AND lower(performed_by)=lower(:p)
        ORDER BY created_at DESC LIMIT 1'''), {"c": str(payload.client_id), "w": str(payload.warehouse_id), "id": str(session_id), "p": worker}).mappings().first()
    if conflict:
        raise HTTPException(409, detail=f"У сотрудника {worker} уже открыта другая приёмка: {conflict.get('reference') or str(conflict['id'])[:8]}")
    db.execute(text("UPDATE fc_pvz_receiving_sessions SET reference=:r,performed_by=:p WHERE id=:id"), {"r": reference, "p": worker, "id": str(session_id)})
    # Existing draft scans follow the reassigned worker too. This is important when adopting a legacy admin-owned draft.
    db.execute(text("UPDATE fc_pvz_receiving_scans SET performed_by=:p WHERE session_id=:id"), {"p": worker, "id": str(session_id)})
    db.commit()
    return {"ok": True, "session": _public(db, _session(db, session_id, payload.client_id, payload.warehouse_id))}


@router.get("/active")
def active_session(client_id: uuid.UUID, warehouse_id: uuid.UUID, performed_by: str = "Администратор", db: Session = Depends(get_db)):
    scope(db, client_id, warehouse_id); _ensure_schema(db)
    row = db.execute(text('''SELECT * FROM fc_pvz_receiving_sessions
        WHERE client_id=:c AND warehouse_id=:w AND status='DRAFT' AND lower(performed_by)=lower(:p)
        ORDER BY created_at DESC LIMIT 1'''), {"c": str(client_id), "w": str(warehouse_id), "p": str(performed_by or "Администратор").strip()}).mappings().first()
    return {"session": _public(db, dict(row)) if row else None}


@router.get("/{session_id}")
def get_session(session_id: uuid.UUID, client_id: uuid.UUID, warehouse_id: uuid.UUID, db: Session = Depends(get_db)):
    scope(db, client_id, warehouse_id); _ensure_schema(db)
    return _public(db, _session(db, session_id, client_id, warehouse_id))


@router.post("/{session_id}/scan")
def add_scan(session_id: uuid.UUID, payload: SessionScan, db: Session = Depends(get_db)):
    scope(db, payload.client_id, payload.warehouse_id); _ensure_schema(db)
    session = _session(db, session_id, payload.client_id, payload.warehouse_id, lock=True)
    if session["status"] != "DRAFT":
        raise HTTPException(409, detail="Эта приёмка уже завершена")
    worker = _assert_worker(session, payload.performed_by)
    # Resolve/validate product using the same local PVZ rules, but do not touch inventory yet.
    probe = PvzReceive(client_id=payload.client_id, warehouse_id=payload.warehouse_id,
        barcode=payload.barcode, article=payload.article, size=payload.size, name=payload.name,
        quantity=1, location_code=payload.location_code, container_code=payload.container_code or "PVZ-CELL-DIRECT",
        reference=session.get("reference") or "", comment=payload.comment, performed_by=payload.performed_by)
    sku = _product(db, probe)
    clean_container = str(payload.container_code or "").strip()
    container_created = False
    if clean_container:
        _, container_created = _ensure_live_storage_box(
            db, payload.client_id, payload.warehouse_id, clean_container, payload.location_code
        )
    box = _get_or_create_session_box(db, session, clean_container, payload.location_code) if clean_container else None
    base_time = datetime.now(timezone.utc)
    ids = []
    for i in range(int(payload.quantity)):
        scan_id = uuid.uuid4(); ids.append(scan_id)
        scanned_at = base_time + timedelta(microseconds=i)
        db.execute(text('''INSERT INTO fc_pvz_receiving_scans(
            id,session_id,box_id,sku_id,barcode,article,size,name,location_code,container_code,comment,source,performed_by,scanned_at)
            VALUES(:id,:s,:box,:sku,:bc,:a,:z,:n,:loc,:cc,:comment,:source,:p,:at)'''), {
            "id": str(scan_id), "s": str(session_id), "box": str(box["id"]) if box else None, "sku": str(sku.id),
            "bc": str(sku.barcode or ""), "a": str(sku.article or ""), "z": str(sku.size or ""), "n": str(sku.name or ""),
            "loc": str(payload.location_code or "ПРИЕМКА").strip() or "ПРИЕМКА", "cc": str(box["container_code"]) if box else "",
            "comment": payload.comment.strip(), "source": payload.source.strip().upper() or "SCAN", "p": worker, "at": scanned_at,
        })
    db.commit()
    current = _public(db, _session(db, session_id, payload.client_id, payload.warehouse_id))
    return {"ok": True, "added": len(ids), "scanned_at": base_time,
            "box_sequence": int(box["sequence"]) if box else None,
            "container": (box.get("container_name") or box.get("container_code")) if box else None,
            "container_created": container_created,
            "storage_kind": "BOX" if box else "LOCATION",
            "location_code": str(payload.location_code or "ПРИЕМКА").strip() or "ПРИЕМКА",
            "barcode": str(sku.barcode or ""), "session": current}


@router.delete("/{session_id}/scans/{scan_id}")
def delete_scan(session_id: uuid.UUID, scan_id: uuid.UUID, client_id: uuid.UUID, warehouse_id: uuid.UUID, db: Session = Depends(get_db)):
    scope(db, client_id, warehouse_id); _ensure_schema(db)
    session = _session(db, session_id, client_id, warehouse_id, lock=True)
    if session["status"] != "DRAFT":
        raise HTTPException(409, detail="Изменять можно только незавершённую приёмку")
    result = db.execute(text("DELETE FROM fc_pvz_receiving_scans WHERE id=:id AND session_id=:s"), {"id": str(scan_id), "s": str(session_id)})
    if not int(result.rowcount or 0):
        raise HTTPException(404, detail="Единица приёмки не найдена")
    db.commit()
    return {"ok": True, "session": _public(db, _session(db, session_id, client_id, warehouse_id))}


@router.post("/{session_id}/cancel")
def cancel_session(session_id: uuid.UUID, payload: SessionAction, db: Session = Depends(get_db)):
    scope(db, payload.client_id, payload.warehouse_id); _ensure_schema(db)
    session = _session(db, session_id, payload.client_id, payload.warehouse_id, lock=True)
    if session["status"] != "DRAFT":
        raise HTTPException(409, detail="Приёмка уже закрыта")
    _assert_worker(session, payload.performed_by)
    db.execute(text("UPDATE fc_pvz_receiving_sessions SET status='CANCELLED',cancelled_at=now() WHERE id=:id"), {"id": str(session_id)})
    db.commit(); return {"ok": True}


def _finalize(db: Session, session: dict, client_id: uuid.UUID, warehouse_id: uuid.UUID, actor: str):
    scans = [dict(r) for r in db.execute(text('''SELECT s.*,b.sequence box_sequence,b.container_name
        FROM fc_pvz_receiving_scans s LEFT JOIN fc_pvz_receiving_boxes b ON b.id=s.box_id
        WHERE s.session_id=:s ORDER BY s.scanned_at,s.id'''), {"s": str(session["id"])}).mappings().all()]
    if not scans:
        raise HTTPException(409, detail="В приёмке нет ни одной отсканированной единицы")

    resolved = {}
    groups = defaultdict(lambda: {"quantity": Decimal(0), "first_at": None, "last_at": None, "comment": ""})
    for row in scans:
        key_place = (row["location_code"], row["container_code"])
        if key_place not in resolved:
            loc = _ensure_location(db, warehouse_id, str(row["location_code"] or "ПРИЕМКА").strip() or "ПРИЕМКА")
            container_code = str(row["container_code"] or "").strip()
            box = _ensure_container(db, client_id, warehouse_id, loc.id, container_code) if container_code else None
            resolved[key_place] = (loc, box)
        loc, box = resolved[key_place]
        key = (str(row["sku_id"]), str(loc.id), str(box.id) if box else "")
        g = groups[key]; g["sku_id"] = row["sku_id"]; g["location"] = loc; g["box"] = box
        g["barcode"] = row["barcode"]; g["article"] = row["article"]; g["size"] = row["size"]; g["name"] = row["name"]
        g["quantity"] += Decimal(1); g["comment"] = row.get("comment") or g["comment"]
        at = row["scanned_at"]; g["first_at"] = at if g["first_at"] is None or at < g["first_at"] else g["first_at"]
        g["last_at"] = at if g["last_at"] is None or at > g["last_at"] else g["last_at"]

    op_lines = [StockLine(sku_id=uuid.UUID(str(g["sku_id"])), quantity=g["quantity"],
        destination_location_id=g["location"].id, destination_container_id=g["box"].id if g["box"] else None,
        reason=f"PVZ_RECEIVING_SESSION:{session['id']}") for g in groups.values()]
    op = execute_operation(db, OperationCreate(client_id=client_id, warehouse_id=warehouse_id,
        operation_type="RECEIVING", idempotency_key=f"PVZ_SESSION:{session['id']}", performed_by=actor,
        reference=(session.get("reference") or f"ПВЗ {session['id']}")[:255],
        comment="Единая приёмка ПВЗ по местам хранения · ЗОНА:MAIN", lines=op_lines), commit=False)

    for g in groups.values():
        reconcile_stock_bucket(db, client_id=client_id, warehouse_id=warehouse_id, sku_id=uuid.UUID(str(g["sku_id"])),
            location_id=g["location"].id, container_id=g["box"].id if g["box"] else None,
            preferred_bucket="STORAGE_ONLY", lock=True)

    batch_id = uuid.uuid4()
    db.execute(text('''INSERT INTO fc_pvz_batches(id,client_id,warehouse_id,kind,status,reference,performed_by,stock_operation_id,receiving_session_id,created_at)
        VALUES(:id,:c,:w,'GOOD','COMPLETED',:ref,:actor,:op,:sid,:created)'''), {
        "id": str(batch_id), "c": str(client_id), "w": str(warehouse_id), "ref": session.get("reference") or "",
        "actor": actor, "op": str(op.id), "sid": str(session["id"]), "created": session.get("created_at") or datetime.now(timezone.utc),
    })
    for g in groups.values():
        db.execute(text('''INSERT INTO fc_pvz_lines(id,batch_id,sku_id,barcode,article,size,name,quantity,location_id,container_id,comment,created_at)
            VALUES(:id,:b,:s,:bc,:a,:z,:n,:q,:loc,:box,:comment,:created)'''), {
            "id": str(uuid.uuid4()), "b": str(batch_id), "s": str(g["sku_id"]), "bc": g["barcode"], "a": g["article"],
            "z": g["size"], "n": g["name"], "q": g["quantity"], "loc": str(g["location"].id),
            "box": str(g["box"].id) if g["box"] else None, "comment": g["comment"], "created": g["first_at"],
        })
    db.execute(text("UPDATE fc_pvz_receiving_sessions SET status='COMPLETED',completed_at=now(),stock_operation_id=:op WHERE id=:id"),
               {"op": str(op.id), "id": str(session["id"])})
    return batch_id, op.id


@router.post("/{session_id}/complete")
def complete_session(session_id: uuid.UUID, payload: SessionAction, db: Session = Depends(get_db)):
    scope(db, payload.client_id, payload.warehouse_id); _ensure_schema(db)
    session = _session(db, session_id, payload.client_id, payload.warehouse_id, lock=True)
    if session["status"] == "COMPLETED":
        return {"ok": True, "replayed": True, "session": _public(db, session)}
    if session["status"] != "DRAFT":
        raise HTTPException(409, detail="Эта приёмка отменена")
    worker = _assert_worker(session, payload.performed_by)
    batch_id, operation_id = _finalize(db, session, payload.client_id, payload.warehouse_id, worker)
    db.commit()
    final = _public(db, _session(db, session_id, payload.client_id, payload.warehouse_id))
    return {"ok": True, "batch_id": str(batch_id), "operation_id": str(operation_id), "session": final,
            "export_url": f"/pvz/receiving-sessions/{session_id}/export.xlsx?client_id={payload.client_id}&warehouse_id={payload.warehouse_id}"}


def _workbook(db: Session, session: dict) -> bytes:
    data = _public(db, session)
    wb = Workbook(); ws = wb.active; ws.title = 'Сводка'
    append_safe(ws,['Приёмка ПВЗ', session.get('reference') or str(session['id'])])
    append_safe(ws,['Статус', session['status']]); append_safe(ws,['Сотрудник', session.get('performed_by') or ''])
    append_safe(ws,['Начало', session.get('created_at')]); append_safe(ws,['Завершение', session.get('completed_at')])
    append_safe(ws,['Коробов', data['box_count']]); append_safe(ws,['Напрямую в ячейки', data['cell_units']]); append_safe(ws,['Единиц', data['units']])

    wb_boxes = wb.create_sheet('Короба')
    append_safe(wb_boxes,['Короб №','Название короба','Код короба','Ячейка / палета','Единиц','Первый скан','Последний скан'])
    for b in data['boxes']:
        append_safe(wb_boxes,[b['sequence'], b.get('container_name') or b.get('container_code'), b.get('container_code'), b.get('location_code'), b.get('units'), b.get('first_scan_at'), b.get('last_scan_at')])

    wb_units = wb.create_sheet('Единицы')
    append_safe(wb_units,['Время принятия','Короб №','Короб / без короба','Ячейка / палета','ШК','Артикул','Размер','Товар','Источник','Сотрудник','Комментарий'])
    for r in data['scans']:
        append_safe(wb_units,[r.get('scanned_at'), r.get('box_sequence'), r.get('container_name') or r.get('container_code') or 'Без короба', r.get('location_code'),
                         r.get('barcode'), r.get('article'), r.get('size'), r.get('name'), r.get('source'), r.get('performed_by'), r.get('comment')])

    agg = defaultdict(int)
    for r in data['scans']:
        place = r.get('container_name') or r.get('container_code') or 'Без короба'
        k = (place, r.get('location_code'), r.get('barcode'), r.get('article'), r.get('size'), r.get('name'))
        agg[k] += 1
    wb_items = wb.create_sheet('Товары')
    append_safe(wb_items,['Короб / без короба','Ячейка / палета','ШК','Артикул','Размер','Товар','Количество'])
    for k, qty in agg.items():
        append_safe(wb_items,[*k, qty])

    for sheet in wb.worksheets:
        if sheet.max_row > 1:
            sheet.freeze_panes = 'A2'
        for col in sheet.columns:
            letter = col[0].column_letter
            sheet.column_dimensions[letter].width = min(42, max(12, max(len(str(x.value or '')) for x in col) + 2))
    return workbook_bytes(wb)


@router.get("/{session_id}/export.xlsx")
def export_session(session_id: uuid.UUID, client_id: uuid.UUID, warehouse_id: uuid.UUID, db: Session = Depends(get_db)):
    scope(db, client_id, warehouse_id); _ensure_schema(db)
    session = _session(db, session_id, client_id, warehouse_id)
    raw = _workbook(db, session)
    return Response(raw, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="FulfillCore_PVZ_receipt_{str(session_id)[:8]}.xlsx"'})
