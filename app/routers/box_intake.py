"""Authenticated, scoped endpoints for box-only receiving and Excel workflows."""
import uuid
from pathlib import Path
from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field
from sqlalchemy import text
from sqlalchemy.orm import Session
from app.core.database import get_db
from app.models import Client, Warehouse, Container
from app.services import box_intake as service, box_intake_edit as edits
from app.services.box_intake_rules import read_sheet
from app.services.workflow_rules import MAX_BYTES
from app.services.excel_safety import append_safe, sanitize_workbook
from app.routers.workflow import excel, ensure_schema as ensure_workflow_schema

router=APIRouter(prefix='/box-intake',tags=['Приёмка и сборка коробов'])


def scope(db,c,w):
    if not db.get(Client,c) or not db.get(Warehouse,w):raise HTTPException(404,detail='Клиент или склад не найден')
    if not db.execute(text("SELECT to_regclass('fc_box_intakes')")).scalar():raise HTTPException(503,detail='Завершите установку r68: схема приёмки коробов отсутствует')
    service.ensure_r145f7_schema(db)


def actor(req):return str(getattr(req.state,'username','') or 'Администратор')[:120]

def checked(fn,*args):
    try:return fn(*args)
    except ValueError as e:raise HTTPException(422,detail=str(e))

def parse(file,kind):
    try:return read_sheet(file.file.read(MAX_BYTES+1),kind)
    except ValueError as e:raise HTTPException(422,detail=str(e))


class Context(BaseModel):
    client_id:uuid.UUID
    warehouse_id:uuid.UUID

class Start(Context):
    id:uuid.UUID
    name:str=Field(default='',max_length=240)
    to_pallets:bool=True
    section:str=Field(default='',max_length=80)
    first_row:int|None=Field(default=1,ge=1,le=900)

class Scan(Context):
    box_code:str=Field(min_length=1,max_length=120)
    receiver:str|None=Field(default=None,max_length=120)

class Extend(Context):
    count:int=Field(ge=1,le=800)
    receiver:str|None=Field(default=None,max_length=120)

class PalletAddress(Context):
    section:str=Field(default='',max_length=80)
    first_row:int|None=Field(default=None,ge=1,le=900)

class Confirm(Context):
    confirm_occupied:bool=False

class TsdReceiptOpen(Context):
    name:str=Field(min_length=1,max_length=240)
    receiver:str=Field(min_length=1,max_length=120)

class TsdPalletSelect(Context):
    pallet_code:str=Field(min_length=1,max_length=120)
    receiver:str=Field(min_length=1,max_length=120)

def operator(req, value=None):
    return (str(value or '').strip() or actor(req))[:120]


@router.post('/receipts')
def create(payload:Start,request:Request,db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id)
    r=checked(service.new_receipt,db,payload.id,payload.client_id,payload.warehouse_id,payload.name,payload.to_pallets,payload.section,payload.first_row,actor(request))
    db.commit();return r


@router.post('/receipts/open-or-create')
def open_or_create_tsd(payload:TsdReceiptOpen,request:Request,db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id)
    r=service.open_or_create_tsd_receipt(db,payload.client_id,payload.warehouse_id,payload.name,operator(request,payload.receiver))
    db.commit();return r

@router.post('/receipts/{rid}/tsd-pallet')
def select_tsd_pallet(rid:uuid.UUID,payload:TsdPalletSelect,request:Request,db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id)
    r=service.select_tsd_pallet(db,rid,payload.client_id,payload.warehouse_id,payload.pallet_code,operator(request,payload.receiver))
    db.commit();return r

@router.get('/receipts')
def receipts(client_id:uuid.UUID,warehouse_id:uuid.UUID,include_deleted:bool=False,db:Session=Depends(get_db)):
    scope(db,client_id,warehouse_id);return {'rows':service.list_receipts(db,client_id,warehouse_id,include_deleted)}

@router.get('/receipts/{rid}')
def receipt(rid:uuid.UUID,client_id:uuid.UUID,warehouse_id:uuid.UUID,db:Session=Depends(get_db)):
    scope(db,client_id,warehouse_id);return service.detail(db,rid,client_id,warehouse_id)

@router.post('/receipts/{rid}/scan')
def scan(rid:uuid.UUID,payload:Scan,request:Request,db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id)
    try:r=service.scan(db,rid,payload.client_id,payload.warehouse_id,payload.box_code,operator(request,payload.receiver))
    except ValueError as e:raise HTTPException(422,detail=str(e))
    db.commit();return r

@router.post('/receipts/{rid}/next-pallet')
def next_pallet(rid:uuid.UUID,payload:Context,request:Request,db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id);service.next_pallet(db,rid,payload.client_id,payload.warehouse_id,actor(request));db.commit();return {'ok':True}

@router.post('/receipts/{rid}/pallets/{pid}/extend')
def extend(rid:uuid.UUID,pid:uuid.UUID,payload:Extend,request:Request,db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id);service.extend_pallet(db,rid,pid,payload.client_id,payload.warehouse_id,payload.count,operator(request,payload.receiver));db.commit();return {'ok':True}

@router.patch('/receipts/{rid}/pallets/{pid}')
def address(rid:uuid.UUID,pid:uuid.UUID,payload:PalletAddress,request:Request,db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id);checked(service.edit_pallet,db,rid,pid,payload.client_id,payload.warehouse_id,payload.section,payload.first_row,actor(request));db.commit();return {'ok':True}

@router.post('/receipts/{rid}/close')
def close(rid:uuid.UUID,payload:Context,request:Request,db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id);service.close_receipt(db,rid,payload.client_id,payload.warehouse_id,actor(request));db.commit();return {'ok':True}

@router.post('/receipts/{rid}/layout')
def layout(rid:uuid.UUID,request:Request,client_id:uuid.UUID=Form(...),warehouse_id:uuid.UUID=Form(...),file:UploadFile=File(...),db:Session=Depends(get_db)):
    scope(db,client_id,warehouse_id);rows=parse(file,'layout');r=service.preview_layout(db,rid,client_id,warehouse_id,rows,actor(request));db.commit();return r

@router.post('/layout/{plan_id}/confirm')
def confirm_layout(plan_id:uuid.UUID,payload:Confirm,request:Request,db:Session=Depends(get_db)):
    ensure_workflow_schema(db);scope(db,payload.client_id,payload.warehouse_id)
    try:
        r=service.commit_layout(db,plan_id,payload.client_id,payload.warehouse_id,payload.confirm_occupied,actor(request));db.commit();return r
    except Exception:
        db.rollback();raise

@router.post('/picks')
def create_pick(request:Request,client_id:uuid.UUID=Form(...),warehouse_id:uuid.UUID=Form(...),number:str=Form(...),file:UploadFile=File(...),db:Session=Depends(get_db)):
    scope(db,client_id,warehouse_id);rows=parse(file,'boxes');r=checked(service.make_box_sheet,db,client_id,warehouse_id,number,rows,actor(request));db.commit();return r

@router.get('/picks')
def picks(client_id:uuid.UUID,warehouse_id:uuid.UUID,db:Session=Depends(get_db)):
    scope(db,client_id,warehouse_id)
    rows=db.execute(text('''SELECT id,number,created_at,result_json->>'requested' requested,result_json->>'found' found
       FROM fc_box_sheets WHERE client_id=:c AND warehouse_id=:w ORDER BY created_at DESC LIMIT 100'''),service.params(client_id,warehouse_id)).mappings().all()
    return {'rows':[dict(r) for r in rows]}


def route_workbook(rows,title,errors=()):
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active;ws.title=title
    ws.append(['Порядок','Секция','Ряд','Палета / ячейка','Место в ряду','Короб','Коробов'])
    for i,r in enumerate(rows,1):
        append_safe(ws,[i,r.get('route_section') or '',r.get('route_row') or '',r['location_code'],r.get('route_position') or '',r.get('container_name') or r.get('box_code') or r.get('container'),1])
    if errors:
        err=wb.create_sheet('Не найдены');err.append(['Строка файла','Короб','Причина'])
        for r in errors:append_safe(err,[r['row'],r['container'],r['error']])
    sanitize_workbook(wb)
    return wb

@router.get('/picks/{sid}/export.xlsx')
def export_pick(sid:uuid.UUID,client_id:uuid.UUID,warehouse_id:uuid.UUID,db:Session=Depends(get_db)):
    scope(db,client_id,warehouse_id);data=service.get_box_sheet(db,sid,client_id,warehouse_id)
    return excel(route_workbook(data['rows'],'Маршрут коробов',data['errors']),'box_pick_'+str(sid)[:8]+'.xlsx')

@router.get('/receipts/{rid}/export.xlsx')
def export_receipt(rid:uuid.UUID,client_id:uuid.UUID,warehouse_id:uuid.UUID,db:Session=Depends(get_db)):
    from app.services.picking_route import route_key
    scope(db,client_id,warehouse_id);data=service.detail(db,rid,client_id,warehouse_id)
    rows=sorted([r for r in data['boxes'] if r['status']=='ACTIVE'],key=route_key)
    wb=route_workbook(rows,'Принятые короба');ws=wb.active;ws.cell(1,8,'Учтено товара, шт.')
    for i,r in enumerate(rows,2):ws.cell(i,8,float(r['quantity']))
    return excel(wb,'box_receipt_'+str(rid)[:8]+'.xlsx')

@router.get('/templates/{kind}.xlsx')
def template(kind:str):
    allowed={'boxes','goods','goods_by_box','box_layout','fbs','transfer'}
    if kind not in allowed:raise HTTPException(404,detail='Шаблон не найден')
    path=Path(__file__).resolve().parents[1]/'static'/'templates'/('r68_'+kind+'.xlsx')
    return FileResponse(path,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',filename=path.name)


class EmptyMove(Context):
    destination_location_id:uuid.UUID

@router.post('/boxes/{bid}/move-empty')
def move_empty(bid:uuid.UUID,payload:EmptyMove,request:Request,db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id)
    result=service.empty_box_change(db,bid,payload.client_id,payload.warehouse_id,actor(request),destination=payload.destination_location_id)
    db.commit();return result

@router.post('/boxes/{bid}/archive-empty')
def archive_empty(bid:uuid.UUID,payload:Context,request:Request,db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id)
    result=service.empty_box_change(db,bid,payload.client_id,payload.warehouse_id,actor(request),archive=True)
    db.commit();return result

class BulkArchive(Context):
    ids:list[uuid.UUID]=Field(min_length=1,max_length=1000)

@router.post('/archive-empty')
def archive_many(payload:BulkArchive,request:Request,db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id);deleted=[];skipped=[]
    for bid in dict.fromkeys(payload.ids):
        try:
            with db.begin_nested():service.empty_box_change(db,bid,payload.client_id,payload.warehouse_id,actor(request),archive=True)
            deleted.append({'id':str(bid)})
        except HTTPException as e:skipped.append({'id':str(bid),'reason':str(e.detail)})
    db.commit();return {'deleted':deleted,'skipped':skipped}


class ReceiptMove(Context):
    pallet_id:uuid.UUID
    allow_expand:bool=False
    expected_location_id:uuid.UUID | None=None

class NewPallet(Context):
    id:uuid.UUID


def edit_request(db,fn,*args,snapshot_context=None):
    try:
        result=checked(fn,*args)
        if snapshot_context:result['receipt_snapshot']=service.detail(db,*snapshot_context)
        db.commit();return result
    except Exception:
        db.rollback();raise


@router.post('/receipts/{rid}/pallets')
def receipt_new_pallet(rid:uuid.UUID,payload:NewPallet,request:Request,db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id)
    return edit_request(db,edits.new_pallet,db,rid,payload.id,payload.client_id,payload.warehouse_id,actor(request),snapshot_context=(rid,payload.client_id,payload.warehouse_id))


@router.post('/receipts/{rid}/boxes/{bid}/archive')
def receipt_box_archive(rid:uuid.UUID,bid:uuid.UUID,payload:Context,request:Request,db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id)
    return edit_request(db,edits.archive_box,db,rid,bid,payload.client_id,payload.warehouse_id,actor(request),snapshot_context=(rid,payload.client_id,payload.warehouse_id))


@router.post('/receipts/{rid}/boxes/{bid}/move')
def receipt_box_move(rid:uuid.UUID,bid:uuid.UUID,payload:ReceiptMove,request:Request,db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id)
    return edit_request(db,edits.move_box,db,rid,bid,payload.client_id,payload.warehouse_id,payload.pallet_id,payload.allow_expand,payload.expected_location_id,actor(request),snapshot_context=(rid,payload.client_id,payload.warehouse_id))


@router.post('/receipts/{rid}/boxes/{bid}/restore')
def receipt_box_restore(rid:uuid.UUID,bid:uuid.UUID,payload:Context,request:Request,db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id)
    return edit_request(db,edits.restore_box,db,rid,bid,payload.client_id,payload.warehouse_id,actor(request),snapshot_context=(rid,payload.client_id,payload.warehouse_id))


@router.post('/receipts/{rid}/pallets/{pid}/select')
def receipt_pallet_select(rid:uuid.UUID,pid:uuid.UUID,payload:Context,request:Request,db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id)
    return edit_request(db,edits.select_pallet,db,rid,pid,payload.client_id,payload.warehouse_id,actor(request),snapshot_context=(rid,payload.client_id,payload.warehouse_id))


@router.post('/receipts/{rid}/pallets/{pid}/archive')
def receipt_pallet_archive(rid:uuid.UUID,pid:uuid.UUID,payload:Context,request:Request,db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id)
    return edit_request(db,edits.remove_pallet,db,rid,pid,payload.client_id,payload.warehouse_id,actor(request),snapshot_context=(rid,payload.client_id,payload.warehouse_id))


@router.post('/receipts/{rid}/archive')
def receipt_archive(rid:uuid.UUID,payload:Context,request:Request,db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id)
    return edit_request(db,edits.remove_receipt,db,rid,payload.client_id,payload.warehouse_id,actor(request),snapshot_context=(rid,payload.client_id,payload.warehouse_id))


@router.post('/picks/{sid}/archive')
def sheet_archive(sid:uuid.UUID,payload:Context,request:Request,db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id)
    return edit_request(db,edits.remove_sheet,db,sid,payload.client_id,payload.warehouse_id,actor(request))
