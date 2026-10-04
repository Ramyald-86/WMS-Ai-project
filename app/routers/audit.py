from __future__ import annotations

import io
import uuid
from datetime import datetime, timezone
from decimal import Decimal
from typing import Any

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import bindparam, text
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.services.stock_core_v2 import stock_truth_for_sku

router = APIRouter(prefix="/audit", tags=["audit"])


def _num(value: Any) -> float:
    try:
        return float(Decimal(str(value or 0)))
    except Exception:
        return 0.0


def _iso(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.astimezone(timezone.utc).isoformat()
    return str(value)



def _history_date(value: object):
    """Return a real date value for Excel filtering while preserving the original timestamp."""
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    raw = str(value).strip()
    try:
        return datetime.fromisoformat(raw.replace("Z", "+00:00")).date()
    except Exception:
        return raw[:10] if len(raw) >= 10 else raw


def _receipt_name(row: dict[str, Any]) -> str:
    """Human receipt name only for inbound rows; blank for reserve/shipping/etc."""
    event_type = str(row.get("event_type") or "").upper()
    if event_type not in {"RECEIVING", "RECEIVING_IMPORT", "PVZ_RECEIVING"}:
        return ""
    comment = str(row.get("comment") or "")
    import re
    match = re.search(r"номер\s+при[её]мки\s*:\s*([^;]+)", comment, flags=re.I)
    if match:
        value = match.group(1).strip()
        if value:
            return value
    ref = str(row.get("reference") or "").strip()
    if event_type == "PVZ_RECEIVING" and re.fullmatch(r"ПВЗ\s+[0-9a-fA-F-]{32,36}", ref):
        return "Приёмка ПВЗ"
    for prefix in ("Приёмка ТСД · ", "Приемка ТСД · ", "Приёмка ТСД - ", "Приемка ТСД - "):
        if ref.startswith(prefix):
            return ref[len(prefix):].strip()
    if ref in {"Приёмка ТСД", "Приемка ТСД", "Приёмка", "Приемка"}:
        return ""
    return ref

def _table_exists(db: Session, name: str) -> bool:
    return bool(db.execute(text("SELECT to_regclass(:name) IS NOT NULL"), {"name": name}).scalar())


def _filters_sql(*, sku_id=None, barcode=None, location_id=None, container_id=None, order_id=None, actor=None):
    clauses: list[str] = []
    params: dict[str, Any] = {}
    if sku_id:
        clauses.append("l.sku_id = :sku_id")
        params["sku_id"] = str(sku_id)
    if barcode:
        clauses.append("s.barcode = :barcode")
        params["barcode"] = str(barcode).strip()
    if location_id:
        clauses.append("(l.source_location_id = :location_id OR l.destination_location_id = :location_id)")
        params["location_id"] = str(location_id)
    if container_id:
        clauses.append("(l.source_container_id = :container_id OR l.destination_container_id = :container_id)")
        params["container_id"] = str(container_id)
    if order_id:
        clauses.append("COALESCE(o.reference,'') ILIKE :order_id")
        params["order_id"] = f"%{str(order_id).strip()}%"
    if actor:
        clauses.append("COALESCE(o.performed_by,'') ILIKE :actor")
        params["actor"] = f"%{str(actor).strip()}%"
    return clauses, params


def _physical_events(
    db: Session,
    *,
    client_id: uuid.UUID,
    warehouse_id: uuid.UUID,
    sku_id: uuid.UUID | None = None,
    barcode: str | None = None,
    location_id: uuid.UUID | None = None,
    container_id: uuid.UUID | None = None,
    order_id: str | None = None,
    actor: str | None = None,
    operation_type: str | None = None,
    limit: int = 500,
) -> list[dict[str, Any]]:
    clauses, params = _filters_sql(
        sku_id=sku_id, barcode=barcode, location_id=location_id,
        container_id=container_id, order_id=order_id, actor=actor,
    )
    clauses.insert(0, "o.client_id = :client_id")
    clauses.insert(1, "o.warehouse_id = :warehouse_id")
    params.update({"client_id": str(client_id), "warehouse_id": str(warehouse_id), "limit": limit})
    if operation_type:
        clauses.append("o.operation_type = :operation_type")
        params["operation_type"] = str(operation_type).upper()
    sql = f"""
        SELECT o.id AS operation_id, o.operation_type, o.status, o.performed_by, o.reference,
               o.idempotency_key, o.comment, o.created_at, l.id AS line_id, l.sku_id, l.quantity, l.reason,
               tw.data->>'first_item_at' AS tsd_box_first_at, tw.data->>'last_item_at' AS tsd_box_last_at,
               ((tw.data->'lines')->(CAST(l.sku_id AS text)))->>'first_scanned_at' AS tsd_first_scanned_at,
               ((tw.data->'lines')->(CAST(l.sku_id AS text)))->>'last_scanned_at' AS tsd_last_scanned_at,
               s.sku_code, s.article, s.name AS sku_name, s.size, s.barcode,
               l.source_location_id, sl.code AS source_location,
               l.destination_location_id, dl.code AS destination_location,
               l.source_container_id, COALESCE(sc.name, sc.barcode) AS source_container,
               l.destination_container_id, COALESCE(dc.name, dc.barcode) AS destination_container
        FROM inventory_operations o
        JOIN inventory_operation_lines l ON l.operation_id=o.id
        JOIN skus s ON s.id=l.sku_id
        LEFT JOIN locations sl ON sl.id=l.source_location_id
        LEFT JOIN locations dl ON dl.id=l.destination_location_id
        LEFT JOIN containers sc ON sc.id=l.source_container_id
        LEFT JOIN containers dc ON dc.id=l.destination_container_id
        LEFT JOIN fc_tsd_work tw ON o.idempotency_key=('tsd-work-' || tw.id)
          AND tw.client_id=CAST(o.client_id AS text) AND tw.warehouse_id=CAST(o.warehouse_id AS text)
        WHERE {' AND '.join(clauses)}
        ORDER BY o.created_at DESC, l.id
        LIMIT :limit
    """
    rows = db.execute(text(sql), params).mappings().all()
    out: list[dict[str, Any]] = []
    for r in rows:
        source = r["source_location_id"] is not None
        destination = r["destination_location_id"] is not None
        if destination and not source:
            physical_delta = _num(r["quantity"])
        elif source and not destination:
            physical_delta = -_num(r["quantity"])
        else:
            physical_delta = 0.0
        event_type = str(r["operation_type"] or "STOCK")
        reason = str(r.get("reason") or "")
        idem = str(r.get("idempotency_key") or "")
        comment = str(r.get("comment") or "")
        # r120: PVZ goods are real stock receipts, but must be visibly distinct from
        # ordinary/manual receiving in both UI history and the warehouse Excel.
        if event_type == "RECEIVING" and (idem.upper().startswith("PVZ:") or reason.upper().startswith("PVZ_RECEIVING:") or "ПРИЁМКА ПВЗ" in comment.upper()):
            event_type = "PVZ_RECEIVING"
        # FC_K9R38_HISTORY_WRITE_OFF_LABEL: warehouse history names direct task write-offs explicitly.
        if reason.startswith("EXTERNAL_PICK_TASK_WRITE_OFF:"):
            event_type = "EXTERNAL_PICK_WRITE_OFF"
        server_created_at=_iso(r["created_at"])
        factual_tsd_at=(r.get("tsd_last_scanned_at") or r.get("tsd_box_last_at")) if idem.startswith("tsd-work-") else None
        out.append({
            "kind": "STOCK",
            "event_type": event_type,
            "created_at": _iso(factual_tsd_at) if factual_tsd_at else server_created_at,
            "server_created_at": server_created_at,
            "tsd_first_scanned_at": _iso(r.get("tsd_first_scanned_at") or r.get("tsd_box_first_at")) if idem.startswith("tsd-work-") else None,
            "tsd_last_scanned_at": _iso(factual_tsd_at) if factual_tsd_at else None,
            "operation_id": str(r["operation_id"]),
            "line_id": str(r["line_id"]),
            "sku_id": str(r["sku_id"]),
            "sku_code": r["sku_code"], "article": r["article"], "name": r["sku_name"],
            "size": r["size"], "barcode": r["barcode"],
            "quantity": _num(r["quantity"]), "physical_delta": physical_delta, "reserve_delta": 0,
            "source_location_id": str(r["source_location_id"]) if r["source_location_id"] else None,
            "source_location": r["source_location"],
            "source_container_id": str(r["source_container_id"]) if r["source_container_id"] else None,
            "source_container": r["source_container"],
            "destination_location_id": str(r["destination_location_id"]) if r["destination_location_id"] else None,
            "destination_location": r["destination_location"],
            "destination_container_id": str(r["destination_container_id"]) if r["destination_container_id"] else None,
            "destination_container": r["destination_container"],
            "actor": r["performed_by"], "reference": r["reference"], "comment": r["comment"],
            "idempotency_key": r.get("idempotency_key"),
            "reason": r["reason"], "status": r["status"], "order_id": None,
        })
    return out


def _reservation_events(
    db: Session,
    *,
    client_id: uuid.UUID,
    warehouse_id: uuid.UUID,
    sku_id: uuid.UUID | None = None,
    barcode: str | None = None,
    location_id: uuid.UUID | None = None,
    container_id: uuid.UUID | None = None,
    order_id: str | None = None,
    limit: int = 500,
) -> list[dict[str, Any]]:
    if not _table_exists(db, "fc_stock_reservations"):
        return []
    where = ["r.client_id=:client_id", "r.warehouse_id=:warehouse_id"]
    params: dict[str, Any] = {"client_id": str(client_id), "warehouse_id": str(warehouse_id), "limit": limit}
    if sku_id:
        where.append("r.sku_id=:sku_id"); params["sku_id"] = str(sku_id)
    if barcode:
        where.append("s.barcode=:barcode"); params["barcode"] = str(barcode).strip()
    if location_id:
        where.append("r.source_location_id=:location_id"); params["location_id"] = str(location_id)
    if container_id:
        where.append("r.source_container_id=:container_id"); params["container_id"] = str(container_id)
    if order_id:
        where.append("(r.order_id ILIKE :order_id OR COALESCE(r.external_order_id,'') ILIKE :order_id)"); params["order_id"] = f"%{str(order_id).strip()}%"
    sql = f"""
      SELECT r.*, s.sku_code, s.article, s.name AS sku_name, s.size, s.barcode,
             l.code AS source_location, COALESCE(c.name,c.barcode) AS source_container,
             mo.local_status AS order_status
      FROM fc_stock_reservations r
      LEFT JOIN skus s ON CAST(s.id AS text)=r.sku_id
      LEFT JOIN locations l ON CAST(l.id AS text)=r.source_location_id
      LEFT JOIN containers c ON CAST(c.id AS text)=r.source_container_id
      LEFT JOIN marketplace_orders mo ON CAST(mo.id AS text)=r.order_id
      WHERE {' AND '.join(where)}
      ORDER BY r.created_at DESC
      LIMIT :limit
    """
    rows = db.execute(text(sql), params).mappings().all()
    events: list[dict[str, Any]] = []
    for r in rows:
        base = {
            "kind": "RESERVE", "operation_id": None, "line_id": None,
            "sku_id": str(r["sku_id"]), "sku_code": r["sku_code"], "article": r["article"],
            "name": r["sku_name"], "size": r["size"], "barcode": r["barcode"],
            "quantity": _num(r["quantity"]), "physical_delta": 0,
            "source_location_id": r["source_location_id"], "source_location": r["source_location"],
            "source_container_id": r["source_container_id"], "source_container": r["source_container"],
            "destination_location_id": None, "destination_location": None,
            "destination_container_id": None, "destination_container": None,
            "actor": "SYSTEM", "reference": r["external_order_id"], "comment": r["note"],
            "reason": r["reservation_mode"], "status": r["status"], "order_id": r["order_id"],
            "order_status": r["order_status"],
        }
        created = dict(base, event_type="RESERVE_CREATED", created_at=_iso(r["created_at"]), reserve_delta=_num(r["quantity"]))
        events.append(created)
        if r["source_location_id"]:
            events.append(dict(base, event_type="RESERVE_BOUND", created_at=_iso(r["updated_at"]), reserve_delta=0))
        if r["consumed_at"]:
            events.append(dict(base, event_type="RESERVE_CONSUMED", created_at=_iso(r["consumed_at"]), reserve_delta=-_num(r["quantity"])))
        if r["released_at"]:
            events.append(dict(base, event_type="RESERVE_RELEASED", created_at=_iso(r["released_at"]), reserve_delta=-_num(r["quantity"])))
        if r["receiving_at"]:
            events.append(dict(base, event_type="RETURN_TO_RECEIVING", created_at=_iso(r["receiving_at"]), reserve_delta=-_num(r["quantity"])))
    return events


def _ledger(db: Session, **kwargs) -> list[dict[str, Any]]:
    limit = int(kwargs.get("limit") or 500)
    physical = _physical_events(db, **kwargs)
    reserve_kwargs = {k: v for k, v in kwargs.items() if k != "actor" and k != "operation_type"}
    reserve = _reservation_events(db, **reserve_kwargs)
    rows = physical + reserve
    rows.sort(key=lambda x: x.get("created_at") or "", reverse=True)
    return rows[:limit]



_EVENT_TYPE_RU = {
    "RECEIVING": "Приёмка",
    "PVZ_RECEIVING": "Приёмка ПВЗ",
    "RECEIVING_IMPORT": "Приёмка из Excel",
    "MOVE": "Перемещение",
    "MOVE_CONTAINER": "Перемещение короба",
    "SHIPPING": "Списание / отгрузка",
    "ADJUSTMENT": "Ручная корректировка",
    "RECONCILIATION": "Техническая сверка",
    "RESERVE_CREATED": "Заказ: резерв создан",
    "RESERVE_BOUND": "Сборочный лист: резерв закреплён",
    "RESERVE_CONSUMED": "Упаковка: резерв списан",
    "RESERVE_RELEASED": "Отмена: резерв освобождён",
    "RETURN_TO_RECEIVING": "Возврат в ПРИЕМКА",
}

def _event_type_ru(value: object) -> str:
    code = str(value or "").strip().upper()
    return _EVENT_TYPE_RU.get(code, code or "Операция")

_STATUS_RU = {
    "COMPLETED": "Проведена",
    "CANCELLED": "Отменена",
    "ACTIVE": "Активна",
    "SALE_RESERVED": "Резерв продажи",
    "PACKING_WRITE_OFF": "Ожидает закрытия списания",
    "CONSUMED": "Закрыта",
    "RELEASED": "Освобождена",
}

def _status_ru(value: object) -> str:
    code = str(value or "COMPLETED").strip().upper()
    return _STATUS_RU.get(code, code or "Проведена")

def _counts_in_physical_stock(row: dict[str, Any]) -> bool:
    if str(row.get("kind") or "").upper() != "STOCK":
        return False
    return str(row.get("status") or "COMPLETED").strip().upper() != "CANCELLED"

@router.get("/ledger")
def audit_ledger(
    client_id: uuid.UUID, warehouse_id: uuid.UUID,
    sku_id: uuid.UUID | None = None, barcode: str | None = None,
    location_id: uuid.UUID | None = None, container_id: uuid.UUID | None = None,
    order_id: str | None = None, actor: str | None = None, operation_type: str | None = None,
    limit: int = Query(default=500, ge=1, le=3000), db: Session = Depends(get_db),
):
    return _ledger(db, client_id=client_id, warehouse_id=warehouse_id, sku_id=sku_id, barcode=barcode,
                   location_id=location_id, container_id=container_id, order_id=order_id, actor=actor,
                   operation_type=operation_type, limit=limit)


@router.get("/stock-summary")
def audit_stock_summary(
    client_id: uuid.UUID, warehouse_id: uuid.UUID, barcode: str | None = None,
    limit: int = Query(default=1000, ge=1, le=5000), db: Session = Depends(get_db),
):
    params: dict[str, Any] = {"client_id": str(client_id), "warehouse_id": str(warehouse_id), "limit": limit}
    bc = " AND s.barcode=:barcode" if barcode else ""
    if barcode: params["barcode"] = str(barcode).strip()
    sql = f"""
      WITH physical AS (
        SELECT b.sku_id, SUM(b.quantity) qty FROM inventory_balances b
        WHERE b.client_id=:client_id AND b.warehouse_id=:warehouse_id GROUP BY b.sku_id
      ), ops AS (
        SELECT l.sku_id,
          SUM(CASE WHEN o.operation_type<>'RECONCILIATION' AND l.destination_location_id IS NOT NULL THEN l.quantity ELSE 0 END)
          - SUM(CASE WHEN o.operation_type<>'RECONCILIATION' AND l.source_location_id IS NOT NULL THEN l.quantity ELSE 0 END) ledger_qty
        FROM inventory_operations o JOIN inventory_operation_lines l ON l.operation_id=o.id
        WHERE o.client_id=:client_id AND o.warehouse_id=:warehouse_id AND COALESCE(o.status,'COMPLETED')<>'CANCELLED'
        GROUP BY l.sku_id
      ), ord AS (
        SELECT local_sku_id AS sku_id,
          COUNT(*) FILTER (WHERE local_status='NEW') AS new_orders,
          COUNT(*) FILTER (WHERE local_status='PICKING') AS picking_orders
        FROM marketplace_orders WHERE client_id=:client_id AND local_sku_id IS NOT NULL GROUP BY local_sku_id
      )
      SELECT s.id sku_id,s.sku_code,s.article,s.name,s.size,s.barcode,
             COALESCE(p.qty,0) physical_qty, COALESCE(o.ledger_qty,0) ledger_qty,
             COALESCE(ord.new_orders,0) new_orders, COALESCE(ord.picking_orders,0) picking_orders
      FROM skus s LEFT JOIN physical p ON p.sku_id=s.id LEFT JOIN ops o ON o.sku_id=s.id LEFT JOIN ord ON ord.sku_id=s.id
      WHERE s.client_id=:client_id {bc} AND (COALESCE(p.qty,0)<>0 OR COALESCE(ord.new_orders,0)+COALESCE(ord.picking_orders,0)>0)
      ORDER BY s.article NULLS LAST,s.size NULLS LAST LIMIT :limit
    """
    rows = db.execute(text(sql), params).mappings().all()
    out=[]
    seen_identities: set[tuple[str, ...]] = set()
    for r in rows:
        try:
            truth = stock_truth_for_sku(
                db, client_id=client_id, warehouse_id=warehouse_id, sku_id=r["sku_id"]
            )
            identity = tuple(sorted(str(x) for x in truth.get("identity_sku_ids") or [r["sku_id"]]))
            if identity in seen_identities:
                continue
            seen_identities.add(identity)
            physical = _num(truth.get("physical"))
            res = _num(truth.get("reserved"))
            hold = _num(truth.get("accounting_pending"))
            free = _num(truth.get("available_to_sell"))
            ledger = _num(truth.get("recorded_net"))
            new_orders = int(_num(truth.get("order_new")))
            picking_orders = int(_num(truth.get("order_picking")))
            gap = _num(truth.get("reservation_gap"))
        except Exception:
            physical=_num(r["physical_qty"]); res=0.0; hold=0.0; free=max(0.0,physical); ledger=_num(r["ledger_qty"]); new_orders=int(r["new_orders"] or 0); picking_orders=int(r["picking_orders"] or 0); gap=0.0
        out.append({**{k:(str(v) if k=="sku_id" else v) for k,v in dict(r).items()},
                    "physical_qty":physical,"reserved_qty":res,"accounting_pending_qty":hold,"free_qty":free,
                    "ledger_qty":ledger,"history_difference":physical-ledger,
                    "new_orders":new_orders,"picking_orders":picking_orders,
                    "active_orders":new_orders+picking_orders,"reservation_gap_qty":gap})
    return out


@router.get("/export.xlsx")
def audit_export_xlsx(
    client_id: uuid.UUID, warehouse_id: uuid.UUID, barcode: str | None = None,
    location_id: uuid.UUID | None = None, container_id: uuid.UUID | None = None,
    order_id: str | None = None, operation_type: str | None = None,
    db: Session = Depends(get_db),
):
    # The old export silently cut the warehouse history to the newest 3000 events.
    # That made a receiving-only Excel total smaller than the actual stock even when
    # older receipts were valid. 100k comfortably covers warehouse-scale history while
    # remaining bounded for browser downloads.
    rows = _ledger(db, client_id=client_id, warehouse_id=warehouse_id, barcode=barcode,
                   location_id=location_id, container_id=container_id, order_id=order_id,
                   actor=None, operation_type=operation_type, limit=100000)
    wb=Workbook(); ws=wb.active; ws.title="Журнал склада"
    headers=["Дата/время","Дата","Выгружено на сервер","Тип","Статус","Учитывается в физическом остатке","ШК","Артикул","Размер","Товар","Физ. изменение","Резерв изменение","Кол-во","Откуда","Куда","Заказ/основание","Название приёмки","Сотрудник","Комментарий"]
    ws.append(headers)
    for cell in ws[1]: cell.font=Font(bold=True); cell.fill=PatternFill("solid", fgColor="D9E1F2")
    for r in rows:
        src=" · ".join(x for x in [r.get("source_location"),r.get("source_container")] if x)
        dst=" · ".join(x for x in [r.get("destination_location"),r.get("destination_container")] if x)
        ws.append([r.get("created_at"),_history_date(r.get("created_at")),r.get("server_created_at") or r.get("created_at"),_event_type_ru(r.get("event_type")),_status_ru(r.get("status")),
                   "Да" if _counts_in_physical_stock(r) else "Нет",r.get("barcode"),r.get("article"),r.get("size"),r.get("name"),
                   r.get("physical_delta"),r.get("reserve_delta"),r.get("quantity"),src,dst,r.get("reference") or r.get("order_id"),_receipt_name(r),r.get("actor"),r.get("comment") or r.get("reason")])
        ws.cell(ws.max_row,2).number_format="dd.mm.yyyy"
        for col in (7,8,9): ws.cell(ws.max_row,col).number_format='@'
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    for col,width in {"A":27,"B":13,"C":27,"D":24,"E":18,"F":28,"G":20,"H":16,"I":12,"J":36,"K":15,"L":15,"M":12,"N":24,"O":24,"P":25,"Q":30,"R":18,"S":38}.items(): ws.column_dimensions[col].width=width
    stream=io.BytesIO(); wb.save(stream); stream.seek(0)
    name=f"FulfillCore_warehouse_history_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(stream,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":f'attachment; filename="{name}"'})
