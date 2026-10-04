# FC_R145F17_PVZ_RECEIPT_EXCEL_SPLIT: exact receipt Excel export; GOOD and DEFECT downloads stay separate.
import io
import json
import re
import uuid
from datetime import datetime, timezone
from decimal import Decimal
from typing import Literal

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, Response, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from pydantic import BaseModel, Field
from sqlalchemy import select, text
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.database import get_db
from app.models import Client, Container, InventoryOperation, Location, SKU
from app.schemas.inventory import OperationCreate, StockLine
from app.services.inventory import execute_operation
from app.services.stock_core_v2 import reconcile_stock_bucket
from app.services import pvz as service
from app.services.local_catalog import catalog as local_catalog
from app.services.local_catalog_rules import barcode as normalize_product_barcode
from app.services.excel_safety import append_safe, workbook_bytes
from app.routers.receiving import _ensure_location, _ensure_container
from app.routers.warehouse_ui import scope

router = APIRouter(prefix='/pvz', tags=['ПВЗ'])
portal_router = APIRouter(prefix='/client', tags=['Клиентский кабинет'])


class PvzReceive(BaseModel):
    client_id: uuid.UUID
    warehouse_id: uuid.UUID
    barcode: str = Field(min_length=1,max_length=120)
    article: str = Field(default='',max_length=180)
    size: str = Field(default='',max_length=100)
    name: str = Field(default='',max_length=240)
    quantity: int = Field(ge=1,le=100000)
    location_code: str = Field(default='ПРИЕМКА',max_length=120)
    container_code: str = Field(default='',max_length=160)
    reference: str = Field(default='',max_length=300)
    comment: str = Field(default='',max_length=1000)
    performed_by: str = Field(default='Администратор',max_length=160)


class PvzDefect(BaseModel):
    client_id: uuid.UUID
    warehouse_id: uuid.UUID
    barcode: str = Field(default='',max_length=120)
    article: str = Field(min_length=1,max_length=180)
    size: str = Field(default='',max_length=100)
    name: str = Field(default='',max_length=240)
    quantity: int = Field(ge=1,le=100000)
    defect_reason: str = Field(min_length=1,max_length=1000)
    reference: str = Field(default='',max_length=300)
    comment: str = Field(default='',max_length=1000)
    performed_by: str = Field(default='Администратор',max_length=160)


def _product(db: Session, payload: PvzReceive) -> SKU:
    try:
        barcode = normalize_product_barcode(payload.barcode)
    except ValueError as exc:
        raise HTTPException(422, detail=str(exc))
    rows=db.execute(select(SKU).where(SKU.client_id==payload.client_id,SKU.barcode==barcode)).scalars().all()
    if len(rows)>1:
        raise HTTPException(409,detail='Этот ШК привязан к нескольким SKU. Исправьте Базу товаров')
    known = local_catalog(db, payload.client_id, barcode)
    card = known[0] if known else {}
    if card.get('wb_conflict'):
        raise HTTPException(409, detail=card.get('warning') or 'ШК встречается в нескольких карточках WB')
    resolved_article = str(card.get('article') or payload.article or '').strip()
    resolved_size = str(card.get('size') or payload.size or '').strip()
    resolved_name = str(card.get('name') or payload.name or resolved_article).strip()
    if rows:
        sku=rows[0]
        if resolved_article and not str(sku.article or '').strip(): sku.article=resolved_article
        if resolved_size and not str(sku.size or '').strip(): sku.size=resolved_size
        if resolved_name and not str(sku.name or '').strip(): sku.name=resolved_name
        sku.is_active=True
        return sku
    # Unknown barcodes must be consciously registered.  Do not silently invent an
    # article or size from the barcode: this new SKU becomes part of the shared local product base.
    if not resolved_article or not resolved_size:
        raise HTTPException(422, detail='ШК не найден в Базе товаров. Укажите артикул и размер вручную (для безразмерного — «Без размера»).')
    sku=SKU(id=uuid.uuid4(),client_id=payload.client_id,
            sku_code='LOCAL-'+uuid.uuid5(payload.client_id, barcode).hex,
            barcode=barcode,article=resolved_article,size=resolved_size,
            name=resolved_name or resolved_article,unit='PCS',is_active=True)
    db.add(sku);db.flush();return sku


@router.post('/receive')
def receive(payload:PvzReceive,db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id)
    service.install_schema(db)
    sku=_product(db,payload)
    location=_ensure_location(db,payload.warehouse_id,payload.location_code.strip() or 'ПРИЕМКА')
    container=_ensure_container(db,payload.client_id,payload.warehouse_id,location.id,payload.container_code.strip())
    batch_id=uuid.uuid4();operation_key=f'PVZ:{batch_id}'
    op=execute_operation(db,OperationCreate(client_id=payload.client_id,warehouse_id=payload.warehouse_id,
        operation_type='RECEIVING',idempotency_key=operation_key,performed_by=payload.performed_by,
        reference=payload.reference or f'ПВЗ {batch_id}',comment='Приёмка ПВЗ · товар доступен для сборки · ЗОНА:MAIN',
        lines=[StockLine(sku_id=sku.id,quantity=Decimal(payload.quantity),destination_location_id=location.id,
                         destination_container_id=container.id if container else None,reason=f'PVZ_RECEIVING:{batch_id}')]),commit=False)
    reconcile_stock_bucket(db,client_id=payload.client_id,warehouse_id=payload.warehouse_id,sku_id=sku.id,
        location_id=location.id,container_id=container.id if container else None,preferred_bucket='STORAGE_ONLY',lock=True)
    db.execute(text('''INSERT INTO fc_pvz_batches(id,client_id,warehouse_id,kind,status,reference,performed_by,stock_operation_id)
        VALUES(:id,:c,:w,'GOOD','COMPLETED',:ref,:actor,:op)'''),{'id':str(batch_id),'c':str(payload.client_id),'w':str(payload.warehouse_id),'ref':payload.reference,'actor':payload.performed_by,'op':str(op.id)})
    db.execute(text('''INSERT INTO fc_pvz_lines(id,batch_id,sku_id,barcode,article,size,name,quantity,location_id,container_id,comment)
        VALUES(:id,:b,:s,:bc,:a,:z,:n,:q,:loc,:box,:comment)'''),{'id':str(uuid.uuid4()),'b':str(batch_id),'s':str(sku.id),'bc':str(sku.barcode or ''),'a':str(sku.article or ''),'z':str(sku.size or ''),'n':str(sku.name or ''),'q':payload.quantity,'loc':str(location.id),'box':str(container.id) if container else None,'comment':payload.comment})
    db.commit()
    return {'ok':True,'batch_id':str(batch_id),'operation_id':str(op.id),'barcode':sku.barcode,'quantity':payload.quantity,
            'location':location.code,'container':(container.name or container.barcode) if container else '',
            'message':'ПВЗ принят на Основной склад и доступен для сборки'}


@router.post('/defect')
def defect(payload:PvzDefect,db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id);service.install_schema(db)
    sku_id=None
    if payload.barcode.strip():
        sku=db.execute(select(SKU).where(SKU.client_id==payload.client_id,SKU.barcode==payload.barcode.strip())).scalars().first()
        sku_id=str(sku.id) if sku else None
    batch_id=uuid.uuid4()
    db.execute(text('''INSERT INTO fc_pvz_batches(id,client_id,warehouse_id,kind,status,reference,performed_by)
        VALUES(:id,:c,:w,'DEFECT','COMPLETED',:ref,:actor)'''),{'id':str(batch_id),'c':str(payload.client_id),'w':str(payload.warehouse_id),'ref':payload.reference,'actor':payload.performed_by})
    db.execute(text('''INSERT INTO fc_pvz_lines(id,batch_id,sku_id,barcode,article,size,name,quantity,defect_reason,comment)
        VALUES(:id,:b,:s,:bc,:a,:z,:n,:q,:reason,:comment)'''),{'id':str(uuid.uuid4()),'b':str(batch_id),'s':sku_id,'bc':payload.barcode.strip(),'a':payload.article.strip(),'z':payload.size.strip(),'n':payload.name.strip(),'q':payload.quantity,'reason':payload.defect_reason.strip(),'comment':payload.comment.strip()})
    db.commit()
    return {'ok':True,'batch_id':str(batch_id),'quantity':payload.quantity,'message':'Брак записан отдельно и НЕ добавлен в складские остатки'}


@router.get('/history')
def history(client_id:uuid.UUID,warehouse_id:uuid.UUID,kind:str='',limit:int=200,db:Session=Depends(get_db)):
    scope(db,client_id,warehouse_id);service.install_schema(db)
    params={'c':str(client_id),'w':str(warehouse_id),'limit':min(max(int(limit),1),1000)}
    clause='';
    if kind in ('GOOD','DEFECT'): params['kind']=kind;clause=' AND b.kind=:kind '
    rows=db.execute(text(f'''SELECT CAST(b.id AS text) id,COALESCE(to_jsonb(b)->>'receiving_session_id','') receiving_session_id,b.kind,b.status,b.reference,b.performed_by,b.created_at,
      COALESCE(SUM(l.quantity),0) quantity,COUNT(l.id) lines,
      STRING_AGG(DISTINCT NULLIF(l.defect_reason,''),'; ') defect_reason
      FROM fc_pvz_batches b LEFT JOIN fc_pvz_lines l ON l.batch_id=b.id
      WHERE CAST(b.client_id AS text)=:c AND CAST(b.warehouse_id AS text)=:w {clause}
      GROUP BY b.id,b.kind,b.status,b.reference,b.performed_by,b.created_at ORDER BY b.created_at DESC LIMIT :limit'''),params).mappings().all()
    return {'rows':[dict(r) for r in rows]}


def _history_rows(db, client_id, warehouse_id, *, kind: str = '', batch_id: uuid.UUID | None = None):
    params = {'c': str(client_id), 'w': str(warehouse_id)}
    clauses = ['CAST(b.client_id AS text)=:c', 'CAST(b.warehouse_id AS text)=:w']
    clean_kind = str(kind or '').strip().upper()
    if clean_kind:
        if clean_kind not in ('GOOD', 'DEFECT'):
            raise HTTPException(422, detail='Тип выгрузки должен быть GOOD или DEFECT')
        params['kind'] = clean_kind
        clauses.append('b.kind=:kind')
    if batch_id is not None:
        params['batch_id'] = str(batch_id)
        clauses.append('CAST(b.id AS text)=:batch_id')
    sql = '''SELECT b.created_at,b.kind,b.reference,b.performed_by,l.barcode,l.article,l.size,l.name,l.quantity,
      COALESCE(loc.code,'') location,COALESCE(c.name,c.barcode,'') container,l.defect_reason,l.comment
      FROM fc_pvz_batches b JOIN fc_pvz_lines l ON l.batch_id=b.id
      LEFT JOIN locations loc ON loc.id=l.location_id LEFT JOIN containers c ON c.id=l.container_id
      WHERE %s ORDER BY b.created_at DESC,l.article,l.size,l.created_at''' % ' AND '.join(clauses)
    return db.execute(text(sql), params).all()


def _history_workbook(db, client_id, warehouse_id, *, kind: str = ''):
    rows = _history_rows(db, client_id, warehouse_id, kind=kind)
    clean_kind = str(kind or '').strip().upper()
    wb=Workbook();ws=wb.active
    ws.title='Приёмки ПВЗ' if clean_kind=='GOOD' else ('Брак ПВЗ' if clean_kind=='DEFECT' else 'ПВЗ')
    append_safe(ws, ['Дата','Тип','Основание','Сотрудник','ШК','Артикул','Размер','Товар','Количество','Ячейка','Короб','Причина брака','Комментарий'])
    for r in rows:
        append_safe(ws, r)
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width=min(42,max(12,max(len(str(x.value or '')) for x in col)+2))
    return workbook_bytes(wb)


def _batch_workbook(db, client_id, warehouse_id, batch_id: uuid.UUID):
    batch=db.execute(text('''SELECT CAST(id AS text) id,kind,status,reference,performed_by,created_at,COALESCE(to_jsonb(fc_pvz_batches)->>'receiving_session_id','') receiving_session_id
      FROM fc_pvz_batches WHERE CAST(id AS text)=:id AND CAST(client_id AS text)=:c AND CAST(warehouse_id AS text)=:w'''),
      {'id':str(batch_id),'c':str(client_id),'w':str(warehouse_id)}).mappings().first()
    if not batch:
        raise HTTPException(404,detail='Приёмка ПВЗ не найдена')
    rows=_history_rows(db,client_id,warehouse_id,batch_id=batch_id)
    wb=Workbook();summary=wb.active;summary.title='Сводка'
    append_safe(summary,['Название приёмки',batch.get('reference') or str(batch_id)])
    append_safe(summary,['Тип','Брак' if batch.get('kind')=='DEFECT' else 'ПВЗ'])
    append_safe(summary,['Сотрудник',batch.get('performed_by') or ''])
    append_safe(summary,['Дата',batch.get('created_at')])
    append_safe(summary,['Статус',batch.get('status') or ''])
    append_safe(summary,['Единиц',sum(float(r[8] or 0) for r in rows)])
    append_safe(summary,['Позиций',len(rows)])
    ws=wb.create_sheet('Состав')
    append_safe(ws,['Дата','Тип','Основание','Сотрудник','ШК','Артикул','Размер','Товар','Количество','Ячейка','Короб','Причина брака','Комментарий'])
    for r in rows:
        append_safe(ws,r)
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for sheet in wb.worksheets:
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width=min(42,max(12,max(len(str(x.value or '')) for x in col)+2))
    return workbook_bytes(wb),dict(batch)


@router.get('/history/export.xlsx')
def export_history(client_id:uuid.UUID,warehouse_id:uuid.UUID,kind:str='',db:Session=Depends(get_db)):
    scope(db,client_id,warehouse_id);service.install_schema(db)
    clean_kind=str(kind or '').strip().upper()
    data=_history_workbook(db,client_id,warehouse_id,kind=clean_kind)
    filename='FulfillCore_PVZ_history.xlsx'
    if clean_kind=='GOOD':
        filename='FulfillCore_PVZ_receipts.xlsx'
    elif clean_kind=='DEFECT':
        filename='FulfillCore_PVZ_defects.xlsx'
    return Response(data,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="{filename}"'})


@router.get('/history/{batch_id}/export.xlsx')
def export_one_receipt(batch_id:uuid.UUID,client_id:uuid.UUID,warehouse_id:uuid.UUID,db:Session=Depends(get_db)):
    scope(db,client_id,warehouse_id);service.install_schema(db)
    data,batch=_batch_workbook(db,client_id,warehouse_id,batch_id)
    prefix='defect' if batch.get('kind')=='DEFECT' else 'receipt'
    return Response(data,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="FulfillCore_PVZ_{prefix}_{str(batch_id)[:8]}.xlsx"'})


@router.post('/labels/upload')
def upload_labels(client_id:uuid.UUID=Form(...),warehouse_id:uuid.UUID=Form(...),performed_by:str=Form('Администратор'),file:UploadFile=File(...),db:Session=Depends(get_db)):
    scope(db,client_id,warehouse_id);service.install_schema(db)
    raw=file.file.read(50*1024*1024+1)
    if len(raw)>50*1024*1024: raise HTTPException(413,detail='PDF больше 50 МБ. Разделите файл на части')
    try:pages=service.analyze_pdf(db,client_id,raw)
    except Exception as e: raise HTTPException(422,detail=f'Не удалось разобрать PDF: {e}')
    batch_id=uuid.uuid4();sha=__import__('hashlib').sha256(raw).hexdigest()
    db.execute(text('''INSERT INTO fc_pvz_label_batches(id,client_id,warehouse_id,original_filename,pdf_bytes,pdf_sha256,page_count,performed_by)
      VALUES(:id,:c,:w,:name,:pdf,:sha,:count,:actor)'''),{'id':str(batch_id),'c':str(client_id),'w':str(warehouse_id),'name':file.filename or 'labels.pdf','pdf':raw,'sha':sha,'count':len(pages),'actor':performed_by})
    for p in pages:
        db.execute(text('''INSERT INTO fc_pvz_label_pages(batch_id,page_no,article,size,barcode_value,label_type,extracted_text,page_key)
          VALUES(:b,:p,:a,:z,:bc,:t,:x,:k)'''),{'b':str(batch_id),'p':p['page_no'],'a':p['article'],'z':p['size'],'bc':p['barcode_value'],'t':p['label_type'],'x':p['text'][:4000],'k':p['page_key']})
    db.commit();groups=service.group_label_pages(pages,service.stock_by_article_size(db,client_id,warehouse_id))
    return {'ok':True,'batch_id':str(batch_id),'filename':file.filename,'page_count':len(pages),'groups':groups,
            'unknown_pages':sum(1 for p in pages if p['label_type']=='UNKNOWN'),'message':'PDF разделён по артикулу/размеру. Неопределённые страницы можно исправить вручную.'}


@router.get('/labels/{batch_id}')
def label_batch(batch_id:uuid.UUID,client_id:uuid.UUID,warehouse_id:uuid.UUID,db:Session=Depends(get_db)):
    scope(db,client_id,warehouse_id);service.install_schema(db)
    b=db.execute(text('SELECT original_filename,page_count,created_at FROM fc_pvz_label_batches WHERE id=:id AND client_id=:c AND warehouse_id=:w'),{'id':str(batch_id),'c':str(client_id),'w':str(warehouse_id)}).mappings().first()
    if not b: raise HTTPException(404,detail='PDF не найден')
    rows=[dict(r) for r in db.execute(text('SELECT page_no,article,size,barcode_value,label_type,extracted_text FROM fc_pvz_label_pages WHERE batch_id=:b ORDER BY page_no'),{'b':str(batch_id)}).mappings().all()]
    return {'batch_id':str(batch_id),**dict(b),'groups':service.group_label_pages(rows,service.stock_by_article_size(db,client_id,warehouse_id)),
            'pages':rows}


class PageUpdate(BaseModel):
    client_id:uuid.UUID;warehouse_id:uuid.UUID
    article:str=Field(default='',max_length=180);size:str=Field(default='',max_length=100);barcode_value:str=Field(default='',max_length=160)
    label_type:Literal['BARCODE','CHZ','UNKNOWN']


@router.patch('/labels/{batch_id}/pages/{page_no}')
def update_page(batch_id:uuid.UUID,page_no:int,payload:PageUpdate,db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id);service.install_schema(db)
    r=db.execute(text('''UPDATE fc_pvz_label_pages p SET article=:a,size=:z,barcode_value=:bc,label_type=:t
      FROM fc_pvz_label_batches b WHERE p.batch_id=b.id AND p.batch_id=:id AND p.page_no=:p
      AND b.client_id=:c AND b.warehouse_id=:w RETURNING p.id'''),{'a':payload.article.strip(),'z':payload.size.strip(),'bc':payload.barcode_value.strip(),'t':payload.label_type,'id':str(batch_id),'p':page_no,'c':str(payload.client_id),'w':str(payload.warehouse_id)}).first()
    if not r: raise HTTPException(404,detail='Страница PDF не найдена')
    db.commit();return {'ok':True}


class PrintSpec(BaseModel):
    client_id:uuid.UUID;warehouse_id:uuid.UUID
    article:str=Field(default='',max_length=180);size:str=Field(default='',max_length=100)
    barcode_copies:int=Field(default=0,ge=0,le=5000);chz_count:int=Field(default=0,ge=0,le=5000);chz_start:int=Field(default=1,ge=1,le=5000)
    performed_by:str=Field(default='Администратор',max_length=160)


@router.post('/labels/{batch_id}/print.pdf')
def print_labels(batch_id:uuid.UUID,payload:PrintSpec,db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id);service.install_schema(db)
    batch=db.execute(text('SELECT pdf_bytes FROM fc_pvz_label_batches WHERE id=:id AND client_id=:c AND warehouse_id=:w'),{'id':str(batch_id),'c':str(payload.client_id),'w':str(payload.warehouse_id)}).first()
    if not batch: raise HTTPException(404,detail='PDF не найден')
    rows=db.execute(text('''SELECT page_no,label_type FROM fc_pvz_label_pages WHERE batch_id=:b AND article=:a AND size=:z ORDER BY page_no'''),{'b':str(batch_id),'a':payload.article,'z':payload.size}).mappings().all()
    barcode_pages=[r['page_no'] for r in rows if r['label_type']=='BARCODE'];chz_pages=[r['page_no'] for r in rows if r['label_type']=='CHZ']
    try:data=service.build_pdf(bytes(batch[0]),barcode_pages,payload.barcode_copies,chz_pages,payload.chz_count,payload.chz_start)
    except ValueError as e: raise HTTPException(422,detail=str(e))
    db.execute(text('''INSERT INTO fc_pvz_print_jobs(id,batch_id,client_id,warehouse_id,article,size,barcode_copies,chz_count,chz_start,performed_by)
      VALUES(:id,:b,:c,:w,:a,:z,:bc,:chz,:start,:actor)'''),{'id':str(uuid.uuid4()),'b':str(batch_id),'c':str(payload.client_id),'w':str(payload.warehouse_id),'a':payload.article,'z':payload.size,'bc':payload.barcode_copies,'chz':payload.chz_count,'start':payload.chz_start,'actor':payload.performed_by})
    db.commit();return Response(data,media_type='application/pdf',headers={'Content-Disposition':'inline; filename="FulfillCore_PVZ_labels.pdf"'})


@router.get('/labels/{batch_id}/page/{page_no}.pdf')
def preview_page(batch_id:uuid.UUID,page_no:int,client_id:uuid.UUID,warehouse_id:uuid.UUID,db:Session=Depends(get_db)):
    scope(db,client_id,warehouse_id);service.install_schema(db)
    row=db.execute(text('SELECT pdf_bytes,page_count FROM fc_pvz_label_batches WHERE id=:id AND client_id=:c AND warehouse_id=:w'),{'id':str(batch_id),'c':str(client_id),'w':str(warehouse_id)}).first()
    if not row or page_no<1 or page_no>int(row[1]): raise HTTPException(404,detail='Страница не найдена')
    from app.vendor.pypdf import PdfReader,PdfWriter
    reader=PdfReader(io.BytesIO(bytes(row[0])));writer=PdfWriter();writer.add_page(reader.pages[page_no-1]);out=io.BytesIO();writer.write(out)
    return Response(out.getvalue(),media_type='application/pdf')


class PortalAccess(BaseModel):
    client_id:uuid.UUID
    username:str=Field(min_length=3,max_length=120)
    password:str=Field(min_length=8,max_length=200)


@router.post('/client-access')
def client_access(payload:PortalAccess,db:Session=Depends(get_db)):
    if not db.get(Client,payload.client_id): raise HTTPException(404,detail='Клиент не найден')
    service.install_schema(db)
    username=payload.username.strip();salt,pwh=service.password_parts(payload.password)
    existing=db.execute(text('SELECT id FROM fc_client_portal_users WHERE client_id=:c ORDER BY created_at LIMIT 1'),{'c':str(payload.client_id)}).first()
    conflict=db.execute(text('SELECT id,client_id FROM fc_client_portal_users WHERE lower(username)=lower(:u)'),{'u':username}).mappings().first()
    if conflict and (not existing or str(conflict['id'])!=str(existing[0])): raise HTTPException(409,detail='Такой логин уже занят')
    if existing:
        db.execute(text('UPDATE fc_client_portal_users SET username=:u,password_salt=:s,password_hash=:h,is_active=TRUE,updated_at=now() WHERE id=:id'),{'u':username,'s':salt,'h':pwh,'id':str(existing[0])})
        user_id=existing[0]
    else:
        user_id=uuid.uuid4();db.execute(text('INSERT INTO fc_client_portal_users(id,client_id,username,password_salt,password_hash) VALUES(:id,:c,:u,:s,:h)'),{'id':str(user_id),'c':str(payload.client_id),'u':username,'s':salt,'h':pwh})
    db.commit();return {'ok':True,'username':username,'url':'/client/login'}


def _portal_session(request:Request,db:Session):
    service.install_schema(db)
    data=service.read_portal_token(settings.app_secret_key,request.cookies.get('fc_client_session'))
    if not data: raise HTTPException(401,detail='Требуется вход клиента')
    row=db.execute(text('SELECT id,client_id,username,is_active FROM fc_client_portal_users WHERE id=:id'),{'id':str(data['uid'])}).mappings().first()
    if not row or not row['is_active'] or str(row['client_id'])!=str(data['cid']): raise HTTPException(401,detail='Сессия клиента недействительна')
    return dict(row)


class PortalLogin(BaseModel):
    username:str;password:str


@portal_router.post('/auth/login')
def portal_login(payload:PortalLogin,response:Response,request:Request,db:Session=Depends(get_db)):
    service.install_schema(db)
    row=db.execute(text('SELECT id,client_id,username,password_salt,password_hash,is_active FROM fc_client_portal_users WHERE lower(username)=lower(:u)'),{'u':payload.username.strip()}).mappings().first()
    if not row or not row['is_active'] or not service.verify_password(payload.password,row['password_salt'],row['password_hash']): raise HTTPException(401,detail='Неверный логин или пароль')
    token=service.make_portal_token(settings.app_secret_key,str(row['id']),str(row['client_id']),str(row['username']))
    response.set_cookie('fc_client_session',token,max_age=60*60*24*7,httponly=True,samesite='lax',secure=request.url.scheme=='https',path='/')
    return {'ok':True,'username':row['username']}


@portal_router.post('/auth/logout')
def portal_logout(response:Response):
    response.delete_cookie('fc_client_session',path='/');return {'ok':True}


@portal_router.get('/api/session')
def portal_session(request:Request,db:Session=Depends(get_db)):
    user=_portal_session(request,db);client=db.get(Client,uuid.UUID(str(user['client_id'])))
    return {'username':user['username'],'client_id':str(user['client_id']),'client_name':getattr(client,'name',None) or str(user['client_id'])}


# FC_R145F11_SELLER_ASSEMBLY_HOLDS: active warehouse picking tasks reserve MAIN without changing physical stock until final close.
def _portal_pick_holds(db:Session,client_id)->dict[tuple[str,str],float]:
    try:
        if not db.execute(text("SELECT to_regclass('public.fc_wh_pick_tasks')")).scalar():
            return {}
        rows=db.execute(text("""
            SELECT t.warehouse_id,l.sku_id,COALESCE(SUM(l.allocated_qty),0) held
            FROM fc_wh_pick_tasks t JOIN fc_wh_pick_lines l ON l.task_id=t.id
            WHERE t.client_id=:c AND t.status IN ('DRAFT','ASSIGNED','IN_PROGRESS','READY_FOR_ADMIN')
              AND l.sku_id IS NOT NULL
            GROUP BY t.warehouse_id,l.sku_id
        """),{'c':str(client_id)}).mappings().all()
        return {(str(r['warehouse_id']),str(r['sku_id'])):float(r.get('held') or 0) for r in rows}
    except Exception:
        db.rollback()
        return {}

def _portal_stock_rows(db:Session,client_id):
    rows=[dict(r) for r in service.stock_snapshot(db,client_id)]
    holds=_portal_pick_holds(db,client_id)
    for row in rows:
        held=float(holds.get((str(row.get('warehouse_id')),str(row.get('sku_id'))),0) or 0)
        row['assembly_reserved']=held
        row['main_free']=max(0.0,float(row.get('storage_only') or 0)-held)
    return rows

@portal_router.get('/api/stocks')
def portal_stocks(request:Request,db:Session=Depends(get_db)):
    user=_portal_session(request,db);return {'rows':_portal_stock_rows(db,user['client_id']),'generated_at':datetime.now(timezone.utc).isoformat()}


@portal_router.get('/api/places')
def portal_places(sku_id:uuid.UUID,request:Request,warehouse_id:uuid.UUID|None=None,db:Session=Depends(get_db)):
    user=_portal_session(request,db)
    sku=db.get(SKU,sku_id)
    if not sku or str(sku.client_id)!=str(user['client_id']):
        raise HTTPException(404,detail='Товар не найден')
    return {'rows':service.stock_locations(db,user['client_id'],sku_id,warehouse_id)}


# FC_R145F8_SELLER_DEFECT_BOX_AND_SAFE_STOCK_XLSX
def _portal_excel_value(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            value = value.astimezone(timezone.utc).replace(tzinfo=None)
        return value
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, uuid.UUID):
        return str(value)
    if isinstance(value, str):
        return re.sub(r'[\x00-\x08\x0B\x0C\x0E-\x1F]', '', value)
    return value


@portal_router.get('/api/defects')
def portal_defects(request:Request,db:Session=Depends(get_db)):
    user=_portal_session(request,db)
    rows=db.execute(text('''SELECT b.created_at,b.reference,b.performed_by,l.barcode,l.article,l.size,l.name,l.quantity,l.defect_reason,l.comment,
      COALESCE(NULLIF(c.name,''),c.barcode,'') container_name
      FROM fc_pvz_batches b JOIN fc_pvz_lines l ON l.batch_id=b.id
      LEFT JOIN containers c ON c.id=l.container_id
      WHERE b.client_id=:c AND b.kind='DEFECT' AND b.deleted_at IS NULL AND l.deleted_at IS NULL
      ORDER BY b.created_at DESC,l.created_at DESC LIMIT 5000'''),{'c':str(user['client_id'])}).mappings().all()
    return {'rows':[dict(r) for r in rows]}


@portal_router.get('/api/stocks.xlsx')
def portal_stock_xlsx(request:Request,db:Session=Depends(get_db)):
    user=_portal_session(request,db);rows=_portal_stock_rows(db,user['client_id'])
    wb=Workbook();ws=wb.active;ws.title='Остатки';ws.append(['Склад','ШК','Артикул','Размер','Товар','Физически','Основной','В сборке','Свободно основной','FBS','Резерв'])
    for r in rows:
        ws.append([_portal_excel_value(r.get(k)) for k in ['warehouse_name','barcode','article','size','name','physical','storage_only','assembly_reserved','main_free','fbs_available','reserved']])
    wp=wb.create_sheet('Места хранения');wp.append(['Склад','ШК','Артикул','Размер','Товар','Ячейка / палета','Короб','Физически','Основной','FBS','Резерв'])
    for r in service.stock_locations_all(db,user['client_id']):
        wp.append([_portal_excel_value(r.get(k)) for k in ['warehouse_name','barcode','article','size','name','location','container','physical','storage_only','fbs_available','reserved']])
    wd=wb.create_sheet('Брак')
    defects=db.execute(text('''SELECT b.created_at,b.reference,b.performed_by,l.barcode,l.article,l.size,l.name,l.quantity,l.defect_reason,l.comment,
      COALESCE(NULLIF(c.name,''),c.barcode,'') container_name
      FROM fc_pvz_batches b JOIN fc_pvz_lines l ON l.batch_id=b.id
      LEFT JOIN containers c ON c.id=l.container_id
      WHERE b.client_id=:c AND b.kind='DEFECT' AND b.deleted_at IS NULL AND l.deleted_at IS NULL
      ORDER BY b.created_at DESC,l.created_at DESC'''),{'c':str(user['client_id'])}).mappings().all()
    wd.append(['Дата','Основание','Сотрудник','ШК','Артикул','Размер','Товар','Количество','Короб брака','Причина брака','Комментарий'])
    for r in defects:
        wd.append([_portal_excel_value(r.get(k)) for k in ['created_at','reference','performed_by','barcode','article','size','name','quantity','container_name','defect_reason','comment']])
    for sheet in (ws,wp,wd):
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width=min(45,max(12,max(len(str(c.value or '')) for c in col)+2))
    data=workbook_bytes(wb);return Response(data,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="FulfillCore_client_stocks.xlsx"'})
