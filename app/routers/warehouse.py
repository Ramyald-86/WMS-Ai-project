# FC_R145F11_ADMIN_RECEIPT_GROUP_PICK_ASSIGN_MODES
from __future__ import annotations

import io
import json
import re
import threading
import uuid
from datetime import datetime, timezone
from decimal import Decimal
from typing import Any

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field
from sqlalchemy import and_, bindparam, func, or_, select, text
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.models import Container, InventoryBalance, InventoryOperation, InventoryOperationLine, Location, SKU
from app.schemas.inventory import OperationCreate, StockLine
from app.services.inventory import execute_operation
from app.services.stock_core_v2 import move_bucket, reconcile_stock_bucket, ensure_stock_core_v2_schema
from app.services.warehouse_availability import lock_scope, held_quantity, available_buckets
from app.services.workflow_rules import natural
from app.services.box_routes import enrich as route_enrich, freeze_task, snapshots as route_snapshots
from app.services.picking_route import route_key
from app.services.container_lifecycle import archive_consumed_empty_containers
from app.services.excel_safety import append_safe, workbook_bytes
from app.services.receipt_history_dates import history_date_key

router = APIRouter(prefix="/warehouse", tags=["Склад · основной/FBS"])

# FC_K9R62_WAREHOUSE_SPLIT_RECEIPTS_TSD_ADMIN
_SCHEMA_READY = False
_SCHEMA_LOCK = threading.Lock()


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _qty(value: Any) -> Decimal:
    try:
        return max(Decimal("0"), Decimal(str(value or 0)))
    except Exception:
        return Decimal("0")


def _qf(value: Any) -> float:
    return float(_qty(value))


def _natural(value: str | None) -> tuple[Any, ...]:
    parts = re.split(r"(\d+)", str(value or "").upper())
    return natural(value)


def _ensure_schema(db: Session) -> None:
    global _SCHEMA_READY
    if _SCHEMA_READY:
        return
    with _SCHEMA_LOCK:
        if _SCHEMA_READY:
            return
        statements = [
            """
            CREATE TABLE IF NOT EXISTS fc_wh_pick_tasks (
                id VARCHAR(36) PRIMARY KEY,
                client_id VARCHAR(36) NOT NULL,
                warehouse_id VARCHAR(36) NOT NULL,
                task_number VARCHAR(120) NOT NULL,
                name VARCHAR(240) NULL,
                source_file_name VARCHAR(255) NULL,
                source VARCHAR(32) NOT NULL DEFAULT 'ADMIN',
                seller_username VARCHAR(160) NULL,
                parent_task_id VARCHAR(36) NULL,
                pick_mode VARCHAR(16) NOT NULL DEFAULT 'CELL',
                picker VARCHAR(120) NULL,
                status VARCHAR(32) NOT NULL DEFAULT 'DRAFT',
                created_by VARCHAR(120) NULL,
                assigned_by VARCHAR(120) NULL,
                closed_by VARCHAR(120) NULL,
                created_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP,
                assigned_at TIMESTAMPTZ NULL,
                submitted_at TIMESTAMPTZ NULL,
                closed_at TIMESTAMPTZ NULL,
                cancelled_at TIMESTAMPTZ NULL,
                close_operation_id VARCHAR(36) NULL,
                note TEXT NULL
            )
            """,
            "CREATE UNIQUE INDEX IF NOT EXISTS uq_fc_wh_pick_task_number ON fc_wh_pick_tasks(client_id,warehouse_id,task_number)",
            "CREATE INDEX IF NOT EXISTS ix_fc_wh_pick_task_scope ON fc_wh_pick_tasks(client_id,warehouse_id,status,created_at DESC)",
            "CREATE INDEX IF NOT EXISTS ix_fc_wh_pick_task_picker ON fc_wh_pick_tasks(client_id,warehouse_id,LOWER(picker),status,assigned_at DESC)",
            "ALTER TABLE fc_wh_pick_tasks ADD COLUMN IF NOT EXISTS source VARCHAR(32) NOT NULL DEFAULT 'ADMIN'",
            "ALTER TABLE fc_wh_pick_tasks ADD COLUMN IF NOT EXISTS seller_username VARCHAR(160) NULL",
            "ALTER TABLE fc_wh_pick_tasks ADD COLUMN IF NOT EXISTS parent_task_id VARCHAR(36) NULL",
            "ALTER TABLE fc_wh_pick_tasks ADD COLUMN IF NOT EXISTS pick_mode VARCHAR(16) NOT NULL DEFAULT 'CELL'",
            "CREATE INDEX IF NOT EXISTS ix_fc_wh_pick_task_parent ON fc_wh_pick_tasks(parent_task_id,created_at DESC)",
            "CREATE INDEX IF NOT EXISTS ix_fc_wh_pick_task_source ON fc_wh_pick_tasks(client_id,warehouse_id,source,created_at DESC)",
            """
            CREATE TABLE IF NOT EXISTS fc_wh_pick_lines (
                id VARCHAR(36) PRIMARY KEY,
                task_id VARCHAR(36) NOT NULL,
                sku_id VARCHAR(36) NULL,
                location_id VARCHAR(36) NULL,
                container_id VARCHAR(36) NULL,
                requested_qty NUMERIC(18,3) NOT NULL DEFAULT 0,
                allocated_qty NUMERIC(18,3) NOT NULL DEFAULT 0,
                scan_qty NUMERIC(18,3) NOT NULL DEFAULT 0,
                manual_qty NUMERIC(18,3) NOT NULL DEFAULT 0,
                missing_qty NUMERIC(18,3) NOT NULL DEFAULT 0,
                deferred BOOLEAN NOT NULL DEFAULT FALSE,
                line_status VARCHAR(32) NOT NULL DEFAULT 'PENDING',
                sort_order INTEGER NOT NULL DEFAULT 0,
                request_barcode VARCHAR(255) NULL,
                request_article VARCHAR(255) NULL,
                request_size VARCHAR(255) NULL,
                created_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP,
                updated_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
            """,
            "CREATE INDEX IF NOT EXISTS ix_fc_wh_pick_lines_task ON fc_wh_pick_lines(task_id,sort_order,id)",
            "CREATE INDEX IF NOT EXISTS ix_fc_wh_pick_lines_bucket ON fc_wh_pick_lines(sku_id,location_id,container_id,task_id)",
            """
            CREATE TABLE IF NOT EXISTS fc_wh_pick_audit (
                id VARCHAR(36) PRIMARY KEY,
                task_id VARCHAR(36) NOT NULL,
                line_id VARCHAR(36) NULL,
                actor VARCHAR(120) NULL,
                action VARCHAR(64) NOT NULL,
                before_json JSONB NULL,
                after_json JSONB NULL,
                reason TEXT NULL,
                created_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
            """,
            "CREATE INDEX IF NOT EXISTS ix_fc_wh_pick_audit_task ON fc_wh_pick_audit(task_id,created_at,id)",
            """CREATE TABLE IF NOT EXISTS fc_wh_receipt_fbs (
                line_id VARCHAR(36) PRIMARY KEY,quantity NUMERIC(18,3) NOT NULL,
                reloadable_quantity NUMERIC(18,3) NOT NULL DEFAULT 0,
                actor VARCHAR(120),updated_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP)""",
            "ALTER TABLE fc_wh_receipt_fbs ADD COLUMN IF NOT EXISTS reloadable_quantity NUMERIC(18,3) NOT NULL DEFAULT 0",
            """CREATE TABLE IF NOT EXISTS fc_receipt_fbs_reset_releases (
                reset_line_id VARCHAR(36) NOT NULL,receipt_line_id VARCHAR(36) NOT NULL,
                quantity NUMERIC(18,3) NOT NULL,created_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP,
                PRIMARY KEY(reset_line_id,receipt_line_id))""",
            "CREATE INDEX IF NOT EXISTS ix_fc_receipt_fbs_reset_releases_receipt ON fc_receipt_fbs_reset_releases(receipt_line_id)",

        ]
        try:
            for stmt in statements:
                db.execute(text(stmt))
            db.commit()
            _SCHEMA_READY = True
        except Exception:
            db.rollback()
            raise


def _audit(db: Session, *, task_id: str, action: str, actor: str | None = None,
           line_id: str | None = None, before: dict | None = None, after: dict | None = None,
           reason: str | None = None) -> None:
    db.execute(text("""
        INSERT INTO fc_wh_pick_audit(id,task_id,line_id,actor,action,before_json,after_json,reason)
        VALUES(:id,:task_id,:line_id,:actor,:action,CAST(:before AS jsonb),CAST(:after AS jsonb),:reason)
    """), {
        "id": str(uuid.uuid4()), "task_id": task_id, "line_id": line_id,
        "actor": (actor or "")[:120] or None, "action": action[:64],
        "before": json.dumps(before, ensure_ascii=False) if before is not None else None,
        "after": json.dumps(after, ensure_ascii=False) if after is not None else None,
        "reason": reason,
    })


def _zone_from_comment(comment: str | None) -> str:
    value = str(comment or "").upper()
    if "ЗОНА:MAIN" in value or "ОСНОВНОЙ СКЛАД" in value or "STORAGE_ONLY" in value:
        return "MAIN"
    if "ЗОНА:FBS" in value:
        return "FBS"
    if "ВНУТРЕННИЙ УЧЁТ" in value:
        return "MAIN"
    return "LEGACY"


@router.get("/summary")
def warehouse_summary(client_id: uuid.UUID, warehouse_id: uuid.UUID, db: Session = Depends(get_db)):
    physical = _qty(db.execute(
        select(func.coalesce(func.sum(InventoryBalance.quantity), 0)).where(
            InventoryBalance.client_id == client_id, InventoryBalance.warehouse_id == warehouse_id, InventoryBalance.quantity > 0
        )
    ).scalar_one() or 0)
    try:
        main = _qty(db.execute(text("""
            SELECT COALESCE(SUM(storage_only),0) FROM fc_stocks_v2
            WHERE client_id=:client AND warehouse_id=:warehouse
        """), {"client": str(client_id), "warehouse": str(warehouse_id)}).scalar_one() or 0)
    except Exception:
        db.rollback(); main = Decimal("0")
    main = min(physical, max(Decimal("0"), main))
    return {"physical": float(physical), "main": float(main), "fbs": float(max(Decimal("0"), physical-main))}


# r76: all receipt formats share one history and one allocation service.
from app.services import receipt_fbs


@router.get("/receipts/recent")
def receipts_recent(client_id: uuid.UUID, warehouse_id: uuid.UUID,
        kind: str = '', limit: int = Query(100,ge=1,le=500), offset: int = Query(0,ge=0),
        query: str = Query('',max_length=200), db: Session = Depends(get_db)):
    return receipt_fbs.recent(db,client_id,warehouse_id,kind,limit,offset,query)


def _receipt_detail(db,op):
    return receipt_fbs.detail(db,op.client_id,op.warehouse_id,op.id)


def _admin_receipt_group_key(row:dict[str,Any])->str:
    idem=str(row.get("idempotency_key") or "").strip()
    upper=idem.upper()
    # FC_R145F35_RECEIPT_GROUP_PARITY
    parts=upper.split(':')
    if len(parts)>=4 and parts[0] in ('R145F27','R145F35','R145F37') and parts[1]=='MANUAL' and parts[2]:return 'MANUAL_DRAFT|'+parts[2]
    if len(parts)>=3 and parts[0] in ('R145F19','R145F35') and parts[1]=='TSD' and parts[2]:return 'TSD_DRAFT|'+parts[2]
    strong=upper.startswith(("PVZ_SESSION:","BULK-RECEIVE-","STRUCTURED-RECEIPT-","PALLET-RECEIPT:","PALLET-RECEIPT-"))
    if strong:return "IDEM|"+upper
    created=row.get("created_at")
    try:day=created.date().isoformat()
    except Exception:day=str(created or "")[:10]
    ref=str(row.get("reference") or "").strip()
    if ref:return f"REF|{str(row.get('type') or '').upper()}|{day}|{ref.casefold()}"
    return "OP|"+str(row.get("id") or "")


def _group_recent_receipts(rows:list[dict[str,Any]])->list[dict[str,Any]]:
    groups={}
    for raw in rows:
        row=dict(raw);key=_admin_receipt_group_key(row);cur=groups.get(key)
        if cur is None:
            cur={**row,"receipt_group_key":key,"operation_ids":[str(row.get("id"))],"operation_count":1,"_actors":[]}
            if row.get("performed_by"):cur["_actors"].append(str(row.get("performed_by")))
            groups[key]=cur;continue
        cur["operation_ids"].append(str(row.get("id")));cur["operation_count"]+=1
        for field in ("lines","units","transferred","tracked_transferred","reloadable"):
            cur[field]=float(cur.get(field) or 0)+float(row.get(field) or 0)
        actor=str(row.get("performed_by") or "").strip()
        if actor and actor not in cur["_actors"]:cur["_actors"].append(actor)
        candidate_date = history_date_key(row.get("created_at"))
        current_date = history_date_key(cur.get("created_at"))
        if candidate_date[0] and (not current_date[0] or candidate_date < current_date):
            cur["created_at"] = row["created_at"]
        if str(row.get("status") or "").upper()!="CANCELLED":cur["status"]=row.get("status") or "COMPLETED"
    result=[]
    for cur in groups.values():
        actors=cur.pop("_actors",[]);cur["performed_by"]=', '.join(actors[:4])+(f" +{len(actors)-4}" if len(actors)>4 else '')
        cur["id"]=cur["operation_ids"][0]
        result.append(cur)
    result.sort(key=lambda x:(history_date_key(x.get("created_at")),str(x.get("id") or "")),reverse=True)
    return result


@router.get("/receipts/recent-grouped")
def receipts_recent_grouped(client_id:uuid.UUID,warehouse_id:uuid.UUID,
        kind:str='',limit:int=Query(100,ge=1,le=500),offset:int=Query(0,ge=0),query:str=Query('',max_length=200),db:Session=Depends(get_db)):
    # Fetch enough raw fragments so one TSD/manual receipt persisted as many operations still becomes one card.
    raw_limit=min(10000,max(1000,limit*20))
    data=receipt_fbs.recent(db,client_id,warehouse_id,kind,raw_limit,0,query)
    grouped=_group_recent_receipts([dict(x) for x in (data.get("rows") or [])])
    page=grouped[offset:offset+limit]
    return {"rows":page,"total":len(grouped),"offset":offset,"limit":limit,"raw_total":data.get("total",0),"truncated":int(data.get("total",0) or 0)>raw_limit}


def _group_operation_scope(db:Session,client_id:uuid.UUID,warehouse_id:uuid.UUID,operation_ids:list[uuid.UUID])->list[str]:
    ids=[str(x) for x in operation_ids]
    if not ids or len(ids)>200:raise HTTPException(422,detail="Выберите от 1 до 200 частей приёмки")
    rows=db.execute(text("""SELECT CAST(id AS text) id FROM inventory_operations
        WHERE CAST(client_id AS text)=:c AND CAST(warehouse_id AS text)=:w
          AND operation_type IN ('RECEIVING','RECEIVING_IMPORT') AND CAST(id AS text) IN :ids""").bindparams(bindparam('ids',expanding=True)),{"c":str(client_id),"w":str(warehouse_id),"ids":ids}).scalars().all()
    allowed={str(x) for x in rows}
    if allowed!=set(ids):raise HTTPException(404,detail="Часть приёмки не найдена в выбранном складе")
    return ids


def _group_receipt_detail(db:Session,client_id:uuid.UUID,warehouse_id:uuid.UUID,operation_ids:list[uuid.UUID]):
    ids=_group_operation_scope(db,client_id,warehouse_id,operation_ids)
    ops=[dict(r) for r in db.execute(text("""SELECT CAST(o.id AS text) id,o.created_at,o.operation_type,o.status,o.idempotency_key,o.reference,o.comment,o.performed_by
        FROM inventory_operations o WHERE CAST(o.client_id AS text)=:c AND CAST(o.warehouse_id AS text)=:w AND CAST(o.id AS text) IN :ids ORDER BY o.created_at,o.id""").bindparams(bindparam('ids',expanding=True)),{"c":str(client_id),"w":str(warehouse_id),"ids":ids}).mappings().all()]
    lines=[dict(r) for r in db.execute(text("""SELECT CAST(l.operation_id AS text) operation_id,l.quantity,
        COALESCE(s.barcode,'') barcode,COALESCE(s.article,'') article,COALESCE(s.size,'') size,COALESCE(s.name,'') name,
        COALESCE(loc.code,'') location,COALESCE(NULLIF(c.name,''),c.barcode,'') container,
        o.created_at,o.reference,o.performed_by
        FROM inventory_operation_lines l JOIN inventory_operations o ON o.id=l.operation_id
        LEFT JOIN skus s ON s.id=l.sku_id
        LEFT JOIN locations loc ON loc.id=l.destination_location_id
        LEFT JOIN containers c ON c.id=l.destination_container_id
        WHERE CAST(o.client_id AS text)=:c AND CAST(o.warehouse_id AS text)=:w AND CAST(o.id AS text) IN :ids
        ORDER BY o.created_at,o.id,l.id""").bindparams(bindparam('ids',expanding=True)),{"c":str(client_id),"w":str(warehouse_id),"ids":ids}).mappings().all()]
    refs=[str(x.get("reference") or "").strip() for x in ops if str(x.get("reference") or "").strip()]
    title=refs[0] if refs else 'Приёмка'
    actors=[]
    for x in ops:
        a=str(x.get("performed_by") or "").strip()
        if a and a not in actors:actors.append(a)
    return {"title":title,"operation_ids":ids,"operation_count":len(ids),"operations":ops,"lines":lines,"units":sum(float(x.get("quantity") or 0) for x in lines),"positions":len({(x.get("barcode"),x.get("article"),x.get("size")) for x in lines}),"performed_by":', '.join(actors)}


def _resolve_receipt_group_ids(db:Session,client_id:uuid.UUID,warehouse_id:uuid.UUID,receipt_group_key:str|None,operation_id:list[uuid.UUID])->list[uuid.UUID]:
    if receipt_group_key:
        data=receipt_fbs.recent(db,client_id,warehouse_id,'',10000,0,'')
        groups=_group_recent_receipts([dict(x) for x in (data.get("rows") or [])])
        row=next((x for x in groups if str(x.get("receipt_group_key") or "")==str(receipt_group_key)),None)
        if not row: raise HTTPException(404,detail="Объединённая приёмка не найдена")
        return [uuid.UUID(str(x)) for x in (row.get("operation_ids") or [])]
    if not operation_id: raise HTTPException(422,detail="Не указана приёмка")
    return operation_id


@router.get("/receipts/group-detail")
def receipt_group_detail(client_id:uuid.UUID,warehouse_id:uuid.UUID,receipt_group_key:str|None=Query(default=None,max_length=500),operation_id:list[uuid.UUID]=Query(default=[]),db:Session=Depends(get_db)):
    ids=_resolve_receipt_group_ids(db,client_id,warehouse_id,receipt_group_key,operation_id)
    return _group_receipt_detail(db,client_id,warehouse_id,ids)


@router.get("/receipts/group-detail.xlsx")
def receipt_group_detail_xlsx(client_id:uuid.UUID,warehouse_id:uuid.UUID,receipt_group_key:str|None=Query(default=None,max_length=500),operation_id:list[uuid.UUID]=Query(default=[]),db:Session=Depends(get_db)):
    ids=_resolve_receipt_group_ids(db,client_id,warehouse_id,receipt_group_key,operation_id)
    data=_group_receipt_detail(db,client_id,warehouse_id,ids)
    wb=Workbook();summary=wb.active;summary.title='Приёмка';summary.append(['Название','Операций','Позиций','Единиц','Сотрудники']);append_safe(summary,[data['title'],data['operation_count'],data['positions'],data['units'],data['performed_by']])
    ws=wb.create_sheet('Состав');ws.append(['Дата','Сотрудник','Артикул','Размер','ШК','Товар','Количество','Ячейка','Короб','Операция'])
    for r in data['lines']:append_safe(ws,[r.get('created_at'),r.get('performed_by') or '',r.get('article') or '',r.get('size') or '',r.get('barcode') or '',r.get('name') or '',r.get('quantity') or 0,r.get('location') or '',r.get('container') or '',r.get('operation_id') or ''])
    for sheet in wb.worksheets:
        for c in sheet[1]:c.font=Font(bold=True)
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=min(45,max(12,max(len(str(c.value or '')) for c in col)+2))
    raw=workbook_bytes(wb)
    return StreamingResponse(io.BytesIO(raw),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="receipt_group_composition.xlsx"'})


@router.get("/receipts/{operation_id}")
def receipt_detail(operation_id: uuid.UUID,client_id: uuid.UUID,warehouse_id: uuid.UUID,db: Session=Depends(get_db)):
    return receipt_fbs.detail(db,client_id,warehouse_id,operation_id)


class ReceiptFbsLine(BaseModel):
    line_id: uuid.UUID
    quantity: int = Field(gt=0,le=10000000,strict=True)


class ReceiptToFbsRequest(BaseModel):
    request_id: uuid.UUID
    fingerprint: str = Field(min_length=64,max_length=64)
    lines: list[ReceiptFbsLine] = Field(min_length=1,max_length=10000)
    wb_warehouse_id: int | None = Field(default=None,gt=0)
    performed_by: str | None = Field(default=None,max_length=120)


@router.post("/receipts/{operation_id}/to-fbs")
def receipt_to_fbs(operation_id:uuid.UUID,payload:ReceiptToFbsRequest,
        client_id:uuid.UUID,warehouse_id:uuid.UUID,db:Session=Depends(get_db)):
    try:
        result=receipt_fbs.transfer(db,client_id,warehouse_id,operation_id,payload.model_dump(mode='json'),payload.performed_by)
        db.commit()
    except Exception:
        db.rollback();raise
    return result


class ReceiptPublishRequest(BaseModel):
    wb_warehouse_id:int|None=Field(default=None,gt=0)


@router.post("/receipts/{operation_id}/publish-fbs")
def receipt_publish_fbs(operation_id:uuid.UUID,payload:ReceiptPublishRequest,
        client_id:uuid.UUID,warehouse_id:uuid.UUID,db:Session=Depends(get_db)):
    try:
        data=receipt_fbs.detail(db,client_id,warehouse_id,operation_id)
        if data['receipt']['status']!='COMPLETED':raise HTTPException(409,detail='Приёмка отменена')
        ids={x['sku_id'] for x in data['lines'] if x['transferred']>0 or x['initially_fbs']}
        if not ids:raise HTTPException(409,detail='Сначала переведите товар этой приёмки в FBS')
        from app.services import stock_actions as stock
        stock.locks(db,str(client_id),str(warehouse_id),ids)
        if receipt_fbs.metadata(db,receipt_fbs.get_operation(db,client_id,warehouse_id,operation_id))['status']!='COMPLETED':
            raise HTTPException(409,detail='Приёмка отменена')
        result=receipt_fbs.queue_publication(db,client_id,warehouse_id,ids,payload.wb_warehouse_id)
        db.commit();return {'ok':True,'publication':result}
    except Exception:db.rollback();raise


@router.get("/receipts/{operation_id}/export.xlsx")
def receipt_export(operation_id:uuid.UUID,client_id:uuid.UUID,warehouse_id:uuid.UUID,db:Session=Depends(get_db)):
    d=receipt_fbs.detail(db,client_id,warehouse_id,operation_id,include_marketplace=False)
    wb=Workbook();ws=wb.active;ws.title='Приёмка'
    append_safe(ws,['Приёмка',d['receipt']['reference']]);append_safe(ws,['Статус',d['receipt']['status']])
    if d['receipt'].get('type')=='TSD':
        append_safe(ws,['Фактически начато ТСД',str(d['receipt'].get('factual_started_at') or '')])
        append_safe(ws,['Фактически принят короб',str(d['receipt'].get('factual_received_at') or '')])
        append_safe(ws,['Выгружено на сервер',str(d['receipt'].get('posted_at') or '')])
    ws.append(['Ячейка','Короб','Артикул','Размер','ШК','Принято','Первый скан','Последний скан','Сейчас в FBS из приёмки','Повторно после обнуления','Доступно для загрузки'])
    header_row=ws.max_row
    for r in d['lines']:
        append_safe(ws,[r['location'],r['container'],r['article'],r['size'],r['barcode'],r['quantity_received'],r.get('first_scanned_at') or '',r.get('last_scanned_at') or '',r['transferred'],r.get('reloadable_quantity',0),r['available_quantity']])
        for col in (2,3,4,5):ws.cell(ws.max_row,col).number_format='@'
    ws.freeze_panes=f'A{header_row+1}';ws.auto_filter.ref=f'A{header_row}:K{ws.max_row}'
    for col,width in [('A',22),('B',32),('C',24),('D',16),('E',24),('F',14),('G',24),('H',24),('I',25),('J',29),('K',25)]:ws.column_dimensions[col].width=width
    hist=wb.create_sheet('Переводы FBS');hist.append(['Дата','Сотрудник','ШК','Ячейка','Короб','Количество'])
    for e in reversed(d['events']):
        for r in ((e.get('result_json') or {}).get('details') or []):
            append_safe(hist,[str(e['created_at']),e['actor'],r['barcode'],r['location'],r['container'],r['quantity']])
            hist.cell(hist.max_row,3).number_format='@'
    data=workbook_bytes(wb)
    return StreamingResponse(io.BytesIO(data),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="receipt.xlsx"'})


# ---------------- warehouse pick tasks ----------------

def _header(value: Any) -> str:
    return re.sub(r"[^A-ZА-Я0-9]", "", str(value or "").strip().upper().replace("Ё", "Е"))


_ALIASES = {
    "barcode": {"ШК", "ШКТОВАРА", "ШТРИХКОД", "ШТРИХКОДТОВАРА", "BARCODE", "EAN", "EAN13", "GTIN"},
    "article": {"АРТИКУЛ", "АРТИКУЛТОВАРА", "АРТИКУЛПРОДАВЦА", "АРТИКУЛПОСТАВЩИКА", "ARTICLE", "VENDORCODE"},
    "size": {"РАЗМЕР", "РАЗМЕРТОВАРА", "SIZE", "TECHSIZE"},
    "qty": {"КОЛИЧЕСТВО", "КОЛВО", "КОЛИЧЕСТВОШТ", "QTY", "QUANTITY", "НУЖНО", "СОБРАТЬ", "СОБРАТЬШТ"},
}


def _parse_pick_file(content: bytes) -> list[dict[str, Any]]:
    from app.services.box_intake_rules import read_sheet
    try: return read_sheet(content, 'goods')
    except ValueError as e: raise HTTPException(422, detail=str(e))


def _find_pick_sku(db: Session, client_id: uuid.UUID, row: dict[str, Any]) -> SKU | None:
    barcode = str(row.get("barcode") or "").strip()
    if barcode:
        matches = db.execute(select(SKU).where(SKU.client_id == client_id, SKU.barcode == barcode, SKU.is_active.is_(True)).limit(2)).scalars().all()
        if len(matches)>1:raise HTTPException(409,detail='ШК неоднозначен в справочнике: '+barcode)
        return matches[0] if matches else None
    article = str(row.get("article") or "").strip()
    size = str(row.get("size") or "").strip()
    if article:
        stmt = select(SKU).where(SKU.client_id == client_id, func.lower(func.coalesce(SKU.article, "")) == article.lower())
        if size:
            stmt = stmt.where(func.lower(func.coalesce(SKU.size, "")) == size.lower())
        matches = db.execute(stmt.limit(3)).scalars().all()
        if len(matches) == 1:
            return matches[0]
    return None


def _task_row(db: Session, task_id: str) -> dict[str, Any]:
    _ensure_schema(db)
    row = db.execute(text("SELECT * FROM fc_wh_pick_tasks WHERE id=:id"), {"id": task_id}).mappings().first()
    if not row:
        raise HTTPException(status_code=404, detail="Сборочное задание не найдено")
    return dict(row)



def _mutable_task(db, task_id, *, tsd=False):
    _ensure_schema(db)
    ensure_stock_core_v2_schema(db)
    initial=_task_row(db,task_id)
    lock_scope(db,initial['client_id'],initial['warehouse_id'])
    task=dict(db.execute(text('SELECT * FROM fc_wh_pick_tasks WHERE id=:id FOR UPDATE'),{'id':task_id}).mappings().one())
    if tsd and task['status'] not in ('ASSIGNED','IN_PROGRESS'):
        raise HTTPException(409,detail='Задание уже передано администратору, закрыто или снято со сборщика. Обновите ТСД')
    return task


def _task_lines(db: Session, task_id: str) -> list[dict[str, Any]]:
    rows = db.execute(text("""
        SELECT l.*, s.barcode,s.article,s.size,s.name,
               loc.code AS location_code,loc.barcode AS location_barcode,
               c.name AS container_name,c.barcode AS container_barcode
        FROM fc_wh_pick_lines l
        LEFT JOIN skus s ON CAST(s.id AS varchar)=l.sku_id
        LEFT JOIN locations loc ON CAST(loc.id AS varchar)=l.location_id
        LEFT JOIN containers c ON CAST(c.id AS varchar)=l.container_id
        WHERE l.task_id=:task_id
        ORDER BY l.sort_order,l.id
    """), {"task_id": task_id}).mappings().all()
    return route_snapshots(db,"MAIN",task_id,rows)


def _pick_payload(db: Session, task_id: str) -> dict[str, Any]:
    task = _task_row(db, task_id)
    lines = _task_lines(db, task_id)
    total = picked = scan = manual = missing = skipped = shortage = 0.0
    tasks = []
    for row in lines:
        alloc = _qf(row.get("allocated_qty")); req = _qf(row.get("requested_qty"))
        sc = _qf(row.get("scan_qty")); man = _qf(row.get("manual_qty")); miss = _qf(row.get("missing_qty"))
        handled = min(alloc, sc + man + miss); remaining = max(0.0, alloc - handled)
        deferred = bool(row.get("deferred")) and remaining > 0
        if alloc <= 0:
            shortage += req; continue
        total += alloc; scan += sc; manual += man; picked += sc + man; missing += miss
        if deferred: skipped += remaining
        tasks.append({
            **{k:v for k,v in row.items() if k.startswith("route_") or k=="_route_order"},
            "task_key": str(row["id"]), "line_id": str(row["id"]),
            "location_id": str(row.get("location_id") or ""), "location": str(row.get("location_code") or ""), "location_barcode": str(row.get("location_barcode") or ""),
            "container_id": str(row.get("container_id") or "") or None, "container": str(row.get("container_name") or row.get("container_barcode") or ""), "container_barcode": str(row.get("container_barcode") or ""),
            "sku_id": str(row.get("sku_id") or ""), "barcode": str(row.get("barcode") or ""), "barcodes": [str(row.get("barcode"))] if row.get("barcode") else [],
            "article": str(row.get("article") or ""), "size": str(row.get("size") or ""), "name": str(row.get("name") or ""),
            "required": alloc, "picked": sc + man, "scan": sc, "manual": man, "missing": miss,
            "remaining": remaining, "deferred": deferred, "complete": remaining <= 0.000001,
        })
    pick_mode=str(task.get("pick_mode") or "CELL").upper()
    if pick_mode=="ARTICLE":
        tasks.sort(key=lambda x: (_natural(x["article"]),_natural(x["size"]),x.get("_route_order",10**9),_natural(x["location"]),_natural(x["container"])))
    else:
        tasks.sort(key=lambda x: (x.get("_route_order",10**9),_natural(x["location"]),_natural(x["container"]),_natural(x["article"]),_natural(x["size"])))
    current = next((x for x in tasks if not x["complete"] and not x["deferred"]), None)
    same_location = [x for x in tasks if current and x["location_id"] == current["location_id"] and (not current.get("route_row") or x.get("route_row")==current["route_row"])]
    current_location = None
    if current:
        current_location = {
            "location_id": current["location_id"], "location": current["location"],
            "route_row":current.get("route_row"), "route_label":current.get("route_label"),
            "required": sum(x["required"] for x in same_location), "picked": sum(x["picked"] for x in same_location),
            "missing": sum(x["missing"] for x in same_location), "remaining": sum(x["remaining"] for x in same_location if not x["deferred"]),
            "task_count": len(same_location),
            "items": [{"task_key": x["task_key"], "article": x["article"], "size": x["size"], "container": x["container"], "required": x["required"], "picked": x["picked"], "missing": x["missing"], "remaining": x["remaining"]} for x in same_location],
        }
    return {
        "session": {"id": str(task["id"]), "kind": "WAREHOUSE", "task_number": task["task_number"], "name": task.get("name") or "", "client_id": task["client_id"], "warehouse_id": task["warehouse_id"], "picker": task.get("picker") or "", "status": task["status"], "created_at": task["created_at"], "assigned_at": task.get("assigned_at"), "assigned_by": task.get("assigned_by"), "source": task.get("source") or "ADMIN", "seller_username": task.get("seller_username") or "", "parent_task_id": task.get("parent_task_id") or "", "pick_mode": pick_mode},
        "progress": {"total": total, "picked": picked, "scan": scan, "manual": manual, "missing": missing, "skipped": skipped, "shortage": shortage, "done": picked + missing + skipped, "remaining": max(0.0, total-picked-missing-skipped)},
        "task_count": len(tasks), "tasks_done": sum(1 for x in tasks if x["complete"]),
        "current": current, "current_location": current_location, "tasks": tasks,
        "can_submit": current is None and str(task["status"]).upper() in {"ASSIGNED", "IN_PROGRESS"},
    }


@router.get("/picks/template.xlsx")
def pick_template():
    from pathlib import Path
    from fastapi.responses import FileResponse
    path=Path(__file__).resolve().parents[1]/'static'/'templates'/'r68_goods.xlsx'
    return FileResponse(path,filename='r68_goods.xlsx',media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@router.post("/picks/import")
def import_pick_task(
    client_id: uuid.UUID = Form(...), warehouse_id: uuid.UUID = Form(...),
    task_number: str = Form(...), name: str | None = Form(None), created_by: str | None = Form(None),
    source: str = Form("ADMIN"), seller_username: str | None = Form(None),
    pick_mode: str = Form("CELL"), seller_only: bool = Form(False),
    file: UploadFile = File(...), db: Session = Depends(get_db),
):
    _ensure_schema(db)
    ensure_stock_core_v2_schema(db)
    task_number = str(task_number or "").strip()
    if len(task_number)>120: raise HTTPException(422,detail="Номер задания не длиннее 120 символов")
    if not task_number:
        raise HTTPException(status_code=400, detail="Укажите номер/название сборочного задания")
    content = file.file.read(8*1024*1024+1)
    if len(content)>8*1024*1024: raise HTTPException(422,detail="Excel больше 8 МБ")
    requests = _parse_pick_file(content)
    source=str(source or "ADMIN").strip().upper()[:32] or "ADMIN"
    pick_mode=str(pick_mode or "CELL").strip().upper()
    if pick_mode not in {"CELL","ARTICLE"}: raise HTTPException(422,detail="Способ сборки: CELL или ARTICLE")
    if seller_only:
        bad=[r.get("row") for r in requests if not str(r.get("barcode") or "").strip()]
        if bad: raise HTTPException(422,detail=f"Кабинет селлера принимает только формат ШК + количество. Нет ШК в строках: {', '.join(map(str,bad[:20]))}")
        for r in requests:
            r["article"]="";r["size"]="";r["location"]="";r["container"]=""
    # Exact box/cell requests allocate first so automatic requests cannot consume their stock.
    requests.sort(key=lambda r:(-int(bool(r.get('container')))-int(bool(r.get('location'))),not bool(r.get('container')),r['row']))
    db.execute(text("SELECT pg_advisory_xact_lock(hashtext(:key))"), {"key": f"fc:r62:wh-pick:{client_id}:{warehouse_id}"})
    existing = db.execute(text("SELECT id FROM fc_wh_pick_tasks WHERE client_id=:c AND warehouse_id=:w AND task_number=:n"), {"c":str(client_id),"w":str(warehouse_id),"n":task_number}).scalar()
    if existing:
        db.rollback(); raise HTTPException(status_code=409, detail="Задание с таким номером уже существует")
    task_id=str(uuid.uuid4())
    db.execute(text("""INSERT INTO fc_wh_pick_tasks(id,client_id,warehouse_id,task_number,name,source_file_name,source,seller_username,pick_mode,status,created_by)
        VALUES(:id,:c,:w,:n,:name,:file,:source,:seller,:mode,'DRAFT',:by)"""), {"id":task_id,"c":str(client_id),"w":str(warehouse_id),"n":task_number[:120],"name":str(name or "")[:240] or None,"file":str(file.filename or "")[:255] or None,"source":source,"seller":str(seller_username or "")[:160] or None,"mode":pick_mode,"by":str(created_by or "")[:120] or None})
    _audit(db,task_id=task_id,action="CREATE",actor=created_by,after={"task_number":task_number,"rows":len(requests)},reason="Импорт Excel")
    line_no=0; requested_total=allocated_total=shortage_total=0
    errors=[]
    source_buckets=[{**x,'storage_only':x['main']} for x in route_enrich(db,client_id,warehouse_id,available_buckets(db,client_id,warehouse_id))]
    for req in requests:
        sku=_find_pick_sku(db,client_id,req); requested=Decimal(req["qty"]); requested_total+=int(requested)
        if not sku:
            errors.append({"row":req["row"],"article":req["article"],"size":req["size"],"barcode":req["barcode"],"error":"Товар не найден"})
            db.execute(text("""INSERT INTO fc_wh_pick_lines(id,task_id,sku_id,requested_qty,allocated_qty,line_status,sort_order,request_barcode,request_article,request_size)
                VALUES(:id,:task,:sku,:requested,0,'SHORTAGE',:sort,:barcode,:article,:size)"""), {"id":str(uuid.uuid4()),"task":task_id,"sku":None,"requested":requested,"sort":line_no,"barcode":req["barcode"],"article":req["article"],"size":req["size"]})
            shortage_total+=int(requested); line_no+=1; continue
        # Main-warehouse stock only: use explicit STORAGE_ONLY bucket; do not take FBS stock.
        buckets=[x for x in source_buckets if x['sku_id']==str(sku.id)]
        if req.get('container'):
            from app.services.box_intake import find_box
            exact=find_box(db,client_id,warehouse_id,req['container'])
            if not exact:raise HTTPException(422,detail=f"Строка {req['row']}: короб {req['container']} не найден")
            buckets=[b for b in buckets if b.get('container_id')==exact['id']]
        if req.get('location'):
            value=req['location'].casefold()
            buckets=[b for b in buckets if value in (str(b.get('location_code') or '').casefold(),str(b.get('location_barcode') or '').casefold())]
        buckets.sort(key=route_key)
        remaining=requested
        for b in buckets:
            if remaining<=0: break
            free=max(Decimal("0"),_qty(b.get("storage_only"))-_qty(b.get("held"))); take=min(remaining,free)
            if take<=0: continue
            line_no+=1; lid=str(uuid.uuid4())
            db.execute(text("""INSERT INTO fc_wh_pick_lines(id,task_id,sku_id,location_id,container_id,requested_qty,allocated_qty,line_status,sort_order,request_barcode,request_article,request_size)
                VALUES(:id,:task,:sku,:loc,:cont,:q,:q,'PENDING',:sort,:barcode,:article,:size)"""), {"id":lid,"task":task_id,"sku":str(sku.id),"loc":str(b["location_id"]),"cont":str(b.get("container_id") or "") or None,"q":take,"sort":line_no,"barcode":req["barcode"],"article":req["article"],"size":req["size"]})
            b['held']=_qty(b.get('held'))+take
            allocated_total+=int(take); remaining-=take
        if remaining>0:
            line_no+=1; shortage_total+=int(remaining)
            db.execute(text("""INSERT INTO fc_wh_pick_lines(id,task_id,sku_id,requested_qty,allocated_qty,line_status,sort_order,request_barcode,request_article,request_size)
                VALUES(:id,:task,:sku,:requested,0,'SHORTAGE',:sort,:barcode,:article,:size)"""), {"id":str(uuid.uuid4()),"task":task_id,"sku":str(sku.id),"requested":remaining,"sort":line_no,"barcode":req["barcode"],"article":req["article"],"size":req["size"]})
    freeze_task(db,task_id)
    db.commit()
    return {"ok":True,"task_id":task_id,"task_number":task_number,"requested":requested_total,"allocated":allocated_total,"reserved":allocated_total,"shortage":shortage_total,"errors":errors[:100],"status":"DRAFT","source":source,"pick_mode":pick_mode}


@router.get("/picks")
def list_pick_tasks(client_id: uuid.UUID, warehouse_id: uuid.UUID, limit:int=Query(default=100,ge=1,le=500), db:Session=Depends(get_db)):
    _ensure_schema(db)
    rows=db.execute(text("""
      SELECT t.*,
        COALESCE(SUM(l.requested_qty),0) requested,
        COALESCE(SUM(l.allocated_qty),0) allocated,
        COALESCE(SUM(l.scan_qty),0) scan,
        COALESCE(SUM(l.manual_qty),0) manual,
        COALESCE(SUM(l.missing_qty),0) missing,
        COALESCE(SUM(CASE WHEN l.deferred THEN GREATEST(l.allocated_qty-l.scan_qty-l.manual_qty-l.missing_qty,0) ELSE 0 END),0) skipped,
        COALESCE(SUM(CASE WHEN l.allocated_qty=0 THEN l.requested_qty ELSE 0 END),0) shortage
      FROM fc_wh_pick_tasks t LEFT JOIN fc_wh_pick_lines l ON l.task_id=t.id
      WHERE t.client_id=:c AND t.warehouse_id=:w
      GROUP BY t.id ORDER BY t.created_at DESC LIMIT :limit
    """),{"c":str(client_id),"w":str(warehouse_id),"limit":limit}).mappings().all()
    return {"rows":[{**dict(r),**{k:_qf(r.get(k)) for k in ["requested","allocated","scan","manual","missing","skipped","shortage"]}} for r in rows]}


@router.get("/picks/{task_id}")
def pick_detail(task_id:str, db:Session=Depends(get_db)):
    _ensure_schema(db); task=_task_row(db,task_id); lines=_task_lines(db,task_id)
    audits=[dict(x) for x in db.execute(text("SELECT * FROM fc_wh_pick_audit WHERE task_id=:id ORDER BY created_at DESC,id DESC LIMIT 500"),{"id":task_id}).mappings().all()]
    payload=_pick_payload(db,task_id)
    return {"task":task,"lines":[{**r,"scan_qty":_qf(r.get("scan_qty")),"manual_qty":_qf(r.get("manual_qty")),"missing_qty":_qf(r.get("missing_qty")),"allocated_qty":_qf(r.get("allocated_qty")),"requested_qty":_qf(r.get("requested_qty")),"skipped_qty":max(0.0,_qf(r.get("allocated_qty"))-_qf(r.get("scan_qty"))-_qf(r.get("manual_qty"))-_qf(r.get("missing_qty"))) if bool(r.get("deferred")) else 0.0} for r in lines],"audit":audits,"progress":payload["progress"]}


class AssignWarehousePick(BaseModel):
    picker:str=Field(min_length=1,max_length=120)
    assigned_by:str|None=Field(default=None,max_length=120)
    pick_mode:str=Field(default="CELL",max_length=16)


@router.post("/picks/{task_id}/assign")
def assign_wh_pick(task_id:str,payload:AssignWarehousePick,db:Session=Depends(get_db)):
    task=_mutable_task(db,task_id); status=str(task["status"]).upper()
    if status not in {"DRAFT","ASSIGNED"}: raise HTTPException(409,detail="Назначать можно только новое/неначатое задание")
    other=db.execute(text("SELECT id,task_number FROM fc_wh_pick_tasks WHERE client_id=:c AND warehouse_id=:w AND LOWER(picker)=LOWER(:p) AND status IN ('ASSIGNED','IN_PROGRESS') AND id<>:id LIMIT 1"),{"c":task["client_id"],"w":task["warehouse_id"],"p":payload.picker.strip(),"id":task_id}).mappings().first()
    if other: raise HTTPException(409,detail=f"У сборщика уже есть активное складское задание {other['task_number']}")
    mode=str(payload.pick_mode or "CELL").strip().upper()
    if mode not in {"CELL","ARTICLE"}: raise HTTPException(422,detail="Способ сборки: по ячейкам или по артикулу")
    before={"picker":task.get("picker"),"status":task["status"],"pick_mode":task.get("pick_mode")}; after={"picker":payload.picker.strip(),"status":"ASSIGNED","pick_mode":mode}
    db.execute(text("UPDATE fc_wh_pick_tasks SET picker=:p,pick_mode=:mode,status='ASSIGNED',assigned_by=:by,assigned_at=CURRENT_TIMESTAMP WHERE id=:id"),{"p":payload.picker.strip(),"mode":mode,"by":payload.assigned_by,"id":task_id})
    _audit(db,task_id=task_id,action="ASSIGN",actor=payload.assigned_by,before=before,after=after); db.commit(); return _pick_payload(db,task_id)


class SplitAssignWarehousePick(BaseModel):
    pickers:list[str]=Field(min_length=2,max_length=20)
    assigned_by:str|None=Field(default=None,max_length=120)
    pick_mode:str=Field(default="CELL",max_length=16)


def _split_group_key(row:dict[str,Any],mode:str):
    if mode=="ARTICLE":
        return (str(row.get("article") or row.get("request_article") or row.get("barcode") or row.get("request_barcode") or row.get("sku_id") or "").casefold(),)
    return (str(row.get("route_section") or ""),str(row.get("route_row") or ""),str(row.get("location_id") or ""))


@router.post("/picks/{task_id}/split-assign")
def split_assign_wh_pick(task_id:str,payload:SplitAssignWarehousePick,db:Session=Depends(get_db)):
    task=_mutable_task(db,task_id);status=str(task["status"]).upper()
    if status!="DRAFT": raise HTTPException(409,detail="Распределять между несколькими сборщиками можно только новое задание")
    mode=str(payload.pick_mode or "CELL").strip().upper()
    if mode not in {"CELL","ARTICLE"}: raise HTTPException(422,detail="Способ сборки: по ячейкам или по артикулу")
    pickers=[]
    for raw in payload.pickers:
        name=str(raw or "").strip()
        if name and name.casefold() not in {x.casefold() for x in pickers}: pickers.append(name[:120])
    if len(pickers)<2: raise HTTPException(422,detail="Укажите минимум двух разных сборщиков")
    for picker in pickers:
        other=db.execute(text("SELECT task_number FROM fc_wh_pick_tasks WHERE client_id=:c AND warehouse_id=:w AND LOWER(picker)=LOWER(:p) AND status IN ('ASSIGNED','IN_PROGRESS') LIMIT 1"),{"c":task["client_id"],"w":task["warehouse_id"],"p":picker}).scalar()
        if other: raise HTTPException(409,detail=f"У сборщика {picker} уже есть активное задание {other}")
    raw_lines=_task_lines(db,task_id)
    if any(_qf(x.get("scan_qty"))+_qf(x.get("manual_qty"))+_qf(x.get("missing_qty"))>0 for x in raw_lines):
        raise HTTPException(409,detail="Задание уже начали собирать; сначала снимите/завершите текущую работу")
    lines=[x for x in raw_lines if _qf(x.get("allocated_qty"))>0]
    if not lines: raise HTTPException(409,detail="В задании нет распределённого товара")
    groups={}
    for row in lines: groups.setdefault(_split_group_key(row,mode),[]).append(row)
    ordered=list(groups.values())
    if mode=="ARTICLE": ordered.sort(key=lambda g:(_natural(g[0].get("article") or g[0].get("request_article")),_natural(g[0].get("size") or g[0].get("request_size"))))
    else: ordered.sort(key=lambda g:min(int(x.get("_route_order") or 10**9) for x in g))
    bins=[{"picker":p,"qty":0.0,"groups":[]} for p in pickers]
    for group in ordered:
        target=min(bins,key=lambda b:(b["qty"],pickers.index(b["picker"])))
        target["groups"].append(group);target["qty"]+=sum(_qf(x.get("allocated_qty")) for x in group)
    children=[]
    for idx,bucket in enumerate((b for b in bins if b["groups"]),1):
        cid=str(uuid.uuid4());base=f"{task['task_number']}-{idx}"
        number=base[:120];suffix=1
        while db.execute(text("SELECT 1 FROM fc_wh_pick_tasks WHERE client_id=:c AND warehouse_id=:w AND task_number=:n"),{"c":task["client_id"],"w":task["warehouse_id"],"n":number}).scalar():
            suffix+=1;number=(base[:110]+f"-{suffix}")[:120]
        db.execute(text("""INSERT INTO fc_wh_pick_tasks(id,client_id,warehouse_id,task_number,name,source_file_name,source,seller_username,parent_task_id,pick_mode,picker,status,created_by,assigned_by,assigned_at,note)
            VALUES(:id,:c,:w,:n,:name,:file,:source,:seller,:parent,:mode,:picker,'ASSIGNED',:created,:by,CURRENT_TIMESTAMP,:note)"""),{
            "id":cid,"c":task["client_id"],"w":task["warehouse_id"],"n":number,"name":((task.get("name") or task["task_number"])+f" · {bucket['picker']}")[:240],"file":task.get("source_file_name"),"source":task.get("source") or "ADMIN","seller":task.get("seller_username"),"parent":task_id,"mode":mode,"picker":bucket["picker"],"created":task.get("created_by"),"by":payload.assigned_by,"note":f"Часть задания {task['task_number']}"})
        sort_order=0
        for group in bucket["groups"]:
            for row in group:
                sort_order+=1
                db.execute(text("""INSERT INTO fc_wh_pick_lines(id,task_id,sku_id,location_id,container_id,requested_qty,allocated_qty,scan_qty,manual_qty,missing_qty,deferred,line_status,sort_order,request_barcode,request_article,request_size)
                    VALUES(:id,:task,:sku,:loc,:cont,:requested,:allocated,0,0,0,FALSE,'PENDING',:sort,:barcode,:article,:size)"""),{
                    "id":str(uuid.uuid4()),"task":cid,"sku":row.get("sku_id"),"loc":row.get("location_id"),"cont":row.get("container_id"),"requested":row.get("allocated_qty") or 0,"allocated":row.get("allocated_qty") or 0,"sort":sort_order,"barcode":row.get("request_barcode") or row.get("barcode") or "","article":row.get("request_article") or row.get("article") or "","size":row.get("request_size") or row.get("size") or ""})
        freeze_task(db,cid)
        _audit(db,task_id=cid,action="CREATE_SPLIT",actor=payload.assigned_by,after={"parent_task_id":task_id,"picker":bucket["picker"],"pick_mode":mode,"quantity":bucket["qty"]},reason="Часть общего сборочного задания")
        _audit(db,task_id=cid,action="ASSIGN",actor=payload.assigned_by,after={"picker":bucket["picker"],"status":"ASSIGNED","pick_mode":mode})
        children.append({"id":cid,"task_number":number,"picker":bucket["picker"],"quantity":bucket["qty"]})
    db.execute(text("UPDATE fc_wh_pick_tasks SET status='SPLIT',pick_mode=:mode,assigned_by=:by,assigned_at=CURRENT_TIMESTAMP WHERE id=:id"),{"mode":mode,"by":payload.assigned_by,"id":task_id})
    _audit(db,task_id=task_id,action="SPLIT_ASSIGN",actor=payload.assigned_by,before={"status":"DRAFT"},after={"status":"SPLIT","pick_mode":mode,"children":children},reason="Задание распределено между сборщиками")
    db.commit()
    return {"ok":True,"parent_task_id":task_id,"pick_mode":mode,"children":children,"assigned":len(children)}


class UnassignWarehousePick(BaseModel):
    actor:str|None=Field(default=None,max_length=120)
    reason:str|None=Field(default=None,max_length=500)


@router.post("/picks/{task_id}/unassign")
def unassign_wh_pick(task_id:str,payload:UnassignWarehousePick,db:Session=Depends(get_db)):
    task=_mutable_task(db,task_id); status=str(task["status"]).upper()
    if status not in {"ASSIGNED","IN_PROGRESS"}: raise HTTPException(409,detail="Снять сборщика можно только с активного незавершённого задания")
    before={"picker":task.get("picker"),"status":task["status"]}; after={"picker":None,"status":"DRAFT"}
    db.execute(text("UPDATE fc_wh_pick_tasks SET picker=NULL,status='DRAFT',assigned_by=NULL,assigned_at=NULL WHERE id=:id"),{"id":task_id})
    _audit(db,task_id=task_id,action="UNASSIGN",actor=payload.actor,before=before,after=after,reason=payload.reason or "Сборщик снят администратором"); db.commit(); return _pick_payload(db,task_id)


class EditWarehouseLine(BaseModel):
    scan_qty:float=Field(ge=0)
    manual_qty:float=Field(ge=0)
    missing_qty:float=Field(ge=0)
    reason:str=Field(min_length=1,max_length=500)
    actor:str|None=Field(default=None,max_length=120)


@router.patch("/picks/{task_id}/lines/{line_id}")
def edit_wh_pick_line(task_id:str,line_id:str,payload:EditWarehouseLine,db:Session=Depends(get_db)):
    task=_mutable_task(db,task_id)
    if str(task["status"]).upper() in {"CLOSED","CANCELLED"}: raise HTTPException(409,detail="Закрытое/отменённое задание редактировать нельзя")
    row=db.execute(text("SELECT * FROM fc_wh_pick_lines WHERE id=:line AND task_id=:task FOR UPDATE"),{"line":line_id,"task":task_id}).mappings().first()
    if not row: raise HTTPException(404,detail="Позиция не найдена")
    alloc=_qf(row.get("allocated_qty")); total=payload.scan_qty+payload.manual_qty+payload.missing_qty
    if any(not Decimal(str(q)).is_finite() or q != int(q) for q in (payload.scan_qty,payload.manual_qty,payload.missing_qty)):
        raise HTTPException(422,detail="Факт должен быть целым неотрицательным количеством")
    if total>0 and not row.get('location_id'):
        raise HTTPException(409,detail="У строки нехватки нет адреса. Создайте дополнительную сборку из актуального остатка")
    if not payload.reason.strip(): raise HTTPException(422,detail="Укажите причину исправления")
    before={"scan":_qf(row.get("scan_qty")),"manual":_qf(row.get("manual_qty")),"missing":_qf(row.get("missing_qty")),"skipped":max(0.0,alloc-_qf(row.get("scan_qty"))-_qf(row.get("manual_qty"))-_qf(row.get("missing_qty"))) if row.get("deferred") else 0}
    deferred=total<alloc-1e-6; line_status="DONE" if not deferred else "DEFERRED"
    db.execute(text("UPDATE fc_wh_pick_lines SET scan_qty=:s,manual_qty=:m,missing_qty=:x,deferred=:d,line_status=:st,updated_at=CURRENT_TIMESTAMP WHERE id=:id"),{"s":payload.scan_qty,"m":payload.manual_qty,"x":payload.missing_qty,"d":deferred,"st":line_status,"id":line_id})
    after={"scan":payload.scan_qty,"manual":payload.manual_qty,"missing":payload.missing_qty,"skipped":max(0.0,alloc-total)}
    _audit(db,task_id=task_id,line_id=line_id,action="ADMIN_EDIT",actor=payload.actor,before=before,after=after,reason=payload.reason); db.commit(); return pick_detail(task_id,db)


class ActorRequest(BaseModel):
    actor:str|None=Field(default=None,max_length=120)
    reason:str|None=Field(default=None,max_length=500)


@router.post("/picks/{task_id}/close")
def close_wh_pick(task_id:str,payload:ActorRequest,db:Session=Depends(get_db)):
    task=_mutable_task(db,task_id)
    if str(task["status"]).upper()=="CLOSED":
        return {"ok":True,"task_id":task_id,"operation_id":task.get("close_operation_id"),"already_closed":True}
    if str(task["status"]).upper()!="READY_FOR_ADMIN": raise HTTPException(409,detail="Задание ещё не передано администратору с ТСД")
    lines=_task_lines(db,task_id)
    unresolved=[]; stock_lines=[]; buckets=[]
    consumed={}  # Aggregate repeated task lines at the same physical bucket.
    for r in lines:
        alloc=_qty(r.get("allocated_qty")); scan=_qty(r.get("scan_qty")); manual=_qty(r.get("manual_qty")); missing=_qty(r.get("missing_qty")); handled=scan+manual+missing
        if alloc<=0: continue
        if handled+Decimal("0.000001")<alloc:
            unresolved.append(str(r["id"])); continue
        loc=r.get("location_id"); cont=r.get("container_id") or None
        bucket=reconcile_stock_bucket(db,client_id=uuid.UUID(str(task["client_id"])),warehouse_id=uuid.UUID(str(task["warehouse_id"])),sku_id=uuid.UUID(str(r["sku_id"])),location_id=uuid.UUID(str(loc)),container_id=uuid.UUID(str(cont)) if cont else None,preferred_bucket="STORAGE_ONLY",lock=True)
        bucket_key=(str(r['sku_id']),str(loc),str(cont or ''))
        other=held_quantity(db,task['client_id'],task['warehouse_id'],r['sku_id'],loc,cont,exclude_task=task_id)
        usable=max(Decimal('0'),min(_qty(bucket.get('storage_only')),_qty(bucket.get('physical'))-_qty(bucket.get('fbs_available'))-_qty(bucket.get('reserved')))-other)
        consumed[bucket_key]=consumed.get(bucket_key,Decimal('0'))+handled
        if usable+Decimal("0.000001")<consumed[bucket_key]:
            db.rollback(); raise HTTPException(409,detail=f"Основного остатка уже недостаточно в {r.get('location_code') or loc} / {r.get('container_name') or ''}. Обновите факт или остаток перед закрытием.")
        picked=scan+manual
        if picked>0:
            stock_lines.append(StockLine(sku_id=uuid.UUID(str(r["sku_id"])),quantity=picked,source_location_id=uuid.UUID(str(loc)),source_container_id=uuid.UUID(str(cont)) if cont else None,reason=f"WAREHOUSE_PICK_COLLECTED:{task['task_number']}"))
        if missing>0:
            stock_lines.append(StockLine(sku_id=uuid.UUID(str(r["sku_id"])),quantity=missing,source_location_id=uuid.UUID(str(loc)),source_container_id=uuid.UUID(str(cont)) if cont else None,reason=f"WAREHOUSE_PICK_NOT_IN_BOX:{task['task_number']}"))
        buckets.append((r,handled))
    if unresolved:
        db.rollback(); raise HTTPException(409,detail=f"Есть пропущенные/неразобранные позиции: {len(unresolved)}. Отредактируйте факт перед закрытием.")
    op=None
    if stock_lines:
        op=execute_operation(db,OperationCreate(client_id=uuid.UUID(str(task["client_id"])),warehouse_id=uuid.UUID(str(task["warehouse_id"])),operation_type="SHIPPING",idempotency_key=f"wh-pick-close-{task_id}",performed_by=payload.actor,reference=str(task["task_number"]),comment=f"Сборка со склада · {task['task_number']} · администратор закрыл",lines=stock_lines),commit=False)
        for r,_handled in buckets:
            reconcile_stock_bucket(db,client_id=uuid.UUID(str(task["client_id"])),warehouse_id=uuid.UUID(str(task["warehouse_id"])),sku_id=uuid.UUID(str(r["sku_id"])),location_id=uuid.UUID(str(r["location_id"])),container_id=uuid.UUID(str(r["container_id"])) if r.get("container_id") else None,preferred_bucket="STORAGE_ONLY",lock=True)
    db.execute(text("UPDATE fc_wh_pick_tasks SET status='CLOSED',closed_by=:by,closed_at=CURRENT_TIMESTAMP,close_operation_id=:op WHERE id=:id"),{"by":payload.actor,"op":str(getattr(op,"id","") or "") or None,"id":task_id})
    emptied = archive_consumed_empty_containers(
        db, task["client_id"], task["warehouse_id"],
        [r.get("container_id") for r in lines if r.get("container_id")],
        actor=payload.actor, reason="WAREHOUSE_PICK_CLOSED_EMPTY_BOX",
    )
    _audit(db,task_id=task_id,action="CLOSE",actor=payload.actor,after={"operation_id":str(getattr(op,"id","") or ""),"posted":sum(float(x.quantity) for x in stock_lines),"archived_empty_boxes":emptied.get("count",0)},reason=payload.reason or "Warehouse pick closed")
    db.commit()
    return {"ok":True,"task_id":task_id,"operation_id":str(getattr(op,"id","") or "") or None,"posted":sum(float(x.quantity) for x in stock_lines),"archived_empty_boxes":emptied.get("count",0)}


@router.post("/picks/import-ready-writeoff")
def import_ready_pick_and_writeoff(
    client_id: uuid.UUID = Form(...), warehouse_id: uuid.UUID = Form(...),
    task_number: str = Form(...), name: str | None = Form(None), created_by: str | None = Form(None),
    file: UploadFile = File(...), db: Session = Depends(get_db),
):
    """Import a factual MAIN-warehouse file and immediately post it only when fully available.

    This deliberately reuses the normal warehouse allocation engine. If even one requested
    unit cannot be allocated from free MAIN stock, nothing is written off; the created DRAFT
    remains available for review as an ordinary picking task.
    """
    result=import_pick_task(client_id=client_id,warehouse_id=warehouse_id,task_number=task_number,name=name,created_by=created_by,source='ADMIN',seller_username=None,pick_mode='CELL',seller_only=False,file=file,db=db)
    if float(result.get("shortage") or 0)>0:
        return {**result,"written_off":False,"message":"Файл не списан: есть нехватка. Создан обычный сборочный лист для проверки."}
    task_id=str(result["task_id"])
    # The file is explicitly declared as the physical fact: every allocated unit is collected.
    db.execute(text("""UPDATE fc_wh_pick_lines SET scan_qty=0,manual_qty=allocated_qty,missing_qty=0,
        deferred=FALSE,line_status=CASE WHEN allocated_qty>0 THEN 'DONE' ELSE line_status END,updated_at=CURRENT_TIMESTAMP
        WHERE task_id=:id AND allocated_qty>0"""),{"id":task_id})
    db.execute(text("UPDATE fc_wh_pick_tasks SET status='READY_FOR_ADMIN',submitted_at=CURRENT_TIMESTAMP WHERE id=:id"),{"id":task_id})
    _audit(db,task_id=task_id,action="READY_FILE_FACT",actor=created_by,after={"allocated":result.get("allocated",0)},reason="Готовый файл: факт подтверждён без ТСД")
    posted=close_wh_pick(task_id,ActorRequest(actor=created_by,reason="Готовый файл Основного склада: прямое списание факта"),db)
    return {**result,**posted,"written_off":True,"message":"Готовый файл проведён: свободный остаток Основного склада списан."}


@router.post("/picks/{task_id}/cancel")
def cancel_wh_pick(task_id:str,payload:ActorRequest,db:Session=Depends(get_db)):
    task=_mutable_task(db,task_id)
    if str(task["status"]).upper() in {"CLOSED","CANCELLED"}: return {"ok":True,"status":task["status"]}
    if str(task["status"]).upper()=="SPLIT":
        db.execute(text("UPDATE fc_wh_pick_tasks SET status='CANCELLED',cancelled_at=CURRENT_TIMESTAMP WHERE parent_task_id=:id AND status IN ('DRAFT','ASSIGNED','IN_PROGRESS','READY_FOR_ADMIN')"),{"id":task_id})
    db.execute(text("UPDATE fc_wh_pick_tasks SET status='CANCELLED',cancelled_at=CURRENT_TIMESTAMP WHERE id=:id"),{"id":task_id})
    _audit(db,task_id=task_id,action="CANCEL",actor=payload.actor,before={"status":task["status"]},after={"status":"CANCELLED"},reason=payload.reason or "Отмена до финального списания"); db.commit(); return {"ok":True,"status":"CANCELLED"}


@router.get("/picks/{task_id}/report.xlsx")
def wh_pick_report(task_id:str,db:Session=Depends(get_db)):
    detail=pick_detail(task_id,db); task=detail["task"]; lines=detail["lines"]; audit=detail["audit"]
    wb=Workbook(); fact=wb.active; fact.title="Факт сборки"; missing=wb.create_sheet("Не найдено"); skipped=wb.create_sheet("Пропущено"); source=wb.create_sheet("Исходный запрос"); hist=wb.create_sheet("История изменений")
    route=wb.create_sheet("Маршрут сборки",0)
    route.append(["№","Секция","Ряд","Палета / ячейка","Место в ряду","Короб","Артикул","Размер","ШК","План"] )
    for ix,r in enumerate((r for r in lines if _qf(r.get("allocated_qty"))>0),1):
        append_safe(route,[ix,r.get("route_section") or "",r.get("route_row") or "",r.get("location_code") or "",r.get("route_position") or "",r.get("container_name") or r.get("container_barcode") or "",r.get("article") or "",r.get("size") or "",r.get("barcode") or "",_qf(r.get("allocated_qty"))])
    common=["Ячейка","Короб","Артикул","Размер","ШК","План","Скан","Вручную","Нет в коробе"]
    for ws in (fact,missing,skipped): ws.append(common)
    source.append(["Артикул","Размер","ШК","Запрошено","Выделено"]); hist.append(["Дата","Кто","Действие","Позиция","Было","Стало","Причина"])
    for r in lines:
        row=[r.get("location_code") or "",r.get("container_name") or r.get("container_barcode") or "",r.get("article") or r.get("request_article") or "",r.get("size") or r.get("request_size") or "",r.get("barcode") or r.get("request_barcode") or "",r.get("allocated_qty") or 0,r.get("scan_qty") or 0,r.get("manual_qty") or 0,r.get("missing_qty") or 0]
        if _qf(r.get("scan_qty"))+_qf(r.get("manual_qty"))>0: append_safe(fact,row)
        if _qf(r.get("missing_qty"))>0: append_safe(missing,row)
        if _qf(r.get("skipped_qty"))>0: append_safe(skipped,row)
        append_safe(source,[r.get("article") or r.get("request_article") or "",r.get("size") or r.get("request_size") or "",r.get("barcode") or r.get("request_barcode") or "",r.get("requested_qty") or 0,r.get("allocated_qty") or 0])
    for a in audit: append_safe(hist,[str(a.get("created_at") or ""),a.get("actor") or "",a.get("action") or "",a.get("line_id") or "",json.dumps(a.get("before_json"),ensure_ascii=False) if a.get("before_json") is not None else "",json.dumps(a.get("after_json"),ensure_ascii=False) if a.get("after_json") is not None else "",a.get("reason") or ""])
    for ws in wb.worksheets:
        for c in ws[1]: c.font=Font(bold=True); c.fill=PatternFill("solid",fgColor="D9EAD3")
        ws.freeze_panes="A2"
        for row in ws.iter_rows():
            for c in row: c.alignment=Alignment(vertical="top",wrap_text=True)
        for col in range(1,ws.max_column+1): ws.column_dimensions[chr(64+col) if col<=26 else "A"].width=18
    data=workbook_bytes(wb)
    return StreamingResponse(io.BytesIO(data),media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":'attachment; filename="warehouse_pick_report.xlsx"'})


# ---------------- TSD facade for warehouse tasks ----------------
class TsdAction(BaseModel):
    task_key:str
    code:str|None=None
    quantity:float|None=Field(default=None,gt=0,le=100000)
    confirm:bool=False
    actor:str|None=Field(default=None,max_length=120)


def _touch_task(db:Session,task_id:str):
    db.execute(text("UPDATE fc_wh_pick_tasks SET status='IN_PROGRESS' WHERE id=:id AND status='ASSIGNED'"),{"id":task_id})


@router.get("/tsd/assigned")
def assigned_wh_pick(client_id:uuid.UUID,warehouse_id:uuid.UUID,picker:str,db:Session=Depends(get_db)):
    _ensure_schema(db)
    row=db.execute(text("""SELECT id FROM fc_wh_pick_tasks WHERE client_id=:c AND warehouse_id=:w AND LOWER(picker)=LOWER(:p) AND status IN ('ASSIGNED','IN_PROGRESS') ORDER BY assigned_at DESC NULLS LAST,created_at DESC LIMIT 1"""),{"c":str(client_id),"w":str(warehouse_id),"p":picker.strip()}).mappings().first()
    return _pick_payload(db,str(row["id"])) if row else {"session":None,"assignment_required":True}


@router.post("/picks/{task_id}/verify-location")
def wh_verify_location(task_id:str,payload:TsdAction,db:Session=Depends(get_db)):
    state=_pick_payload(db,task_id); cur=state.get("current")
    if not cur or str(cur["task_key"])!=str(payload.task_key): raise HTTPException(409,detail="Позиция уже изменилась. Обновите задание")
    scanned=re.sub(r"\s+","",str(payload.code or "").strip().upper()); accepted={re.sub(r"\s+","",str(x or "").strip().upper()) for x in [cur.get("location"),cur.get("location_barcode")] if x}
    if scanned not in accepted: raise HTTPException(409,detail=f"Неверная ячейка. Нужна {cur.get('location')}")
    return {"ok":True}


@router.post("/picks/{task_id}/verify-container")
def wh_verify_container(task_id:str,payload:TsdAction,db:Session=Depends(get_db)):
    state=_pick_payload(db,task_id); cur=state.get("current")
    if not cur or str(cur["task_key"])!=str(payload.task_key): raise HTTPException(409,detail="Позиция уже изменилась. Обновите задание")
    if not cur.get("container"): return {"ok":True}
    scanned=re.sub(r"\s+","",str(payload.code or "").strip().upper()); accepted={re.sub(r"\s+","",str(x or "").strip().upper()) for x in [cur.get("container"),cur.get("container_barcode")] if x}
    if scanned not in accepted: raise HTTPException(409,detail=f"Неверный короб. Нужен {cur.get('container')}")
    return {"ok":True}


def _lock_line(db:Session,task_id:str,line_id:str):
    row=db.execute(text("SELECT * FROM fc_wh_pick_lines WHERE task_id=:task AND id=:line FOR UPDATE"),{"task":task_id,"line":line_id}).mappings().first()
    if not row: raise HTTPException(404,detail="Позиция задания не найдена")
    return dict(row)


@router.post("/picks/{task_id}/scan")
def wh_scan(task_id:str,payload:TsdAction,db:Session=Depends(get_db)):
    task=_mutable_task(db,task_id,tsd=True); row=_lock_line(db,task_id,payload.task_key)
    if bool(row.get("deferred")): db.execute(text("UPDATE fc_wh_pick_lines SET deferred=FALSE WHERE id=:id"),{"id":row["id"]})
    sku=db.get(SKU,uuid.UUID(str(row["sku_id"]))); expected=re.sub(r"\s+","",str(getattr(sku,"barcode","") or "").strip().upper()); scanned=re.sub(r"\s+","",str(payload.code or "").strip().upper())
    if not expected or scanned!=expected: db.rollback(); raise HTTPException(409,detail=f"Неверный товар. Нужен {getattr(sku,'article','') or ''} {getattr(sku,'size','') or ''}")
    alloc=_qf(row["allocated_qty"]); current=_qf(row["scan_qty"])+_qf(row["manual_qty"])+_qf(row["missing_qty"])
    if current>=alloc-1e-6: db.rollback(); raise HTTPException(409,detail="Позиция уже обработана")
    before={"scan":_qf(row["scan_qty"]),"manual":_qf(row["manual_qty"]),"missing":_qf(row["missing_qty"])}; new_scan=_qf(row["scan_qty"])+1
    done=new_scan+_qf(row["manual_qty"])+_qf(row["missing_qty"])>=alloc-1e-6
    db.execute(text("UPDATE fc_wh_pick_lines SET scan_qty=:q,deferred=FALSE,line_status=:st,updated_at=CURRENT_TIMESTAMP WHERE id=:id"),{"q":new_scan,"st":"DONE" if done else "PARTIAL","id":row["id"]}); _touch_task(db,task_id)
    _audit(db,task_id=task_id,line_id=str(row["id"]),action="TSD_SCAN",actor=task.get("picker"),before=before,after={**before,"scan":new_scan}); db.commit(); return _pick_payload(db,task_id)


@router.post("/picks/{task_id}/manual")
def wh_manual(task_id:str,payload:TsdAction,db:Session=Depends(get_db)):
    task=_mutable_task(db,task_id,tsd=True); row=_lock_line(db,task_id,payload.task_key); amount=float(payload.quantity or 1)
    if not Decimal(str(amount)).is_finite() or amount!=int(amount): raise HTTPException(422,detail="Количество должно быть целым")
    alloc=_qf(row["allocated_qty"]); current=_qf(row["scan_qty"])+_qf(row["manual_qty"])+_qf(row["missing_qty"]); remaining=max(0.0,alloc-current)
    if amount>remaining+1e-6: db.rollback(); raise HTTPException(409,detail=f"Осталось подтвердить только {remaining:g} шт.")
    before={"scan":_qf(row["scan_qty"]),"manual":_qf(row["manual_qty"]),"missing":_qf(row["missing_qty"])}; new_manual=_qf(row["manual_qty"])+amount
    done=current+amount>=alloc-1e-6
    db.execute(text("UPDATE fc_wh_pick_lines SET manual_qty=:q,deferred=FALSE,line_status=:st,updated_at=CURRENT_TIMESTAMP WHERE id=:id"),{"q":new_manual,"st":"DONE" if done else "PARTIAL","id":row["id"]}); _touch_task(db,task_id)
    _audit(db,task_id=task_id,line_id=str(row["id"]),action="TSD_MANUAL",actor=task.get("picker"),before=before,after={**before,"manual":new_manual},reason="Ручное подтверждение факта"); db.commit(); return _pick_payload(db,task_id)


@router.post("/picks/{task_id}/skip")
def wh_skip(task_id:str,payload:TsdAction,db:Session=Depends(get_db)):
    task=_mutable_task(db,task_id,tsd=True); row=_lock_line(db,task_id,payload.task_key)
    remaining=max(0.0,_qf(row["allocated_qty"])-_qf(row["scan_qty"])-_qf(row["manual_qty"])-_qf(row["missing_qty"]))
    if remaining<=1e-6: db.rollback(); return _pick_payload(db,task_id)
    db.execute(text("UPDATE fc_wh_pick_lines SET deferred=TRUE,line_status='DEFERRED',updated_at=CURRENT_TIMESTAMP WHERE id=:id"),{"id":row["id"]}); _touch_task(db,task_id)
    _audit(db,task_id=task_id,line_id=str(row["id"]),action="TSD_SKIP",actor=task.get("picker"),before={"skipped":0},after={"skipped":remaining},reason="Позиция пропущена сборщиком"); db.commit(); return _pick_payload(db,task_id)


@router.post("/picks/{task_id}/missing")
def wh_missing(task_id:str,payload:TsdAction,db:Session=Depends(get_db)):
    if not payload.confirm: raise HTTPException(400,detail='Нажмите подтверждение «Вы уверены?»')
    task=_mutable_task(db,task_id,tsd=True); row=_lock_line(db,task_id,payload.task_key)
    remaining=max(0.0,_qf(row["allocated_qty"])-_qf(row["scan_qty"])-_qf(row["manual_qty"])-_qf(row["missing_qty"]))
    before={"missing":_qf(row["missing_qty"]),"skipped":remaining if row.get("deferred") else 0}; new_missing=_qf(row["missing_qty"])+remaining
    db.execute(text("UPDATE fc_wh_pick_lines SET missing_qty=:q,deferred=FALSE,line_status='DONE',updated_at=CURRENT_TIMESTAMP WHERE id=:id"),{"q":new_missing,"id":row["id"]}); _touch_task(db,task_id)
    _audit(db,task_id=task_id,line_id=str(row["id"]),action="TSD_NOT_IN_BOX",actor=task.get("picker"),before=before,after={"missing":new_missing,"skipped":0},reason="Подтверждено: единицы нет в указанном коробе. Остаток пока не списан."); db.commit(); return _pick_payload(db,task_id)


@router.post("/picks/{task_id}/submit")
def wh_submit(task_id:str,payload:ActorRequest,db:Session=Depends(get_db)):
    task=_mutable_task(db,task_id)
    if str(task['status']).upper()=='CLOSED':return {'ok':True,'status':'CLOSED','already_closed':True,'operation_id':task.get('close_operation_id')}
    state=_pick_payload(db,task_id)
    if state.get("current") is not None: raise HTTPException(409,detail="Сначала обработайте или пропустите текущие позиции")
    if str(task["status"]).upper() not in {"ASSIGNED","IN_PROGRESS"}: raise HTTPException(409,detail="Задание нельзя завершить в текущем статусе")
    db.execute(text("UPDATE fc_wh_pick_tasks SET status='READY_FOR_ADMIN',submitted_at=CURRENT_TIMESTAMP WHERE id=:id"),{"id":task_id}); _audit(db,task_id=task_id,action="SUBMIT_ADMIN",actor=payload.actor or task.get("picker"),before={"status":task["status"]},after={"status":"READY_FOR_ADMIN","skipped":state["progress"]["skipped"]}); return close_wh_pick(task_id,ActorRequest(actor=payload.actor or task.get('picker'),reason='Сборщик завершил сборочный лист'),db)
