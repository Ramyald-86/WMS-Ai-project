"""Authenticated mobile receiving and inventory, same catalogue and stock ledger."""
import io
import json
import uuid
from datetime import datetime
from typing import Literal
from fastapi import APIRouter, Depends, Request, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field, model_validator
from sqlalchemy import text
from sqlalchemy.orm import Session
from app.core.database import get_db
from app.routers.warehouse_ui import scope
from app.services import tsd_stocktake as service

router=APIRouter(prefix='/tsd/mobile/work', tags=['Приёмка и инвентаризация ТСД'])


def owner(request):
    value=getattr(request.state,'username',None)
    if not value:raise HTTPException(401,detail='Войдите в WMS')
    return str(value)


class Context(BaseModel):
    client_id:uuid.UUID
    warehouse_id:uuid.UUID


class Start(Context):
    request_id:uuid.UUID
    mode:Literal['COUNT','RECEIVE']
    picker:str=Field(min_length=1,max_length=120)
    code:str=Field(min_length=1,max_length=255)
    target_type:Literal['AUTO','BOX','LOCATION']='AUTO'
    section_code:str=Field(default='',max_length=255)
    placement_code:str=Field(default='',max_length=255)
    place_empty_box:bool=False
    row:int|None=Field(default=None,ge=1,le=999,strict=True)
    receipt_number:str=Field(default='',max_length=160)

    @model_validator(mode='after')
    def clean(self):
        self.picker=self.picker.strip()
        if not self.picker:raise ValueError('Укажите сотрудника')
        return self


class ReceiptBatch(Context):
    receipt_number:str=Field(min_length=1,max_length=160)
    picker:str=Field(min_length=1,max_length=120)

    @model_validator(mode='after')
    def clean(self):
        self.receipt_number=self.receipt_number.strip()
        self.picker=self.picker.strip()
        if not self.receipt_number:raise ValueError('\u0423\u043a\u0430\u0436\u0438\u0442\u0435 \u043d\u043e\u043c\u0435\u0440 \u0438\u043b\u0438 \u043d\u0430\u0437\u0432\u0430\u043d\u0438\u0435 \u043f\u0440\u0438\u0451\u043c\u043a\u0438')
        if not self.picker:raise ValueError('\u0423\u043a\u0430\u0436\u0438\u0442\u0435 \u043f\u0440\u0438\u0451\u043c\u0449\u0438\u043a\u0430')
        return self


class Product(BaseModel):
    article:str=Field(min_length=1,max_length=120)
    size:str=Field(min_length=1,max_length=80)
    name:str=Field(default='',max_length=200)
    barcode_confirmed:bool=False


class Edit(Context):
    event_id:uuid.UUID
    version:int=Field(ge=0,strict=True)
    action:Literal['SCAN','SET','RESET','REMOVE','UNDO','CHANGE_PLACE','CANCEL','POST']
    barcode:str|None=Field(default=None,max_length=80)
    sku_id:uuid.UUID|None=None
    quantity:int|None=Field(default=None,ge=0,le=1000000,strict=True)
    product:Product|None=None
    place_code:str|None=Field(default=None,max_length=255)
    target_type:Literal['AUTO','BOX','LOCATION']='AUTO'
    section_code:str=Field(default='',max_length=255)
    placement_code:str=Field(default='',max_length=255)
    place_empty_box:bool=False
    row:int|None=Field(default=None,ge=1,le=999,strict=True)
    receipt_number:str=Field(default='',max_length=160)
    occurred_at:datetime|None=None


@router.post('/start')
def start(payload:Start, request:Request, db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id)
    try:
        result=service.start(db,payload.model_dump(mode='json'),owner(request));db.commit();return result
    except Exception:db.rollback();raise


@router.get('/drafts')
def drafts(client_id:uuid.UUID,warehouse_id:uuid.UUID,picker:str,request:Request,db:Session=Depends(get_db)):
    scope(db,client_id,warehouse_id)
    # Show the full active receipt for visibility, but keep every row attributed
    # to its receiver. COUNT remains scoped to the selected receiver.
    rows=db.execute(text('''SELECT * FROM fc_tsd_work WHERE client_id=:c AND warehouse_id=:w
        AND owner=:owner AND status='DRAFT' AND (mode='RECEIVE' OR picker=:picker)
        ORDER BY updated_at DESC'''),
        {'c':str(client_id),'w':str(warehouse_id),'owner':owner(request),'picker':picker.strip()}).mappings().all()
    return {'rows':[service.public(r) for r in rows]}


def _receipt_drafts(db, client_id, warehouse_id, actor, number, picker=None, include_posted=False):
    params={'c':str(client_id),'w':str(warehouse_id),'o':actor,'n':number}
    picker=str(picker or '').strip()
    picker_sql=''
    if picker:
        picker_sql=' AND picker=:picker'
        params['picker']=picker
    status_sql="status IN ('DRAFT','POSTED')" if include_posted else "status='DRAFT'"
    return db.execute(text('''SELECT id,owner,picker,mode,status,data,created_at,updated_at FROM fc_tsd_work
        WHERE client_id=:c AND warehouse_id=:w AND owner=:o AND mode='RECEIVE' AND '''+status_sql+'''
          AND COALESCE(data->>'receipt_number','')=:n'''+picker_sql+''' ORDER BY created_at,id'''),params).mappings().all()


@router.get('/receipt-batch/export.xlsx')
def receipt_batch_export(client_id:uuid.UUID,warehouse_id:uuid.UUID,receipt_number:str,request:Request,picker:str|None=None,db:Session=Depends(get_db)):
    scope(db,client_id,warehouse_id)
    actor=owner(request); number=str(receipt_number or '').strip(); picker=str(picker or '').strip()
    if not number: raise HTTPException(422,detail='\u0423\u043a\u0430\u0436\u0438\u0442\u0435 \u043d\u043e\u043c\u0435\u0440 \u0438\u043b\u0438 \u043d\u0430\u0437\u0432\u0430\u043d\u0438\u0435 \u043f\u0440\u0438\u0451\u043c\u043a\u0438')
    rows=_receipt_drafts(db,client_id,warehouse_id,actor,number,picker or None,include_posted=True)
    if not rows:
        detail=('\u041f\u043e \u043f\u0440\u0438\u0451\u043c\u0449\u0438\u043a\u0443 '+picker+' \u043d\u0435\u0442 \u0434\u0430\u043d\u043d\u044b\u0445 \u0432 \u044d\u0442\u043e\u0439 \u043f\u0440\u0438\u0451\u043c\u043a\u0435') if picker else '\u041f\u043e \u044d\u0442\u043e\u0439 \u043f\u0440\u0438\u0451\u043c\u043a\u0435 \u043d\u0435\u0442 \u0434\u0430\u043d\u043d\u044b\u0445 \u0434\u043b\u044f \u0432\u044b\u0433\u0440\u0443\u0437\u043a\u0438'
        raise HTTPException(404,detail=detail)
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
    def decoded(value):
        return value if isinstance(value,dict) else json.loads(value or '{}')
    def safe_place(data):
        saved=data.get('place') or {}; warning=''
        try: current=service.current_place(db,client_id,warehouse_id,saved,allow_box_move=True)
        except HTTPException as exc:
            current=saved; warning=str(exc.detail)
        return saved,current,warning
    wb=Workbook(); summary=wb.active; summary.title='\u0421\u0432\u043e\u0434\u043a\u0430'
    total_units=0; total_lines=0; pickers=set(); box_rows=[]; item_rows=[]; by_picker={}
    for r in rows:
        data=decoded(r['data']); saved,current,warning=safe_place(data); lines=list((data.get('lines') or {}).values())
        units=sum(max(0,int(x.get('quantity') or 0)) for x in lines); total_units+=units
        active_items=[x for x in lines if int(x.get('quantity') or 0)>0]; total_lines+=len(active_items)
        receiver=str(r['picker'] or '').strip(); pickers.add(receiver)
        stats=by_picker.setdefault(receiver,{'boxes':0,'lines':0,'units':0,'first':'','last':'','statuses':set()})
        stats['boxes']+=1; stats['lines']+=len(active_items); stats['units']+=units; stats['statuses'].add(str(r['status'] or ''))
        first=str(data.get('first_item_at') or ''); last=str(data.get('last_item_at') or '')
        if first and (not stats['first'] or first<stats['first']): stats['first']=first
        if last and (not stats['last'] or last>stats['last']): stats['last']=last
        box_rows.append([number,str(r['id']),receiver,str(r['status'] or ''),saved.get('container') or saved.get('location') or '',saved.get('barcode') or '',saved.get('location') or '',current.get('location') or '',units,len(active_items),first,last,str(r['updated_at'] or ''),warning])
        for x in active_items:
            item_rows.append([number,receiver,str(r['status'] or ''),current.get('container') or saved.get('container') or '',current.get('location') or '',x.get('barcode') or '',x.get('article') or '',x.get('size') or '',x.get('name') or '',int(x.get('quantity') or 0),x.get('first_scanned_at') or '',x.get('last_scanned_at') or '',str(r['id'])])
    summary.append(['\u041f\u043e\u043a\u0430\u0437\u0430\u0442\u0435\u043b\u044c','\u0417\u043d\u0430\u0447\u0435\u043d\u0438\u0435'])
    summary.append(['\u041d\u043e\u043c\u0435\u0440 / \u043d\u0430\u0437\u0432\u0430\u043d\u0438\u0435 \u043f\u0440\u0438\u0451\u043c\u043a\u0438',number])
    summary.append(['\u0420\u0435\u0436\u0438\u043c \u0432\u044b\u0433\u0440\u0443\u0437\u043a\u0438',('\u0422\u043e\u043b\u044c\u043a\u043e \u043f\u0440\u0438\u0451\u043c\u0449\u0438\u043a: '+picker) if picker else '\u0412\u0441\u044f \u043f\u0440\u0438\u0451\u043c\u043a\u0430'])
    summary.append(['\u041a\u043e\u0440\u043e\u0431\u043e\u0432',len(rows)])
    summary.append(['\u0422\u043e\u0432\u0430\u0440\u043d\u044b\u0445 \u0441\u0442\u0440\u043e\u043a',total_lines])
    summary.append(['\u041f\u0440\u0438\u043d\u044f\u0442\u043e \u0435\u0434\u0438\u043d\u0438\u0446',total_units])
    summary.append(['\u041f\u0440\u0438\u0451\u043c\u0449\u0438\u043a\u043e\u0432',len([x for x in pickers if x])])
    summary.append(['\u041f\u0440\u0438\u0451\u043c\u0449\u0438\u043a\u0438',', '.join(sorted(x for x in pickers if x))])
    summary.append(['\u0423\u0447\u0451\u0442\u043d\u0430\u044f \u0437\u0430\u043f\u0438\u0441\u044c WMS',actor])
    summary.append(['\u041f\u0440\u0438\u043c\u0435\u0447\u0430\u043d\u0438\u0435','\u0412\u044b\u0433\u0440\u0443\u0437\u043a\u0430 \u0432\u043a\u043b\u044e\u0447\u0430\u0435\u0442 \u0430\u043a\u0442\u0438\u0432\u043d\u044b\u0435 \u0438 \u0443\u0436\u0435 \u0437\u0430\u0432\u0435\u0440\u0448\u0451\u043d\u043d\u044b\u0435 \u0447\u0430\u0441\u0442\u0438 \u044d\u0442\u043e\u0439 \u043f\u0440\u0438\u0451\u043c\u043a\u0438. \u041a\u0430\u0436\u0434\u044b\u0439 \u043f\u0440\u0438\u0451\u043c\u0449\u0438\u043a \u0443\u0447\u0438\u0442\u044b\u0432\u0430\u0435\u0442\u0441\u044f \u043e\u0442\u0434\u0435\u043b\u044c\u043d\u043e; \u043f\u0435\u0440\u0435\u043d\u043e\u0441 \u043a\u043e\u0440\u043e\u0431\u0430 \u043d\u0430 \u043f\u0430\u043b\u0435\u0442\u0443 \u043d\u0435 \u043c\u0435\u043d\u044f\u0435\u0442 \u0430\u0432\u0442\u043e\u0440\u0430 \u043f\u0440\u0438\u0451\u043c\u043a\u0438.'])
    staff=wb.create_sheet('\u041f\u043e \u043f\u0440\u0438\u0451\u043c\u0449\u0438\u043a\u0430\u043c')
    staff.append(['\u041f\u0440\u0438\u0451\u043c\u043a\u0430','\u041f\u0440\u0438\u0451\u043c\u0449\u0438\u043a','\u0421\u0442\u0430\u0442\u0443\u0441','\u041a\u043e\u0440\u043e\u0431\u043e\u0432','\u0422\u043e\u0432\u0430\u0440\u043d\u044b\u0445 \u0441\u0442\u0440\u043e\u043a','\u0415\u0434\u0438\u043d\u0438\u0446','\u041f\u0435\u0440\u0432\u044b\u0439 \u0441\u043a\u0430\u043d','\u041f\u043e\u0441\u043b\u0435\u0434\u043d\u0438\u0439 \u0441\u043a\u0430\u043d'])
    for receiver in sorted(by_picker):
        st=by_picker[receiver]; staff.append([number,receiver,', '.join(sorted(x for x in st['statuses'] if x)),st['boxes'],st['lines'],st['units'],st['first'],st['last']])
    boxes=wb.create_sheet('\u041a\u043e\u0440\u043e\u0431\u0430')
    boxes.append(['\u041f\u0440\u0438\u0451\u043c\u043a\u0430','ID \u0447\u0435\u0440\u043d\u043e\u0432\u0438\u043a\u0430','\u041f\u0440\u0438\u0451\u043c\u0449\u0438\u043a','\u0421\u0442\u0430\u0442\u0443\u0441','\u041a\u043e\u0440\u043e\u0431','\u0428\u041a \u043a\u043e\u0440\u043e\u0431\u0430','\u0410\u0434\u0440\u0435\u0441 \u043f\u0440\u0438 \u0441\u043a\u0430\u043d\u0438\u0440\u043e\u0432\u0430\u043d\u0438\u0438','\u0422\u0435\u043a\u0443\u0449\u0438\u0439 \u0430\u0434\u0440\u0435\u0441','\u0415\u0434\u0438\u043d\u0438\u0446','\u0422\u043e\u0432\u0430\u0440\u043d\u044b\u0445 \u0441\u0442\u0440\u043e\u043a','\u041f\u0435\u0440\u0432\u044b\u0439 \u0442\u043e\u0432\u0430\u0440','\u041f\u043e\u0441\u043b\u0435\u0434\u043d\u0438\u0439 \u0442\u043e\u0432\u0430\u0440','\u041e\u0431\u043d\u043e\u0432\u043b\u0435\u043d\u043e','\u041f\u0440\u043e\u0432\u0435\u0440\u043a\u0430 \u0442\u0435\u043a\u0443\u0449\u0435\u0433\u043e \u043c\u0435\u0441\u0442\u0430'])
    for row in box_rows: boxes.append(row)
    items=wb.create_sheet('\u0422\u043e\u0432\u0430\u0440\u044b')
    items.append(['\u041f\u0440\u0438\u0451\u043c\u043a\u0430','\u041f\u0440\u0438\u0451\u043c\u0449\u0438\u043a','\u0421\u0442\u0430\u0442\u0443\u0441','\u041a\u043e\u0440\u043e\u0431','\u0422\u0435\u043a\u0443\u0449\u0438\u0439 \u0430\u0434\u0440\u0435\u0441','\u0428\u041a \u0442\u043e\u0432\u0430\u0440\u0430','\u0410\u0440\u0442\u0438\u043a\u0443\u043b','\u0420\u0430\u0437\u043c\u0435\u0440','\u0422\u043e\u0432\u0430\u0440','\u041a\u043e\u043b\u0438\u0447\u0435\u0441\u0442\u0432\u043e','\u041f\u0435\u0440\u0432\u044b\u0439 \u0441\u043a\u0430\u043d','\u041f\u043e\u0441\u043b\u0435\u0434\u043d\u0438\u0439 \u0441\u043a\u0430\u043d','ID \u0447\u0435\u0440\u043d\u043e\u0432\u0438\u043a\u0430'])
    for row in item_rows: items.append(row)
    scans=wb.create_sheet('\u0421\u043a\u0430\u043d\u044b')
    scans.append(['\u041f\u043e\u043b\u0443\u0447\u0435\u043d\u043e \u0441\u0435\u0440\u0432\u0435\u0440\u043e\u043c','\u0424\u0430\u043a\u0442\u0438\u0447\u0435\u0441\u043a\u043e\u0435 \u0432\u0440\u0435\u043c\u044f','\u041f\u0440\u0438\u0451\u043c\u0449\u0438\u043a','\u041a\u043e\u0440\u043e\u0431','\u0414\u0435\u0439\u0441\u0442\u0432\u0438\u0435','\u0428\u041a \u0442\u043e\u0432\u0430\u0440\u0430','\u0410\u0440\u0442\u0438\u043a\u0443\u043b','\u0420\u0430\u0437\u043c\u0435\u0440','\u041a\u043e\u043b\u0438\u0447\u0435\u0441\u0442\u0432\u043e / \u0437\u043d\u0430\u0447\u0435\u043d\u0438\u0435','ID \u0447\u0435\u0440\u043d\u043e\u0432\u0438\u043a\u0430'])
    event_params={'c':str(client_id),'w':str(warehouse_id),'o':actor,'n':number}; event_picker=''
    if picker:
        event_picker=' AND w.picker=:picker'; event_params['picker']=picker
    evrows=db.execute(text('''SELECT e.work_id,e.action,e.payload,e.created_at,w.picker,w.data
        FROM fc_tsd_work_events e JOIN fc_tsd_work w ON w.id=e.work_id
        WHERE w.client_id=:c AND w.warehouse_id=:w AND w.owner=:o AND w.mode='RECEIVE' AND w.status IN ('DRAFT','POSTED')
          AND COALESCE(w.data->>'receipt_number','')=:n'''+event_picker+''' AND e.action IN ('SCAN','SET')
        ORDER BY e.created_at,e.id'''),event_params).mappings().all()
    for e in evrows:
        payload=decoded(e['payload']); data=decoded(e['data']); lines=data.get('lines') or {}; sid=str(payload.get('sku_id') or ''); barcode=str(payload.get('barcode') or '')
        item=(lines.get(sid) or {}) if sid else next((x for x in lines.values() if str(x.get('barcode') or '')==barcode),{})
        place=data.get('place') or {}
        scans.append([str(e['created_at'] or ''),payload.get('effective_at') or payload.get('occurred_at') or str(e['created_at'] or ''),str(e['picker'] or ''),place.get('container') or place.get('location') or '',str(e['action'] or ''),barcode or item.get('barcode') or '',item.get('article') or '',item.get('size') or '',1 if str(e['action'])=='SCAN' else payload.get('quantity'),str(e['work_id'])])
    fill=PatternFill('solid',fgColor='D9EAD3')
    for sheet in (summary,staff,boxes,items,scans):
        sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
        for c in sheet[1]: c.font=Font(bold=True); c.fill=fill; c.alignment=Alignment(vertical='center',wrap_text=True)
        for rr in sheet.iter_rows(min_row=2):
            for c in rr: c.alignment=Alignment(vertical='top',wrap_text=True)
        for i in range(1,sheet.max_column+1):
            width=max((len(str(sheet.cell(rr,i).value or '')) for rr in range(1,min(sheet.max_row,300)+1)),default=10)+2
            sheet.column_dimensions[get_column_letter(i)].width=min(55,max(12,width))
    for sheet,col in ((boxes,6),(items,6),(scans,6)):
        for rr in range(2,sheet.max_row+1): sheet.cell(rr,col).number_format='@'
    out=io.BytesIO(); wb.save(out); out.seek(0)
    return StreamingResponse(out,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="tsd_receipt_current.xlsx"'})


@router.get('/history/list')
def history_list(client_id:uuid.UUID,warehouse_id:uuid.UUID,limit:int=200,db:Session=Depends(get_db)):
    scope(db,client_id,warehouse_id)
    return {'rows':service.history_rows(db,client_id,warehouse_id,limit)}


@router.get('/history/export.xlsx')
def history_export(client_id:uuid.UUID,warehouse_id:uuid.UUID,db:Session=Depends(get_db)):
    scope(db,client_id,warehouse_id)
    rows=service.history_rows(db,client_id,warehouse_id,1000)
    events=service.history_events(db,client_id,warehouse_id,100000)
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
    def cell(value):
        if value is None:return ''
        if isinstance(value,(dict,list,tuple,set)):
            return json.dumps(value,ensure_ascii=False,sort_keys=True,default=str)
        return value if isinstance(value,(str,int,float,bool)) else str(value)
    wb=Workbook();ws=wb.active;ws.title='История ТСД'
    ws.append(['Фактически начато','Фактически принято','Выгружено на сервер','Номер приёмки','Тип','Статус','Сотрудник','Учётная запись WMS','Место / короб','Адрес','Количество','Основание','ID'])
    labels={'RECEIVE':'Приёмка товара','COUNT':'Инвентаризация','MOVE':'Перемещение короб → короб','PICK':'Сборка'}
    for r in rows:
        ws.append([cell(str(r.get('first_item_at') or '')),cell(str(r.get('last_item_at') or '')),cell(str(r.get('updated_at') or '')),cell(r.get('receipt_number')),cell(labels.get(r.get('kind'),r.get('kind'))),cell(r.get('status')),cell(r.get('picker')),cell(r.get('owner')),cell(r.get('place')),cell(r.get('address')),cell(r.get('units')),cell(r.get('reference')),cell(r.get('id'))])
    ev=wb.create_sheet('Действия ТСД');ev.append(['Дата','Тип','Действие','Сотрудник','Учётная запись WMS','Сессия / задание','Детали'])
    action_labels={'SCAN':'Скан товара','SET':'Изменено количество','RESET':'Сброшено','REMOVE':'Удалена строка','UNDO':'Отменён последний скан','CHANGE_PLACE':'Пересканирован короб / место','CANCEL':'Черновик удалён','POST':'Завершено','TSD_SCAN':'Скан при сборке','TSD_MANUAL':'Подтверждено вручную','TSD_SKIP':'Позиция пропущена','TSD_NOT_IN_BOX':'Нет в коробе','SUBMIT_ADMIN':'Сборка завершена'}
    for r in events:
        ev.append([cell(str(r.get('created_at') or '')),cell(labels.get(r.get('kind'),r.get('kind'))),cell(action_labels.get(r.get('action'),r.get('action'))),cell(r.get('picker')),cell(r.get('owner')),cell(r.get('id')),cell(r.get('details'))])
    staff=wb.create_sheet('\u0418\u0442\u043e\u0433\u0438 \u043f\u043e \u043f\u0440\u0438\u0451\u043c\u0449\u0438\u043a\u0430\u043c')
    staff.append(['\u041d\u043e\u043c\u0435\u0440 \u043f\u0440\u0438\u0451\u043c\u043a\u0438','\u041f\u0440\u0438\u0451\u043c\u0449\u0438\u043a','\u0421\u0442\u0430\u0442\u0443\u0441','\u041a\u043e\u0440\u043e\u0431\u043e\u0432','\u0415\u0434\u0438\u043d\u0438\u0446','\u041f\u0435\u0440\u0432\u044b\u0439 \u0441\u043a\u0430\u043d','\u041f\u043e\u0441\u043b\u0435\u0434\u043d\u0438\u0439 \u0441\u043a\u0430\u043d','\u0412\u044b\u0433\u0440\u0443\u0436\u0435\u043d\u043e \u043d\u0430 \u0441\u0435\u0440\u0432\u0435\u0440'])
    grouped={}
    for r in rows:
        if str(r.get('kind') or '')!='RECEIVE': continue
        key=(str(r.get('receipt_number') or ''),str(r.get('picker') or ''),str(r.get('status') or ''))
        g=grouped.setdefault(key,{'boxes':0,'units':0,'first':'','last':'','posted':''})
        g['boxes']+=1; g['units']+=float(r.get('units') or 0)
        first=str(r.get('first_item_at') or ''); last=str(r.get('last_item_at') or ''); posted=str(r.get('updated_at') or r.get('created_at') or '')
        if first and (not g['first'] or first<g['first']):g['first']=first
        if last and (not g['last'] or last>g['last']):g['last']=last
        if posted and posted>g['posted']:g['posted']=posted
    for (number,receiver,status),g in sorted(grouped.items()):
        staff.append([number,receiver,status,g['boxes'],g['units'],g['first'],g['last'],g['posted']])
    fill=PatternFill('solid',fgColor='D9EAD3')
    for sheet in (ws,ev,staff):
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]: cell.font=Font(bold=True);cell.fill=fill;cell.alignment=Alignment(vertical='center',wrap_text=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row: cell.alignment=Alignment(vertical='top',wrap_text=True)
        for i in range(1,sheet.max_column+1):
            vals=[str(sheet.cell(r,i).value or '') for r in range(1,min(sheet.max_row,250)+1)]
            sheet.column_dimensions[get_column_letter(i)].width=min(55,max(12,max((len(v) for v in vals),default=10)+2))
    detail=wb.create_sheet('Товары по приёмкам')
    detail.append(['Номер приёмки','Первый скан','Последний скан','Выгружено на сервер','Сотрудник','Учётная запись WMS','Короб / место','Ячейка','ШК','Артикул','Размер','Товар','Количество','ID черновика'])
    for r in rows:
        if str(r.get('kind') or '')!='RECEIVE': continue
        for x in (r.get('items') or []):
            detail.append([cell(r.get('receipt_number')),cell(str(x.get('first_scanned_at') or r.get('first_item_at') or '')),cell(str(x.get('last_scanned_at') or r.get('last_item_at') or '')),cell(str(r.get('updated_at') or r.get('created_at') or '')),cell(r.get('picker')),cell(r.get('owner')),cell(r.get('place')),cell(r.get('address')),cell(x.get('barcode')),cell(x.get('article')),cell(x.get('size')),cell(x.get('name')),cell(x.get('quantity')),cell(r.get('id'))])
    scans=wb.create_sheet('Фактические сканы')
    scans.append(['Фактическое время','Получено сервером','Номер приёмки','Короб / место','ШК','Артикул','Размер','Действие','Кол-во / значение','Сотрудник','Учётная запись WMS','ID черновика'])
    for e in reversed(events):
        if str(e.get('kind') or '')!='RECEIVE' or str(e.get('action') or '') not in ('SCAN','SET'): continue
        scans.append([cell(str(e.get('occurred_at') or e.get('created_at') or '')),cell(str(e.get('created_at') or '')),cell(e.get('receipt_number')),cell(e.get('place')),cell(e.get('barcode')),cell(e.get('article')),cell(e.get('size')),cell(action_labels.get(e.get('action'),e.get('action'))),cell(e.get('quantity')),cell(e.get('picker')),cell(e.get('owner')),cell(e.get('id'))])
        scans.cell(scans.max_row,5).number_format='@'
    scans.freeze_panes='A2';scans.auto_filter.ref=scans.dimensions
    for c in scans[1]: c.font=Font(bold=True);c.fill=fill;c.alignment=Alignment(vertical='center',wrap_text=True)
    for i in range(1,scans.max_column+1):
        scans.column_dimensions[get_column_letter(i)].width=min(45,max(12,max((len(str(scans.cell(rr,i).value or '')) for rr in range(1,min(scans.max_row,300)+1)),default=10)+2))
    detail.freeze_panes='A2';detail.auto_filter.ref=detail.dimensions
    for c in detail[1]: c.font=Font(bold=True);c.fill=fill;c.alignment=Alignment(vertical='center',wrap_text=True)
    for i in range(1,detail.max_column+1):
        detail.column_dimensions[get_column_letter(i)].width=min(45,max(12,max((len(str(detail.cell(rr,i).value or '')) for rr in range(1,min(detail.max_row,300)+1)),default=10)+2))
    out=io.BytesIO();wb.save(out);out.seek(0)
    return StreamingResponse(out,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="tsd_history.xlsx"'})


@router.post('/receipt-batch/post')
def receipt_batch_post(payload:ReceiptBatch, request:Request, db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id)
    actor=owner(request); number=payload.receipt_number; receiver=payload.picker
    try:
        # A receipt number may be shared by several receivers, but each receiver owns
        # a separate result. Finalization never closes another receiver's boxes.
        sql="""SELECT id,version,data,picker FROM fc_tsd_work
            WHERE client_id=:c AND warehouse_id=:w AND owner=:o AND mode='RECEIVE' AND status='DRAFT'
              AND COALESCE(data->>'receipt_number','')=:n AND picker=:picker
            ORDER BY created_at,id FOR UPDATE"""
        rows=db.execute(text(sql),{'c':str(payload.client_id),'w':str(payload.warehouse_id),'o':actor,'n':number,'picker':receiver}).mappings().all()
        if not rows:
            raise HTTPException(409,detail='\u0423 \u043f\u0440\u0438\u0451\u043c\u0449\u0438\u043a\u0430 '+receiver+' \u043d\u0435\u0442 \u0430\u043a\u0442\u0438\u0432\u043d\u044b\u0445 \u043a\u043e\u0440\u043e\u0431\u043e\u0432 \u043f\u043e \u044d\u0442\u043e\u0439 \u043f\u0440\u0438\u0451\u043c\u043a\u0435')
        ready=[]; empty=[]
        for r in rows:
            data=r['data'] if isinstance(r['data'],dict) else json.loads(r['data'] or '{}')
            lines=(data.get('lines') or {}).values()
            row_units=sum(max(0,int(x.get('quantity') or 0)) for x in lines)
            (ready if row_units>0 else empty).append(r)
        if not ready:
            raise HTTPException(409,detail='\u0423 \u0432\u0430\u0441 \u0432 \u044d\u0442\u043e\u0439 \u043f\u0440\u0438\u0451\u043c\u043a\u0435 \u043f\u043e\u043a\u0430 \u043d\u0435\u0442 \u043e\u0442\u0441\u043a\u0430\u043d\u0438\u0440\u043e\u0432\u0430\u043d\u043d\u043e\u0433\u043e \u0442\u043e\u0432\u0430\u0440\u0430.')
        posted=[]; units=0
        for r in ready:
            result=service.edit(db,r['id'],payload.client_id,payload.warehouse_id,actor,{
                'event_id':str(uuid.uuid4()),'version':int(r['version']),'action':'POST','barcode':None,
                'sku_id':None,'quantity':None,'product':None,'place_code':None,'target_type':'AUTO',
                'section_code':'','placement_code':'','place_empty_box':False,'row':None})
            posted.append(result); units+=int(result.get('units') or 0)
        for r in empty:
            db.execute(text("""UPDATE fc_tsd_work SET status='CANCELLED',version=version+1,updated_at=now(),
                data=jsonb_set(COALESCE(data,'{}'::jsonb),'{batch_empty_cancelled}','true'::jsonb,true)
                WHERE id=:id AND status='DRAFT'"""),{'id':str(r['id'])})
        db.commit()
        return {'ok':True,'receipt_number':number,'picker':receiver,'boxes':len(posted),'units':units,'rows':posted,'ignored_empty':len(empty),'pickers':[receiver]}
    except Exception:
        db.rollback();raise


@router.get('/{work_id}')
def detail(work_id:uuid.UUID,client_id:uuid.UUID,warehouse_id:uuid.UUID,request:Request,db:Session=Depends(get_db)):
    scope(db,client_id,warehouse_id)
    return service.public(service.get(db,work_id,client_id,warehouse_id,owner(request)))


@router.post('/{work_id}/event')
def event(work_id:uuid.UUID,payload:Edit,request:Request,db:Session=Depends(get_db)):
    scope(db,payload.client_id,payload.warehouse_id)
    try:
        req=payload.model_dump(mode='json');req.pop('client_id');req.pop('warehouse_id')
        result=service.edit(db,work_id,payload.client_id,payload.warehouse_id,owner(request),req)
        db.commit();return result
    except Exception:db.rollback();raise
