from __future__ import annotations

import io
import re
import uuid
from datetime import datetime, timezone
from decimal import Decimal
from typing import Literal

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel, Field
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.services.printing import queue_print_job
from app.services.excel_safety import workbook_bytes
from app.vendor.pypdf import PdfReader, PdfWriter

router = APIRouter(prefix="/fbo", tags=["FBO shipment foundation"])

# FC_R145F8_FBO_FOUNDATION_NO_STOCK_EFFECTS
# This module deliberately owns its own fc_fbo_* tables and never writes
# inventory_operations / inventory balances. It is a packing / label workspace only.

PDF_LIMIT = 60 * 1024 * 1024
KINDS = {"BARCODE", "CHZ"}


class BatchCreate(BaseModel):
    client_id: uuid.UUID
    warehouse_id: uuid.UUID
    name: str = Field(min_length=1, max_length=200)
    source: str = Field(default="Собрано со склада", max_length=120)
    comment: str = Field(default="", max_length=1000)
    performed_by: str = Field(default="Администратор", max_length=160)


class BoxCreate(BaseModel):
    name: str = Field(min_length=1, max_length=180)
    performed_by: str = Field(default="Администратор", max_length=160)


class ItemAdd(BaseModel):
    barcode: str = Field(min_length=1, max_length=160)
    box_name: str = Field(min_length=1, max_length=180)
    quantity: Decimal = Field(default=Decimal("1"), gt=0, le=100000)
    performed_by: str = Field(default="Администратор", max_length=160)


class PrintGroup(BaseModel):
    page_no: int = Field(ge=1)
    copies: int = Field(ge=0, le=500)


class PrintRequest(BaseModel):
    print_station: str = Field(min_length=1, max_length=80)
    groups: list[PrintGroup] = Field(default_factory=list)


class PrintAllRequest(BaseModel):
    print_station: str = Field(min_length=1, max_length=80)


def _schema(db: Session) -> None:
    db.execute(text("""
    CREATE TABLE IF NOT EXISTS fc_fbo_batches (
      id uuid PRIMARY KEY,
      client_id uuid NOT NULL,
      warehouse_id uuid NOT NULL,
      name text NOT NULL,
      source text NOT NULL DEFAULT '',
      comment text NOT NULL DEFAULT '',
      status text NOT NULL DEFAULT 'ACTIVE',
      performed_by text NOT NULL DEFAULT '',
      created_at timestamptz NOT NULL DEFAULT now(),
      updated_at timestamptz NOT NULL DEFAULT now()
    )"""))
    db.execute(text("""
    CREATE TABLE IF NOT EXISTS fc_fbo_boxes (
      id uuid PRIMARY KEY,
      batch_id uuid NOT NULL REFERENCES fc_fbo_batches(id) ON DELETE CASCADE,
      name text NOT NULL,
      performed_by text NOT NULL DEFAULT '',
      created_at timestamptz NOT NULL DEFAULT now(),
      UNIQUE(batch_id, name)
    )"""))
    db.execute(text("""
    CREATE TABLE IF NOT EXISTS fc_fbo_items (
      id uuid PRIMARY KEY,
      batch_id uuid NOT NULL REFERENCES fc_fbo_batches(id) ON DELETE CASCADE,
      box_id uuid NOT NULL REFERENCES fc_fbo_boxes(id) ON DELETE CASCADE,
      sku_id uuid,
      barcode text NOT NULL,
      article text NOT NULL DEFAULT '',
      size text NOT NULL DEFAULT '',
      product_name text NOT NULL DEFAULT '',
      quantity numeric(14,3) NOT NULL,
      performed_by text NOT NULL DEFAULT '',
      created_at timestamptz NOT NULL DEFAULT now()
    )"""))
    db.execute(text("""
    CREATE TABLE IF NOT EXISTS fc_fbo_label_batches (
      id uuid PRIMARY KEY,
      batch_id uuid NOT NULL REFERENCES fc_fbo_batches(id) ON DELETE CASCADE,
      kind text NOT NULL,
      file_name text NOT NULL DEFAULT '',
      pdf_bytes bytea NOT NULL,
      page_count integer NOT NULL,
      performed_by text NOT NULL DEFAULT '',
      created_at timestamptz NOT NULL DEFAULT now()
    )"""))
    db.execute(text("CREATE INDEX IF NOT EXISTS ix_fc_fbo_batches_scope ON fc_fbo_batches(client_id,warehouse_id,status,created_at DESC)"))
    db.execute(text("CREATE INDEX IF NOT EXISTS ix_fc_fbo_items_batch ON fc_fbo_items(batch_id,box_id,created_at)"))
    db.execute(text("CREATE INDEX IF NOT EXISTS ix_fc_fbo_labels_batch ON fc_fbo_label_batches(batch_id,kind,created_at DESC)"))
    db.commit()


def _scope_batch(db: Session, batch_id: uuid.UUID) -> dict:
    _schema(db)
    row = db.execute(text("SELECT * FROM fc_fbo_batches WHERE id=:id"), {"id": str(batch_id)}).mappings().first()
    if not row:
        raise HTTPException(404, detail="FBO-партия не найдена")
    return dict(row)


def _batch_payload(db: Session, row: dict) -> dict:
    bid = str(row["id"])
    counts = db.execute(text("""
      SELECT COUNT(DISTINCT b.id)::int boxes,
             COUNT(i.id)::int lines,
             COALESCE(SUM(i.quantity),0) units
      FROM fc_fbo_batches x
      LEFT JOIN fc_fbo_boxes b ON b.batch_id=x.id
      LEFT JOIN fc_fbo_items i ON i.batch_id=x.id
      WHERE x.id=:id GROUP BY x.id
    """), {"id": bid}).mappings().first() or {}
    latest = {r["kind"]: str(r["id"]) for r in db.execute(text("""
      SELECT DISTINCT ON (kind) id,kind FROM fc_fbo_label_batches
      WHERE batch_id=:id ORDER BY kind,created_at DESC
    """), {"id": bid}).mappings().all()}
    out = dict(row)
    out["id"] = bid
    out["client_id"] = str(out["client_id"])
    out["warehouse_id"] = str(out["warehouse_id"])
    out["boxes"] = int(counts.get("boxes") or 0)
    out["lines"] = int(counts.get("lines") or 0)
    out["units"] = float(counts.get("units") or 0)
    out["barcode_label_batch_id"] = latest.get("BARCODE")
    out["chz_label_batch_id"] = latest.get("CHZ")
    return out


def _safe_xlsx(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        if v.tzinfo is not None:
            v = v.astimezone(timezone.utc).replace(tzinfo=None)
        return v
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, uuid.UUID):
        return str(v)
    if isinstance(v, str):
        # openpyxl rejects C0 control characters.
        return re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]", "", v)
    return v


def _page_bytes(pdf: bytes, page_no: int) -> bytes:
    reader = PdfReader(io.BytesIO(pdf))
    if page_no < 1 or page_no > len(reader.pages):
        raise HTTPException(404, detail="Страница PDF не найдена")
    writer = PdfWriter()
    writer.add_page(reader.pages[page_no - 1])
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def _label_groups(pdf: bytes) -> list[dict]:
    reader = PdfReader(io.BytesIO(pdf))
    groups: dict[str, dict] = {}
    for idx, page in enumerate(reader.pages, start=1):
        try:
            raw = " ".join((page.extract_text() or "").split())
        except Exception:
            raw = ""
        codes = re.findall(r"(?<!\d)\d{8,24}(?!\d)", raw)
        code = codes[0] if codes else ""
        key = code or raw[:120] or f"PAGE:{idx}"
        item = groups.setdefault(key, {
            "key": key,
            "code": code,
            "title": code or (raw[:90] if raw else f"Страница {idx}"),
            "page_no": idx,
            "page_nos": [],
            "found": 0,
        })
        item["page_nos"].append(idx)
        item["found"] += 1
    return list(groups.values())


def _label_row(db: Session, label_batch_id: uuid.UUID) -> dict:
    _schema(db)
    row = db.execute(text("""SELECT l.*,b.client_id,b.warehouse_id FROM fc_fbo_label_batches l
      JOIN fc_fbo_batches b ON b.id=l.batch_id WHERE l.id=:id"""), {"id": str(label_batch_id)}).mappings().first()
    if not row:
        raise HTTPException(404, detail="Загрузка PDF не найдена")
    return dict(row)


@router.post("/batches")
def create_batch(payload: BatchCreate, db: Session = Depends(get_db)):
    _schema(db)
    # Validate the physical scope without changing stock.
    if not db.execute(text("SELECT 1 FROM clients WHERE id=:id"), {"id": str(payload.client_id)}).first():
        raise HTTPException(404, detail="Клиент не найден")
    if not db.execute(text("SELECT 1 FROM warehouses WHERE id=:id"), {"id": str(payload.warehouse_id)}).first():
        raise HTTPException(404, detail="Склад не найден")
    bid = uuid.uuid4()
    db.execute(text("""
      INSERT INTO fc_fbo_batches(id,client_id,warehouse_id,name,source,comment,performed_by)
      VALUES(:id,:c,:w,:n,:s,:m,:p)
    """), {"id": str(bid), "c": str(payload.client_id), "w": str(payload.warehouse_id),
           "n": payload.name.strip(), "s": payload.source.strip(), "m": payload.comment.strip(), "p": payload.performed_by.strip()})
    db.commit()
    return _batch_payload(db, _scope_batch(db, bid))


@router.get("/batches")
def list_batches(client_id: uuid.UUID, warehouse_id: uuid.UUID, status: str = "ACTIVE", db: Session = Depends(get_db)):
    _schema(db)
    rows = db.execute(text("""
      SELECT * FROM fc_fbo_batches
      WHERE client_id=:c AND warehouse_id=:w AND (:status='' OR status=:status)
      ORDER BY created_at DESC
    """), {"c": str(client_id), "w": str(warehouse_id), "status": str(status or "").upper()}).mappings().all()
    return {"rows": [_batch_payload(db, dict(r)) for r in rows]}


@router.get("/batches/{batch_id}")
def batch_detail(batch_id: uuid.UUID, db: Session = Depends(get_db)):
    row = _scope_batch(db, batch_id)
    boxes = [dict(r) for r in db.execute(text("""
      SELECT b.id,b.name,b.performed_by,b.created_at,COUNT(i.id)::int lines,COALESCE(SUM(i.quantity),0) units
      FROM fc_fbo_boxes b LEFT JOIN fc_fbo_items i ON i.box_id=b.id
      WHERE b.batch_id=:id GROUP BY b.id ORDER BY b.created_at,b.name
    """), {"id": str(batch_id)}).mappings().all()]
    items = [dict(r) for r in db.execute(text("""
      SELECT i.id,i.barcode,i.article,i.size,i.product_name,i.quantity,i.performed_by,i.created_at,
             b.name box_name
      FROM fc_fbo_items i JOIN fc_fbo_boxes b ON b.id=i.box_id
      WHERE i.batch_id=:id ORDER BY b.created_at,b.name,i.created_at,i.id
    """), {"id": str(batch_id)}).mappings().all()]
    for x in boxes:
        x["id"] = str(x["id"]); x["units"] = float(x.get("units") or 0)
    for x in items:
        x["id"] = str(x["id"]); x["quantity"] = float(x.get("quantity") or 0)
    return {"batch": _batch_payload(db, row), "boxes": boxes, "items": items}


@router.post("/batches/{batch_id}/boxes")
def create_box(batch_id: uuid.UUID, payload: BoxCreate, db: Session = Depends(get_db)):
    _scope_batch(db, batch_id)
    name = payload.name.strip()
    row = db.execute(text("SELECT * FROM fc_fbo_boxes WHERE batch_id=:b AND lower(name)=lower(:n)"), {"b": str(batch_id), "n": name}).mappings().first()
    if row:
        return {"id": str(row["id"]), "name": row["name"], "created": False}
    box_id = uuid.uuid4()
    db.execute(text("INSERT INTO fc_fbo_boxes(id,batch_id,name,performed_by) VALUES(:id,:b,:n,:p)"),
               {"id": str(box_id), "b": str(batch_id), "n": name, "p": payload.performed_by.strip()})
    db.commit()
    return {"id": str(box_id), "name": name, "created": True}


@router.post("/batches/{batch_id}/items")
def add_item(batch_id: uuid.UUID, payload: ItemAdd, db: Session = Depends(get_db)):
    batch = _scope_batch(db, batch_id)
    box = db.execute(text("SELECT id,name FROM fc_fbo_boxes WHERE batch_id=:b AND lower(name)=lower(:n)"),
                     {"b": str(batch_id), "n": payload.box_name.strip()}).mappings().first()
    if not box:
        raise HTTPException(422, detail="Сначала создайте или выберите короб FBO")
    bc = payload.barcode.strip()
    sku = db.execute(text("""
      SELECT id,barcode,article,size,name FROM skus
      WHERE client_id=:c AND barcode=:bc AND is_active=TRUE
      ORDER BY created_at DESC LIMIT 1
    """), {"c": str(batch["client_id"]), "bc": bc}).mappings().first()
    if not sku:
        raise HTTPException(404, detail="ШК товара не найден в Базе товаров")
    item_id = uuid.uuid4()
    db.execute(text("""
      INSERT INTO fc_fbo_items(id,batch_id,box_id,sku_id,barcode,article,size,product_name,quantity,performed_by)
      VALUES(:id,:b,:box,:sku,:bc,:a,:z,:n,:q,:p)
    """), {"id": str(item_id), "b": str(batch_id), "box": str(box["id"]), "sku": str(sku["id"]),
           "bc": sku["barcode"] or bc, "a": sku["article"] or "", "z": sku["size"] or "", "n": sku["name"] or "",
           "q": str(payload.quantity), "p": payload.performed_by.strip()})
    db.commit()
    return {"ok": True, "id": str(item_id), "box": box["name"], "article": sku["article"], "size": sku["size"], "barcode": sku["barcode"], "quantity": float(payload.quantity)}


@router.delete("/batches/{batch_id}/items/{item_id}")
def delete_item(batch_id: uuid.UUID, item_id: uuid.UUID, db: Session = Depends(get_db)):
    _scope_batch(db, batch_id)
    result = db.execute(text("DELETE FROM fc_fbo_items WHERE id=:id AND batch_id=:b"), {"id": str(item_id), "b": str(batch_id)})
    if result.rowcount == 0:
        raise HTTPException(404, detail="Позиция FBO не найдена")
    db.commit()
    return {"ok": True}


@router.post("/batches/{batch_id}/labels/{kind}")
async def upload_labels(batch_id: uuid.UUID, kind: Literal["BARCODE", "CHZ"], file: UploadFile = File(...), performed_by: str = Form(default="Администратор"), db: Session = Depends(get_db)):
    _scope_batch(db, batch_id)
    data = await file.read(PDF_LIMIT + 1)
    if len(data) > PDF_LIMIT:
        raise HTTPException(413, detail="PDF больше 60 МБ")
    if not data.startswith(b"%PDF"):
        raise HTTPException(422, detail="Нужен PDF-файл")
    try:
        reader = PdfReader(io.BytesIO(data))
        count = len(reader.pages)
    except Exception as exc:
        raise HTTPException(422, detail=f"Не удалось прочитать PDF: {exc}")
    if not count:
        raise HTTPException(422, detail="В PDF нет страниц")
    lid = uuid.uuid4()
    db.execute(text("""
      INSERT INTO fc_fbo_label_batches(id,batch_id,kind,file_name,pdf_bytes,page_count,performed_by)
      VALUES(:id,:b,:k,:n,:pdf,:pc,:p)
    """), {"id": str(lid), "b": str(batch_id), "k": kind, "n": file.filename or f"{kind}.pdf", "pdf": data, "pc": count, "p": performed_by.strip()})
    db.commit()
    return {"id": str(lid), "kind": kind, "file_name": file.filename, "page_count": count, "groups": _label_groups(data)}


@router.get("/label-batches/{label_batch_id}")
def label_info(label_batch_id: uuid.UUID, db: Session = Depends(get_db)):
    row = _label_row(db, label_batch_id)
    return {"id": str(row["id"]), "batch_id": str(row["batch_id"]), "kind": row["kind"], "file_name": row["file_name"], "page_count": row["page_count"], "groups": _label_groups(bytes(row["pdf_bytes"]))}


@router.get("/label-batches/{label_batch_id}/pages/{page_no}.pdf")
def label_page(label_batch_id: uuid.UUID, page_no: int, db: Session = Depends(get_db)):
    row = _label_row(db, label_batch_id)
    data = _page_bytes(bytes(row["pdf_bytes"]), page_no)
    return Response(data, media_type="application/pdf", headers={"Content-Disposition": f'inline; filename="FBO_{row["kind"]}_{page_no}.pdf"'})


@router.post("/label-batches/{label_batch_id}/print")
def print_selected(label_batch_id: uuid.UUID, payload: PrintRequest, db: Session = Depends(get_db)):
    row = _label_row(db, label_batch_id)
    station = payload.print_station.strip().upper()
    queued = 0
    for group in payload.groups:
        if group.copies <= 0:
            continue
        page = _page_bytes(bytes(row["pdf_bytes"]), group.page_no)
        for copy_no in range(group.copies):
            queue_print_job(db, station_code=station, payload=page,
                            file_name=f"FBO-{row['kind']}-{label_batch_id}-{group.page_no}-{copy_no+1}.pdf",
                            document_type=f"FBO_{row['kind']}_PDF",
                            client_id=row["client_id"], object_type="FBO_LABEL_BATCH", object_id=str(label_batch_id),
                            content_type="application/pdf")
            queued += 1
    db.commit()
    return {"ok": True, "queued": queued, "print_station": station}


@router.post("/label-batches/{label_batch_id}/print-all")
def print_all(label_batch_id: uuid.UUID, payload: PrintAllRequest, db: Session = Depends(get_db)):
    row = _label_row(db, label_batch_id)
    station = payload.print_station.strip().upper()
    queued = 0
    for page_no in range(1, int(row["page_count"]) + 1):
        page = _page_bytes(bytes(row["pdf_bytes"]), page_no)
        queue_print_job(db, station_code=station, payload=page,
                        file_name=f"FBO-{row['kind']}-{label_batch_id}-{page_no}.pdf",
                        document_type=f"FBO_{row['kind']}_PDF",
                        client_id=row["client_id"], object_type="FBO_LABEL_BATCH", object_id=str(label_batch_id),
                        content_type="application/pdf")
        queued += 1
    db.commit()
    return {"ok": True, "queued": queued, "print_station": station}


@router.get("/batches/{batch_id}/composition.xlsx")
def composition_xlsx(batch_id: uuid.UUID, db: Session = Depends(get_db)):
    batch = _scope_batch(db, batch_id)
    rows = db.execute(text("""
      SELECT b.name box_name,i.article,i.size,i.barcode,i.product_name,i.quantity,i.performed_by,i.created_at
      FROM fc_fbo_items i JOIN fc_fbo_boxes b ON b.id=i.box_id
      WHERE i.batch_id=:id ORDER BY b.created_at,b.name,i.created_at
    """), {"id": str(batch_id)}).mappings().all()
    wb = Workbook(); ws = wb.active; ws.title = "Раскладка по коробам"
    ws.append(["Партия FBO","Короб","Артикул","Размер","ШК","Товар","Количество","Сотрудник","Добавлено"])
    for c in ws[1]: c.font = Font(bold=True)
    for r in rows:
        ws.append([_safe_xlsx(batch["name"]),_safe_xlsx(r["box_name"]),_safe_xlsx(r["article"]),_safe_xlsx(r["size"]),_safe_xlsx(r["barcode"]),_safe_xlsx(r["product_name"]),_safe_xlsx(r["quantity"]),_safe_xlsx(r["performed_by"]),_safe_xlsx(r["created_at"])])
    summary = wb.create_sheet("Сводка")
    summary.append(["Короб","Артикул","Размер","ШК","Количество"])
    for c in summary[1]: c.font = Font(bold=True)
    sums = db.execute(text("""
      SELECT b.name box_name,i.article,i.size,i.barcode,SUM(i.quantity) quantity
      FROM fc_fbo_items i JOIN fc_fbo_boxes b ON b.id=i.box_id
      WHERE i.batch_id=:id GROUP BY b.name,i.article,i.size,i.barcode ORDER BY b.name,i.article,i.size
    """), {"id": str(batch_id)}).mappings().all()
    for r in sums: summary.append([_safe_xlsx(r["box_name"]),_safe_xlsx(r["article"]),_safe_xlsx(r["size"]),_safe_xlsx(r["barcode"]),_safe_xlsx(r["quantity"])])
    for sheet in (ws, summary):
        sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(45, max(12, max(len(str(c.value or "")) for c in col)+2))
    data = workbook_bytes(wb)
    stamp = datetime.now().strftime("%Y-%m-%d")
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="FBO_boxes_{stamp}.xlsx"'})
