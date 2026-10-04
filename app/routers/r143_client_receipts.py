# FC_R145F11_SELLER_RECEIPT_SELECT_AND_PICK_UPLOAD
from __future__ import annotations

import io
import json
import re
import uuid
from datetime import datetime, timezone

from fastapi import APIRouter, Depends, Request, Query, File, Form, UploadFile, HTTPException
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.routers.pvz import _portal_session
from app.services.excel_safety import workbook_bytes

router = APIRouter(prefix="/client", tags=["Кабинет селлера · приёмки"])

# FC_R145F8_SELLER_RECEIPT_GROUP_BY_NAME: one seller row per human receipt name/day; Excel remains detailed.


def _excel_dt(value):
    if hasattr(value, "tzinfo") and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    return value


def _excel_value(value):
    value = _excel_dt(value)
    if value is None:
        return ""
    if isinstance(value, str):
        return re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]", "", value)
    return value


def _kind(row) -> str:
    idem = str(row.get("idempotency_key") or "").upper()
    ref = str(row.get("reference") or "").upper()
    comment = str(row.get("comment") or "").upper()
    op_type = str(row.get("operation_type") or "").upper()
    if idem.startswith("PVZ_SESSION:") or "ПВЗ" in ref or "ПВЗ" in comment:
        return "ПВЗ"
    if idem.startswith("TSD-") or idem.startswith("TSD_") or "TSD" in idem:
        return "ТСД"
    if op_type == "RECEIVING_IMPORT" or "EXCEL" in comment or "EXCEL" in ref:
        return "Excel"
    return "Ручная"


def _strong_receipt_idempotency(idempotency_key: str) -> bool:
    upper = str(idempotency_key or "").strip().upper()
    return upper.startswith(
        (
            "PVZ_SESSION:",
            "BULK-RECEIVE-",
            "STRUCTURED-RECEIPT-",
            "PALLET-RECEIPT:",
            "PALLET-RECEIPT-",
        )
    )


def _receipt_group_key(row: dict) -> str:
    """Return a stable UI grouping key for one human receiving session.

    Strong server session/batch keys win. Older/TSD receipts that were persisted as
    several inventory operations are compacted by warehouse + receiving type +
    calendar date + the explicit receipt/reference name. We intentionally do not
    group unnamed UUID-only operations so unrelated receipts can never be merged.
    """
    warehouse_id = str(row.get("warehouse_id") or "")
    idem = str(row.get("idempotency_key") or "").strip()
    if idem and _strong_receipt_idempotency(idem):
        return f"{warehouse_id}|IDEM|{idem.upper()}"

    reference = str(row.get("reference") or "").strip()
    created = row.get("created_at")
    if hasattr(created, "date"):
        day = created.date().isoformat()
    else:
        day = str(created or "")[:10]
    if reference:
        return f"{warehouse_id}|REF|{_kind(row)}|{day}|{reference.casefold()}"

    return f"{warehouse_id}|OP|{row.get('id')}"


def _operation_rows(db: Session, client_id):
    return [
        dict(r)
        for r in db.execute(
            text(
                """
                SELECT o.id,o.created_at,o.warehouse_id,w.name warehouse_name,
                       o.operation_type,o.status,o.idempotency_key,o.performed_by,
                       o.reference,o.comment,
                       COUNT(l.id)::int line_count,
                       COALESCE(SUM(l.quantity),0) quantity
                FROM inventory_operations o
                JOIN warehouses w ON w.id=o.warehouse_id
                JOIN inventory_operation_lines l ON l.operation_id=o.id
                WHERE o.client_id=:client_id
                  AND o.operation_type IN ('RECEIVING','RECEIVING_IMPORT')
                  AND o.status <> 'CANCELLED'
                GROUP BY o.id,o.created_at,o.warehouse_id,w.name,o.operation_type,o.status,
                         o.idempotency_key,o.performed_by,o.reference,o.comment
                ORDER BY o.created_at DESC,o.id DESC
                """
            ),
            {"client_id": str(client_id)},
        ).mappings().all()
    ]


def _receipt_rows(db: Session, client_id):
    grouped: dict[str, dict] = {}
    for raw in _operation_rows(db, client_id):
        row = dict(raw)
        row["id"] = str(row["id"])
        row["warehouse_id"] = str(row["warehouse_id"])
        row["quantity"] = float(row.get("quantity") or 0)
        row["line_count"] = int(row.get("line_count") or 0)
        row["type"] = _kind(row)
        key = _receipt_group_key(row)

        current = grouped.get(key)
        if current is None:
            current = {
                **row,
                "receipt_key": key,
                "operation_ids": [row["id"]],
                "operation_count": 1,
                "performed_by_values": [str(row.get("performed_by") or "").strip()] if row.get("performed_by") else [],
            }
            grouped[key] = current
            continue

        current["operation_ids"].append(row["id"])
        current["operation_count"] = int(current.get("operation_count") or 1) + 1
        current["quantity"] = float(current.get("quantity") or 0) + row["quantity"]
        current["line_count"] = int(current.get("line_count") or 0) + row["line_count"]
        if row.get("created_at") and (not current.get("created_at") or row["created_at"] < current["created_at"]):
            current["created_at"] = row["created_at"]
        if not current.get("reference") and row.get("reference"):
            current["reference"] = row["reference"]
        if not current.get("comment") and row.get("comment"):
            current["comment"] = row["comment"]
        actor = str(row.get("performed_by") or "").strip()
        if actor and actor not in current["performed_by_values"]:
            current["performed_by_values"].append(actor)

    result = list(grouped.values())
    for row in result:
        actors = row.pop("performed_by_values", [])
        if actors:
            row["performed_by"] = ", ".join(actors[:3]) + (f" +{len(actors)-3}" if len(actors) > 3 else "")
        row["id"] = row["operation_ids"][0]
    result.sort(key=lambda r: (r.get("created_at") or datetime.min, str(r.get("id") or "")), reverse=True)
    return result


def _detail_rows(db: Session, client_id):
    return [
        dict(r)
        for r in db.execute(
            text(
                """
                SELECT o.id operation_id,o.created_at,o.warehouse_id,w.name warehouse_name,
                       o.operation_type,o.status,o.idempotency_key,o.reference,o.comment,o.performed_by,
                       s.barcode,s.article,s.size,s.name,l.quantity,
                       loc.code location_code,
                       COALESCE(NULLIF(c.name,''),c.barcode) container_name
                FROM inventory_operations o
                JOIN warehouses w ON w.id=o.warehouse_id
                JOIN inventory_operation_lines l ON l.operation_id=o.id
                JOIN skus s ON s.id=l.sku_id
                LEFT JOIN locations loc ON loc.id=l.destination_location_id
                LEFT JOIN containers c ON c.id=l.destination_container_id
                WHERE o.client_id=:client_id
                  AND o.operation_type IN ('RECEIVING','RECEIVING_IMPORT')
                  AND o.status <> 'CANCELLED'
                ORDER BY o.created_at DESC,o.id,l.id
                """
            ),
            {"client_id": str(client_id)},
        ).mappings().all()
    ]


@router.get("/api/receipts")
def portal_receipts(request: Request, limit: int | None = None, db: Session = Depends(get_db)):
    # ``limit`` is accepted only for backward compatibility with older cached UI.
    # The seller history must always show every receipt for the authenticated client.
    _ = limit
    user = _portal_session(request, db)
    rows = _receipt_rows(db, user["client_id"])
    return {
        "rows": rows,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "total": len(rows),
        "complete": True,
    }




def _seller_warehouse_allowed(db: Session, client_id, warehouse_id) -> bool:
    return bool(db.execute(text("""
        SELECT 1 WHERE EXISTS(
          SELECT 1 FROM inventory_balances WHERE CAST(client_id AS text)=:c AND CAST(warehouse_id AS text)=:w LIMIT 1
        ) OR EXISTS(
          SELECT 1 FROM inventory_operations WHERE CAST(client_id AS text)=:c AND CAST(warehouse_id AS text)=:w LIMIT 1
        )
    """),{"c":str(client_id),"w":str(warehouse_id)}).scalar())


@router.get("/api/picks/template.xlsx")
def portal_pick_template(request: Request, db: Session = Depends(get_db)):
    _portal_session(request, db)
    wb=Workbook();ws=wb.active;ws.title="Сборочный лист"
    ws.append(["ШК","Количество"]);ws["A1"].font=Font(bold=True);ws["B1"].font=Font(bold=True)
    ws.append(["2000000000000",1]);ws.column_dimensions["A"].width=24;ws.column_dimensions["B"].width=14
    ws.freeze_panes="A2"
    data=workbook_bytes(wb)
    return Response(data,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":'attachment; filename="FulfillCore_seller_pick_template.xlsx"'})


@router.post("/api/picks/upload")
def portal_pick_upload(
    request: Request,
    warehouse_id: uuid.UUID = Form(...),
    task_number: str = Form(...),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    user=_portal_session(request,db);client_id=uuid.UUID(str(user["client_id"]));username=str(user.get("username") or "seller")
    if not _seller_warehouse_allowed(db,client_id,warehouse_id):
        raise HTTPException(403,detail="Этот склад недоступен в кабинете селлера")
    if not str(file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(422,detail="Загрузите Excel .xlsx в формате ШК + количество")
    from app.routers.warehouse import import_pick_task
    result=import_pick_task(
        client_id=client_id,warehouse_id=warehouse_id,task_number=str(task_number or "").strip(),name=str(task_number or "").strip(),
        created_by=f"SELLER:{username}",source="SELLER",seller_username=username,pick_mode="CELL",seller_only=True,file=file,db=db,
    )
    return {**result,"message":"Задание передано складу. Выделенное количество сразу исключено из свободного MAIN-остатка резервом сборки; физическое списание произойдёт один раз после закрытия сборки."}


@router.get("/api/picks")
def portal_pick_list(request: Request, db: Session = Depends(get_db)):
    user=_portal_session(request,db);client_id=str(user["client_id"])
    from app.routers.warehouse import _ensure_schema
    _ensure_schema(db)
    rows=db.execute(text("""
      SELECT t.id,t.warehouse_id,w.name warehouse_name,t.task_number,t.name,t.status,t.pick_mode,t.picker,t.created_at,t.assigned_at,t.closed_at,
        COALESCE(SUM(l.requested_qty),0) requested,COALESCE(SUM(l.allocated_qty),0) allocated,
        COALESCE(SUM(l.scan_qty+l.manual_qty),0) picked,COALESCE(SUM(l.missing_qty),0) missing,
        COALESCE(SUM(CASE WHEN l.allocated_qty=0 THEN l.requested_qty ELSE 0 END),0) shortage,
        (SELECT COUNT(*) FROM fc_wh_pick_tasks ch WHERE ch.parent_task_id=t.id) child_count,
        (SELECT COUNT(*) FROM fc_wh_pick_tasks ch WHERE ch.parent_task_id=t.id AND ch.status='CLOSED') child_closed,
        (SELECT COUNT(*) FROM fc_wh_pick_tasks ch WHERE ch.parent_task_id=t.id AND ch.status IN ('ASSIGNED','IN_PROGRESS','READY_FOR_ADMIN')) child_active
      FROM fc_wh_pick_tasks t
      LEFT JOIN warehouses w ON CAST(w.id AS text)=t.warehouse_id
      LEFT JOIN fc_wh_pick_lines l ON l.task_id=t.id
      WHERE t.client_id=:c AND COALESCE(t.source,'ADMIN')='SELLER' AND t.parent_task_id IS NULL
      GROUP BY t.id,w.name ORDER BY t.created_at DESC LIMIT 200
    """),{"c":client_id}).mappings().all()
    out=[]
    for raw in rows:
        r=dict(raw);status=str(r.get("status") or "DRAFT").upper();children=int(r.get("child_count") or 0)
        if status=="SPLIT" and children:
            if int(r.get("child_closed") or 0)==children: display="CLOSED"
            elif int(r.get("child_active") or 0)>0: display="IN_PROGRESS"
            else: display="SPLIT"
        else: display=status
        r["display_status"]=display
        for k in ("requested","allocated","picked","missing","shortage"): r[k]=float(r.get(k) or 0)
        out.append(r)
    return {"rows":out,"total":len(out)}


@router.get("/api/receipts.xlsx")
def portal_receipts_xlsx(request: Request, limit: int | None = None, receipt_key: list[str] | None = Query(default=None), db: Session = Depends(get_db)):
    # Excel remains intentionally detailed: one summary row per receiving session and
    # the complete cell/container-level composition on the second sheet.
    _ = limit
    user = _portal_session(request, db)
    client_id = user["client_id"]
    receipts = _receipt_rows(db, client_id)
    requested_keys={str(x) for x in (receipt_key or []) if str(x).strip()}
    if requested_keys:
        receipts=[r for r in receipts if str(r.get("receipt_key") or "") in requested_keys]
    selected_operation_ids={str(op) for receipt in receipts for op in (receipt.get("operation_ids") or [])}
    detail = [r for r in _detail_rows(db, client_id) if str(r.get("operation_id")) in selected_operation_ids]

    group_by_operation_id = {}
    for receipt in receipts:
        for operation_id in receipt.get("operation_ids") or []:
            group_by_operation_id[str(operation_id)] = receipt

    wb = Workbook()
    summary = wb.active
    summary.title = "Приёмки"
    summary.append(["Дата/время", "Дата", "Склад", "Тип", "Название / основание", "Сотрудник", "Операций", "Строк", "Единиц", "Статус"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    for row in receipts:
        created = row.get("created_at")
        summary.append([
            _excel_value(created),
            created.date() if hasattr(created, "date") else created,
            _excel_value(row.get("warehouse_name") or ""),
            _excel_value(row.get("type") or ""),
            _excel_value(row.get("reference") or ""),
            _excel_value(row.get("performed_by") or ""),
            int(row.get("operation_count") or 1),
            int(row.get("line_count") or 0),
            float(row.get("quantity") or 0),
            _excel_value(row.get("status") or ""),
        ])

    lines = wb.create_sheet("Состав")
    lines.append(["Дата/время", "Дата", "Склад", "Тип", "Название / основание", "Сотрудник", "Артикул", "Размер", "ШК", "Товар", "Количество", "Адрес", "Короб"])
    for cell in lines[1]:
        cell.font = Font(bold=True)
    for row in detail:
        op_id = str(row.get("operation_id"))
        parent = group_by_operation_id.get(op_id, {})
        created = row.get("created_at")
        lines.append([
            _excel_value(created),
            created.date() if hasattr(created, "date") else created,
            _excel_value(row.get("warehouse_name") or ""),
            _excel_value(parent.get("type") or _kind(row)),
            _excel_value(parent.get("reference") or row.get("reference") or ""),
            _excel_value(row.get("performed_by") or parent.get("performed_by") or ""),
            _excel_value(row.get("article") or ""),
            _excel_value(row.get("size") or ""),
            _excel_value(row.get("barcode") or ""),
            _excel_value(row.get("name") or ""),
            float(row.get("quantity") or 0),
            _excel_value(row.get("location_code") or ""),
            _excel_value(row.get("container_name") or ""),
        ])

    for ws in (summary, lines):
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            letter = col[0].column_letter
            ws.column_dimensions[letter].width = min(45, max(12, max(len(str(c.value or "")) for c in col) + 2))

    data = workbook_bytes(wb)
    stamp = datetime.now().strftime("%Y-%m-%d")
    return Response(
        data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="FulfillCore_receipts_{stamp}.xlsx"'},
    )


@router.post("/api/receipts-selection.xlsx")
def portal_receipts_selection_xlsx(request: Request, receipt_keys: str = Form(...), db: Session = Depends(get_db)):
    _portal_session(request, db)
    try:
        keys=json.loads(receipt_keys)
    except Exception:
        raise HTTPException(422,detail="Некорректный список выбранных приёмок")
    if not isinstance(keys,list) or not keys:
        raise HTTPException(422,detail="Выберите хотя бы одну приёмку")
    keys=[str(x) for x in keys if str(x).strip()]
    if len(keys)>1000:
        raise HTTPException(422,detail="За один раз можно выгрузить до 1000 приёмок")
    return portal_receipts_xlsx(request=request,limit=None,receipt_key=keys,db=db)
