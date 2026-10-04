import io
import re
import uuid
from datetime import datetime, timezone
from decimal import Decimal

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import bindparam, func, or_, select, text
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.database import get_db
from app.models import Container, InventoryBalance, InventoryOperation, InventoryOperationLine, Location, MarketplaceProduct, SKU
from app.schemas.inventory import (
    AdjustCreate,
    CatalogAdjustCreate,
    BalanceRead,
    MoveContainerCreate,
    MoveCreate,
    OperationCreate,
    OperationLineRead,
    OperationRead,
    ReceiveCreate,
    ShipCreate,
    StockLine,
)
from app.services.inventory import container_move_preview, execute_operation, move_entire_container
from app.services.stock_reservations import active_reserved_totals
from app.services.stock_core_v2 import active_order_floor_for_identity, stock_truth_for_skus
from app.services.container_lifecycle import archive_consumed_empty_containers

router = APIRouter(tags=["Остатки и операции"])


def _archive_empty_source_boxes(db: Session, payload: OperationCreate) -> int:
    ids = [line.source_container_id for line in (payload.lines or []) if line.source_container_id]
    if not ids:
        return 0
    result = archive_consumed_empty_containers(
        db, payload.client_id, payload.warehouse_id, ids,
        actor=payload.performed_by, reason=f"OPERATION_{payload.operation_type}_EMPTY_BOX",
    )
    if result.get("count"):
        db.commit()
    return int(result.get("count", 0))


def _operation_read(op: InventoryOperation, lines: list[InventoryOperationLine]) -> OperationRead:
    return OperationRead(
        id=op.id,
        client_id=op.client_id,
        warehouse_id=op.warehouse_id,
        operation_type=op.operation_type,
        status=op.status,
        idempotency_key=op.idempotency_key,
        performed_by=op.performed_by,
        reference=op.reference,
        comment=op.comment,
        created_at=op.created_at,
        lines=[OperationLineRead.model_validate(line) for line in lines],
    )


def _serialize_operation(db: Session, op: InventoryOperation) -> OperationRead:
    # Scalar helper is kept for write endpoints that return one operation.
    lines = db.execute(
        select(InventoryOperationLine).where(InventoryOperationLine.operation_id == op.id)
    ).scalars().all()
    return _operation_read(op, lines)


# FC_BOX_SAFE_ARCHIVE_09415D_20260903
def _fc_box_json(container: Container) -> dict:
    return {
        "id": str(container.id),
        "client_id": str(container.client_id),
        "warehouse_id": str(container.warehouse_id),
        "location_id": str(container.location_id) if container.location_id else None,
        "barcode": str(container.barcode or ""),
        "name": str(getattr(container, "name", None) or container.barcode or ""),
        "container_type": str(container.container_type or "BOX"),
        "status": str(container.status or "ACTIVE"),
    }


def _fc_box_qty(db: Session, container_id: uuid.UUID, client_id: uuid.UUID, warehouse_id: uuid.UUID) -> Decimal:
    value = db.execute(select(func.coalesce(func.sum(InventoryBalance.quantity), 0)).where(
        InventoryBalance.client_id == client_id,
        InventoryBalance.warehouse_id == warehouse_id,
        InventoryBalance.container_id == container_id,
    )).scalar_one()
    return Decimal(value or 0)


def _fc_box_technical_barcode(db: Session, client_id: uuid.UUID, warehouse_id: uuid.UUID, code: str) -> str:
    clean = str(code or "").strip()
    occupied = db.execute(select(Container.id).where(Container.barcode == clean).limit(1)).scalar_one_or_none()
    if not occupied:
        return clean
    candidate = "FCBOX-" + uuid.uuid5(uuid.NAMESPACE_URL, f"fulfillcore-box:{client_id}:{warehouse_id}:{clean}").hex[:24].upper()
    if db.execute(select(Container.id).where(Container.barcode == candidate).limit(1)).scalar_one_or_none():
        candidate = "FCBOX-" + uuid.uuid4().hex[:24].upper()
    return candidate


@router.post("/inventory/containers/ensure")
def ensure_scoped_box(payload: dict, db: Session = Depends(get_db)):
    try:
        client_id = uuid.UUID(str(payload.get("client_id")))
        warehouse_id = uuid.UUID(str(payload.get("warehouse_id")))
        location_id = uuid.UUID(str(payload.get("location_id")))
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="Некорректный клиент, склад или ячейка")
    code = str(payload.get("name") or payload.get("barcode") or "").strip()
    if not code:
        raise HTTPException(status_code=400, detail="Укажите код короба")
    location = db.get(Location, location_id)
    if not location or location.warehouse_id != warehouse_id:
        raise HTTPException(status_code=409, detail="Ячейка относится к другому складу")
    name_column = getattr(Container, "name", Container.barcode)
    container = db.execute(select(Container).where(
        Container.client_id == client_id,
        Container.warehouse_id == warehouse_id,
        or_(Container.barcode == code, name_column == code),
    ).with_for_update()).scalars().first()
    if container:
        if str(container.status or "").upper() != "ACTIVE":
            container.status = "ACTIVE"
            container.location_id = location_id
        elif container.location_id is None:
            container.location_id = location_id
        db.commit()
        db.refresh(container)
        return _fc_box_json(container)
    technical = _fc_box_technical_barcode(db, client_id, warehouse_id, code)
    kwargs = dict(
        id=uuid.uuid4(), client_id=client_id, warehouse_id=warehouse_id, location_id=location_id,
        barcode=technical, container_type=str(payload.get("container_type") or "BOX"), status="ACTIVE",
    )
    if hasattr(Container, "name"):
        kwargs["name"] = code
    container = Container(**kwargs)
    db.add(container)
    db.commit()
    db.refresh(container)
    return _fc_box_json(container)


@router.post("/inventory/containers/{container_id}/archive")
def archive_empty_container(
    container_id: uuid.UUID,
    client_id: uuid.UUID = Query(...),
    warehouse_id: uuid.UUID = Query(...),
    db: Session = Depends(get_db),
):
    container = db.get(Container, container_id)
    if not container or container.client_id != client_id or container.warehouse_id != warehouse_id:
        raise HTTPException(status_code=404, detail="Короб не найден на выбранном складе")
    quantity = _fc_box_qty(db, container.id, client_id, warehouse_id)
    if quantity > 0:
        raise HTTPException(status_code=409, detail=f"В коробе находится {quantity} шт. Сначала переместите или спишите товар.")
    container.status = "ARCHIVED"
    db.commit()
    return {"ok": True, "id": str(container.id), "status": "ARCHIVED"}


@router.post("/inventory/containers/archive-empty-bulk")
def archive_empty_containers(
    payload: dict,
    client_id: uuid.UUID = Query(...),
    warehouse_id: uuid.UUID = Query(...),
    db: Session = Depends(get_db),
):
    raw_ids = payload.get("ids") or []
    ids = []
    for value in raw_ids:
        try:
            ids.append(uuid.UUID(str(value)))
        except (TypeError, ValueError):
            continue
    if not ids:
        return {"deleted": [], "skipped": []}
    rows = db.execute(select(Container).where(
        Container.id.in_(ids),
        Container.client_id == client_id,
        Container.warehouse_id == warehouse_id,
    )).scalars().all()
    by_id = {str(row.id): row for row in rows}
    deleted, skipped = [], []
    for container_id in ids:
        container = by_id.get(str(container_id))
        if not container:
            skipped.append({"id": str(container_id), "reason": "не найден на выбранном складе"})
            continue
        quantity = _fc_box_qty(db, container.id, client_id, warehouse_id)
        if quantity > 0:
            skipped.append({"id": str(container.id), "reason": f"в коробе {quantity} шт."})
            continue
        container.status = "ARCHIVED"
        deleted.append({"id": str(container.id)})
    db.commit()
    return {"deleted": deleted, "skipped": skipped}


@router.get("/inventory", response_model=list[BalanceRead])
def list_inventory(
    client_id: uuid.UUID,
    warehouse_id: uuid.UUID,
    sku_id: uuid.UUID | None = None,
    location_id: uuid.UUID | None = None,
    container_id: uuid.UUID | None = None,
    only_positive: bool = True,
    db: Session = Depends(get_db),
):
    stmt = select(InventoryBalance).where(
        InventoryBalance.client_id == client_id,
        InventoryBalance.warehouse_id == warehouse_id,
    )
    if sku_id:
        stmt = stmt.where(InventoryBalance.sku_id == sku_id)
    if location_id:
        stmt = stmt.where(InventoryBalance.location_id == location_id)
    if container_id:
        stmt = stmt.where(InventoryBalance.container_id == container_id)
    if only_positive:
        stmt = stmt.where(InventoryBalance.quantity > 0)
    return db.execute(stmt.order_by(InventoryBalance.updated_at.desc())).scalars().all()


# FC_K9R23_INVENTORY_QUICK_FIRST_PAINT
@router.get("/inventory/quick")
def inventory_quick(
    client_id: uuid.UUID,
    warehouse_id: uuid.UUID,
    db: Session = Depends(get_db),
):
    """Return positive local stock in one joined query for immediate Inventory first paint.

    This endpoint intentionally does not calculate order reserve/free-to-sell truth.  The
    UI marks every row as pending until /inventory/catalog finishes its canonical batch
    reconciliation.  Its only job is to make the physical stock list visible immediately
    instead of leaving a blank table while the heavier WB-card projection is calculated.
    """
    raw = db.execute(
        select(InventoryBalance, SKU, Location, Container)
        .join(SKU, SKU.id == InventoryBalance.sku_id)
        .join(Location, Location.id == InventoryBalance.location_id)
        .outerjoin(Container, Container.id == InventoryBalance.container_id)
        .where(
            InventoryBalance.client_id == client_id,
            InventoryBalance.warehouse_id == warehouse_id,
            InventoryBalance.quantity > 0,
            SKU.client_id == client_id,
        )
        .order_by(SKU.article.asc().nullslast(), SKU.size.asc().nullslast(), SKU.barcode.asc().nullslast(), Location.code.asc())
    ).all()

    grouped: dict[tuple[str, str, str], dict] = {}
    for balance, sku, location, container in raw:
        barcode = str(sku.barcode or sku.sku_code or "").strip()
        article = str(sku.article or sku.sku_code or barcode or sku.id).strip()
        size = str(sku.size or "").strip()
        key = (barcode or str(sku.id), article, size)
        row = grouped.get(key)
        if row is None:
            row = {
                "product_id": f"quick:{sku.id}",
                "local_only": True,
                "quick_local": True,
                "nm_id": None,
                "article": article,
                "title": str(sku.name or article or barcode or f"SKU {sku.id}"),
                "brand": None,
                "subject_name": None,
                "photo": None,
                "size": size or None,
                "barcodes": [barcode] if barcode else [],
                "barcode": barcode or None,
                "sku_id": str(sku.id),
                "sku_ids": [],
                "quantity": 0.0,
                "locations": [],
            }
            grouped[key] = row
        sid = str(sku.id)
        if sid not in row["sku_ids"]:
            row["sku_ids"].append(sid)
        qty = float(balance.quantity or 0)
        row["quantity"] += qty
        row["locations"].append({
            "balance_id": str(balance.id),
            "location_id": str(location.id),
            "location": str(location.code or ""),
            "container_id": str(container.id) if container else None,
            "container": (getattr(container, "name", None) or container.barcode) if container else None,
            "container_barcode": container.barcode if container else None,
            "quantity": qty,
        })

    rows = list(grouped.values())
    rows.sort(key=lambda x: (str(x.get("article") or ""), str(x.get("size") or ""), str(x.get("barcode") or "")))
    return {"rows": rows, "count": len(rows), "products": len(rows), "truth_pending": True}


@router.post("/operations/execute", response_model=OperationRead)
def execute(payload: OperationCreate, db: Session = Depends(get_db)):
    op = execute_operation(db, payload)
    _archive_empty_source_boxes(db, payload)
    return _serialize_operation(db, op)


@router.post("/operations/receive", response_model=OperationRead)
def receive(payload: ReceiveCreate, db: Session = Depends(get_db)):
    op = execute_operation(
        db,
        OperationCreate(
            client_id=payload.client_id,
            warehouse_id=payload.warehouse_id,
            operation_type="RECEIVING",
            idempotency_key=payload.idempotency_key,
            performed_by=payload.performed_by,
            reference=payload.reference,
            comment=payload.comment,
            lines=[
                StockLine(
                    sku_id=payload.sku_id,
                    quantity=payload.quantity,
                    destination_location_id=payload.destination_location_id,
                    destination_container_id=payload.destination_container_id,
                )
            ],
        ),
    )
    return _serialize_operation(db, op)


@router.post("/operations/move", response_model=OperationRead)
def move(payload: MoveCreate, db: Session = Depends(get_db)):
    operation_payload = OperationCreate(
        client_id=payload.client_id,
        warehouse_id=payload.warehouse_id,
        operation_type="MOVE",
        idempotency_key=payload.idempotency_key,
        performed_by=payload.performed_by,
        reference=payload.reference,
        comment=payload.comment,
        lines=[
            StockLine(
                sku_id=payload.sku_id,
                quantity=payload.quantity,
                source_location_id=payload.source_location_id,
                source_container_id=payload.source_container_id,
                destination_location_id=payload.destination_location_id,
                destination_container_id=payload.destination_container_id,
            )
        ],
    )
    op = execute_operation(db, operation_payload)
    _archive_empty_source_boxes(db, operation_payload)
    return _serialize_operation(db, op)


@router.post("/operations/ship", response_model=OperationRead)
def ship(payload: ShipCreate, db: Session = Depends(get_db)):
    operation_payload = OperationCreate(
        client_id=payload.client_id,
        warehouse_id=payload.warehouse_id,
        operation_type="SHIPPING",
        idempotency_key=payload.idempotency_key,
        performed_by=payload.performed_by,
        reference=payload.reference,
        comment=payload.comment,
        lines=[
            StockLine(
                sku_id=payload.sku_id,
                quantity=payload.quantity,
                source_location_id=payload.source_location_id,
                source_container_id=payload.source_container_id,
            )
        ],
    )
    op = execute_operation(db, operation_payload)
    _archive_empty_source_boxes(db, operation_payload)
    return _serialize_operation(db, op)


@router.post("/operations/adjust", response_model=OperationRead)
def adjust(payload: AdjustCreate, db: Session = Depends(get_db)):
    if payload.delta > 0:
        line = StockLine(
            sku_id=payload.sku_id,
            quantity=payload.delta,
            destination_location_id=payload.location_id,
            destination_container_id=payload.container_id,
            reason=payload.reason,
        )
    else:
        line = StockLine(
            sku_id=payload.sku_id,
            quantity=abs(payload.delta),
            source_location_id=payload.location_id,
            source_container_id=payload.container_id,
            reason=payload.reason,
        )
    operation_payload = OperationCreate(
        client_id=payload.client_id,
        warehouse_id=payload.warehouse_id,
        operation_type="ADJUSTMENT",
        idempotency_key=payload.idempotency_key,
        performed_by=payload.performed_by,
        reference=payload.reference,
        comment=payload.comment,
        lines=[line],
    )
    op = execute_operation(db, operation_payload)
    _archive_empty_source_boxes(db, operation_payload)
    return _serialize_operation(db, op)


# FC_K9R56_CONTAINER_MOVE_PREVIEW
@router.get("/operations/move-container/preview")
def preview_move_container(
    client_id: uuid.UUID = Query(...),
    warehouse_id: uuid.UUID = Query(...),
    container_id: uuid.UUID = Query(...),
    destination_location_id: uuid.UUID | None = Query(default=None),
    db: Session = Depends(get_db),
):
    return container_move_preview(
        db,
        client_id=client_id,
        warehouse_id=warehouse_id,
        container_id=container_id,
        destination_location_id=destination_location_id,
    )


@router.post("/operations/move-container", response_model=OperationRead)
def move_container(payload: MoveContainerCreate, db: Session = Depends(get_db)):
    op = move_entire_container(
        db,
        client_id=payload.client_id,
        warehouse_id=payload.warehouse_id,
        container_id=payload.container_id,
        destination_location_id=payload.destination_location_id,
        idempotency_key=payload.idempotency_key,
        performed_by=payload.performed_by,
        reference=payload.reference,
        comment=payload.comment,
    )
    return _serialize_operation(db, op)


@router.get("/operations", response_model=list[OperationRead])
def list_operations(
    client_id: uuid.UUID,
    warehouse_id: uuid.UUID,
    operation_type: str | None = Query(default=None),
    limit: int = Query(default=100, ge=1, le=1000),
    db: Session = Depends(get_db),
):
    stmt = select(InventoryOperation).where(
        InventoryOperation.client_id == client_id,
        InventoryOperation.warehouse_id == warehouse_id,
    )
    if operation_type:
        stmt = stmt.where(InventoryOperation.operation_type == operation_type.upper())
    ops = db.execute(stmt.order_by(InventoryOperation.created_at.desc()).limit(limit)).scalars().all()
    # FC_K9R4_OPERATION_HISTORY_BATCH: this endpoint is loaded in the background on
    # almost every warehouse context. The old serializer executed one extra SELECT for
    # every operation (301 SQL statements at limit=300), which could exhaust the pool
    # while another screen was loading. Fetch all lines in one indexed query instead.
    if not ops:
        return []
    op_ids = [op.id for op in ops]
    all_lines = db.execute(
        select(InventoryOperationLine)
        .where(InventoryOperationLine.operation_id.in_(op_ids))
        .order_by(InventoryOperationLine.operation_id, InventoryOperationLine.id)
    ).scalars().all()
    by_operation: dict[uuid.UUID, list[InventoryOperationLine]] = {}
    for line in all_lines:
        by_operation.setdefault(line.operation_id, []).append(line)
    return [_operation_read(op, by_operation.get(op.id, [])) for op in ops]



def _catalog_products(db: Session, client_id: uuid.UUID) -> list[MarketplaceProduct]:
    """Latest WB product-card projection per nmID for a client.

    Credentials can be disconnected/reconnected. The warehouse catalog must not disappear
    just because a new API credential row was created, so inventory is based on the latest
    imported card across all WB credential generations for the client.
    """
    # FC_K9R4_LATEST_WB_CARD_SQL: credentials are often re-created and every import
    # leaves older MarketplaceProduct generations behind. Loading all historical JSON
    # cards and deduplicating them in Python made the stock page progressively slower.
    # PostgreSQL DISTINCT ON returns only the newest row per nmID before ORM hydration.
    rows = db.execute(
        select(MarketplaceProduct)
        .where(MarketplaceProduct.client_id == client_id, MarketplaceProduct.provider == "WB")
        .distinct(MarketplaceProduct.external_product_id)
        .order_by(
            MarketplaceProduct.external_product_id,
            MarketplaceProduct.updated_at.desc(),
            MarketplaceProduct.created_at.desc(),
        )
    ).scalars().all()
    return rows


def _size_name(size: dict) -> str | None:
    value = size.get("techSize") or size.get("wbSize") or size.get("name")
    text = str(value or "").strip()
    return text or None


def _size_barcodes(size: dict) -> list[str]:
    return [str(x).strip() for x in (size.get("skus") or []) if str(x or "").strip()]


def _catalog_sku_maps(db: Session, client_id: uuid.UUID, products: list[MarketplaceProduct]):
    vendor_codes = {str(p.vendor_code).strip() for p in products if str(p.vendor_code or "").strip()}
    barcodes = {code for p in products for size in (p.sizes or []) if isinstance(size, dict) for code in _size_barcodes(size)}
    clauses = []
    if vendor_codes:
        clauses.extend([SKU.article.in_(vendor_codes), SKU.sku_code.in_(vendor_codes)])
    if barcodes:
        clauses.append(SKU.barcode.in_(barcodes))
    skus = db.execute(select(SKU).where(SKU.client_id == client_id, or_(*clauses))).scalars().all() if clauses else []
    by_barcode: dict[str, list[SKU]] = {}
    by_article_size: dict[tuple[str, str], list[SKU]] = {}
    for sku in skus:
        if sku.barcode:
            by_barcode.setdefault(str(sku.barcode).strip(), []).append(sku)
        article = str(sku.article or sku.sku_code or "").strip()
        size = str(sku.size or "").strip()
        if article:
            by_article_size.setdefault((article, size), []).append(sku)
    return by_barcode, by_article_size


def _catalog_row_skus(product: MarketplaceProduct, size_name: str | None, codes: list[str], by_barcode: dict[str, list[SKU]], by_article_size: dict[tuple[str, str], list[SKU]]) -> list[SKU]:
    # WB barcode is authoritative. Article+size is only a fallback when the WB card
    # has no matching local barcode, so two distinct barcodes cannot be merged merely
    # because they share the same article and size.
    matches: list[SKU] = []
    for code in codes:
        matches.extend(by_barcode.get(str(code).strip(), []))
    if not matches:
        article_key = str(product.vendor_code or "").strip()
        if article_key:
            matches.extend(by_article_size.get((article_key, str(size_name or "").strip()), []))
    unique: dict[str, SKU] = {}
    for sku in matches:
        unique[str(sku.id)] = sku
    return list(unique.values())


def _catalog_balance_map(db: Session, client_id: uuid.UUID, warehouse_id: uuid.UUID, sku_ids: set[uuid.UUID]):
    result: dict[uuid.UUID, list[dict]] = {sku_id: [] for sku_id in sku_ids}
    if not sku_ids:
        return result
    rows = db.execute(
        select(InventoryBalance, Location, Container).join(Location, Location.id == InventoryBalance.location_id)
        .outerjoin(Container, Container.id == InventoryBalance.container_id)
        .where(InventoryBalance.client_id == client_id, InventoryBalance.warehouse_id == warehouse_id, InventoryBalance.sku_id.in_(sku_ids), InventoryBalance.quantity != 0)
        .order_by(Location.code.asc(), Container.name.asc().nullslast(), Container.barcode.asc().nullslast())
    ).all()
    ids = [str(x) for x in sku_ids]
    reserve_stmt = text("""
        SELECT r.sku_id,r.source_location_id,COALESCE(r.source_container_id::text,'') AS container_key,
          COALESCE(SUM(CASE WHEN o.local_status IN ('NEW','PICKING') THEN r.quantity ELSE 0 END),0) AS reserved,
          COALESCE(SUM(CASE WHEN COALESCE(o.local_status,'') NOT IN ('NEW','PICKING')
            AND NOT (r.status='PACKING_WRITE_OFF' AND io.id IS NOT NULL) THEN r.quantity ELSE 0 END),0) AS accounting_pending
        FROM fc_stock_reservations r
        LEFT JOIN marketplace_orders o ON CAST(o.id AS varchar)=r.order_id
        LEFT JOIN inventory_operations io
          ON CAST(io.client_id AS varchar)=CAST(r.client_id AS varchar) AND COALESCE(io.status,'COMPLETED')='COMPLETED'
         AND io.idempotency_key=('stock-core-v2-pack-' || CAST(o.id AS varchar) || '-' || CAST(r.id AS varchar))
        WHERE r.client_id=:client AND r.warehouse_id=:warehouse AND r.sku_id IN :ids AND r.source_location_id IS NOT NULL
          AND r.status IN ('ACTIVE','CONSUMING','CANCELLING','PACKING_WRITE_OFF')
        GROUP BY r.sku_id,r.source_location_id,COALESCE(r.source_container_id::text,'')
    """).bindparams(bindparam("ids", expanding=True))
    reserve_map: dict[tuple[str,str,str], tuple[float,float]] = {}
    try:
        for row in db.execute(reserve_stmt,{"client":str(client_id),"warehouse":str(warehouse_id),"ids":ids}).mappings():
            reserve_map[(str(row["sku_id"]),str(row["source_location_id"]),str(row["container_key"] or ""))]=(float(row["reserved"] or 0),float(row["accounting_pending"] or 0))
    except Exception:
        reserve_map={}
    # FC_K9R37_EXTERNAL_PICK_BUCKET_TRUTH: units allocated by an uploaded external FBS
    # picking list remain physically in their source cell until picked, but they are no
    # longer free and must not be offered to WB/orders/a second external list.
    external_pick_map: dict[tuple[str,str,str], float] = {}
    try:
        ext_stmt = text("""
            SELECT sku_id,location_id,COALESCE(container_id,'') AS container_key,COALESCE(SUM(quantity),0) AS quantity
            FROM fc_external_pick_allocations
            WHERE client_id=:client AND warehouse_id=:warehouse AND sku_id IN :ids AND status='ACTIVE'
            GROUP BY sku_id,location_id,COALESCE(container_id,'')
        """).bindparams(bindparam("ids", expanding=True))
        for row in db.execute(ext_stmt,{"client":str(client_id),"warehouse":str(warehouse_id),"ids":ids}).mappings():
            external_pick_map[(str(row["sku_id"]),str(row["location_id"]),str(row["container_key"] or ""))]=max(0.0,float(row["quantity"] or 0))
    except Exception:
        external_pick_map={}

    # FC_K9R35_BUCKET_STORAGE_TRUTH: show the exact part of each physical bucket that is
    # intentionally excluded from WB FBS.  Without this, a location could say "free 34"
    # while five of those units were STORAGE_ONLY and the WB target correctly used 29.
    storage_map: dict[tuple[str,str,str], float] = {}
    try:
        storage_stmt = text("""
            SELECT CAST(sku_id AS varchar) AS sku_id,CAST(location_id AS varchar) AS location_id,
                   COALESCE(container_key,'') AS container_key,COALESCE(SUM(storage_only),0) AS storage_only
            FROM fc_stocks_v2
            WHERE client_id=:client AND warehouse_id=:warehouse AND CAST(sku_id AS varchar) IN :ids
            GROUP BY CAST(sku_id AS varchar),CAST(location_id AS varchar),COALESCE(container_key,'')
        """).bindparams(bindparam("ids", expanding=True))
        for row in db.execute(storage_stmt,{"client":str(client_id),"warehouse":str(warehouse_id),"ids":ids}).mappings():
            storage_map[(str(row["sku_id"]),str(row["location_id"]),str(row["container_key"] or ""))]=max(0.0,float(row["storage_only"] or 0))
    except Exception:
        storage_map={}
    for balance, location, container in rows:
        physical=float(balance.quantity or 0)
        key=(str(balance.sku_id),str(location.id),str(container.id) if container else "")
        reserved,pending=reserve_map.get(key,(0.0,0.0))
        external_pick=max(0.0,external_pick_map.get(key,0.0))
        storage=min(physical,max(0.0,storage_map.get(key,0.0)))
        result.setdefault(balance.sku_id,[]).append({
            "balance_id":str(balance.id),"sku_id":str(balance.sku_id),"location_id":str(location.id),"location":location.code,
            "container_id":str(container.id) if container else None,"container":(container.name or container.barcode) if container else None,
            "container_barcode":container.barcode if container else None,"quantity":physical,
            "reserved_quantity":reserved,"external_pick_reserved_quantity":external_pick,
            "accounting_pending_quantity":pending,"storage_only_quantity":storage,
            "free_quantity":max(0.0,physical-storage-reserved-pending-external_pick),
        })
    return result


def _catalog_stock_state_map(db: Session, client_id: uuid.UUID, warehouse_id: uuid.UUID, sku_ids: set[uuid.UUID]):
    """NEW/PICKING are reserve; terminal live rows are a separate safety hold."""
    result: dict[str, dict] = {str(sku_id): {"fbs_pool":0.0,"storage_only":0.0,"bucket_reserved":0.0,"sale_reserved":0.0,"address_reserved":0.0,"accounting_pending":0.0,"accounted_terminal":0.0,"live_reserved":0.0} for sku_id in sku_ids}
    if not sku_ids:
        return result
    ids=[str(x) for x in sku_ids]
    stock_stmt=text("""SELECT sku_id,COALESCE(SUM(fbs_available),0) AS fbs_pool,COALESCE(SUM(storage_only),0) AS storage_only,COALESCE(SUM(reserved),0) AS bucket_reserved FROM fc_stocks_v2 WHERE client_id=:client AND warehouse_id=:warehouse AND sku_id IN :ids GROUP BY sku_id""").bindparams(bindparam("ids",expanding=True))
    try:
        for row in db.execute(stock_stmt,{"client":str(client_id),"warehouse":str(warehouse_id),"ids":ids}).mappings():
            result.setdefault(str(row["sku_id"]),{}).update({"fbs_pool":float(row["fbs_pool"] or 0),"storage_only":float(row["storage_only"] or 0),"bucket_reserved":float(row["bucket_reserved"] or 0)})
    except Exception:
        pass
    reserve_stmt=text("""
        SELECT r.sku_id,
          COALESCE(SUM(CASE WHEN o.local_status='NEW' THEN r.quantity ELSE 0 END),0) AS sale_reserved,
          COALESCE(SUM(CASE WHEN o.local_status='PICKING' THEN r.quantity ELSE 0 END),0) AS address_reserved,
          COALESCE(SUM(CASE WHEN COALESCE(o.local_status,'') NOT IN ('NEW','PICKING')
             AND NOT (r.status='PACKING_WRITE_OFF' AND io.id IS NOT NULL) THEN r.quantity ELSE 0 END),0) AS accounting_pending,
          COALESCE(SUM(CASE WHEN COALESCE(o.local_status,'') NOT IN ('NEW','PICKING')
             AND r.status='PACKING_WRITE_OFF' AND io.id IS NOT NULL THEN r.quantity ELSE 0 END),0) AS accounted_terminal
        FROM fc_stock_reservations r
        LEFT JOIN marketplace_orders o ON CAST(o.id AS varchar)=r.order_id
        LEFT JOIN inventory_operations io
          ON CAST(io.client_id AS varchar)=CAST(r.client_id AS varchar) AND COALESCE(io.status,'COMPLETED')='COMPLETED'
         AND io.idempotency_key=('stock-core-v2-pack-' || CAST(o.id AS varchar) || '-' || CAST(r.id AS varchar))
        WHERE r.client_id=:client AND r.warehouse_id=:warehouse AND r.sku_id IN :ids
          AND r.status IN ('SALE_RESERVED','ACTIVE','CONSUMING','CANCELLING','PACKING_WRITE_OFF')
        GROUP BY r.sku_id
    """).bindparams(bindparam("ids",expanding=True))
    try:
        for row in db.execute(reserve_stmt,{"client":str(client_id),"warehouse":str(warehouse_id),"ids":ids}).mappings():
            result.setdefault(str(row["sku_id"]),{}).update({"sale_reserved":float(row["sale_reserved"] or 0),"address_reserved":float(row["address_reserved"] or 0),"accounting_pending":float(row["accounting_pending"] or 0),"accounted_terminal":float(row["accounted_terminal"] or 0)})
    except Exception:
        pass
    return result


def _catalog_movement_map(db: Session, client_id: uuid.UUID, warehouse_id: uuid.UUID, sku_ids: set[uuid.UUID]):
    """Batch movement ledger totals keyed by local SKU id.

    This is display/audit data only. Current physical stock remains InventoryBalance;
    WB availability remains the canonical free calculation with live safety holds.
    """
    result: dict[str, dict[str, float]] = {str(x): {
        "received": 0.0, "shipped": 0.0, "adjustment_in": 0.0,
        "adjustment_out": 0.0, "inbound_total": 0.0, "outbound_total": 0.0,
    } for x in sku_ids}
    if not sku_ids:
        return result
    stmt = text("""
        SELECT iol.sku_id,
          COALESCE(SUM(CASE WHEN io.operation_type IN ('RECEIVING','RECEIVING_IMPORT')
              AND iol.source_location_id IS NULL AND iol.destination_location_id IS NOT NULL
              THEN iol.quantity ELSE 0 END),0) AS received,
          COALESCE(SUM(CASE WHEN io.operation_type='SHIPPING'
              AND iol.source_location_id IS NOT NULL AND iol.destination_location_id IS NULL
              THEN iol.quantity ELSE 0 END),0) AS shipped,
          COALESCE(SUM(CASE WHEN io.operation_type='ADJUSTMENT'
              AND iol.source_location_id IS NULL AND iol.destination_location_id IS NOT NULL
              THEN iol.quantity ELSE 0 END),0) AS adjustment_in,
          COALESCE(SUM(CASE WHEN io.operation_type='ADJUSTMENT'
              AND iol.source_location_id IS NOT NULL AND iol.destination_location_id IS NULL
              THEN iol.quantity ELSE 0 END),0) AS adjustment_out,
          COALESCE(SUM(CASE WHEN io.operation_type<>'RECONCILIATION' AND iol.source_location_id IS NULL AND iol.destination_location_id IS NOT NULL
              THEN iol.quantity ELSE 0 END),0) AS inbound_total,
          COALESCE(SUM(CASE WHEN io.operation_type<>'RECONCILIATION' AND iol.source_location_id IS NOT NULL AND iol.destination_location_id IS NULL
              THEN iol.quantity ELSE 0 END),0) AS outbound_total
        FROM inventory_operation_lines iol
        JOIN inventory_operations io ON io.id=iol.operation_id
        WHERE io.client_id=:client AND io.warehouse_id=:warehouse AND COALESCE(io.status,'COMPLETED')='COMPLETED'
          AND iol.sku_id IN :ids
        GROUP BY iol.sku_id
    """).bindparams(bindparam('ids', expanding=True))
    for row in db.execute(stmt, {
        "client": str(client_id), "warehouse": str(warehouse_id),
        # Keep UUID parameters typed so PostgreSQL can use inventory_operation_lines.sku_id index.
        "ids": list(sku_ids),
    }).mappings():
        result[str(row["sku_id"])] = {name: float(row[name] or 0) for name in (
            "received","shipped","adjustment_in","adjustment_out","inbound_total","outbound_total"
        )}
    return result


@router.get("/inventory/catalog")
def inventory_catalog(
    client_id: uuid.UUID,
    warehouse_id: uuid.UUID,
    search: str | None = Query(default=None),
    wb_warehouse_id: int | None = Query(default=None),
    limit: int = Query(default=10000, ge=1, le=20000),
    db: Session = Depends(get_db),
):
    """Warehouse stock view based on WB product cards, including zero stock.

    Every WB size becomes a warehouse row. If there is no physical InventoryBalance for
    that size yet, quantity is returned as 0 instead of hiding the product.
    """
    products = _catalog_products(db, client_id)
    by_barcode, by_article_size = _catalog_sku_maps(db, client_id, products)
    # FC_K9R35_LOCAL_ONLY_CANONICAL_TRUTH: include every SKU that physically exists on
    # the warehouse, even when the latest local WB-card cache does not contain it.  This
    # makes physical / reserve / STORAGE_ONLY / free totals complete instead of dropping
    # "live-mapped but uncached" rows from the warehouse math.
    local_stock_skus = db.execute(
        select(SKU).join(InventoryBalance, InventoryBalance.sku_id == SKU.id).where(
            SKU.client_id == client_id, InventoryBalance.client_id == client_id,
            InventoryBalance.warehouse_id == warehouse_id, InventoryBalance.quantity > 0,
        ).distinct()
    ).scalars().all()
    candidate_sku_ids = ({sku.id for group in by_barcode.values() for sku in group}
                         | {sku.id for group in by_article_size.values() for sku in group}
                         | {sku.id for sku in local_stock_skus})
    balance_map = _catalog_balance_map(db, client_id, warehouse_id, candidate_sku_ids)
    # FC_K9R1_INVENTORY_CANONICAL_PARITY
    # Use the exact same alias-aware canonical stock truth as the address/location UI.
    # Older catalog code summed only the SKU rows discovered through the current WB card;
    # a historical duplicate SKU with the same exact barcode could therefore own the live
    # reservation while the inventory row displayed reserve=0.  stock_truth_for_skus()
    # resolves those canonical aliases in one batch and keeps the page fast.
    # FC_K9R48_INVENTORY_TRUTH_FAST_SPLIT: the stock screen must not hold reserve/free
    # truth behind the historical movement ledger. Movement totals are audit-only and
    # can be expensive on a warehouse with a long operation history. Load only the
    # canonical physical/storage/reservation truth in the critical request; the UI
    # hydrates movement totals separately in the background.
    canonical_truth_map = stock_truth_for_skus(
        db, client_id=client_id, warehouse_id=warehouse_id, sku_ids=candidate_sku_ids,
        wb_warehouse_id=wb_warehouse_id, include_movement=False,
    )
    # Legacy fbs_available is retained for diagnostics only. Its reservation columns are
    # deliberately not used for the visible reserve/free calculation below.
    stock_state_map = _catalog_stock_state_map(db, client_id, warehouse_id, candidate_sku_ids)
    from app.services.wb_sales_control import paused_ids as sale_pauses
    paused_sku_ids = sale_pauses(db,client_id,warehouse_id,candidate_sku_ids)
    needle = (search or "").strip().lower()
    rows: list[dict] = []
    represented_sku_ids: set[str] = set()
    represented_identity_groups: set[tuple[str, ...]] = set()
    for product in products:
        sizes = [x for x in (product.sizes or []) if isinstance(x, dict)] or [{}]
        for size in sizes:
            size_name = _size_name(size)
            codes = _size_barcodes(size) or [str(x) for x in (product.barcodes or []) if x]
            row_skus = _catalog_row_skus(product, size_name, codes, by_barcode, by_article_size)
            sku = row_skus[0] if row_skus else None
            locations = [place for item in row_skus for place in balance_map.get(item.id, [])]
            states = [stock_state_map.get(str(item.id), {}) for item in row_skus]

            # A WB size can point at one or several local SKU rows.  Collapse duplicate
            # canonical identities before summing so the inventory page and the location
            # page cannot disagree merely because an old duplicate SKU owns the reserve.
            truth_groups: dict[tuple[str, ...], dict] = {}
            for item in row_skus:
                truth = canonical_truth_map.get(str(item.id), {})
                identity = tuple(sorted(str(x) for x in (truth.get("identity_sku_ids") or [str(item.id)])))
                truth_groups.setdefault(identity, truth)
            canonical_states = list(truth_groups.values())
            represented_sku_ids.update(str(item.id) for item in row_skus)
            represented_identity_groups.update(truth_groups.keys())
            # FC_K9R35_CACHED_ROW_PHYSICAL_CANONICAL: current stock must come from the same
            # canonical identity as reserve/storage/free.  Otherwise a historical alias SKU
            # can own physical units that vanish from the cached WB row while the identity is
            # marked represented and therefore also suppressed from local-only rows.
            qty = sum(float(x.get("physical") or 0) for x in canonical_states)
            if not canonical_states:
                qty = sum(float(x.get("quantity") or 0) for x in locations)

            ledger_sale_reserved = sum(float(x.get("ledger_sale_reserved") or 0) for x in canonical_states)
            ledger_picking_reserved = sum(float(x.get("ledger_picking_reserved") or 0) for x in canonical_states)
            accounting_pending = sum(float(x.get("accounting_pending") or 0) for x in canonical_states)
            accounted_terminal = sum(float(x.get("accounted_terminal") or 0) for x in canonical_states)

            # Active NEW/PICKING orders remain the reserve safety floor.  Use one union
            # across every SKU/barcode represented by this WB size so one order is never
            # counted twice when the card exposes barcode aliases.
            order_floor = active_order_floor_for_identity(
                db, client_id=client_id, warehouse_id=warehouse_id,
                sku_ids=[item.id for item in row_skus], barcodes=codes,
                wb_warehouse_id=wb_warehouse_id,
            )
            order_new = float(order_floor.get("new") or 0)
            order_picking = float(order_floor.get("picking") or 0)
            sale_reserved = max(ledger_sale_reserved, order_new)
            address_reserved = max(ledger_picking_reserved, order_picking)
            total_reserved = sale_reserved + address_reserved
            external_pick_reserved = sum(float(x.get("external_pick_reserved") or 0) for x in canonical_states)
            reservation_gap = max(0.0, (order_new + order_picking) - (ledger_sale_reserved + ledger_picking_reserved))
            live_reserved = total_reserved + accounting_pending + external_pick_reserved

            # fbs_available is an old diagnostic field only.  The visible truth below
            # comes from physical InventoryBalance + canonical reserve/holds.
            fbs_pool = sum(float(x.get("fbs_pool") or 0) for x in states)
            storage_only = min(qty, max(0.0, sum(float(x.get("storage_only") or 0) for x in canonical_states)))
            available_to_sell = max(0.0, qty - storage_only - live_reserved)

            # Movement totals use the same canonical identity groups too, so a duplicate
            # SKU cannot make the card say "reserve from one row / movement from another".
            received = sum(float(x.get("received") or 0) for x in canonical_states)
            shipped = sum(float(x.get("shipped") or 0) for x in canonical_states)
            adjustment_in = sum(float(x.get("adjustment_in") or 0) for x in canonical_states)
            adjustment_out = sum(float(x.get("adjustment_out") or 0) for x in canonical_states)
            inbound_total = sum(float(x.get("inbound_total") or 0) for x in canonical_states)
            outbound_total = sum(float(x.get("outbound_total") or 0) for x in canonical_states)
            photo = None
            if product.photos:
                first = product.photos[0] if isinstance(product.photos[0], dict) else {}
                photo = first.get("big") or first.get("c516x688") or first.get("square") or first.get("tm")
            row = {
                "product_id": str(product.id),
                "nm_id": product.external_product_id,
                "article": product.vendor_code,
                "title": product.title,
                "brand": product.brand,
                "subject_name": product.subject_name,
                "photo": photo,
                "size": size_name,
                "chrt_id": size.get("chrtID") or size.get("chrtId"),
                "barcodes": codes,
                "barcode": codes[0] if codes else (sku.barcode if sku else None),
                "sku_id": str(sku.id) if sku else None,
                "sku_ids": [str(x.id) for x in row_skus],
                "quantity": qty,
                "movement_deferred": True,
                "received_quantity": None,
                "shipped_quantity": None,
                "adjustment_in_quantity": None,
                "adjustment_out_quantity": None,
                "recorded_inbound_quantity": None,
                "recorded_outbound_quantity": None,
                "opening_balance_estimate": None,
                "sale_reserved_quantity": sale_reserved,
                "picking_reserved_quantity": address_reserved,
                "reserved_quantity": total_reserved,
                "external_pick_reserved_quantity": external_pick_reserved,
                "ledger_sale_reserved_quantity": ledger_sale_reserved,
                "ledger_picking_reserved_quantity": ledger_picking_reserved,
                "order_new_quantity": order_new,
                "order_picking_quantity": order_picking,
                "reservation_gap_quantity": reservation_gap,
                "accounting_pending_quantity": accounting_pending,
                "accounted_terminal_quantity": accounted_terminal,
                "live_unavailable_quantity": live_reserved,
                "available_to_sell": available_to_sell,
                "fbs_pool_quantity": fbs_pool,
                "storage_only_quantity": storage_only,
                "legacy_fbs_pool_mismatch": bool(abs(float(fbs_pool) - max(0.0, qty - storage_only - live_reserved)) > 0.0001),
                "locations": locations,
                "wb_sales_paused": any(str(x.id) in paused_sku_ids for x in row_skus),
            }
            hay = " ".join([
                str(row.get("nm_id") or ""), str(row.get("article") or ""), str(row.get("title") or ""),
                str(row.get("brand") or ""), str(row.get("size") or ""), " ".join(codes),
                " ".join(f"{x.get('location','')} {x.get('container','')} {x.get('container_barcode','')}" for x in locations),
            ]).lower()
            if needle and needle not in hay:
                continue
            rows.append(row)
            if len(rows) >= limit:
                break
        if len(rows) >= limit:
            break

    # Append physical local rows that are not represented by the cached WB-card projection.
    # They still receive the same canonical truth fields, so UI/export totals remain complete.
    local_by_id = {str(sku.id): sku for sku in local_stock_skus}
    emitted_local_groups: set[tuple[str, ...]] = set()
    for source_sku in local_stock_skus:
        if len(rows) >= limit:
            break
        truth = canonical_truth_map.get(str(source_sku.id), {})
        identity = tuple(sorted(str(x) for x in (truth.get("identity_sku_ids") or [str(source_sku.id)])))
        if identity in represented_identity_groups or identity in emitted_local_groups:
            continue
        group_skus = [local_by_id[x] for x in identity if x in local_by_id]
        if not group_skus:
            group_skus = [source_sku]
        locations = [place for item in group_skus for place in balance_map.get(item.id, [])]
        qty = float(truth.get("physical") or sum(float(x.get("quantity") or 0) for x in locations))
        if qty <= 0:
            continue
        codes = list(dict.fromkeys(str(item.barcode or "").strip() for item in group_skus if str(item.barcode or "").strip()))
        primary = next((item for item in group_skus if str(item.barcode or "").strip()), group_skus[0])
        storage_only = min(qty, max(0.0, float(truth.get("storage_only") or 0)))
        sale_reserved = max(0.0, float(truth.get("sale_reserved") or 0))
        picking_reserved = max(0.0, float(truth.get("picking_reserved") or 0))
        accounting_pending = max(0.0, float(truth.get("accounting_pending") or 0))
        total_reserved = sale_reserved + picking_reserved
        external_pick_reserved = max(0.0, float(truth.get("external_pick_reserved") or 0))
        available_to_sell = max(0.0, qty - storage_only - total_reserved - accounting_pending - external_pick_reserved)
        # Historical movement is deliberately deferred on the list page; current stock
        # and WB safety truth above are already authoritative.
        row = {
            "product_id": f"local:{primary.id}", "local_only": True, "nm_id": None,
            "article": primary.article or primary.sku_code or primary.barcode,
            "title": getattr(primary, "name", None) or primary.article or primary.sku_code or primary.barcode,
            "brand": None, "subject_name": None, "photo": None, "size": primary.size,
            "chrt_id": None, "barcodes": codes, "barcode": codes[0] if codes else str(primary.barcode or ""),
            "sku_id": str(primary.id), "sku_ids": [str(x.id) for x in group_skus], "quantity": qty,
            "movement_deferred": True,
            "received_quantity": None, "shipped_quantity": None,
            "adjustment_in_quantity": None, "adjustment_out_quantity": None,
            "recorded_inbound_quantity": None, "recorded_outbound_quantity": None,
            "opening_balance_estimate": None,
            "sale_reserved_quantity": sale_reserved, "picking_reserved_quantity": picking_reserved,
            "reserved_quantity": total_reserved, "external_pick_reserved_quantity": external_pick_reserved,
            "ledger_sale_reserved_quantity": float(truth.get("ledger_sale_reserved") or 0),
            "ledger_picking_reserved_quantity": float(truth.get("ledger_picking_reserved") or 0),
            "order_new_quantity": float(truth.get("order_new") or 0),
            "order_picking_quantity": float(truth.get("order_picking") or 0),
            "reservation_gap_quantity": float(truth.get("reservation_gap") or 0),
            "accounting_pending_quantity": accounting_pending,
            "accounted_terminal_quantity": float(truth.get("accounted_terminal") or 0),
            "live_unavailable_quantity": total_reserved + accounting_pending + external_pick_reserved,
            "available_to_sell": available_to_sell, "fbs_pool_quantity": available_to_sell,
            "storage_only_quantity": storage_only, "legacy_fbs_pool_mismatch": False,
            "locations": locations,
            "wb_sales_paused": any(str(x.id) in paused_sku_ids for x in group_skus),
        }
        hay = " ".join([
            str(row.get("article") or ""), str(row.get("title") or ""), str(row.get("size") or ""), " ".join(codes),
            " ".join(f"{x.get('location','')} {x.get('container','')} {x.get('container_barcode','')}" for x in locations),
        ]).lower()
        if needle and needle not in hay:
            continue
        rows.append(row)
        emitted_local_groups.add(identity)

    rows.sort(key=lambda x: (str(x.get("article") or ""), str(x.get("size") or ""), str(x.get("barcode") or "")))
    return {"rows": rows, "count": len(rows), "products": len(products), "local_only_truth_rows": len(emitted_local_groups)}


# FC_K9R48_INVENTORY_MOVEMENT_BACKGROUND: audit/history totals are intentionally
# separated from the reserve/free critical path. A slow historical ledger can no longer
# keep the entire Inventory screen in the unsafe "truth pending" state.
@router.post("/inventory/movement-summary")
def inventory_movement_summary(payload: dict, db: Session = Depends(get_db)):
    try:
        client_id = uuid.UUID(str(payload.get("client_id")))
        warehouse_id = uuid.UUID(str(payload.get("warehouse_id")))
    except Exception:
        raise HTTPException(status_code=400, detail="Некорректный клиент или склад")
    raw_ids = payload.get("sku_ids") or []
    if not isinstance(raw_ids, list):
        raise HTTPException(status_code=400, detail="sku_ids должен быть списком")
    ids: list[uuid.UUID] = []
    for value in raw_ids[:20000]:
        try:
            ids.append(uuid.UUID(str(value)))
        except Exception:
            continue
    ids = list(dict.fromkeys(ids))
    if not ids:
        return {"rows": [], "count": 0}
    movement = _catalog_movement_map(db, client_id, warehouse_id, set(ids))
    rows = []
    for sid in ids:
        item = movement.get(str(sid), {})
        rows.append({
            "sku_id": str(sid),
            "received_quantity": float(item.get("received") or 0),
            "shipped_quantity": float(item.get("shipped") or 0),
            "adjustment_in_quantity": float(item.get("adjustment_in") or 0),
            "adjustment_out_quantity": float(item.get("adjustment_out") or 0),
            "recorded_inbound_quantity": float(item.get("inbound_total") or 0),
            "recorded_outbound_quantity": float(item.get("outbound_total") or 0),
        })
    return {"rows": rows, "count": len(rows)}


# FC_K9R35_XLSX_EXPORT: inventory user exports are real .xlsx files, not SpreadsheetML
# renamed to .xls/.xlsx.  The browser sends only the already visible report projection; this
# endpoint serializes it with openpyxl and does not alter stock, reservations or WB state.
@router.post("/inventory/export.xlsx")
def inventory_export_xlsx(payload: dict):
    headers = payload.get("headers") or []
    rows = payload.get("rows") or []
    chrt_headers = payload.get("chrt_headers") or []
    chrt_rows = payload.get("chrt_rows") or []
    notes = payload.get("notes") or []
    if not isinstance(headers, list) or not isinstance(rows, list):
        raise HTTPException(status_code=400, detail="Некорректные данные выгрузки")
    if len(rows) > 20000 or len(chrt_rows) > 20000 or len(headers) > 40 or len(chrt_headers) > 40:
        raise HTTPException(status_code=400, detail="Слишком большой отчёт для одной выгрузки")

    def safe_cell(value):
        if value is None:
            return None
        if isinstance(value, (int, float)):
            return value
        text_value = str(value)
        # Prevent spreadsheet formula injection from article/name/barcode text.
        if text_value[:1] in ("=", "+", "-", "@"):  # numeric negatives arrive as numbers
            text_value = "'" + text_value
        return text_value[:32767]

    wb = Workbook()
    ws = wb.active
    ws.title = "Остатки"
    ws.append([safe_cell(x) for x in headers])
    for row in rows:
        if isinstance(row, list):
            ws.append([safe_cell(x) for x in row[:len(headers)]])

    sheets = [ws]
    if chrt_headers:
        ws2 = wb.create_sheet("Сверка WB chrtId")
        ws2.append([safe_cell(x) for x in chrt_headers])
        for row in chrt_rows:
            if isinstance(row, list):
                ws2.append([safe_cell(x) for x in row[:len(chrt_headers)]])
        sheets.append(ws2)
    if notes:
        ws3 = wb.create_sheet("Пояснения")
        ws3.append(["Показатель", "Что означает"])
        for row in notes:
            if isinstance(row, list) and row:
                ws3.append([safe_cell(row[0]), safe_cell(row[1] if len(row) > 1 else "")])
        sheets.append(ws3)

    header_fill = PatternFill("solid", fgColor="D9EAD3")
    for sheet in sheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions if sheet.max_row > 1 else None
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for col_idx in range(1, sheet.max_column + 1):
            values = [str(sheet.cell(r, col_idx).value or "") for r in range(1, min(sheet.max_row, 200) + 1)]
            width = min(55, max(10, max((len(v) for v in values), default=10) + 2))
            sheet.column_dimensions[get_column_letter(col_idx)].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    raw_name = str(payload.get("filename") or "Остатки_FulfillCore.xlsx")
    base = re.sub(r'[\\/:*?"<>|]+', '_', raw_name).strip() or "Остатки_FulfillCore.xlsx"
    if not base.lower().endswith(".xlsx"):
        base += ".xlsx"
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    # FC_K9R36_XLSX_HEADER_ASCII_SAFE: Starlette encodes ordinary response headers as
    # latin-1.  A Cyrillic X-FulfillCore-Filename therefore raised UnicodeEncodeError
    # after the workbook had already been built and surfaced as HTTP 500 in the browser.
    # The browser-side saveBlob() already owns the human-readable Russian filename, so
    # keep the transport header strictly ASCII and never fail a valid export at response time.
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="inventory.xlsx"'},
    )



# FC_K9R62_PLACE_LIFECYCLE_HISTORY
# Reconstruct zero/refill moments from the authoritative physical operation ledger.
# No second stock truth is stored: current InventoryBalance + immutable operation deltas
# are enough to show when an exact cell/box became empty and when stock was placed there again.
def _place_lifecycle_history(
    db: Session, *, client_id: uuid.UUID, warehouse_id: uuid.UUID, sku_ids: set[uuid.UUID]
) -> tuple[dict[tuple[str, str], dict], list[dict]]:
    # FC_K9R63_PRODUCT_HISTORY: classify the net result of each atomic operation.
    from app.services.place_history import place_history, product_place_summary
    events = place_history(db, client_id=client_id, warehouse_id=warehouse_id, sku_ids=sku_ids)
    return product_place_summary(events), events


def _decorate_locations_with_lifecycle(locations: list[dict], lifecycle: dict[tuple[str, str], dict]) -> list[dict]:
    out=[]
    for place in locations:
        key=(str(place.get('location_id') or ''),str(place.get('container_id') or ''))
        out.append({**place, **lifecycle.get(key,{})})
    return out


def _inventory_history_rows(db: Session, *, client_id: uuid.UUID, warehouse_id: uuid.UUID, sku_ids: set[uuid.UUID], lifecycle_rows: list[dict]) -> list[dict]:
    history=[]
    if sku_ids:
        op_rows=db.execute(
            select(InventoryOperationLine,InventoryOperation)
            .join(InventoryOperation,InventoryOperation.id==InventoryOperationLine.operation_id)
            .where(InventoryOperation.client_id==client_id,InventoryOperation.warehouse_id==warehouse_id,InventoryOperationLine.sku_id.in_(sku_ids))
            .order_by(InventoryOperation.created_at.desc()).limit(120)
        ).all()
        location_ids={x for line,_ in op_rows for x in (line.source_location_id,line.destination_location_id) if x}
        container_ids={x for line,_ in op_rows for x in (line.source_container_id,line.destination_container_id) if x}
        loc_map={x.id:x.code for x in db.execute(select(Location).where(Location.id.in_(location_ids))).scalars().all()} if location_ids else {}
        cont_map={x.id:(x.name or x.barcode) for x in db.execute(select(Container).where(Container.id.in_(container_ids))).scalars().all()} if container_ids else {}
        for line,op in op_rows:
            history.append({'operation_id':str(op.id),'type':op.operation_type,'status':op.status,'quantity':float(line.quantity),'created_at':op.created_at,'performed_by':op.performed_by,'reference':op.reference,'comment':op.comment,'source':{'location':loc_map.get(line.source_location_id),'container':cont_map.get(line.source_container_id)} if line.source_location_id else None,'destination':{'location':loc_map.get(line.destination_location_id),'container':cont_map.get(line.destination_container_id)} if line.destination_location_id else None,'reason':line.reason})
    history.extend(lifecycle_rows)
    history.sort(key=lambda x: x.get('created_at') or datetime.min.replace(tzinfo=timezone.utc), reverse=True)
    return history[:180]


def _inventory_local_detail(
    db: Session, *, client_id: uuid.UUID, warehouse_id: uuid.UUID, raw_product_id: str,
    barcode: str | None, wb_warehouse_id: int | None,
) -> dict:
    raw=raw_product_id.split(':',1)[1] if ':' in raw_product_id else raw_product_id
    try:
        sku_id=uuid.UUID(raw)
    except Exception:
        raise HTTPException(status_code=404, detail='Локальный товар не найден')
    sku=db.get(SKU,sku_id)
    if not sku or sku.client_id!=client_id:
        raise HTTPException(status_code=404, detail='Локальный товар не найден')
    initial=stock_truth_for_skus(db,client_id=client_id,warehouse_id=warehouse_id,sku_ids={sku.id},wb_warehouse_id=wb_warehouse_id)
    truth=initial.get(str(sku.id),{})
    identity_ids=set()
    for value in truth.get('identity_sku_ids') or [str(sku.id)]:
        try: identity_ids.add(uuid.UUID(str(value)))
        except Exception: pass
    if not identity_ids: identity_ids={sku.id}
    skus=db.execute(select(SKU).where(SKU.client_id==client_id,SKU.id.in_(identity_ids))).scalars().all()
    if not skus: skus=[sku]
    codes=list(dict.fromkeys(str(x.barcode or '').strip() for x in skus if str(x.barcode or '').strip()))
    selected_barcode=str(barcode or sku.barcode or (codes[0] if codes else '')).strip() or None
    balance_map=_catalog_balance_map(db,client_id,warehouse_id,identity_ids)
    locations=[place for item in skus for place in balance_map.get(item.id,[])]
    lifecycle,lifecycle_rows=_place_lifecycle_history(db,client_id=client_id,warehouse_id=warehouse_id,sku_ids=identity_ids)
    locations=_decorate_locations_with_lifecycle(locations,lifecycle)
    history=_inventory_history_rows(db,client_id=client_id,warehouse_id=warehouse_id,sku_ids=identity_ids,lifecycle_rows=lifecycle_rows)
    physical=float(truth.get('physical') or sum(float(x.get('quantity') or 0) for x in locations))
    storage=min(physical,max(0.0,float(truth.get('storage_only') or 0)))
    return {
        'product': {'id':raw_product_id,'nm_id':None,'article':sku.article or sku.sku_code or selected_barcode,'title':sku.name or sku.article or sku.sku_code or selected_barcode,'brand':None,'photo':None,'local_only':True},
        'size':sku.size,'barcodes':codes,'barcode':selected_barcode,'sku_id':str(sku.id),'sku_ids':[str(x) for x in identity_ids],
        'quantity':physical,'received_quantity':float(truth.get('received') or 0),'shipped_quantity':float(truth.get('shipped') or 0),
        'adjustment_in_quantity':float(truth.get('adjustment_in') or 0),'adjustment_out_quantity':float(truth.get('adjustment_out') or 0),
        'recorded_inbound_quantity':float(truth.get('inbound_total') or 0),'recorded_outbound_quantity':float(truth.get('outbound_total') or 0),
        'opening_balance_estimate':float(truth.get('opening_balance_estimate') or 0),
        'sale_reserved_quantity':float(truth.get('sale_reserved') or 0),'picking_reserved_quantity':float(truth.get('picking_reserved') or 0),
        'reserved_quantity':float(truth.get('reserved') or 0),'external_pick_reserved_quantity':float(truth.get('external_pick_reserved') or 0),
        'ledger_sale_reserved_quantity':float(truth.get('ledger_sale_reserved') or 0),'ledger_picking_reserved_quantity':float(truth.get('ledger_picking_reserved') or 0),
        'order_new_quantity':float(truth.get('order_new') or 0),'order_picking_quantity':float(truth.get('order_picking') or 0),
        'reservation_gap_quantity':float(truth.get('reservation_gap') or 0),'accounting_pending_quantity':float(truth.get('accounting_pending') or 0),
        'accounted_terminal_quantity':float(truth.get('accounted_terminal') or 0),'live_unavailable_quantity':float(truth.get('live_reserved') or 0),
        'storage_only_quantity':storage,'available_to_sell':float(truth.get('available_to_sell') or 0),
        'locations':locations,'wb_sales_paused':False,'history':history,'place_history':lifecycle_rows,
    }


@router.get("/inventory/catalog/detail")
def inventory_catalog_detail(client_id: uuid.UUID, warehouse_id: uuid.UUID, product_id: str, barcode: str | None = Query(default=None), wb_warehouse_id: int | None = Query(default=None), db: Session = Depends(get_db)):
    raw_product_id=str(product_id or '').strip()
    if raw_product_id.startswith(('local:','quick:')):
        return _inventory_local_detail(db,client_id=client_id,warehouse_id=warehouse_id,raw_product_id=raw_product_id,barcode=barcode,wb_warehouse_id=wb_warehouse_id)
    try:
        product_uuid=uuid.UUID(raw_product_id)
    except Exception:
        raise HTTPException(status_code=404,detail="Товар не найден")
    product=db.get(MarketplaceProduct,product_uuid)
    if not product or product.client_id!=client_id or product.provider!="WB":
        raise HTTPException(status_code=404,detail="Товар WB не найден")
    sizes=[x for x in (product.sizes or []) if isinstance(x,dict)] or [{}]
    selected_size=next((x for x in sizes if barcode and barcode in _size_barcodes(x)),None) or sizes[0]
    codes=_size_barcodes(selected_size) or [str(x) for x in (product.barcodes or []) if x]
    selected_barcode=barcode or (codes[0] if codes else None); size_name=_size_name(selected_size)
    by_barcode,by_article_size=_catalog_sku_maps(db,client_id,[product])
    row_skus=_catalog_row_skus(product,size_name,codes,by_barcode,by_article_size); sku=row_skus[0] if row_skus else None
    sku_ids={x.id for x in row_skus}; balance_map=_catalog_balance_map(db,client_id,warehouse_id,sku_ids)
    locations=[place for item in row_skus for place in balance_map.get(item.id,[])]
    # FC_K9R1_INVENTORY_DETAIL_CANONICAL_PARITY
    # Product detail must use the same canonical alias groups as the catalog and
    # location contents. Historical duplicate SKU rows with one exact barcode must
    # not make one screen own the reserve while another screen displays zero.
    truth_map = stock_truth_for_skus(
        db, client_id=client_id, warehouse_id=warehouse_id, sku_ids=sku_ids,
        wb_warehouse_id=wb_warehouse_id,
    ) if sku_ids else {}
    truth_groups: dict[tuple[str, ...], dict] = {}
    for item in row_skus:
        truth = truth_map.get(str(item.id), {})
        identity = tuple(sorted(str(x) for x in (truth.get("identity_sku_ids") or [str(item.id)])))
        truth_groups.setdefault(identity, truth)
    canonical_states = list(truth_groups.values())

    qty = sum(float(x.get("quantity") or 0) for x in locations)
    ledger_sale_reserved = sum(float(x.get("ledger_sale_reserved") or 0) for x in canonical_states)
    ledger_picking_reserved = sum(float(x.get("ledger_picking_reserved") or 0) for x in canonical_states)
    accounting_pending = sum(float(x.get("accounting_pending") or 0) for x in canonical_states)
    accounted_terminal = sum(float(x.get("accounted_terminal") or 0) for x in canonical_states)
    order_floor = active_order_floor_for_identity(
        db, client_id=client_id, warehouse_id=warehouse_id,
        sku_ids=[x.id for x in row_skus], barcodes=codes,
        wb_warehouse_id=wb_warehouse_id,
    )
    order_new = float(order_floor.get("new") or 0)
    order_picking = float(order_floor.get("picking") or 0)
    sale_reserved = max(ledger_sale_reserved, order_new)
    picking_reserved = max(ledger_picking_reserved, order_picking)
    reservation_gap = max(0.0, (order_new + order_picking) - (ledger_sale_reserved + ledger_picking_reserved))
    external_pick_reserved = max(0.0, sum(float(x.get("external_pick_reserved") or 0) for x in canonical_states))
    live_reserved = sale_reserved + picking_reserved + accounting_pending + external_pick_reserved
    storage_only = min(qty, max(0.0, sum(float(x.get("storage_only") or 0) for x in canonical_states)))
    available = max(0.0, qty - storage_only - live_reserved)
    received = sum(float(x.get("received") or 0) for x in canonical_states)
    shipped = sum(float(x.get("shipped") or 0) for x in canonical_states)
    adjustment_in = sum(float(x.get("adjustment_in") or 0) for x in canonical_states)
    adjustment_out = sum(float(x.get("adjustment_out") or 0) for x in canonical_states)
    inbound_total = sum(float(x.get("inbound_total") or 0) for x in canonical_states)
    outbound_total = sum(float(x.get("outbound_total") or 0) for x in canonical_states)
    from app.services.wb_sales_control import paused_ids as sale_pauses
    wb_paused=bool(sale_pauses(db,client_id,warehouse_id,sku_ids))
    lifecycle,lifecycle_rows=_place_lifecycle_history(db,client_id=client_id,warehouse_id=warehouse_id,sku_ids=sku_ids)
    locations=_decorate_locations_with_lifecycle(locations,lifecycle)
    history=_inventory_history_rows(db,client_id=client_id,warehouse_id=warehouse_id,sku_ids=sku_ids,lifecycle_rows=lifecycle_rows)
    photo=None
    if product.photos:
        first=product.photos[0] if isinstance(product.photos[0],dict) else {}; photo=first.get("big") or first.get("c516x688") or first.get("square") or first.get("tm")
    return {"product":{"id":str(product.id),"nm_id":product.external_product_id,"article":product.vendor_code,"title":product.title,"brand":product.brand,"photo":photo},"size":size_name,"barcodes":codes,"barcode":selected_barcode,"sku_id":str(sku.id) if sku else None,"sku_ids":[str(x.id) for x in row_skus],"quantity":qty,"received_quantity":received,"shipped_quantity":shipped,"adjustment_in_quantity":adjustment_in,"adjustment_out_quantity":adjustment_out,"recorded_inbound_quantity":inbound_total,"recorded_outbound_quantity":outbound_total,"opening_balance_estimate":qty-(inbound_total-outbound_total),"sale_reserved_quantity":sale_reserved,"picking_reserved_quantity":picking_reserved,"reserved_quantity":sale_reserved+picking_reserved,"external_pick_reserved_quantity":external_pick_reserved,"ledger_sale_reserved_quantity":ledger_sale_reserved,"ledger_picking_reserved_quantity":ledger_picking_reserved,"order_new_quantity":order_new,"order_picking_quantity":order_picking,"reservation_gap_quantity":reservation_gap,"accounting_pending_quantity":accounting_pending,"accounted_terminal_quantity":accounted_terminal,"live_unavailable_quantity":live_reserved,"storage_only_quantity":storage_only,"available_to_sell":available,"locations":locations,"wb_sales_paused":wb_paused,"history":history,"place_history":lifecycle_rows}


def _ensure_catalog_sku(db: Session, payload: CatalogAdjustCreate, product: MarketplaceProduct) -> SKU:
    sku = db.execute(select(SKU).where(SKU.client_id == payload.client_id, SKU.barcode == payload.barcode)).scalars().first()
    if sku:
        return sku
    selected_size = next((x for x in (product.sizes or []) if isinstance(x, dict) and payload.barcode in _size_barcodes(x)), None) or {}
    size_name = _size_name(selected_size)
    chrt = selected_size.get("chrtID") or selected_size.get("chrtId") or payload.barcode[-12:]
    base_code = f"WB-{product.external_product_id}-{chrt}"[:128]
    code = base_code
    suffix = 1
    while db.execute(select(SKU.id).where(SKU.client_id == payload.client_id, SKU.sku_code == code)).scalar_one_or_none():
        suffix += 1
        code = f"{base_code[:120]}-{suffix}"
    sku = SKU(
        client_id=payload.client_id,
        sku_code=code,
        article=product.vendor_code or str(product.external_product_id),
        name=product.title or product.vendor_code or str(product.external_product_id),
        size=size_name,
        barcode=payload.barcode,
        unit="PCS",
        is_active=True,
    )
    db.add(sku)
    db.flush()
    return sku


@router.post("/inventory/catalog/adjust")
def inventory_catalog_adjust(payload: CatalogAdjustCreate, db: Session = Depends(get_db)):
    from fastapi import HTTPException
    product = db.get(MarketplaceProduct, payload.product_id)
    if not product or product.client_id != payload.client_id or product.provider != "WB":
        raise HTTPException(status_code=404, detail="Товар WB не найден")
    location = db.get(Location, payload.location_id)
    if not location or location.warehouse_id != payload.warehouse_id:
        raise HTTPException(status_code=404, detail="Ячейка не найдена на выбранном складе")
    if payload.container_id:
        container = db.get(Container, payload.container_id)
        if not container or container.warehouse_id != payload.warehouse_id or container.client_id != payload.client_id:
            raise HTTPException(status_code=404, detail="Короб не найден на выбранном складе")
    sku = _ensure_catalog_sku(db, payload, product)
    if payload.delta > 0:
        line = StockLine(
            sku_id=sku.id, quantity=payload.delta,
            destination_location_id=payload.location_id, destination_container_id=payload.container_id,
            reason=payload.reason,
        )
    else:
        line = StockLine(
            sku_id=sku.id, quantity=abs(payload.delta),
            source_location_id=payload.location_id, source_container_id=payload.container_id,
            reason=payload.reason,
        )
    operation_payload = OperationCreate(
        client_id=payload.client_id, warehouse_id=payload.warehouse_id,
        operation_type="ADJUSTMENT", idempotency_key=str(uuid.uuid4()),
        performed_by=payload.performed_by, reference=f"WB {product.external_product_id}",
        comment="Ручная корректировка из складской карточки", lines=[line],
    )
    op = execute_operation(db, operation_payload)
    _archive_empty_source_boxes(db, operation_payload)
    return {"ok": True, "sku_id": str(sku.id), "operation": _serialize_operation(db, op)}
