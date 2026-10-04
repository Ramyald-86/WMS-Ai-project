from __future__ import annotations

import json
import re
import threading
import uuid
from datetime import datetime, timezone
from io import BytesIO
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from pydantic import BaseModel, Field
from sqlalchemy import func, or_, select, text
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.models import Container, InventoryBalance, Location, MarketplaceOrder, SKU
from app.schemas.inventory import OperationCreate, StockLine
from app.services.inventory import execute_operation

router = APIRouter(prefix="/tsd/api", tags=["TSD"])
# FC_K9R49_TSD_ASSIGNMENT_MANAGER: desktop assigns/removes current picker work; TSD follows newest assignment.
# FC_K9R26_TSD_ANDROID_APP_RESUME_STATE:
# Android/WebView resume uses existing active-session API; verified cell/box stage is persisted client-side.
# FC_K9R25_TSD_DYNAMIC_BATCH_WEDGE_CAPTURE:
# task sizing follows actual PICKING workload; browser scanner diagnostics support
# focused TextBox injection/polling in addition to key events.
# FC_K9R24_TSD_SCANNER_FEEDBACK_RESULTS:
# scanner source verification supports location + exact box, session payload exposes
# per-location totals, and desktop FulfillCore can read finished TSD results.

_SCHEMA_READY = False
_SCHEMA_LOCK = threading.Lock()


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _qty(value: Any) -> float:
    try:
        return max(0.0, float(value or 0))
    except Exception:
        return 0.0


def _natural(value: str | None) -> tuple[Any, ...]:
    parts = re.split(r"(\d+)", str(value or "").upper())
    return tuple(int(x) if x.isdigit() else x for x in parts)


def _norm_scan(value: str | None) -> str:
    value = str(value or "").strip().upper()
    value = value.replace("–", "-").replace("—", "-")
    return re.sub(r"\s+", "", value)


def _ensure_schema(db: Session) -> None:
    global _SCHEMA_READY
    if _SCHEMA_READY:
        return
    with _SCHEMA_LOCK:
        if _SCHEMA_READY:
            return
        statements = [
            """
            CREATE TABLE IF NOT EXISTS fc_tsd_pick_sessions (
                id VARCHAR(36) PRIMARY KEY,
                client_id VARCHAR(36) NOT NULL,
                warehouse_id VARCHAR(36) NOT NULL,
                picker VARCHAR(120) NOT NULL,
                status VARCHAR(24) NOT NULL DEFAULT 'ACTIVE',
                created_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP,
                finished_at TIMESTAMPTZ NULL,
                cancelled_at TIMESTAMPTZ NULL,
                assignment_source VARCHAR(24) NOT NULL DEFAULT 'LEGACY',
                assigned_by VARCHAR(120) NULL,
                assigned_at TIMESTAMPTZ NULL
            )
            """,
            "ALTER TABLE fc_tsd_pick_sessions ADD COLUMN IF NOT EXISTS assignment_source VARCHAR(24) NOT NULL DEFAULT 'LEGACY'",
            "ALTER TABLE fc_tsd_pick_sessions ADD COLUMN IF NOT EXISTS assigned_by VARCHAR(120) NULL",
            "ALTER TABLE fc_tsd_pick_sessions ADD COLUMN IF NOT EXISTS assigned_at TIMESTAMPTZ NULL",
            "CREATE INDEX IF NOT EXISTS ix_fc_tsd_pick_sessions_scope ON fc_tsd_pick_sessions(client_id, warehouse_id, status, created_at)",
            "CREATE INDEX IF NOT EXISTS ix_fc_tsd_pick_sessions_picker_active ON fc_tsd_pick_sessions(client_id, warehouse_id, picker, status, created_at)",
            """
            CREATE TABLE IF NOT EXISTS fc_tsd_pick_lines (
                id VARCHAR(36) PRIMARY KEY,
                session_id VARCHAR(36) NOT NULL,
                reservation_id VARCHAR(36) NOT NULL,
                order_id VARCHAR(36) NOT NULL,
                sku_id VARCHAR(36) NOT NULL,
                location_id VARCHAR(36) NOT NULL,
                container_id VARCHAR(36) NULL,
                required_qty NUMERIC(18,3) NOT NULL DEFAULT 1,
                picked_qty NUMERIC(18,3) NOT NULL DEFAULT 0,
                missing_qty NUMERIC(18,3) NOT NULL DEFAULT 0,
                result VARCHAR(24) NOT NULL DEFAULT 'PENDING',
                created_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP,
                updated_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
            """,
            "CREATE UNIQUE INDEX IF NOT EXISTS uq_fc_tsd_pick_lines_reservation ON fc_tsd_pick_lines(reservation_id)",
            "CREATE INDEX IF NOT EXISTS ix_fc_tsd_pick_lines_session ON fc_tsd_pick_lines(session_id, location_id, container_id, sku_id)",
        ]
        try:
            for stmt in statements:
                db.execute(text(stmt))
            db.commit()
            _SCHEMA_READY = True
        except Exception:
            db.rollback()
            raise


def _load_objects(db: Session, rows: list[dict[str, Any]]):
    order_ids: list[uuid.UUID] = []
    sku_ids: list[uuid.UUID] = []
    loc_ids: list[uuid.UUID] = []
    cont_ids: list[uuid.UUID] = []
    for row in rows:
        for key, dest in (("order_id", order_ids), ("sku_id", sku_ids), ("location_id", loc_ids), ("container_id", cont_ids)):
            raw = row.get(key)
            if not raw:
                continue
            try:
                dest.append(uuid.UUID(str(raw)))
            except Exception:
                pass
    orders = {str(x.id): x for x in db.execute(select(MarketplaceOrder).where(MarketplaceOrder.id.in_(set(order_ids)))).scalars().all()} if order_ids else {}
    skus = {str(x.id): x for x in db.execute(select(SKU).where(SKU.id.in_(set(sku_ids)))).scalars().all()} if sku_ids else {}
    locs = {str(x.id): x for x in db.execute(select(Location).where(Location.id.in_(set(loc_ids)))).scalars().all()} if loc_ids else {}
    conts = {str(x.id): x for x in db.execute(select(Container).where(Container.id.in_(set(cont_ids)))).scalars().all()} if cont_ids else {}
    return orders, skus, locs, conts


def _order_barcodes(order: MarketplaceOrder | None, sku: SKU | None) -> list[str]:
    out: list[str] = []
    for value in (getattr(sku, "barcode", None), getattr(sku, "sku_code", None)):
        value = str(value or "").strip()
        if value and value not in out:
            out.append(value)
    raw = getattr(order, "barcodes", None) if order else None
    if isinstance(raw, list):
        for value in raw:
            value = str(value or "").strip()
            if value and value not in out:
                out.append(value)
    return out


def _session_row(db: Session, session_id: str) -> dict[str, Any]:
    _ensure_schema(db)
    row = db.execute(text("SELECT * FROM fc_tsd_pick_sessions WHERE id=:id"), {"id": session_id}).mappings().first()
    if not row:
        raise HTTPException(status_code=404, detail="Задание ТСД не найдено")
    return dict(row)


def _session_payload(db: Session, session_id: str) -> dict[str, Any]:
    session = _session_row(db, session_id)
    rows = [dict(x) for x in db.execute(text("""
        SELECT id, session_id, reservation_id, order_id, sku_id, location_id, container_id,
               required_qty, picked_qty, missing_qty, result, created_at, updated_at
        FROM fc_tsd_pick_lines
        WHERE session_id=:session_id
        ORDER BY created_at, id
    """), {"session_id": session_id}).mappings().all()]
    orders, skus, locs, conts = _load_objects(db, rows)
    groups: dict[tuple[str, str, str], dict[str, Any]] = {}
    total = picked = missing = 0.0
    for row in rows:
        required = _qty(row.get("required_qty"))
        got = _qty(row.get("picked_qty"))
        miss = _qty(row.get("missing_qty"))
        total += required; picked += got; missing += miss
        key = (str(row.get("location_id") or ""), str(row.get("container_id") or ""), str(row.get("sku_id") or ""))
        order = orders.get(str(row.get("order_id") or ""))
        sku = skus.get(str(row.get("sku_id") or ""))
        loc = locs.get(str(row.get("location_id") or ""))
        cont = conts.get(str(row.get("container_id") or ""))
        group = groups.setdefault(key, {
            "task_key": "|".join(key),
            "location_id": key[0],
            "location": getattr(loc, "code", None) or "",
            "location_barcode": getattr(loc, "barcode", None) or "",
            "container_id": key[1] or None,
            "container": (getattr(cont, "name", None) or getattr(cont, "barcode", None) or "") if cont else "",
            "container_barcode": getattr(cont, "barcode", None) or "" if cont else "",
            "sku_id": key[2],
            "article": getattr(sku, "article", None) or getattr(order, "article", None) or "",
            "size": getattr(sku, "size", None) or "",
            "name": getattr(sku, "name", None) or getattr(order, "article", None) or "",
            "barcodes": [],
            "required": 0.0,
            "picked": 0.0,
            "missing": 0.0,
            "order_numbers": [],
        })
        group["required"] += required
        group["picked"] += got
        group["missing"] += miss
        for barcode in _order_barcodes(order, sku):
            if barcode not in group["barcodes"]:
                group["barcodes"].append(barcode)
        if order:
            number = getattr(order, "external_order_num", None) or getattr(order, "external_order_id", None)
            if number is not None and str(number) not in group["order_numbers"]:
                group["order_numbers"].append(str(number))
    tasks = list(groups.values())
    tasks.sort(key=lambda x: (_natural(x.get("location")), _natural(x.get("container")), _natural(x.get("article")), _natural(x.get("size"))))
    for task in tasks:
        task["remaining"] = max(0.0, task["required"] - task["picked"] - task["missing"])
        task["complete"] = task["remaining"] <= 0.000001
    current = next((x for x in tasks if not x["complete"]), None)

    # FC_K9R24_TSD_CELL_SUMMARY: show the full workload for the physical cell,
    # not only the currently highlighted SKU.
    current_location = None
    if current:
        same_location = [x for x in tasks if str(x.get("location_id") or "") == str(current.get("location_id") or "")]
        loc_required = sum(_qty(x.get("required")) for x in same_location)
        loc_picked = sum(_qty(x.get("picked")) for x in same_location)
        loc_missing = sum(_qty(x.get("missing")) for x in same_location)
        current_location = {
            "location_id": current.get("location_id"),
            "location": current.get("location"),
            "required": loc_required,
            "picked": loc_picked,
            "missing": loc_missing,
            "remaining": max(0.0, loc_required - loc_picked - loc_missing),
            "task_count": len(same_location),
            "items": [
                {
                    "task_key": x.get("task_key"),
                    "article": x.get("article"),
                    "size": x.get("size"),
                    "container": x.get("container"),
                    "required": _qty(x.get("required")),
                    "picked": _qty(x.get("picked")),
                    "missing": _qty(x.get("missing")),
                    "remaining": max(0.0, _qty(x.get("required")) - _qty(x.get("picked")) - _qty(x.get("missing"))),
                }
                for x in same_location
            ],
        }
    return {
        "session": {
            "id": str(session["id"]), "client_id": str(session["client_id"]), "warehouse_id": str(session["warehouse_id"]),
            "picker": session["picker"], "status": session["status"], "created_at": session["created_at"],
            "finished_at": session.get("finished_at"),
            "assignment_source": session.get("assignment_source") or "LEGACY",
            "assigned_by": session.get("assigned_by"),
            "assigned_at": session.get("assigned_at"),
        },
        "progress": {"total": total, "picked": picked, "missing": missing, "done": picked + missing, "remaining": max(0.0, total - picked - missing)},
        "task_count": len(tasks),
        "tasks_done": sum(1 for x in tasks if x["complete"]),
        "current": current,
        "current_location": current_location,
        "tasks": tasks,
    }


class StartPickingRequest(BaseModel):
    client_id: uuid.UUID
    warehouse_id: uuid.UUID
    picker: str = Field(min_length=1, max_length=120)
    batch_size: int = Field(default=120, ge=1, le=20000)


class AssignPickingRequest(BaseModel):
    client_id: uuid.UUID
    warehouse_id: uuid.UUID
    picker: str = Field(min_length=1, max_length=120)
    batch_size: int = Field(default=120, ge=1, le=20000)
    assigned_by: str | None = Field(default=None, max_length=120)


class TaskActionRequest(BaseModel):
    task_key: str
    code: str | None = None
    confirm: bool = False


class MoveRequest(BaseModel):
    client_id: uuid.UUID
    warehouse_id: uuid.UUID
    picker: str = Field(min_length=1, max_length=120)
    source_container: str = Field(min_length=1, max_length=160)
    product_barcode: str = Field(min_length=1, max_length=160)
    destination_container: str = Field(min_length=1, max_length=160)
    quantity: float = Field(default=1, gt=0, le=100000)


@router.get("/ping")
def tsd_ping():
    return {"ok": True, "module": "FulfillCore TSD", "version": "1.3", "assignment_mode": "DESKTOP_LATEST"}


@router.get("/picking-results")
def picking_results(
    client_id: uuid.UUID,
    warehouse_id: uuid.UUID,
    limit: int = Query(default=30, ge=1, le=200),
    db: Session = Depends(get_db),
):
    """Recent TSD picking sessions for the desktop FulfillCore result panel."""
    _ensure_schema(db)
    rows = db.execute(text("""
        SELECT s.id, s.picker, s.status, s.created_at, s.finished_at, s.cancelled_at,
               s.assignment_source, s.assigned_by, s.assigned_at,
               COALESCE(SUM(l.required_qty),0) AS total_qty,
               COALESCE(SUM(l.picked_qty),0) AS picked_qty,
               COALESCE(SUM(l.missing_qty),0) AS missing_qty,
               COUNT(l.id) AS line_count
        FROM fc_tsd_pick_sessions s
        LEFT JOIN fc_tsd_pick_lines l ON l.session_id=s.id
        WHERE s.client_id=:client_id AND s.warehouse_id=:warehouse_id
        GROUP BY s.id, s.picker, s.status, s.created_at, s.finished_at, s.cancelled_at,
                 s.assignment_source, s.assigned_by, s.assigned_at
        ORDER BY s.created_at DESC
        LIMIT :limit
    """), {
        "client_id": str(client_id),
        "warehouse_id": str(warehouse_id),
        "limit": limit,
    }).mappings().all()
    out = []
    for row in rows:
        item = dict(row)
        item["id"] = str(item["id"])
        item["total_qty"] = _qty(item.get("total_qty"))
        item["picked_qty"] = _qty(item.get("picked_qty"))
        item["missing_qty"] = _qty(item.get("missing_qty"))
        item["done_qty"] = item["picked_qty"] + item["missing_qty"]
        item["remaining_qty"] = max(0.0, item["total_qty"] - item["done_qty"])
        item["line_count"] = int(item.get("line_count") or 0)
        item["report_url"] = f"/tsd/api/picking/{item['id']}/report.xlsx" if str(item.get("status") or "").upper() == "FINISHED" else None
        out.append(item)
    return out


@router.get("/picking/availability")
def picking_availability(
    client_id: uuid.UUID,
    warehouse_id: uuid.UUID,
    db: Session = Depends(get_db),
):
    """Exact TSD workload derived from current WB orders in local PICKING status.

    total_qty includes all active, cell-bound PICKING reservations in this warehouse.
    available_qty excludes reservations already claimed by any TSD session, so a second
    picker cannot receive the same physical unit.
    """
    _ensure_schema(db)
    row = db.execute(text("""
        SELECT
            COALESCE(SUM(r.quantity),0) AS total_qty,
            COALESCE(SUM(CASE WHEN NOT EXISTS (
                SELECT 1 FROM fc_tsd_pick_lines t WHERE t.reservation_id=r.id
            ) THEN r.quantity ELSE 0 END),0) AS available_qty,
            COUNT(*) AS reservation_count
        FROM fc_stock_reservations r
        JOIN marketplace_orders o ON CAST(o.id AS varchar)=r.order_id
        WHERE r.client_id=:client_id AND r.warehouse_id=:warehouse_id
          AND r.status='ACTIVE'
          AND r.source_location_id IS NOT NULL
          AND r.quantity > 0
          AND UPPER(COALESCE(o.provider,''))='WB'
          AND UPPER(COALESCE(o.local_status,''))='PICKING'
    """), {"client_id": str(client_id), "warehouse_id": str(warehouse_id)}).mappings().first() or {}
    total_qty = _qty(row.get("total_qty"))
    available_qty = _qty(row.get("available_qty"))
    return {
        "total_qty": total_qty,
        "available_qty": available_qty,
        "already_claimed_qty": max(0.0, total_qty - available_qty),
        "reservation_count": int(row.get("reservation_count") or 0),
    }


@router.get("/picking/active")
def active_session(
    client_id: uuid.UUID,
    warehouse_id: uuid.UUID,
    picker: str = Query(min_length=1),
    db: Session = Depends(get_db),
):
    _ensure_schema(db)
    row = db.execute(text("""
        SELECT id FROM fc_tsd_pick_sessions
        WHERE client_id=:client_id AND warehouse_id=:warehouse_id AND LOWER(picker)=LOWER(:picker) AND status='ACTIVE'
        ORDER BY created_at DESC LIMIT 1
    """), {"client_id": str(client_id), "warehouse_id": str(warehouse_id), "picker": picker.strip()}).mappings().first()
    return _session_payload(db, str(row["id"])) if row else {"session": None}


# FC_K9R49_TSD_DESKTOP_ASSIGNMENT_LATEST:
# New Android/browser TSD clients do not self-create work. The desktop FulfillCore
# explicitly assigns the current batch to a picker; the TSD always resolves the
# newest ACTIVE DESKTOP assignment and therefore cannot resume an obsolete local session.
@router.get("/picking/assigned")
def assigned_session(
    client_id: uuid.UUID,
    warehouse_id: uuid.UUID,
    picker: str = Query(min_length=1),
    db: Session = Depends(get_db),
):
    _ensure_schema(db)
    row = db.execute(text("""
        SELECT id FROM fc_tsd_pick_sessions
        WHERE client_id=:client_id AND warehouse_id=:warehouse_id AND LOWER(picker)=LOWER(:picker)
          AND status='ACTIVE' AND assignment_source='DESKTOP'
        ORDER BY COALESCE(assigned_at, created_at) DESC, created_at DESC LIMIT 1
    """), {"client_id": str(client_id), "warehouse_id": str(warehouse_id), "picker": picker.strip()}).mappings().first()
    return _session_payload(db, str(row["id"])) if row else {"session": None, "assignment_required": True}


def _create_picking_session(
    db: Session,
    *,
    client_id: uuid.UUID,
    warehouse_id: uuid.UUID,
    picker: str,
    batch_size: int,
    assignment_source: str,
    assigned_by: str | None = None,
) -> dict[str, Any]:
    _ensure_schema(db)
    picker = picker.strip()
    existing = db.execute(text("""
        SELECT id, assignment_source FROM fc_tsd_pick_sessions
        WHERE client_id=:client_id AND warehouse_id=:warehouse_id AND LOWER(picker)=LOWER(:picker) AND status='ACTIVE'
        ORDER BY created_at DESC LIMIT 1
    """), {"client_id": str(client_id), "warehouse_id": str(warehouse_id), "picker": picker}).mappings().first()
    if existing:
        return _session_payload(db, str(existing["id"]))

    session_id = str(uuid.uuid4())
    source = (assignment_source or "LEGACY").strip().upper()[:24]
    by = (assigned_by or "").strip()[:120] or None
    db.execute(text("""
        INSERT INTO fc_tsd_pick_sessions(
            id,client_id,warehouse_id,picker,status,assignment_source,assigned_by,assigned_at
        ) VALUES(:id,:client_id,:warehouse_id,:picker,'ACTIVE',:assignment_source,:assigned_by,CURRENT_TIMESTAMP)
    """), {
        "id": session_id, "client_id": str(client_id), "warehouse_id": str(warehouse_id),
        "picker": picker, "assignment_source": source, "assigned_by": by,
    })

    # Filter by CURRENT WB/PICKING state in SQL before the candidate limit. This prevents
    # old/stale reservation rows from crowding out the actual current picking workload.
    candidate_limit = min(max(batch_size * 5, 300), 50000)
    candidates = [dict(x) for x in db.execute(text("""
        SELECT r.id AS reservation_id, r.order_id, r.sku_id, r.source_location_id AS location_id,
               r.source_container_id AS container_id, r.quantity, r.created_at
        FROM fc_stock_reservations r
        JOIN marketplace_orders o ON CAST(o.id AS varchar)=r.order_id
        WHERE r.client_id=:client_id AND r.warehouse_id=:warehouse_id
          AND r.status='ACTIVE' AND r.source_location_id IS NOT NULL AND r.quantity > 0
          AND UPPER(COALESCE(o.provider,''))='WB'
          AND UPPER(COALESCE(o.local_status,''))='PICKING'
          AND NOT EXISTS (SELECT 1 FROM fc_tsd_pick_lines t WHERE t.reservation_id=r.id)
        ORDER BY r.created_at, r.id
        LIMIT :candidate_limit
    """), {"client_id": str(client_id), "warehouse_id": str(warehouse_id), "candidate_limit": candidate_limit}).mappings().all()]

    orders, skus, locs, conts = _load_objects(db, candidates)
    candidates.sort(key=lambda row: (
        _natural(getattr(locs.get(str(row.get("location_id") or "")), "code", "")),
        _natural((getattr(conts.get(str(row.get("container_id") or "")), "name", None) or getattr(conts.get(str(row.get("container_id") or "")), "barcode", None) or "")),
        _natural(getattr(skus.get(str(row.get("sku_id") or "")), "article", "")),
        str(row.get("order_id") or ""),
    ))

    claimed = 0.0
    for row in candidates:
        if claimed >= batch_size:
            break
        required = max(1.0, _qty(row.get("quantity")))
        result = db.execute(text("""
            INSERT INTO fc_tsd_pick_lines(
                id,session_id,reservation_id,order_id,sku_id,location_id,container_id,required_qty
            ) VALUES(:id,:session_id,:reservation_id,:order_id,:sku_id,:location_id,:container_id,:required_qty)
            ON CONFLICT (reservation_id) DO NOTHING
        """), {
            "id": str(uuid.uuid4()), "session_id": session_id, "reservation_id": str(row["reservation_id"]),
            "order_id": str(row["order_id"]), "sku_id": str(row["sku_id"]), "location_id": str(row["location_id"]),
            "container_id": str(row.get("container_id") or "") or None, "required_qty": required,
        })
        if result.rowcount:
            claimed += required
    if claimed <= 0:
        db.execute(text("DELETE FROM fc_tsd_pick_sessions WHERE id=:id"), {"id": session_id})
        db.commit()
        raise HTTPException(status_code=409, detail="Нет свободных заказов «На сборке» с закреплённой ячейкой")
    db.commit()
    return _session_payload(db, session_id)


@router.post("/picking/start")
def start_picking(payload: StartPickingRequest, db: Session = Depends(get_db)):
    # Backward-compatible endpoint for older TSD builds. New clients use /assigned.
    return _create_picking_session(
        db,
        client_id=payload.client_id,
        warehouse_id=payload.warehouse_id,
        picker=payload.picker,
        batch_size=payload.batch_size,
        assignment_source="LEGACY",
    )


@router.post("/picking/assign")
def assign_picking(payload: AssignPickingRequest, db: Session = Depends(get_db)):
    _ensure_schema(db)
    picker = payload.picker.strip()
    existing = db.execute(text("""
        SELECT id FROM fc_tsd_pick_sessions
        WHERE client_id=:client_id AND warehouse_id=:warehouse_id AND LOWER(picker)=LOWER(:picker) AND status='ACTIVE'
        ORDER BY created_at DESC LIMIT 1
    """), {"client_id": str(payload.client_id), "warehouse_id": str(payload.warehouse_id), "picker": picker}).mappings().first()
    if existing:
        state = _session_payload(db, str(existing["id"]))
        remaining = _qty((state.get("progress") or {}).get("remaining"))
        raise HTTPException(status_code=409, detail=f"У сборщика уже есть активная сборка ({remaining:g} шт. осталось). Сначала снимите её.")
    return _create_picking_session(
        db,
        client_id=payload.client_id,
        warehouse_id=payload.warehouse_id,
        picker=picker,
        batch_size=payload.batch_size,
        assignment_source="DESKTOP",
        assigned_by=payload.assigned_by,
    )


@router.post("/picking/{session_id}/unassign")
def unassign_picking(session_id: str, db: Session = Depends(get_db)):
    session = _session_row(db, session_id)
    if str(session.get("status") or "").upper() != "ACTIVE":
        return {"ok": True, "already_inactive": True}
    progress = db.execute(text("""
        SELECT COALESCE(SUM(picked_qty),0) AS picked, COALESCE(SUM(missing_qty),0) AS missing
        FROM fc_tsd_pick_lines WHERE session_id=:id
    """), {"id": session_id}).mappings().first() or {}
    touched = _qty(progress.get("picked")) + _qty(progress.get("missing"))
    if touched > 0.000001:
        raise HTTPException(status_code=409, detail=f"Сборка уже начата: обработано {touched:g} шт. Снять её без возврата товара нельзя. Сначала завершите текущую сборку.")
    # Removing an untouched assignment releases only the TSD claim. It never changes
    # WB/local order status or physical stock.
    db.execute(text("DELETE FROM fc_tsd_pick_lines WHERE session_id=:id"), {"id": session_id})
    db.execute(text("""
        UPDATE fc_tsd_pick_sessions
        SET status='CANCELLED',cancelled_at=CURRENT_TIMESTAMP
        WHERE id=:id AND status='ACTIVE'
    """), {"id": session_id})
    db.commit()
    return {"ok": True, "session_id": session_id, "picker": session.get("picker")}


@router.get("/picking/{session_id}")
def get_picking_session(session_id: str, db: Session = Depends(get_db)):
    return _session_payload(db, session_id)


def _require_current(payload: TaskActionRequest, state: dict[str, Any]) -> dict[str, Any]:
    current = state.get("current")
    if not current:
        raise HTTPException(status_code=409, detail="Задание уже полностью обработано")
    if str(current.get("task_key")) != str(payload.task_key):
        raise HTTPException(status_code=409, detail="Позиция уже изменилась. Обновите задание на ТСД")
    return current


@router.post("/picking/{session_id}/verify-location")
def verify_location(session_id: str, payload: TaskActionRequest, db: Session = Depends(get_db)):
    state = _session_payload(db, session_id)
    current = _require_current(payload, state)
    scanned = _norm_scan(payload.code)
    accepted = {_norm_scan(current.get("location")), _norm_scan(current.get("location_barcode"))}
    accepted.discard("")
    if scanned not in accepted:
        raise HTTPException(status_code=409, detail=f"Неверная ячейка. Нужна: {current.get('location') or 'без кода'}")
    return {"ok": True, "task_key": current["task_key"], "location": current.get("location")}


@router.post("/picking/{session_id}/verify-container")
def verify_container(session_id: str, payload: TaskActionRequest, db: Session = Depends(get_db)):
    state = _session_payload(db, session_id)
    current = _require_current(payload, state)
    expected = str(current.get("container") or "").strip()
    if not expected:
        return {"ok": True, "task_key": current["task_key"], "container": None}
    scanned = _norm_scan(payload.code)
    accepted = {_norm_scan(current.get("container")), _norm_scan(current.get("container_barcode"))}
    accepted.discard("")
    if scanned not in accepted:
        raise HTTPException(status_code=409, detail=f"Неверный короб. Нужен: {expected}")
    return {"ok": True, "task_key": current["task_key"], "container": expected}


@router.post("/picking/{session_id}/scan-product")
def scan_product(session_id: str, payload: TaskActionRequest, db: Session = Depends(get_db)):
    state = _session_payload(db, session_id)
    current = _require_current(payload, state)
    scanned = _norm_scan(payload.code)
    accepted = {_norm_scan(x) for x in (current.get("barcodes") or []) if x}
    if scanned not in accepted:
        raise HTTPException(status_code=409, detail=f"Неверный товар. Нужен артикул {current.get('article')}, размер {current.get('size')}")
    loc, cont, sku = current["task_key"].split("|", 2)
    rows = db.execute(text("""
        SELECT id, required_qty, picked_qty, missing_qty
        FROM fc_tsd_pick_lines
        WHERE session_id=:session_id AND location_id=:location_id
          AND COALESCE(container_id,'')=:container_id AND sku_id=:sku_id
          AND (picked_qty + missing_qty) < required_qty
        ORDER BY created_at, id
        FOR UPDATE
    """), {"session_id": session_id, "location_id": loc, "container_id": cont, "sku_id": sku}).mappings().all()
    if not rows:
        db.rollback()
        raise HTTPException(status_code=409, detail="Эта позиция уже полностью обработана")
    row = rows[0]
    new_picked = min(_qty(row["required_qty"]), _qty(row["picked_qty"]) + 1.0)
    remaining = _qty(row["required_qty"]) - new_picked - _qty(row["missing_qty"])
    result = "PICKED" if remaining <= 0.000001 and _qty(row["missing_qty"]) <= 0.000001 else ("PARTIAL" if new_picked > 0 else "PENDING")
    db.execute(text("UPDATE fc_tsd_pick_lines SET picked_qty=:picked,result=:result,updated_at=CURRENT_TIMESTAMP WHERE id=:id"), {"picked": new_picked, "result": result, "id": row["id"]})
    db.commit()
    return _session_payload(db, session_id)


@router.post("/picking/{session_id}/missing")
def mark_missing(session_id: str, payload: TaskActionRequest, db: Session = Depends(get_db)):
    if not payload.confirm:
        raise HTTPException(status_code=400, detail="Требуется подтверждение отсутствия товара")
    state = _session_payload(db, session_id)
    current = _require_current(payload, state)
    loc, cont, sku = current["task_key"].split("|", 2)
    rows = db.execute(text("""
        SELECT id, required_qty, picked_qty, missing_qty
        FROM fc_tsd_pick_lines
        WHERE session_id=:session_id AND location_id=:location_id
          AND COALESCE(container_id,'')=:container_id AND sku_id=:sku_id
          AND (picked_qty + missing_qty) < required_qty
        ORDER BY created_at, id
        FOR UPDATE
    """), {"session_id": session_id, "location_id": loc, "container_id": cont, "sku_id": sku}).mappings().all()
    for row in rows:
        required = _qty(row["required_qty"]); picked = _qty(row["picked_qty"])
        missing = max(0.0, required - picked)
        result = "MISSING" if picked <= 0.000001 else "PARTIAL"
        db.execute(text("UPDATE fc_tsd_pick_lines SET missing_qty=:missing,result=:result,updated_at=CURRENT_TIMESTAMP WHERE id=:id"), {"missing": missing, "result": result, "id": row["id"]})
    db.commit()
    return _session_payload(db, session_id)


@router.post("/picking/{session_id}/finish")
def finish_picking(session_id: str, db: Session = Depends(get_db)):
    state = _session_payload(db, session_id)
    if _qty(state["progress"].get("remaining")) > 0.000001:
        raise HTTPException(status_code=409, detail="Сначала обработайте все позиции задания")
    db.execute(text("UPDATE fc_tsd_pick_sessions SET status='FINISHED',finished_at=CURRENT_TIMESTAMP WHERE id=:id AND status='ACTIVE'"), {"id": session_id})
    db.commit()
    return {"ok": True, "session_id": session_id, "report_url": f"/tsd/api/picking/{session_id}/report.xlsx"}


@router.post("/picking/{session_id}/cancel")
def cancel_picking(session_id: str, db: Session = Depends(get_db)):
    _session_row(db, session_id)
    db.execute(text("DELETE FROM fc_tsd_pick_lines WHERE session_id=:id"), {"id": session_id})
    db.execute(text("UPDATE fc_tsd_pick_sessions SET status='CANCELLED',cancelled_at=CURRENT_TIMESTAMP WHERE id=:id"), {"id": session_id})
    db.commit()
    return {"ok": True}


def _safe_filename(value: str) -> str:
    value = re.sub(r"[^0-9A-Za-zА-Яа-я_-]+", "_", value.strip())
    return value.strip("_")[:60] or "picker"


@router.get("/picking/{session_id}/report.xlsx")
def picking_report(session_id: str, db: Session = Depends(get_db)):
    state = _session_payload(db, session_id)
    session = state["session"]
    wb = Workbook()
    summary = wb.active
    summary.title = "Итог"
    summary.append(["Сборщик", session["picker"]])
    summary.append(["Дата задания", session.get("created_at")])
    summary.append(["Статус", session["status"]])
    summary.append(["Требовалось", state["progress"]["total"]])
    summary.append(["Собрано", state["progress"]["picked"]])
    summary.append(["Нет в наличии", state["progress"]["missing"]])
    for cell in summary[1] + summary[4]:
        cell.font = Font(bold=True)

    headers = ["Ячейка", "Короб", "Артикул", "Размер", "ШК", "Требовалось", "Собрано", "Нет в наличии", "Заказы"]
    picked_ws = wb.create_sheet("Собрано")
    missing_ws = wb.create_sheet("Нет в наличии")
    for ws in (picked_ws, missing_ws):
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"
    for task in state["tasks"]:
        row = [task["location"], task["container"], task["article"], task["size"], ", ".join(task["barcodes"]), task["required"], task["picked"], task["missing"], ", ".join(task["order_numbers"])]
        if _qty(task["picked"]) > 0:
            picked_ws.append(row)
        if _qty(task["missing"]) > 0:
            missing_ws.append(row)
    for ws in (picked_ws, missing_ws):
        widths = [16, 18, 18, 12, 24, 14, 12, 16, 34]
        for idx, width in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + idx)].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    out = BytesIO(); wb.save(out)
    stamp = _now().astimezone().strftime("%Y%m%d_%H%M")
    filename = f"TSD_picking_{stamp}_{_safe_filename(session['picker'])}.xlsx"
    return Response(content=out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{filename}"'})


def _resolve_container(db: Session, client_id: uuid.UUID, warehouse_id: uuid.UUID, code: str) -> Container:
    value = str(code or "").strip()
    container = db.execute(select(Container).where(
        Container.client_id == client_id,
        Container.warehouse_id == warehouse_id,
        or_(func.lower(Container.barcode) == value.lower(), func.lower(Container.name) == value.lower()),
    )).scalars().first()
    if not container:
        raise HTTPException(status_code=404, detail=f"Короб не найден: {value}")
    return container


def _resolve_sku(db: Session, client_id: uuid.UUID, barcode: str) -> SKU:
    value = str(barcode or "").strip()
    sku = db.execute(select(SKU).where(
        SKU.client_id == client_id,
        or_(func.lower(SKU.barcode) == value.lower(), func.lower(SKU.sku_code) == value.lower()),
    )).scalars().first()
    if not sku:
        raise HTTPException(status_code=404, detail=f"Товар по ШК не найден: {value}")
    return sku


@router.post("/move")
def move_box_to_box(payload: MoveRequest, db: Session = Depends(get_db)):
    source = _resolve_container(db, payload.client_id, payload.warehouse_id, payload.source_container)
    destination = _resolve_container(db, payload.client_id, payload.warehouse_id, payload.destination_container)
    if source.id == destination.id:
        raise HTTPException(status_code=409, detail="Исходный и конечный короб совпадают")
    if not source.location_id or not destination.location_id:
        raise HTTPException(status_code=409, detail="Оба короба должны быть размещены в реальных ячейках")
    sku = _resolve_sku(db, payload.client_id, payload.product_barcode)
    available = db.execute(select(func.coalesce(func.sum(InventoryBalance.quantity), 0)).where(
        InventoryBalance.client_id == payload.client_id,
        InventoryBalance.warehouse_id == payload.warehouse_id,
        InventoryBalance.sku_id == sku.id,
        InventoryBalance.container_id == source.id,
        InventoryBalance.location_id == source.location_id,
        InventoryBalance.quantity > 0,
    )).scalar_one()
    if _qty(available) + 0.000001 < payload.quantity:
        raise HTTPException(status_code=409, detail=f"В исходном коробе только {_qty(available):g} шт., запрошено {payload.quantity:g}")
    op = execute_operation(db, OperationCreate(
        client_id=payload.client_id,
        warehouse_id=payload.warehouse_id,
        operation_type="MOVE",
        idempotency_key=f"tsd-move-{uuid.uuid4()}",
        performed_by=payload.picker.strip(),
        reference="TSD BOX-TO-BOX",
        comment=f"ТСД: {source.name or source.barcode} -> {destination.name or destination.barcode}",
        lines=[StockLine(
            sku_id=sku.id,
            quantity=payload.quantity,
            source_location_id=source.location_id,
            source_container_id=source.id,
            destination_location_id=destination.location_id,
            destination_container_id=destination.id,
        )],
    ))
    return {
        "ok": True,
        "operation_id": str(op.id),
        "article": sku.article,
        "size": sku.size,
        "quantity": payload.quantity,
        "source": source.name or source.barcode,
        "destination": destination.name or destination.barcode,
    }
