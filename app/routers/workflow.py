from __future__ import annotations
import io
import json
import threading
import uuid
from collections import defaultdict
from decimal import Decimal

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, Field
from sqlalchemy import func, or_, select, text
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.models import Client, Warehouse, Container, InventoryBalance, InventoryOperation, Location, SKU
from app.schemas.inventory import OperationCreate, StockLine
from app.services.inventory import execute_operation
from app.services.stock_core_v2 import ensure_stock_core_v2_schema, reconcile_stock_bucket, move_bucket
from app.services.warehouse_availability import available_buckets, held_quantity, lock_scope
from app.services.local_catalog import catalog as local_catalog
from app.services.box_routes import enrich as route_enrich, set_box_route, freeze_task
from app.services.picking_route import route_key
from app.services.workflow_rules import (MAX_BYTES, allocate, fingerprint, normalize_package,
    read_terminal_workbook, read_transfer_workbook)
from app.routers.warehouse import _ensure_schema as ensure_picks, _audit
from app.routers.layout import _location_barcode

router=APIRouter(prefix='/workflow',tags=['Приёмка → сборка → FBS'])
_READY=False
_LOCK=threading.Lock()
DDL=[
    '''CREATE TABLE IF NOT EXISTS fc_workflow_plans (
       id VARCHAR(36) PRIMARY KEY, client_id VARCHAR(36) NOT NULL, warehouse_id VARCHAR(36) NOT NULL,
       kind VARCHAR(24) NOT NULL,status VARCHAR(24) NOT NULL DEFAULT 'PREVIEW',
       content_hash VARCHAR(64) NOT NULL,payload_json JSONB NOT NULL,plan_json JSONB NOT NULL,
       actor VARCHAR(120),created_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP,
       confirmed_at TIMESTAMPTZ,operation_id VARCHAR(36),task_id VARCHAR(36))''',
    'CREATE INDEX IF NOT EXISTS ix_fc_workflow_scope ON fc_workflow_plans(client_id,warehouse_id,created_at DESC)',
    '''CREATE TABLE IF NOT EXISTS fc_terminal_receipts (
       client_id VARCHAR(36) NOT NULL,warehouse_id VARCHAR(36) NOT NULL,package_id VARCHAR(120) NOT NULL,
       content_hash VARCHAR(64) NOT NULL,operation_id VARCHAR(36) NOT NULL,plan_id VARCHAR(36) NOT NULL,
       created_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP,
       PRIMARY KEY(client_id,warehouse_id,package_id))''',
]


def ensure_schema(db):
    global _READY
    ensure_picks(db)
    ensure_stock_core_v2_schema(db)
    if _READY: return
    with _LOCK:
        if _READY: return
        try:
            for sql in DDL: db.execute(text(sql))
            db.commit(); _READY=True
        except Exception:
            db.rollback(); raise


def scope(db,c,w):
    if not db.get(Client,c) or not db.get(Warehouse,w):
        raise HTTPException(404,detail='Клиент или склад не найден. Откройте терминал из WMS')


def actor(request):
    return str(getattr(request.state,'username','') or 'Администратор')[:120]


def parsed(fn,*args):
    try: return fn(*args)
    except ValueError as exc: raise HTTPException(422,detail=str(exc))


def obj(v):
    return json.loads(v) if isinstance(v,str) else v


def params(c,w): return {'c':str(c),'w':str(w)}


def plan_response(row):
    result=dict(obj(row['plan_json']))
    return {**result,'plan_id':str(row['id']),'fingerprint':row['content_hash'],
            'status':row['status'],'operation_id':row.get('operation_id'),'task_id':row.get('task_id'),
            'actor':row.get('actor'),'created_at':str(row.get('created_at') or ''),'confirmed_at':str(row.get('confirmed_at') or '')}


def save_plan(db,c,w,kind,payload,plan,by):
    pid=str(uuid.uuid4()); digest=fingerprint({'kind':kind,'payload':payload,'plan':plan})
    db.execute(text('''INSERT INTO fc_workflow_plans(id,client_id,warehouse_id,kind,content_hash,payload_json,plan_json,actor)
        VALUES(:id,:c,:w,:kind,:hash,CAST(:payload AS jsonb),CAST(:plan AS jsonb),:actor)'''),
        {**params(c,w),'id':pid,'kind':kind,'hash':digest,'payload':json.dumps(payload,ensure_ascii=False),
         'plan':json.dumps(plan,ensure_ascii=False),'actor':by})
    db.commit()
    return {**plan,'plan_id':pid,'fingerprint':digest,'status':'PREVIEW'}


@router.get('/context')
def context(client_id:uuid.UUID,warehouse_id:uuid.UUID,db:Session=Depends(get_db)):
    ensure_schema(db); scope(db,client_id,warehouse_id)
    locs=db.execute(select(Location).where(Location.warehouse_id==warehouse_id,Location.is_active==True).order_by(Location.code)).scalars().all()
    skus=db.execute(select(SKU).where(SKU.client_id==client_id,SKU.is_active==True,SKU.barcode.is_not(None))).scalars().all()
    return {'client':{'id':str(client_id),'name':db.get(Client,client_id).name},
            'warehouse':{'id':str(warehouse_id),'name':db.get(Warehouse,warehouse_id).name},
            'locations':[{'id':str(x.id),'code':x.code,'barcode':x.barcode,'name':x.name,'kind':x.location_type} for x in locs],
            'catalog':local_catalog(db,client_id)}


class ReceiptPreview(BaseModel):
    client_id:uuid.UUID
    warehouse_id:uuid.UUID
    package:dict


class Confirm(BaseModel):
    client_id:uuid.UUID
    warehouse_id:uuid.UUID
    fingerprint:str=Field(min_length=64,max_length=64)
    confirm_occupied:bool=False


def receipt_seen(db,c,w,package):
    row=db.execute(text('''SELECT r.*,o.status operation_status FROM fc_terminal_receipts r
       JOIN inventory_operations o ON CAST(o.id AS varchar)=r.operation_id
       WHERE r.client_id=:c AND r.warehouse_id=:w AND r.package_id=:id'''),
       {**params(c,w),'id':package['packageId']}).mappings().first()
    if row and row['content_hash']!=fingerprint(package):
        raise HTTPException(409,detail='Этот номер приёмки уже проведён с другим составом. Создайте новую приёмку; повторное добавление заблокировано')
    if row and str(row.get('operation_status') or '').upper()=='CANCELLED':
        raise HTTPException(409,detail='Эта приёмка отменена в WMS. Повторный импорт её номера запрещён; создайте новую приёмку')
    return row


def receipt_preview(db,c,w,package,lock=False):
    if lock:
        for code in sorted({i['barcode'] for b in package['boxes'] for i in b['items']}):
            db.execute(text('SELECT pg_advisory_xact_lock(hashtext(:key))'), {'key':f'fc:local-sku:{c}:{code}'})
    locations={}; containers={}; skus={}; rows=[]; new_addresses={}; occupied={}; new_products={}
    for b in package['boxes']:
        code=b['locationCode']
        if code not in locations:
            stmt=select(Location).where(Location.warehouse_id==w,or_(Location.code==code,Location.barcode==code))
            matches=db.execute(stmt.with_for_update() if lock else stmt).scalars().all()
            if len(matches)>1: raise HTTPException(409,detail=f'Неоднозначный адрес {code}')
            loc=matches[0] if matches else None
            if loc and (not loc.is_active or loc.location_type not in ('CELL','PALLET','PALLET_PLACE')):
                raise HTTPException(409,detail=f'{code}: выберите действующую ячейку, палету или паллетоместо')
            if not loc and b['autoSection'] is None:
                raise HTTPException(409,detail=f'Адрес {code} не найден. Сначала добавьте его в адресном хранении')
            if loc:
                foreign=db.execute(select(InventoryBalance.id).where(InventoryBalance.location_id==loc.id,InventoryBalance.client_id!=c,InventoryBalance.quantity>0).limit(1)).first()
                if foreign: raise HTTPException(409,detail=f'Адрес {loc.code} занят товаром другого клиента')
            else:
                parent_code=f"SEC-{b['autoSection']}"
                parent=db.execute(select(Location).where(Location.warehouse_id==w,Location.code==parent_code)).scalar_one_or_none()
                if parent and (parent.location_type!='SECTION' or not parent.is_active):
                    raise HTTPException(409,detail=f'{parent_code}: секция недоступна')
                new_addresses[code]={'code':code,'section':b['autoSection'],'pallet':b['autoPallet']}
            locations[code]=loc
        loc=locations[code]
        sscc=b['sscc']; cont=None
        if sscc:
            if sscc not in containers:
                stmt=select(Container).where(or_(Container.barcode==sscc,
                    (Container.name==sscc)&(Container.client_id==c)&(Container.warehouse_id==w)))
                matches=db.execute(stmt.with_for_update() if lock else stmt).scalars().all()
                if len(matches)>1: raise HTTPException(409,detail=f'ШК/имя короба {sscc} неоднозначны. Отсканируйте технический ШК')
                cont=matches[0] if matches else None
                if cont:
                    if str(cont.client_id)!=str(c) or str(cont.warehouse_id)!=str(w):
                        raise HTTPException(409,detail=f'Короб {sscc} принадлежит другому клиенту или складу')
                    if cont.status!='ACTIVE': raise HTTPException(409,detail=f'Короб {sscc} архивный. Используйте действующий короб')
                    if cont.location_id and (not loc or str(cont.location_id)!=str(loc.id)):
                        raise HTTPException(409,detail=f'Короб {sscc} находится по другому адресу. Сначала переместите его в адресном хранении')
                    current=db.execute(select(func.coalesce(func.sum(InventoryBalance.quantity),0)).where(InventoryBalance.container_id==cont.id)).scalar_one()
                    if current>0: occupied[sscc]=float(current)
                containers[sscc]=cont
            cont=containers[sscc]
        for item in b['items']:
            bc=item['barcode']
            if bc not in skus:
                matches=db.execute(select(SKU).where(SKU.client_id==c,SKU.barcode==bc)).scalars().all()
                if len(matches)>1: raise HTTPException(409,detail=f'ШК {bc} привязан к нескольким товарам. Исправьте справочник WMS')
                sku=matches[0] if matches else None
                if sku and not sku.is_active: raise HTTPException(409,detail=f'Товар {bc} архивный')
                skus[bc]=sku
            sku=skus[bc]
            if sku:
                for attr in ('article','size'):
                    if item[attr] and getattr(sku,attr,None) and item[attr].strip().casefold()!=str(getattr(sku,attr)).strip().casefold():
                        raise HTTPException(409,detail=f'{bc}: артикул/размер терминала расходится со справочником WMS. Обновите товар в черновике')
            else:
                if not item['article'] or not item['size']: raise HTTPException(422,detail=f'Новый товар {bc}: укажите артикул и размер (для безразмерного — «Без размера»)')
                identity=(item['article'],item['size'])
                if bc in new_products and tuple(new_products[bc])!=identity: raise HTTPException(409,detail=f'{bc}: разные артикулы/размеры в одной приёмке')
                new_products[bc]=list(identity)
            rows.append({'barcode':bc,'article':(sku.article if sku else item['article']) or '',
                'size':(sku.size if sku else item['size']) or '', 'qty':item['qty'],'purpose':item['stockType'],
                'location_code':loc.code if loc else code,'location_id':str(loc.id) if loc else None,
                'container_name':sscc,'container_id':str(cont.id) if cont else None,
                'sku_id':str(sku.id) if sku else None,'row_on_pallet':b['row'] if sscc and (loc is None or loc.location_type in ('PALLET','PALLET_PLACE')) else None,
                'autoSection':b['autoSection'],'autoPallet':b['autoPallet']})
    return {'kind':'RECEIPT','package_id':package['packageId'],'rows':rows,
            'units':sum(r['qty'] for r in rows),'boxes':sum(bool(b['sscc']) for b in package['boxes']),
            'new_addresses':list(new_addresses.values()),'new_products':len(new_products),
            'occupied_boxes':[{'box':k,'quantity':v} for k,v in sorted(occupied.items())]}


def preview_package(db,c,w,raw,by):
    ensure_schema(db); scope(db,c,w)
    incoming=raw.get('wmsScope') or {}
    if incoming.get('client_id') and (incoming.get('client_id')!=str(c) or incoming.get('warehouse_id')!=str(w)):
        raise HTTPException(409,detail='Файл приёмки относится к другому клиенту или складу')
    package=parsed(normalize_package,raw)
    seen=receipt_seen(db,c,w,package)
    if seen:
        row=db.execute(text('SELECT * FROM fc_workflow_plans WHERE id=:id'),{'id':seen['plan_id']}).mappings().one()
        return plan_response(row)
    plan=receipt_preview(db,c,w,package)
    return save_plan(db,c,w,'RECEIPT',package,plan,by)


@router.post('/receipts/preview')
def preview_receipt(payload:ReceiptPreview,request:Request,db:Session=Depends(get_db)):
    return preview_package(db,payload.client_id,payload.warehouse_id,payload.package,actor(request))


@router.post('/receipts/import')
def import_receipt(request:Request,client_id:uuid.UUID=Form(...),warehouse_id:uuid.UUID=Form(...),file:UploadFile=File(...),db:Session=Depends(get_db)):
    raw=parsed(read_terminal_workbook,file.file.read(MAX_BYTES+1))
    return preview_package(db,client_id,warehouse_id,raw,actor(request))


def receipt_commit(db,c,w,package,plan,plan_id,by,confirm_occupied):
    seen=receipt_seen(db,c,w,package)
    if seen: return str(seen['operation_id'])
    fresh=receipt_preview(db,c,w,package,lock=True)
    if fingerprint(fresh)!=fingerprint(plan):
        raise HTTPException(409,detail='Адреса, короба или справочник изменились после проверки. Обновите предпросмотр приёмки')
    if plan['occupied_boxes'] and not confirm_occupied:
        raise HTTPException(409,detail='Подтвердите добавление товара в непустые короба')
    # All schema preparation happened before scope lock. No helper below commits.
    locs={}; boxes={}; skus={}; stock=defaultdict(int); before={}; purposes=defaultdict(set)
    for r in plan['rows']:
        lc=r['location_code']
        if lc not in locs:
            loc=db.get(Location,uuid.UUID(r['location_id'])) if r['location_id'] else None
            if not loc:
                section_code=f"SEC-{r['autoSection']}"
                section=db.execute(select(Location).where(Location.warehouse_id==w,Location.code==section_code)).scalar_one_or_none()
                if not section:
                    section=Location(id=uuid.uuid4(),warehouse_id=w,parent_id=None,code=section_code,
                      barcode=_location_barcode(w,section_code),name=f"Секция {r['autoSection']}",location_type='SECTION',is_active=True)
                    db.add(section); db.flush()
                loc=Location(id=uuid.uuid4(),warehouse_id=w,parent_id=section.id,code=lc,
                    barcode=_location_barcode(w,lc),name=f"Секция {r['autoSection']} · Паллетоместо {r['autoPallet']}",location_type='PALLET_PLACE',is_active=True)
                db.add(loc); db.flush()
            locs[lc]=loc
        loc=locs[lc]
        code=r['container_name']; cont=None
        if code:
            if code not in boxes:
                cont=db.get(Container,uuid.UUID(r['container_id'])) if r['container_id'] else None
                if not cont:
                    cont=Container(id=uuid.uuid4(),client_id=c,warehouse_id=w,location_id=loc.id,barcode=code,name=code,container_type='BOX',status='ACTIVE')
                    db.add(cont); db.flush()
                elif not cont.location_id: cont.location_id=loc.id; db.flush()
                boxes[code]=cont
            cont=boxes[code]
        bc=r['barcode']
        if bc not in skus:
            sku=db.get(SKU,uuid.UUID(r['sku_id'])) if r['sku_id'] else None
            if not sku:
                sku=SKU(id=uuid.uuid4(),client_id=c,sku_code=f'WB-{bc}',barcode=bc,article=r['article'],size=r['size'],name=r['article'],unit='PCS',is_active=True)
                db.add(sku); db.flush()
            skus[bc]=sku
        key=(skus[bc].id,loc.id,cont.id if cont else None)
        stock[key]+=r['qty']; purposes[key].add(r['purpose'])
    for code,cont in boxes.items():
        side=next((r.get("row_on_pallet") for r in plan["rows"] if r["container_name"]==code),None)
        if side is not None:
            set_box_route(db,container_id=cont.id,client_id=c,warehouse_id=w,row_on_pallet=side,actor=by)
    lines=[]
    for key,qty in sorted(stock.items(),key=lambda x:tuple(str(v or '') for v in x[0])):
        sid,lid,cid=key
        b=reconcile_stock_bucket(db,client_id=c,warehouse_id=w,sku_id=sid,location_id=lid,container_id=cid,preferred_bucket='FBS_AVAILABLE',lock=True)
        before[key]=Decimal(b['physical'])
        lines.append(StockLine(sku_id=sid,quantity=Decimal(qty),destination_location_id=lid,destination_container_id=cid,
            reason='TERMINAL_MAIN:'+','.join(sorted(purposes[key]))))
    op=execute_operation(db,OperationCreate(client_id=c,warehouse_id=w,operation_type='RECEIVING',
       idempotency_key='terminal-'+fingerprint({'warehouse':str(w),'package':package['packageId']})[:40],performed_by=by,reference=package['packageId'],
       comment='Ручная приёмка · ТЕРМИНАЛ · ЗОНА:MAIN',lines=lines),commit=False)
    db.flush()
    for key,qty in stock.items():
        sid,lid,cid=key
        b=reconcile_stock_bucket(db,client_id=c,warehouse_id=w,sku_id=sid,location_id=lid,container_id=cid,preferred_bucket='STORAGE_ONLY',lock=True)
        if Decimal(b['physical'])!=before[key]+qty:
            raise HTTPException(409,detail='Контроль остатка приёмки не пройден. Операция отменена')
    db.execute(text('''INSERT INTO fc_terminal_receipts(client_id,warehouse_id,package_id,content_hash,operation_id,plan_id)
        VALUES(:c,:w,:package,:hash,:op,:plan)'''),{**params(c,w),'package':package['packageId'],
        'hash':fingerprint(package),'op':str(op.id),'plan':plan_id})
    return str(op.id)


@router.get('/transfer/template.xlsx')
def transfer_template():
    wb=Workbook(); ws=wb.active; ws.title='Перевод в FBS'
    ws.append(['ШК товара','Количество','Ячейка','Короб'])
    for row in range(2,502): ws.cell(row,1).number_format='@'
    help=wb.create_sheet('Как заполнить')
    for line in ['ШК товара и Количество обязательны. Количество — целое число больше нуля.',
        'Ячейка и Короб необязательны. Пусто — WMS предложит места сама.',
        'Ячейка: точный код адреса или QR. Короб: точный ШК или уникальное имя.',
        'Если указан только короб, WMS найдёт его адрес. Адреса вне выбранного склада не используются.',
        'Строки с указанными адресами обрабатываются первыми; остаток не используется дважды.',
        'Не удаляйте ведущие нули ШК. Не используйте формулы. Сохраните как .xlsx.',
        'Это перевод учёта MAIN → FBS на том же физическом адресе. Физического списания нет.',
        'Проверьте предложенные места и подтвердите. При нехватке исправьте список и загрузите снова.']:
        help.append([line])
    return excel(wb,'FBS_transfer_template.xlsx')


def excel(wb,name):
    for ws in wb.worksheets:
        ws.freeze_panes='A2'
        for cell in ws[1]: cell.font=Font(bold=True); cell.fill=PatternFill('solid',fgColor='D9EAD3')
        for col in 'ABCDEFGHIJKL': ws.column_dimensions[col].width=23
        # External values must never be Excel formulas.
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value,str) and cell.value.startswith(('=','+','-','@')):
                    cell.data_type='s'
    out=io.BytesIO(); wb.save(out); out.seek(0)
    return StreamingResponse(out,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="{name}"'})


@router.post('/transfer/preview')
def preview_transfer(request:Request,client_id:uuid.UUID=Form(...),warehouse_id:uuid.UUID=Form(...),file:UploadFile=File(...),db:Session=Depends(get_db)):
    ensure_schema(db); scope(db,client_id,warehouse_id)
    requests=parsed(read_transfer_workbook,file.file.read(MAX_BYTES+1))
    buckets=route_enrich(db,client_id,warehouse_id,available_buckets(db,client_id,warehouse_id))
    # A barcode shared by several SKUs or a nonunique box name is not safe for automatic allocation.
    for req in requests:
        matches=[b for b in buckets if b['barcode']==req['barcode']]
        if len({b['sku_id'] for b in matches})>1:
            raise HTTPException(409,detail=f"ШК {req['barcode']} неоднозначен в WMS. Исправьте справочник")
        if req['container']:
            ids={b['container_id'] for b in matches if req['container'].casefold() in (b['container_name'].casefold(),b['container_barcode'].casefold())
                 and (not req['location'] or req['location'].casefold() in (b['location_code'].casefold(),b['location_barcode'].casefold()))}
            if len(ids)>1: raise HTTPException(409,detail=f"Короб {req['container']} неоднозначен. Укажите его ШК или ячейку")
    plan={'kind':'FBS',**allocate(requests,buckets,key=route_key),'file_name':str(file.filename or '')[:255]}
    return save_plan(db,client_id,warehouse_id,'FBS',{'requests':requests},plan,actor(request))


def revalidate_rows(db,c,w,rows):
    current=available_buckets(db,c,w)
    if any('route_row' in r for r in rows): current=route_enrich(db,c,w,current)
    available={b['key']:b for b in current}
    totals=defaultdict(int)
    for row in rows: totals[row['key']]+=row['qty']
    for key,qty in sorted(totals.items()):
        now=available.get(key)
        original=next(r for r in rows if r['key']==key)
        if not now or any(now.get(k)!=original.get(k) for k in (('barcode','location_code','container_barcode','route_section_id','route_row','route_position') if 'route_row' in original else ('barcode','location_code','container_barcode'))):
            raise HTTPException(409,detail='Адрес или товар из плана изменён. Загрузите список и проверьте новый план')
        sid=uuid.UUID(now['sku_id']); lid=uuid.UUID(now['location_id']); cid=uuid.UUID(now['container_id']) if now['container_id'] else None
        b=reconcile_stock_bucket(db,client_id=c,warehouse_id=w,sku_id=sid,location_id=lid,container_id=cid,preferred_bucket='FBS_AVAILABLE',lock=True)
        held=held_quantity(db,c,w,sid,lid,cid)
        free=max(Decimal(0),min(b['storage_only'],b['physical']-b['fbs_available']-b['reserved'])-held)
        if free<qty:
            raise HTTPException(409,detail=f"{now['location_code']} · {now['container_name'] or 'без короба'} · {now['barcode']}: нужно {qty}, доступно {free}. Остатки изменились; обновите план")
    return totals,available


@router.post('/plans/{plan_id}/confirm')
def confirm_plan(plan_id:uuid.UUID,payload:Confirm,request:Request,db:Session=Depends(get_db)):
    ensure_schema(db); scope(db,payload.client_id,payload.warehouse_id)
    c=payload.client_id; w=payload.warehouse_id
    try:
        lock_scope(db,c,w)
        row=db.execute(text('SELECT * FROM fc_workflow_plans WHERE id=:id AND client_id=:c AND warehouse_id=:w FOR UPDATE'),
                       {**params(c,w),'id':str(plan_id)}).mappings().first()
        if not row: raise HTTPException(404,detail='План не найден для выбранного клиента/склада')
        if row['content_hash']!=payload.fingerprint: raise HTTPException(409,detail='План изменился. Проверьте предпросмотр заново')
        if row['status']=='CONFIRMED':
            if row['kind']=='RECEIPT': receipt_seen(db,c,w,obj(row['payload_json']))
            return plan_response(row)
        plan=obj(row['plan_json']); saved=obj(row['payload_json']); op=None
        if row['kind']=='RECEIPT':
            op=receipt_commit(db,c,w,saved,plan,str(plan_id),actor(request),payload.confirm_occupied)
        elif row['kind']=='FBS':
            if plan['shortages']: raise HTTPException(409,detail='Есть нехватка товара. Исправьте Excel; частичный перевод автоматически не выполняется')
            totals,available=revalidate_rows(db,c,w,plan['rows'])
            for key,qty in sorted(totals.items()):
                b=available[key]
                move_bucket(db,client_id=c,warehouse_id=w,sku_id=uuid.UUID(b['sku_id']),location_id=uuid.UUID(b['location_id']),
                  container_id=uuid.UUID(b['container_id']) if b['container_id'] else None,source='STORAGE_ONLY',target='FBS_AVAILABLE',quantity=Decimal(qty))
        else: raise HTTPException(409,detail='Неизвестный вид плана')
        db.execute(text("UPDATE fc_workflow_plans SET status='CONFIRMED',confirmed_at=CURRENT_TIMESTAMP,actor=:actor,operation_id=:op WHERE id=:id"),
            {'id':str(plan_id),'actor':actor(request),'op':op})
        db.commit()
        return {**plan,'plan_id':str(plan_id),'fingerprint':payload.fingerprint,'status':'CONFIRMED','operation_id':op}
    except Exception:
        db.rollback(); raise


@router.get('/plans')
def list_plans(client_id:uuid.UUID,warehouse_id:uuid.UUID,db:Session=Depends(get_db)):
    ensure_schema(db); scope(db,client_id,warehouse_id)
    rows=db.execute(text('''SELECT id,kind,status,actor,created_at,confirmed_at,operation_id,task_id,
        plan_json->>'package_id' package_id,plan_json->>'file_name' file_name,
        COALESCE(plan_json->>'units',plan_json->>'allocated') units
        FROM fc_workflow_plans WHERE client_id=:c AND warehouse_id=:w
        ORDER BY created_at DESC LIMIT 100'''),params(client_id,warehouse_id)).mappings().all()
    result=[]
    for r in rows:
        item=dict(r)
        if item['kind']=='RECEIPT' and item.get('operation_id'):
            status=db.execute(text('SELECT status FROM inventory_operations WHERE CAST(id AS varchar)=:id'),{'id':item['operation_id']}).scalar()
            if status=='CANCELLED': item['status']='CANCELLED'
        result.append(item)
    return {'rows':result}


@router.get('/plans/{plan_id}')
def get_plan(plan_id:uuid.UUID,client_id:uuid.UUID,warehouse_id:uuid.UUID,db:Session=Depends(get_db)):
    ensure_schema(db)
    row=db.execute(text('SELECT * FROM fc_workflow_plans WHERE id=:id AND client_id=:c AND warehouse_id=:w'),
        {**params(client_id,warehouse_id),'id':str(plan_id)}).mappings().first()
    if not row: raise HTTPException(404,detail='План не найден')
    result=plan_response(row)
    if row['kind']=='RECEIPT' and row.get('operation_id'):
        status=db.execute(text('SELECT status FROM inventory_operations WHERE CAST(id AS varchar)=:id'),{'id':row['operation_id']}).scalar()
        if status=='CANCELLED': result['status']='CANCELLED'
    return result


@router.get('/plans/{plan_id}/report.xlsx')
def plan_report(plan_id:uuid.UUID,client_id:uuid.UUID,warehouse_id:uuid.UUID,db:Session=Depends(get_db)):
    plan=get_plan(plan_id,client_id,warehouse_id,db)
    wb=Workbook(); ws=wb.active; ws.title='Адреса и товары'
    ws.append(['Ячейка','Короб','ШК товара','Артикул','Размер','Количество','Строка Excel','Назначение'])
    for r in plan['rows']: ws.append([r.get('location_code'),r.get('container_name'),r['barcode'],r['article'],r['size'],r['qty'],r.get('excel_row'),r.get('purpose')])
    if plan.get('shortages'):
        ws=wb.create_sheet('Нехватка'); ws.append(['Строка','ШК товара','Не хватает','Ячейка','Короб'])
        for r in plan['shortages']: ws.append([r['row'],r['barcode'],r['missing'],r['location'],r['container']])
    ws=wb.create_sheet('Подтверждение'); ws.append(['Параметр','Значение'])
    for k in ('plan_id','status','operation_id','kind','package_id','units','requested','allocated','actor','created_at','confirmed_at'): ws.append([k,plan.get(k)])
    return excel(wb,'warehouse_plan.xlsx')


class ReceiptPick(BaseModel):
    client_id:uuid.UUID
    warehouse_id:uuid.UUID
    task_number:str=Field(min_length=1,max_length=120)


@router.post('/plans/{plan_id}/pick')
def receipt_pick(plan_id:uuid.UUID,payload:ReceiptPick,request:Request,db:Session=Depends(get_db)):
    """Create one exact-address pick from a posted receipt; repeated click is idempotent."""
    ensure_schema(db); c=payload.client_id; w=payload.warehouse_id; scope(db,c,w)
    try:
        lock_scope(db,c,w)
        row=db.execute(text('SELECT * FROM fc_workflow_plans WHERE id=:id AND client_id=:c AND warehouse_id=:w FOR UPDATE'),
            {**params(c,w),'id':str(plan_id)}).mappings().first()
        if not row or row['kind']!='RECEIPT' or row['status']!='CONFIRMED': raise HTTPException(409,detail='Сначала проведите приёмку')
        prior=db.execute(text('SELECT task_id FROM fc_workflow_plans WHERE client_id=:c AND warehouse_id=:w AND operation_id=:op AND task_id IS NOT NULL LIMIT 1'),{**params(c,w),'op':row['operation_id']}).scalar()
        if prior: return {'task_id':prior,'already_created':True}
        op=db.get(InventoryOperation,uuid.UUID(row['operation_id']))
        if not op or str(op.status).upper()=='CANCELLED': raise HTTPException(409,detail='Приёмка отменена')
        # Use actual canonical operation destinations, including IDs created at posting.
        from app.models import InventoryOperationLine
        lines=db.execute(select(InventoryOperationLine).where(InventoryOperationLine.operation_id==op.id)).scalars().all()
        buckets={b['key']:b for b in available_buckets(db,c,w)}; requested=defaultdict(int)
        for line in lines:
            key=':'.join(str(x or '') for x in (line.sku_id,line.destination_location_id,line.destination_container_id))
            requested[key]+=int(line.quantity)
        planned=[]
        for key,q in requested.items():
            if key not in buckets: raise HTTPException(409,detail='Товар приёмки уже перемещён/собран. Создайте сборку по Excel для текущих адресов')
            planned.append({**buckets[key],'qty':q})
        revalidate_rows(db,c,w,planned)
        task_id=str(uuid.uuid4()); task_number=payload.task_number.strip()
        if not task_number: raise HTTPException(422,detail='Введите номер сборки')
        duplicate=db.execute(text('SELECT id FROM fc_wh_pick_tasks WHERE client_id=:c AND warehouse_id=:w AND task_number=:number'),{**params(c,w),'number':task_number}).first()
        if duplicate: raise HTTPException(409,detail='Сборка с таким номером уже существует')
        db.execute(text('''INSERT INTO fc_wh_pick_tasks(id,client_id,warehouse_id,task_number,name,status,created_by,note)
          VALUES(:id,:c,:w,:number,:name,'DRAFT',:actor,:note)'''),{**params(c,w),'id':task_id,'number':task_number,
          'name':'Сборка после приёмки','actor':actor(request),'note':'Приёмка '+str(op.reference or op.id)})
        for ix,r in enumerate(planned):
            db.execute(text('''INSERT INTO fc_wh_pick_lines(id,task_id,sku_id,location_id,container_id,requested_qty,allocated_qty,sort_order,request_barcode,request_article,request_size)
                VALUES(:id,:task,:sku,:loc,:box,:qty,:qty,:sort,:barcode,:article,:size)'''),
                {'id':str(uuid.uuid4()),'task':task_id,'sku':r['sku_id'],'loc':r['location_id'],'box':r['container_id'],
                 'qty':r['qty'],'sort':ix,'barcode':r['barcode'],'article':r['article'],'size':r['size']})
        freeze_task(db,task_id)
        _audit(db,task_id=task_id,action='CREATE_FROM_RECEIPT',actor=actor(request),after={'operation_id':str(op.id),'units':sum(requested.values())},reason='Точные адреса проведённой приёмки')
        db.execute(text('UPDATE fc_workflow_plans SET task_id=:task WHERE id=:id'),{'task':task_id,'id':str(plan_id)})
        db.commit(); return {'task_id':task_id,'task_number':task_number}
    except Exception:
        db.rollback(); raise
