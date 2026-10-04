"""Transit warehouse with supplier/lot custody, history and explicit stock transfers.

FC_R145F12_TRANSIT_EXISTING_BOX_ATTACH
FC_R145F15_TRANSIT_LOCAL_PRODUCT_REGISTRY
FC_R145F16_TRANSIT_SCANNED_BOX_STORAGE_ONLY
New r145f6 batches are TRANSIT_ONLY: they bypass MAIN/FBS until an operator explicitly
transfers goods to MAIN/FBS or marks the physical departure. Legacy batches preserve the
r137 MAIN/FBS write-off behavior so existing active/sent records do not silently change.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import uuid
from datetime import datetime, timezone
from decimal import Decimal

from fastapi import HTTPException
from sqlalchemy import bindparam, func, select, text

from app.models import SKU
from app.services.local_catalog import catalog
from app.services.excel_safety import append_safe, workbook_bytes


BATCH_STATUSES = {"OPEN", "SENT", "CANCELLED"}
BOX_STATUSES = {"OPEN", "CLOSED", "SENT"}


def install_schema(db):
    db.execute(text("CREATE SEQUENCE IF NOT EXISTS fc_transit_box_code_seq START 1"))
    db.execute(text('''CREATE TABLE IF NOT EXISTS fc_transit_batches (
        id varchar(36) PRIMARY KEY,
        client_id varchar(36) NOT NULL,
        warehouse_id varchar(36) NOT NULL,
        number varchar(160) NOT NULL,
        partner varchar(200) NOT NULL,
        status varchar(16) NOT NULL DEFAULT 'OPEN' CHECK(status IN ('OPEN','SENT','CANCELLED')),
        created_by varchar(160) NOT NULL,
        created_at timestamptz NOT NULL DEFAULT now(),
        updated_at timestamptz NOT NULL DEFAULT now(),
        sent_at timestamptz NULL
    )'''))
    db.execute(text('''CREATE INDEX IF NOT EXISTS ix_fc_transit_batches_scope
        ON fc_transit_batches(client_id, warehouse_id, status, updated_at DESC)'''))
    # r137: additive metadata only. Existing transit batches remain untouched; legacy
    # SENT batches simply have NULL write-off fields and can be posted once from the UI.
    db.execute(text("ALTER TABLE fc_transit_batches ADD COLUMN IF NOT EXISTS stock_writeoff_at timestamptz NULL"))
    db.execute(text("ALTER TABLE fc_transit_batches ADD COLUMN IF NOT EXISTS stock_writeoff_summary jsonb NULL"))
    # r145f6: additive supplier/lot/departure metadata. Existing rows stay legacy.
    db.execute(text("ALTER TABLE fc_transit_batches ADD COLUMN IF NOT EXISTS supplier_id varchar(36) NULL"))
    db.execute(text("ALTER TABLE fc_transit_batches ADD COLUMN IF NOT EXISTS lot_number varchar(160) NOT NULL DEFAULT ''"))
    db.execute(text("ALTER TABLE fc_transit_batches ADD COLUMN IF NOT EXISTS destination varchar(200) NOT NULL DEFAULT ''"))
    db.execute(text("ALTER TABLE fc_transit_batches ADD COLUMN IF NOT EXISTS departure_note varchar(500) NOT NULL DEFAULT ''"))
    db.execute(text("ALTER TABLE fc_transit_batches ADD COLUMN IF NOT EXISTS accounting_mode varchar(24) NOT NULL DEFAULT 'LEGACY_MAIN_FBS'"))
    db.execute(text('''CREATE TABLE IF NOT EXISTS fc_transit_suppliers (
        id varchar(36) PRIMARY KEY,
        client_id varchar(36) NOT NULL,
        warehouse_id varchar(36) NOT NULL,
        name varchar(200) NOT NULL,
        note varchar(500) NOT NULL DEFAULT '',
        created_by varchar(160) NOT NULL,
        created_at timestamptz NOT NULL DEFAULT now(),
        updated_at timestamptz NOT NULL DEFAULT now()
    )'''))
    db.execute(text('''CREATE UNIQUE INDEX IF NOT EXISTS ux_fc_transit_suppliers_scope_name
        ON fc_transit_suppliers(client_id,warehouse_id,lower(name))'''))
    db.execute(text('''CREATE INDEX IF NOT EXISTS ix_fc_transit_batches_supplier
        ON fc_transit_batches(client_id,warehouse_id,supplier_id,status,updated_at DESC)'''))
    # r145f15: unknown scanned product barcodes live in an isolated transit-only registry.
    # They do not become MAIN/FBS SKU cards until an explicit Transit -> MAIN/FBS transfer.
    db.execute(text('''CREATE TABLE IF NOT EXISTS fc_transit_products (
        id varchar(36) PRIMARY KEY,
        client_id varchar(36) NOT NULL,
        barcode varchar(80) NOT NULL,
        barcode_norm varchar(80) NOT NULL,
        article varchar(120) NOT NULL DEFAULT '',
        size varchar(80) NOT NULL DEFAULT '',
        name varchar(200) NOT NULL DEFAULT '',
        created_by varchar(160) NOT NULL,
        first_batch_id varchar(36) NULL,
        first_box_id varchar(36) NULL,
        created_at timestamptz NOT NULL DEFAULT now(),
        updated_at timestamptz NOT NULL DEFAULT now(),
        UNIQUE(client_id, barcode_norm)
    )'''))
    db.execute(text('''CREATE INDEX IF NOT EXISTS ix_fc_transit_products_client_updated
        ON fc_transit_products(client_id, updated_at DESC)'''))
    db.execute(text('''CREATE TABLE IF NOT EXISTS fc_transit_boxes (
        id varchar(36) PRIMARY KEY,
        batch_id varchar(36) NOT NULL REFERENCES fc_transit_batches(id) ON DELETE CASCADE,
        barcode varchar(64) NOT NULL UNIQUE,
        sequence integer NOT NULL,
        status varchar(16) NOT NULL DEFAULT 'OPEN' CHECK(status IN ('OPEN','CLOSED','SENT')),
        created_by varchar(160) NOT NULL,
        created_at timestamptz NOT NULL DEFAULT now(),
        updated_at timestamptz NOT NULL DEFAULT now(),
        UNIQUE(batch_id, sequence)
    )'''))
    db.execute(text('''CREATE INDEX IF NOT EXISTS ix_fc_transit_boxes_batch
        ON fc_transit_boxes(batch_id, sequence)'''))
    # r145f12: an existing physical FulfillCore box may be attached to transit by its QR/name.
    db.execute(text("ALTER TABLE fc_transit_boxes ADD COLUMN IF NOT EXISTS source_container_id varchar(36) NULL"))
    db.execute(text("CREATE INDEX IF NOT EXISTS ix_fc_transit_boxes_source_container ON fc_transit_boxes(source_container_id)"))
    db.execute(text('''CREATE TABLE IF NOT EXISTS fc_transit_box_items (
        id varchar(36) PRIMARY KEY,
        box_id varchar(36) NOT NULL REFERENCES fc_transit_boxes(id) ON DELETE CASCADE,
        item_key varchar(96) NOT NULL,
        sku_id varchar(36) NULL,
        barcode varchar(80) NULL,
        article varchar(120) NOT NULL DEFAULT '',
        size varchar(80) NOT NULL DEFAULT '',
        name varchar(200) NOT NULL DEFAULT '',
        source varchar(24) NOT NULL CHECK(source IN ('CATALOG','UNKNOWN_BARCODE','MANUAL')),
        quantity integer NOT NULL DEFAULT 0 CHECK(quantity >= 0),
        created_at timestamptz NOT NULL DEFAULT now(),
        updated_at timestamptz NOT NULL DEFAULT now(),
        UNIQUE(box_id, item_key)
    )'''))
    db.execute(text('''CREATE INDEX IF NOT EXISTS ix_fc_transit_items_box
        ON fc_transit_box_items(box_id, updated_at DESC)'''))
    db.execute(text('''CREATE TABLE IF NOT EXISTS fc_transit_events (
        id varchar(36) PRIMARY KEY,
        batch_id varchar(36) NOT NULL REFERENCES fc_transit_batches(id) ON DELETE CASCADE,
        box_id varchar(36) NULL REFERENCES fc_transit_boxes(id) ON DELETE CASCADE,
        action varchar(32) NOT NULL,
        payload jsonb NOT NULL DEFAULT '{}'::jsonb,
        actor varchar(160) NOT NULL,
        created_at timestamptz NOT NULL DEFAULT now()
    )'''))
    db.execute(text('''CREATE INDEX IF NOT EXISTS ix_fc_transit_events_batch
        ON fc_transit_events(batch_id, created_at DESC)'''))


def _clean(value, limit, label, required=False):
    value = re.sub(r"\s+", " ", str(value or "").strip())
    if required and not value:
        raise HTTPException(422, detail=f"Укажите {label}")
    if len(value) > limit:
        raise HTTPException(422, detail=f"{label.capitalize()} слишком длинный")
    return value


def _uuid(value, label="ID"):
    try:
        return str(uuid.UUID(str(value)))
    except Exception:
        raise HTTPException(422, detail=f"Некорректный {label}")


def _event(db, event_id, batch_id, box_id, action, payload, actor):
    event_id = _uuid(event_id, "ID события")
    exists = db.execute(text("SELECT 1 FROM fc_transit_events WHERE id=:id"), {"id": event_id}).first()
    if exists:
        return False
    db.execute(text('''INSERT INTO fc_transit_events(id,batch_id,box_id,action,payload,actor)
        VALUES(:id,:batch,:box,:action,CAST(:payload AS jsonb),:actor)'''), {
        "id": event_id, "batch": str(batch_id), "box": str(box_id) if box_id else None,
        "action": action, "payload": __import__('json').dumps(payload or {}, ensure_ascii=False, default=str),
        "actor": _clean(actor,160,"пользователя",True)
    })
    return True


def _batch(db, batch_id, client_id=None, warehouse_id=None, lock=False):
    sql = "SELECT * FROM fc_transit_batches WHERE id=:id"
    params = {"id": _uuid(batch_id, "ID транзитной партии")}
    if client_id is not None:
        sql += " AND client_id=:c"; params["c"] = str(client_id)
    if warehouse_id is not None:
        sql += " AND warehouse_id=:w"; params["w"] = str(warehouse_id)
    if lock:
        sql += " FOR UPDATE"
    row = db.execute(text(sql), params).mappings().first()
    if not row:
        raise HTTPException(404, detail="Транзитная партия не найдена у выбранного клиента и склада")
    return dict(row)


def _box(db, box_id=None, barcode=None, client_id=None, warehouse_id=None, lock=False):
    if not box_id and not barcode:
        raise HTTPException(422, detail="Укажите транзитный короб")
    where = "b.id=:value" if box_id else "lower(b.barcode)=lower(:value)"
    params = {"value": _uuid(box_id, "ID транзитного короба") if box_id else _clean(barcode,64,"ШК короба",True)}
    if client_id is not None:
        where += " AND a.client_id=:c"; params["c"] = str(client_id)
    if warehouse_id is not None:
        where += " AND a.warehouse_id=:w"; params["w"] = str(warehouse_id)
    sql = f'''SELECT b.*,a.client_id,a.warehouse_id,a.number batch_number,a.partner,a.status batch_status
        FROM fc_transit_boxes b JOIN fc_transit_batches a ON a.id=b.batch_id WHERE {where}'''
    if lock:
        sql += " FOR UPDATE OF b"
    row = db.execute(text(sql), params).mappings().first()
    if not row:
        raise HTTPException(404, detail="Транзитный короб не найден")
    return dict(row)




def _lock_box_scope(db, box_id, client_id, warehouse_id):
    # All mutations lock in one order: batch -> box. This avoids scan/send deadlocks.
    first=_box(db,box_id=box_id,client_id=client_id,warehouse_id=warehouse_id,lock=False)
    _batch(db,first["batch_id"],client_id,warehouse_id,lock=True)
    return _box(db,box_id=box_id,client_id=client_id,warehouse_id=warehouse_id,lock=True)

def _items(db, box_id):
    return [dict(r) for r in db.execute(text('''SELECT id,item_key,sku_id,barcode,article,size,name,source,quantity,
        created_at,updated_at FROM fc_transit_box_items WHERE box_id=:b AND quantity>0
        ORDER BY updated_at DESC, created_at DESC'''), {"b": str(box_id)}).mappings().all()]


def box_public(db, box):
    rows = _items(db, box["id"])
    return {
        "id": str(box["id"]), "batch_id": str(box["batch_id"]), "barcode": box["barcode"],
        "sequence": int(box["sequence"]), "status": box["status"],
        "source_container_id": str(box.get("source_container_id") or "") or None,
        "units": sum(int(x["quantity"]) for x in rows), "lines": len(rows), "items": rows,
        "created_at": str(box.get("created_at") or ""), "updated_at": str(box.get("updated_at") or "")
    }



def _supplier(db, supplier_id, client_id=None, warehouse_id=None):
    if not supplier_id:
        return None
    sql="SELECT * FROM fc_transit_suppliers WHERE id=:id"
    params={"id":_uuid(supplier_id,"ID поставщика")}
    if client_id is not None:
        sql+=" AND client_id=:c";params["c"]=str(client_id)
    if warehouse_id is not None:
        sql+=" AND warehouse_id=:w";params["w"]=str(warehouse_id)
    row=db.execute(text(sql),params).mappings().first()
    return dict(row) if row else None


def _ensure_supplier(db, client_id, warehouse_id, name, actor):
    name=_clean(name,200,"поставщика",True)
    row=db.execute(text('''SELECT * FROM fc_transit_suppliers
        WHERE client_id=:c AND warehouse_id=:w AND lower(name)=lower(:name)'''),
        {"c":str(client_id),"w":str(warehouse_id),"name":name}).mappings().first()
    if row:
        return dict(row)
    sid=str(uuid.uuid4())
    db.execute(text('''INSERT INTO fc_transit_suppliers(id,client_id,warehouse_id,name,created_by)
        VALUES(:id,:c,:w,:name,:actor) ON CONFLICT DO NOTHING'''),
        {"id":sid,"c":str(client_id),"w":str(warehouse_id),"name":name,"actor":_clean(actor,160,"пользователя",True)})
    row=db.execute(text('''SELECT * FROM fc_transit_suppliers
        WHERE client_id=:c AND warehouse_id=:w AND lower(name)=lower(:name)'''),
        {"c":str(client_id),"w":str(warehouse_id),"name":name}).mappings().first()
    if not row:
        raise HTTPException(409,detail="Не удалось создать поставщика. Повторите операцию")
    return dict(row)


def list_suppliers(db, client_id, warehouse_id, search="", limit=200):
    q=_clean(search,200,"поиск")
    params={"c":str(client_id),"w":str(warehouse_id),"limit":max(1,min(int(limit),500))}
    extra=""
    if q:
        params["q"]="%"+q.casefold()+"%"
        extra=" AND lower(s.name) LIKE :q"
    rows=db.execute(text(f'''SELECT s.*,
        COUNT(DISTINCT a.id)::int total_batches,
        COUNT(DISTINCT a.id) FILTER (WHERE a.status='OPEN')::int active_batches,
        COUNT(DISTINCT a.id) FILTER (WHERE a.status='SENT')::int departed_batches,
        COALESCE(SUM(i.quantity) FILTER (WHERE a.status='OPEN'),0)::int active_units,
        COALESCE(SUM(i.quantity) FILTER (WHERE a.status='SENT'),0)::int departed_units,
        MAX(COALESCE(i.updated_at,a.updated_at,s.updated_at)) last_activity
        FROM fc_transit_suppliers s
        LEFT JOIN fc_transit_batches a ON a.supplier_id=s.id
        LEFT JOIN fc_transit_boxes b ON b.batch_id=a.id
        LEFT JOIN fc_transit_box_items i ON i.box_id=b.id
        WHERE s.client_id=:c AND s.warehouse_id=:w {extra}
        GROUP BY s.id ORDER BY last_activity DESC NULLS LAST,s.name LIMIT :limit'''),params).mappings().all()
    return [{**dict(r),"id":str(r["id"]),"created_at":str(r["created_at"]),"updated_at":str(r["updated_at"]),
             "last_activity":str(r["last_activity"]) if r.get("last_activity") else None} for r in rows]


def supplier_detail(db, supplier_id, client_id, warehouse_id):
    supplier=_supplier(db,supplier_id,client_id,warehouse_id)
    if not supplier:
        raise HTTPException(404,detail="Поставщик не найден")
    batches=list_batches(db,client_id,warehouse_id,limit=300,status=None,search="",supplier_id=supplier["id"])
    products=[dict(r) for r in db.execute(text('''SELECT
        COALESCE(i.barcode,'') barcode,COALESCE(i.article,'') article,COALESCE(i.size,'') size,COALESCE(i.name,'') name,
        COALESCE(SUM(i.quantity) FILTER (WHERE a.status='OPEN'),0)::int active_units,
        COALESCE(SUM(i.quantity) FILTER (WHERE a.status='SENT'),0)::int departed_units,
        COALESCE(SUM(i.quantity),0)::int total_units,MIN(i.created_at) first_received,MAX(i.updated_at) last_received,
        COUNT(DISTINCT a.id)::int lots
        FROM fc_transit_batches a JOIN fc_transit_boxes b ON b.batch_id=a.id
        JOIN fc_transit_box_items i ON i.box_id=b.id
        WHERE a.client_id=:c AND a.warehouse_id=:w AND a.supplier_id=:s AND i.quantity>0
        GROUP BY COALESCE(i.barcode,''),COALESCE(i.article,''),COALESCE(i.size,''),COALESCE(i.name,'')
        ORDER BY last_received DESC'''),{"c":str(client_id),"w":str(warehouse_id),"s":str(supplier["id"])}).mappings().all()]
    product_map={}
    for row in products:
        row["active_units"]=int(row.get("active_units") or 0);row["departed_units"]=int(row.get("departed_units") or 0);row["total_units"]=int(row.get("total_units") or 0)
        row["first_received"]=str(row["first_received"]) if row.get("first_received") else None
        row["last_received"]=str(row["last_received"]) if row.get("last_received") else None
        key=(str(row.get("barcode") or "").casefold(),str(row.get("article") or "").casefold(),str(row.get("size") or "").casefold())
        row.update({"received_units":0,"from_main_units":0,"from_fbs_units":0,"to_main_units":0,"to_fbs_units":0,"deleted_correction_units":0})
        product_map[key]=row
    event_rows=db.execute(text("""SELECT e.created_at,e.action,e.payload FROM fc_transit_events e
        JOIN fc_transit_batches a ON a.id=e.batch_id
        WHERE a.client_id=:c AND a.warehouse_id=:w AND a.supplier_id=:s
          AND e.action IN ('SCAN','MANUAL_ADD','TRANSFER_FROM_MAIN','TRANSFER_FROM_FBS','TRANSFER_TO_MAIN','TRANSFER_TO_FBS','ADMIN_ITEM_REMOVE')
        ORDER BY e.created_at"""),{"c":str(client_id),"w":str(warehouse_id),"s":str(supplier["id"])}).mappings().all()
    for ev in event_rows:
        payload=ev.get("payload") or {};barcode=str(payload.get("barcode") or "");article=str(payload.get("article") or "");size=str(payload.get("size") or "");name=str(payload.get("name") or "")
        key=(barcode.casefold(),article.casefold(),size.casefold())
        row=product_map.get(key)
        if row is None:
            row={"barcode":barcode,"article":article,"size":size,"name":name,"active_units":0,"departed_units":0,"total_units":0,"lots":0,
                 "first_received":None,"last_received":None,"received_units":0,"from_main_units":0,"from_fbs_units":0,"to_main_units":0,"to_fbs_units":0,"deleted_correction_units":0}
            product_map[key]=row;products.append(row)
        qty=int(payload.get("quantity") or 1)
        action=str(ev.get("action") or "")
        when=str(ev.get("created_at") or "")
        if action in {"SCAN","MANUAL_ADD"}:
            row["received_units"]+=qty
            if not row.get("first_received"):row["first_received"]=when
            row["last_received"]=when
        elif action=="TRANSFER_FROM_MAIN":row["from_main_units"]+=qty
        elif action=="TRANSFER_FROM_FBS":row["from_fbs_units"]+=qty
        elif action=="TRANSFER_TO_MAIN":row["to_main_units"]+=qty
        elif action=="TRANSFER_TO_FBS":row["to_fbs_units"]+=qty
        elif action=="ADMIN_ITEM_REMOVE":row["deleted_correction_units"]+=qty
    products.sort(key=lambda x:str(x.get("last_received") or ""),reverse=True)
    return {"supplier":{**supplier,"id":str(supplier["id"]),"created_at":str(supplier["created_at"]),"updated_at":str(supplier["updated_at"])},
            "batches":batches,"products":products,
            "active_units":sum(int(x.get("units") or 0) for x in batches if x.get("status")=="OPEN"),
            "departed_units":sum(int(x.get("units") or 0) for x in batches if x.get("status")=="SENT"),
            "received_units":sum(int(x.get("received_units") or 0) for x in products),
            "from_main_units":sum(int(x.get("from_main_units") or 0) for x in products),
            "from_fbs_units":sum(int(x.get("from_fbs_units") or 0) for x in products),
            "to_main_units":sum(int(x.get("to_main_units") or 0) for x in products),
            "to_fbs_units":sum(int(x.get("to_fbs_units") or 0) for x in products)}


def batch_public(db, batch, include_items=False):
    boxes = [dict(r) for r in db.execute(text('''SELECT b.*,
        COALESCE(SUM(i.quantity),0)::int units, COUNT(i.id) FILTER (WHERE i.quantity>0)::int lines
        FROM fc_transit_boxes b LEFT JOIN fc_transit_box_items i ON i.box_id=b.id
        WHERE b.batch_id=:a GROUP BY b.id ORDER BY b.sequence'''), {"a": str(batch["id"])}).mappings().all()]
    if include_items:
        for b in boxes:
            b["items"] = _items(db, b["id"])
    supplier=_supplier(db,batch.get("supplier_id"),batch.get("client_id"),batch.get("warehouse_id")) if batch.get("supplier_id") else None
    mode=str(batch.get("accounting_mode") or "LEGACY_MAIN_FBS")
    writeoff_at=batch.get("stock_writeoff_at")
    writeoff_summary=batch.get("stock_writeoff_summary") or None
    destination=str(batch.get("destination") or batch.get("partner") or "")
    return {
        "id": str(batch["id"]), "client_id": str(batch["client_id"]), "warehouse_id": str(batch["warehouse_id"]),
        "number": batch["number"], "partner": batch.get("partner") or "", "destination":destination,
        "departure_note":batch.get("departure_note") or "", "lot_number":batch.get("lot_number") or batch.get("number") or "",
        "supplier_id":str(batch.get("supplier_id") or "") or None,"supplier_name":supplier.get("name") if supplier else "",
        "accounting_mode":mode,"transit_only":mode=="TRANSIT_ONLY","status": batch["status"],
        "created_by": batch["created_by"], "created_at": str(batch["created_at"]),
        "updated_at": str(batch["updated_at"]), "sent_at": str(batch["sent_at"]) if batch.get("sent_at") else None,
        "stock_writeoff_at": str(writeoff_at) if writeoff_at else None,
        "stock_writeoff_summary": writeoff_summary,
        "stock_writeoff_done": bool(writeoff_at),
        "needs_stock_writeoff": mode=="LEGACY_MAIN_FBS" and batch["status"]=="SENT" and not bool(writeoff_at),
        "boxes": boxes, "box_count": len(boxes), "units": sum(int(x["units"]) for x in boxes),
        "open_boxes": sum(1 for x in boxes if x["status"] == "OPEN")
    }



def list_batches(db, client_id, warehouse_id, limit=100, status=None, search="", supplier_id=None):
    params = {"c": str(client_id), "w": str(warehouse_id), "limit": max(1, min(int(limit), 300))}
    extra = ""
    if status:
        status = str(status).upper()
        if status not in BATCH_STATUSES:
            raise HTTPException(422, detail="Некорректный статус транзитной партии")
        extra += " AND a.status=:status"; params["status"] = status
    if supplier_id:
        extra += " AND a.supplier_id=:supplier";params["supplier"]=_uuid(supplier_id,"ID поставщика")
    q=_clean(search,200,"поиск")
    if q:
        params["q"]="%"+q.casefold()+"%"
        extra += ''' AND (lower(a.number) LIKE :q OR lower(COALESCE(a.lot_number,'')) LIKE :q
            OR lower(COALESCE(a.destination,'')) LIKE :q OR lower(COALESCE(a.partner,'')) LIKE :q
            OR lower(COALESCE(s.name,'')) LIKE :q OR EXISTS (
                SELECT 1 FROM fc_transit_boxes bx JOIN fc_transit_box_items ix ON ix.box_id=bx.id
                WHERE bx.batch_id=a.id AND (lower(COALESCE(ix.barcode,'')) LIKE :q OR lower(COALESCE(ix.article,'')) LIKE :q
                OR lower(COALESCE(ix.size,'')) LIKE :q OR lower(COALESCE(ix.name,'')) LIKE :q)))'''
    rows = db.execute(text(f'''SELECT a.*,s.name supplier_name,
        COUNT(DISTINCT b.id)::int box_count,
        COALESCE(SUM(i.quantity),0)::int units,
        COUNT(DISTINCT b.id) FILTER (WHERE b.status='OPEN')::int open_boxes
        FROM fc_transit_batches a
        LEFT JOIN fc_transit_suppliers s ON s.id=a.supplier_id
        LEFT JOIN fc_transit_boxes b ON b.batch_id=a.id
        LEFT JOIN fc_transit_box_items i ON i.box_id=b.id
        WHERE a.client_id=:c AND a.warehouse_id=:w {extra}
        GROUP BY a.id,s.name ORDER BY COALESCE(a.sent_at,a.updated_at) DESC LIMIT :limit'''), params).mappings().all()
    result=[]
    for r in rows:
        mode=str(r.get("accounting_mode") or "LEGACY_MAIN_FBS")
        result.append({**dict(r), "id": str(r["id"]), "client_id": str(r["client_id"]), "warehouse_id": str(r["warehouse_id"]),
             "supplier_id":str(r.get("supplier_id") or "") or None,"supplier_name":r.get("supplier_name") or "",
             "lot_number":r.get("lot_number") or r.get("number") or "","destination":r.get("destination") or r.get("partner") or "",
             "accounting_mode":mode,"transit_only":mode=="TRANSIT_ONLY",
             "created_at": str(r["created_at"]), "updated_at": str(r["updated_at"]),
             "sent_at": str(r["sent_at"]) if r["sent_at"] else None,
             "stock_writeoff_at": str(r.get("stock_writeoff_at")) if r.get("stock_writeoff_at") else None,
             "stock_writeoff_done": bool(r.get("stock_writeoff_at")),
             "needs_stock_writeoff": mode=="LEGACY_MAIN_FBS" and r["status"]=="SENT" and not bool(r.get("stock_writeoff_at"))})
    return result



def create_batch(db, client_id, warehouse_id, number, partner, actor, supplier="", lot_number=""):
    supplier=_clean(supplier,200,"поставщика")
    lot_number=_clean(lot_number,160,"номер партии")
    modern=bool(supplier or lot_number)
    if modern:
        supplier=_clean(supplier,200,"поставщика",True)
        lot_number=_clean(lot_number,160,"номер партии",True)
        sup=_ensure_supplier(db,client_id,warehouse_id,supplier,actor)
        number=_clean(number,160,"название партии") or lot_number
        partner=""
        mode="TRANSIT_ONLY"
        supplier_id=sup["id"]
    else:
        number = _clean(number,160,"номер / название партии")
        if not number:
            number = "Транзит " + datetime.now().strftime("%d.%m.%Y %H:%M")
        partner = _clean(partner,200,"партнёра / склад назначения",True)
        mode="LEGACY_MAIN_FBS"
        supplier_id=None
    row = {
        "id": str(uuid.uuid4()), "c": str(client_id), "w": str(warehouse_id), "number": number,
        "partner": partner, "actor": _clean(actor,160,"пользователя",True),"supplier":supplier_id,
        "lot":lot_number,"mode":mode
    }
    db.execute(text('''INSERT INTO fc_transit_batches(id,client_id,warehouse_id,number,partner,created_by,supplier_id,lot_number,accounting_mode)
        VALUES(:id,:c,:w,:number,:partner,:actor,:supplier,:lot,:mode)'''), row)
    _event(db, uuid.uuid4(), row["id"], None, "BATCH_CREATE", {"number":number,"partner":partner,"supplier":supplier,"supplier_id":supplier_id,"lot_number":lot_number,"accounting_mode":mode}, actor)
    return batch_public(db, _batch(db,row["id"],client_id,warehouse_id))


def _new_box_code(db):
    seq_no = int(db.execute(text("SELECT nextval('fc_transit_box_code_seq')")).scalar_one())
    return f"TR-{datetime.now(timezone.utc).strftime('%y%m%d')}-{seq_no:06d}"



def create_boxes(db, batch_id, client_id, warehouse_id, count, actor):
    count = int(count)
    if count < 1 or count > 100:
        raise HTTPException(422, detail="За один раз можно создать от 1 до 100 коробов")
    batch = _batch(db,batch_id,client_id,warehouse_id,lock=True)
    if batch["status"] != "OPEN":
        raise HTTPException(409, detail="Короба можно создавать только в открытой транзитной партии")
    start = db.execute(text("SELECT COALESCE(MAX(sequence),0) FROM fc_transit_boxes WHERE batch_id=:a"), {"a":batch["id"]}).scalar_one()
    created=[]
    for i in range(count):
        code = _new_box_code(db)
        box_id=str(uuid.uuid4()); order=int(start)+i+1
        db.execute(text('''INSERT INTO fc_transit_boxes(id,batch_id,barcode,sequence,created_by)
            VALUES(:id,:a,:code,:sequence,:actor)'''), {"id":box_id,"a":batch["id"],"code":code,"sequence":order,"actor":actor})
        _event(db,uuid.uuid4(),batch["id"],box_id,"BOX_CREATE",{"barcode":code,"sequence":order},actor)
        created.append({"id":box_id,"barcode":code,"sequence":order,"status":"OPEN","units":0,"lines":0,"items":[]})
    db.execute(text("UPDATE fc_transit_batches SET updated_at=now() WHERE id=:id"), {"id":batch["id"]})
    return {"created":created,"batch":batch_public(db,_batch(db,batch["id"],client_id,warehouse_id))}


def attach_existing_box(db, batch_id, client_id, warehouse_id, raw_code, actor):
    """Open or create a scanned transit box without changing MAIN/FBS stock."""
    code=_clean(raw_code,64,"ШК / название короба",True)
    batch=_batch(db,batch_id,client_id,warehouse_id,lock=True)
    if batch["status"]!="OPEN":
        raise HTTPException(409,detail="Короб можно открыть только в активной транзитной партии")

    # Reuse this exact transit box only inside the same batch. Any other transit record,
    # including historical/sent, blocks reusing the physical code.
    current=db.execute(text("""SELECT b.*,a.client_id,a.warehouse_id,a.number batch_number,a.partner,a.status batch_status
        FROM fc_transit_boxes b JOIN fc_transit_batches a ON a.id=b.batch_id
        WHERE lower(b.barcode)=lower(:code)"""),{"code":code}).mappings().first()
    if current:
        current=dict(current)
        if str(current["batch_id"])!=str(batch["id"]):
            raise HTTPException(409,detail=f"Короб {code} уже есть в транзитной истории / другой партии: {current.get('batch_number') or '—'}")
        result=box_public(db,current)
        result["attached_existing"]=bool(current.get("source_container_id"))
        result["transit_only_created"]=False
        result["reused"]=True
        result["stock_effect"]=False
        return result

    # Look for an ordinary FulfillCore container anywhere. Only an ACTIVE empty box from
    # the exact selected client+warehouse may be attached; every other collision blocks.
    row=db.execute(text("""SELECT CAST(c.id AS text) id,CAST(c.client_id AS text) client_id,
        CAST(c.warehouse_id AS text) container_warehouse_id,
        COALESCE(to_jsonb(c)->>'name','') name,COALESCE(to_jsonb(c)->>'barcode','') barcode,
        COALESCE(to_jsonb(c)->>'status','ACTIVE') status,COALESCE(to_jsonb(c)->>'container_type','') container_type,
        CAST(c.location_id AS text) location_id,
        COALESCE(SUM(CASE WHEN ib.client_id=c.client_id AND ib.warehouse_id=c.warehouse_id THEN ib.quantity ELSE 0 END),0) quantity
        FROM containers c LEFT JOIN inventory_balances ib ON ib.container_id=c.id
        WHERE lower(COALESCE(to_jsonb(c)->>'barcode',''))=lower(:code)
           OR lower(COALESCE(to_jsonb(c)->>'name',''))=lower(:code)
        GROUP BY c.id,c.client_id,c.warehouse_id,c.location_id
        ORDER BY CASE WHEN CAST(c.client_id AS text)=:c AND CAST(c.warehouse_id AS text)=:w THEN 0 ELSE 1 END LIMIT 1"""),
        {"c":str(client_id),"w":str(warehouse_id),"code":code}).mappings().first()
    if row:
        row=dict(row)
        if str(row.get("client_id"))!=str(client_id) or str(row.get("container_warehouse_id"))!=str(warehouse_id):
            raise HTTPException(409,detail=f"Код {code} уже используется коробом в другом клиенте/складе. Новый транзитный короб не создан")
        if str(row.get("container_type") or "").upper()=="DEFECT_BOX" or str(row.get("status") or "").upper().startswith("DEFECT"):
            raise HTTPException(409,detail="Короб брака нельзя использовать как транзитный")
        if str(row.get("status") or "ACTIVE").upper()!="ACTIVE":
            raise HTTPException(409,detail=f"Короб {code} не активен в системе")
        if Decimal(str(row.get("quantity") or 0))>0:
            raise HTTPException(409,detail=f"В коробе {code} уже есть складской остаток. Сначала освободите короб, затем принимайте его в транзит")
        occupied=db.execute(text("""SELECT b.id,a.number FROM fc_transit_boxes b JOIN fc_transit_batches a ON a.id=b.batch_id
            WHERE b.source_container_id=:cid AND a.status='OPEN' LIMIT 1"""),{"cid":row["id"]}).mappings().first()
        if occupied:
            raise HTTPException(409,detail=f"Этот физический короб уже используется в активной транзитной партии: {occupied['number']}")
        seq=int(db.execute(text("SELECT COALESCE(MAX(sequence),0)+1 FROM fc_transit_boxes WHERE batch_id=:a"),{"a":batch["id"]}).scalar_one())
        box_id=str(uuid.uuid4())
        db.execute(text("""INSERT INTO fc_transit_boxes(id,batch_id,barcode,sequence,created_by,source_container_id)
            VALUES(:id,:a,:code,:sequence,:actor,:source)"""),{"id":box_id,"a":batch["id"],"code":code,"sequence":seq,
            "actor":_clean(actor,160,"пользователя",True),"source":row["id"]})
        _event(db,uuid.uuid4(),batch["id"],box_id,"BOX_ATTACH_EXISTING",{"barcode":code,"source_container_id":row["id"],"location_id":row.get("location_id")},actor)
        db.execute(text("UPDATE fc_transit_batches SET updated_at=now() WHERE id=:id"),{"id":batch["id"]})
        result=box_public(db,_box(db,box_id=box_id,client_id=client_id,warehouse_id=warehouse_id))
        result["attached_existing"]=True
        result["transit_only_created"]=False
        result["reused"]=False
        result["stock_effect"]=False
        return result

    # A new transit box code must not be an existing cell/pallet/location identifier.
    location=db.execute(text("""SELECT CAST(l.id AS text) id,
        COALESCE(NULLIF(to_jsonb(l)->>'code',''),NULLIF(to_jsonb(l)->>'name',''),NULLIF(to_jsonb(l)->>'barcode',''),CAST(l.id AS text)) label
        FROM locations l
        WHERE lower(COALESCE(to_jsonb(l)->>'code',''))=lower(:code)
           OR lower(COALESCE(to_jsonb(l)->>'name',''))=lower(:code)
           OR lower(COALESCE(to_jsonb(l)->>'barcode',''))=lower(:code)
        LIMIT 1"""),{"code":code}).mappings().first()
    if location:
        raise HTTPException(409,detail=f"Код {code} уже используется как ячейка/палета. Новый транзитный короб не создан")

    # Also reserve every code that ever appeared in transit history, even if an old row
    # was later removed administratively.
    history=db.execute(text("""SELECT a.number,e.action,e.created_at
        FROM fc_transit_events e JOIN fc_transit_batches a ON a.id=e.batch_id
        WHERE lower(COALESCE(e.payload->>'barcode',''))=lower(:code)
           OR lower(COALESCE(e.payload->>'new_barcode',''))=lower(:code)
           OR lower(COALESCE(e.payload->>'old_barcode',''))=lower(:code)
        ORDER BY e.created_at DESC LIMIT 1"""),{"code":code}).mappings().first()
    if history:
        raise HTTPException(409,detail=f"Код {code} уже встречался в истории транзита (партия {history.get('number') or '—'}). Повторно использовать его нельзя")

    # Completely unknown physical box: create ONLY a transit record with the exact scan.
    # No row in containers and no inventory_balances/fc_stocks_v2 mutation is performed.
    seq=int(db.execute(text("SELECT COALESCE(MAX(sequence),0)+1 FROM fc_transit_boxes WHERE batch_id=:a"),{"a":batch["id"]}).scalar_one())
    box_id=str(uuid.uuid4())
    db.execute(text("""INSERT INTO fc_transit_boxes(id,batch_id,barcode,sequence,created_by,source_container_id)
        VALUES(:id,:a,:code,:sequence,:actor,NULL)"""),{"id":box_id,"a":batch["id"],"code":code,"sequence":seq,
        "actor":_clean(actor,160,"пользователя",True)})
    _event(db,uuid.uuid4(),batch["id"],box_id,"BOX_CREATE_SCANNED_TRANSIT_ONLY",{
        "barcode":code,"sequence":seq,"storage_only":True,"inventory_effect":False
    },actor)
    db.execute(text("UPDATE fc_transit_batches SET updated_at=now() WHERE id=:id"),{"id":batch["id"]})
    result=box_public(db,_box(db,box_id=box_id,client_id=client_id,warehouse_id=warehouse_id))
    result["attached_existing"]=False
    result["transit_only_created"]=True
    result["reused"]=False
    result["stock_effect"]=False
    return result

# FC_R145F18_TRANSIT_EMPTY_BOX_DELETE_TSD_DESKTOP
def delete_empty_box(db, box_id, client_id, warehouse_id, event_id, actor):
    """Delete/detach an erroneous EMPTY transit box without touching MAIN/FBS.

    For a box attached from the ordinary warehouse only the transit link is removed;
    the original containers row remains untouched.  A history event with box_id=NULL
    preserves the scanned barcode after the box row is deleted.
    """
    first=_box(db,box_id=box_id,client_id=client_id,warehouse_id=warehouse_id,lock=False)
    batch=_batch(db,first["batch_id"],client_id,warehouse_id,lock=True)
    box=_box(db,box_id=box_id,client_id=client_id,warehouse_id=warehouse_id,lock=True)
    if batch["status"]!="OPEN":
        raise HTTPException(409,detail="Удалять короб можно только из активной транзитной партии")
    units=int(db.execute(text("SELECT COALESCE(SUM(quantity),0) FROM fc_transit_box_items WHERE box_id=:b"),{"b":str(box["id"])}).scalar() or 0)
    if units>0:
        raise HTTPException(409,detail=f"В коробе {box['barcode']} есть товар ({units} шт.). Сначала удалите ошибочные позиции или переведите товар, затем удаляйте короб")
    eid=_uuid(event_id,"ID операции")
    if db.execute(text("SELECT 1 FROM fc_transit_events WHERE id=:id"),{"id":eid}).first():
        return batch_public(db,_batch(db,batch["id"],client_id,warehouse_id),include_items=True)
    payload={"barcode":box.get("barcode"),"sequence":box.get("sequence"),
             "source_container_id":box.get("source_container_id"),"inventory_effect":False}
    # Insert history with box_id NULL so ON DELETE CASCADE cannot erase this audit event.
    _event(db,eid,batch["id"],None,"BOX_DELETE_EMPTY",payload,actor)
    db.execute(text("DELETE FROM fc_transit_boxes WHERE id=:id"),{"id":str(box["id"])})
    db.execute(text("UPDATE fc_transit_batches SET updated_at=now() WHERE id=:id"),{"id":batch["id"]})
    result=batch_public(db,_batch(db,batch["id"],client_id,warehouse_id),include_items=True)
    result["deleted_box"]={"id":str(box["id"]),"barcode":box.get("barcode"),"inventory_effect":False}
    return result


def _transit_barcode(raw_barcode):
    code=str(raw_barcode or '').strip()
    if not code or len(code)>80 or re.search(r'[\s\x00-\x1f\x7f]',code):
        raise HTTPException(422,detail='Сканируйте ШК товара без пробелов, до 80 знаков')
    # Transit accepts supplier/local alphanumeric barcodes too, but never treats a
    # DataMatrix/ЧЗ code as a product identity.
    if code.startswith((']d2',']C1','(01)')) or (len(code)>25 and code.startswith('01')):
        raise HTTPException(422,detail='Нужен ШК товара, не DataMatrix / Честный знак')
    return code


def _existing_sku_identity(db, client_id, code):
    """Return one existing SKU even if it is inactive, so transit never creates a duplicate identity."""
    rows=db.execute(select(SKU).where(
        SKU.client_id==uuid.UUID(str(client_id)),
        func.lower(func.coalesce(SKU.barcode,''))==code.casefold(),
    )).scalars().all()
    if len(rows)>1:
        raise HTTPException(409,detail=f"ШК {code} уже привязан к нескольким SKU. Сначала исправьте карточки товара")
    if not rows:
        return None
    sku=rows[0]
    return {"item_key":f"SKU:{sku.id}","sku_id":str(sku.id),"barcode":code,
            "article":str(sku.article or ""),"size":str(sku.size or ""),"name":str(sku.name or ""),
            "source":"CATALOG","catalog_found":True,"catalog_inactive":not bool(getattr(sku,"is_active",True)),
            "warning":"Использована существующая локальная карточка товара"}


def _guard_transit_product_identifier(db, client_id, warehouse_id, code):
    """A product scan must not silently reuse a storage identifier (box/cell/pallet)."""
    box=db.execute(text('''SELECT COALESCE(name,'') name,COALESCE(barcode,'') barcode,COALESCE(status,'') status
        FROM containers WHERE client_id=:c
          AND (lower(COALESCE(barcode,''))=lower(:code) OR lower(COALESCE(name,''))=lower(:code))
        LIMIT 1'''),{"c":str(client_id),"code":code}).mappings().first()
    if box:
        label=box.get("name") or box.get("barcode") or code
        raise HTTPException(409,detail=f"Код {code} уже используется как короб {label}. Нельзя создавать товар транзита с ШК короба")
    loc=db.execute(text('''SELECT code FROM locations WHERE warehouse_id=:w
        AND lower(COALESCE(code,''))=lower(:code) LIMIT 1'''),{"w":str(warehouse_id),"code":code}).mappings().first()
    if loc:
        raise HTTPException(409,detail=f"Код {code} уже используется как ячейка/палета {loc['code']}. Нельзя создавать товар транзита с этим кодом")
    transit_box=db.execute(text('''SELECT b.barcode,a.number,a.status FROM fc_transit_boxes b
        JOIN fc_transit_batches a ON a.id=b.batch_id
        WHERE a.client_id=:c AND lower(b.barcode)=lower(:code)
        ORDER BY b.updated_at DESC LIMIT 1'''),{"c":str(client_id),"code":code}).mappings().first()
    if transit_box:
        raise HTTPException(409,detail=f"Код {code} уже использовался как транзитный короб в партии {transit_box.get('number') or '—'}")


def _legacy_transit_identity(db, client_id, code):
    """Find the same product barcode in active transit rows or immutable transit event history."""
    row=db.execute(text('''SELECT COALESCE(i.article,'') article,COALESCE(i.size,'') size,COALESCE(i.name,'') name,
        CAST(b.batch_id AS text) batch_id,CAST(i.box_id AS text) box_id
        FROM fc_transit_box_items i
        JOIN fc_transit_boxes b ON b.id=i.box_id
        JOIN fc_transit_batches a ON a.id=b.batch_id
        WHERE a.client_id=:c AND lower(COALESCE(i.barcode,''))=lower(:code)
        ORDER BY i.updated_at DESC LIMIT 1'''),{"c":str(client_id),"code":code}).mappings().first()
    if row:
        return dict(row)
    row=db.execute(text('''SELECT COALESCE(e.payload->>'article','') article,COALESCE(e.payload->>'size','') size,
        COALESCE(e.payload->>'name','') name,CAST(e.batch_id AS text) batch_id,CAST(e.box_id AS text) box_id
        FROM fc_transit_events e JOIN fc_transit_batches a ON a.id=e.batch_id
        WHERE a.client_id=:c AND lower(COALESCE(e.payload->>'barcode',''))=lower(:code)
        ORDER BY e.created_at DESC LIMIT 1'''),{"c":str(client_id),"code":code}).mappings().first()
    return dict(row) if row else None


def _transit_local_identity(db, client_id, warehouse_id, code, batch_id, box_id, actor):
    """Create/reuse a client-wide transit-only product identity without touching SKU/MAIN/FBS."""
    _guard_transit_product_identifier(db,client_id,warehouse_id,code)
    norm=code.casefold()
    existing=db.execute(text('''SELECT * FROM fc_transit_products
        WHERE client_id=:c AND barcode_norm=:n FOR UPDATE'''),{"c":str(client_id),"n":norm}).mappings().first()
    if existing:
        db.execute(text("UPDATE fc_transit_products SET updated_at=now() WHERE id=:id"),{"id":existing["id"]})
        row=dict(existing)
        return {"item_key":f"BARCODE:{code}","sku_id":None,"barcode":code,
                "article":str(row.get("article") or code),"size":str(row.get("size") or ""),
                "name":str(row.get("name") or code),"source":"UNKNOWN_BARCODE","catalog_found":False,
                "transit_local":True,"transit_product_id":str(row["id"]),"transit_local_created":False,
                "transit_local_reused":True,"warning":"Использована существующая карточка товара транзита"}
    legacy=_legacy_transit_identity(db,client_id,code)
    article=str((legacy or {}).get("article") or code)[:120]
    size=str((legacy or {}).get("size") or "")[:80]
    name=str((legacy or {}).get("name") or code)[:200]
    product_id=str(uuid.uuid4())
    first_batch=str((legacy or {}).get("batch_id") or batch_id or "") or None
    first_box=str((legacy or {}).get("box_id") or box_id or "") or None
    # The explicit normalized key makes creation race-safe and also prevents duplicate
    # transit-local product names for the same scanned identifier.
    row=db.execute(text('''INSERT INTO fc_transit_products(
            id,client_id,barcode,barcode_norm,article,size,name,created_by,first_batch_id,first_box_id)
        VALUES(:id,:c,:barcode,:norm,:article,:size,:name,:actor,:batch,:box)
        ON CONFLICT (client_id,barcode_norm) DO UPDATE SET updated_at=now()
        RETURNING *'''),{"id":product_id,"c":str(client_id),"barcode":code,"norm":norm,
        "article":article,"size":size,"name":name,"actor":_clean(actor,160,"пользователя",True),
        "batch":first_batch,"box":first_box}).mappings().one()
    created=str(row["id"])==product_id
    return {"item_key":f"BARCODE:{code}","sku_id":None,"barcode":code,
            "article":str(row.get("article") or code),"size":str(row.get("size") or ""),
            "name":str(row.get("name") or code),"source":"UNKNOWN_BARCODE","catalog_found":False,
            "transit_local":True,"transit_product_id":str(row["id"]),"transit_local_created":created,
            "transit_local_reused":not created or bool(legacy),"history_reused":bool(legacy),
            "warning":"Создан отдельный товар транзита" if created and not legacy else "Использована карточка товара транзита из истории"}


def _catalog_identity(db, client_id, raw_barcode, *, warehouse_id=None, batch_id=None, box_id=None, actor="", allow_transit_local=False):
    code=_transit_barcode(raw_barcode)
    rows = catalog(db, client_id, code)
    if rows:
        row = rows[0]
        if row.get("duplicate_sku") or row.get("wb_conflict"):
            raise HTTPException(409, detail=row.get("warning") or "ШК неоднозначен в локальных карточках")
        sku_ids = row.get("sku_ids") or []
        sku_id = sku_ids[0] if len(sku_ids)==1 else None
        identity = sku_id or f"BARCODE:{code}"
        return {"item_key":f"SKU:{sku_id}" if sku_id else identity,"sku_id":sku_id,"barcode":code,
                "article":str(row.get("article") or ""),"size":str(row.get("size") or ""),
                "name":str(row.get("name") or ""),"source":"CATALOG","catalog_found":True,
                "warning":str(row.get("warning") or "")}
    existing=_existing_sku_identity(db,client_id,code)
    if existing:
        return existing
    if allow_transit_local:
        if warehouse_id is None or batch_id is None or box_id is None:
            raise HTTPException(500,detail="Не передан контекст для создания товара транзита")
        return _transit_local_identity(db,client_id,warehouse_id,code,batch_id,box_id,actor)
    return {"item_key":f"BARCODE:{code}","sku_id":None,"barcode":code,"article":"","size":"","name":"",
            "source":"UNKNOWN_BARCODE","catalog_found":False,"warning":"ШК не найден в локальных карточках"}


def _upsert_item(db, box_id, product, delta=1):
    delta=int(delta)
    if delta == 0:
        return
    existing = db.execute(text("SELECT * FROM fc_transit_box_items WHERE box_id=:b AND item_key=:k FOR UPDATE"),
                          {"b":str(box_id),"k":product["item_key"]}).mappings().first()
    if existing:
        new_qty=int(existing["quantity"])+delta
        if new_qty < 0:
            raise HTTPException(409, detail="Количество в коробе не может быть отрицательным")
        if new_qty == 0:
            db.execute(text("DELETE FROM fc_transit_box_items WHERE id=:id"), {"id":existing["id"]})
        else:
            db.execute(text("UPDATE fc_transit_box_items SET quantity=:q,updated_at=now() WHERE id=:id"), {"q":new_qty,"id":existing["id"]})
        return
    if delta < 0:
        raise HTTPException(409, detail="Этой позиции уже нет в коробе")
    db.execute(text('''INSERT INTO fc_transit_box_items(id,box_id,item_key,sku_id,barcode,article,size,name,source,quantity)
        VALUES(:id,:b,:k,:sku,:barcode,:article,:size,:name,:source,:qty)'''), {
        "id":str(uuid.uuid4()),"b":str(box_id),"k":product["item_key"],"sku":product.get("sku_id"),
        "barcode":product.get("barcode"),"article":product.get("article","")[:120],"size":product.get("size","")[:80],
        "name":product.get("name","")[:200],"source":product["source"],"qty":delta})


def scan_product(db, box_id, client_id, warehouse_id, raw_barcode, event_id, actor):
    box=_lock_box_scope(db,box_id,client_id,warehouse_id)
    if box["batch_status"]!="OPEN" or box["status"]!="OPEN":
        raise HTTPException(409, detail="Сканирование доступно только в открытый транзитный короб")
    # Replay protection: exact event is returned without adding a second unit.
    if db.execute(text("SELECT 1 FROM fc_transit_events WHERE id=:id"), {"id":_uuid(event_id,"ID скана")}).first():
        result=box_public(db,box);result["replayed"]=True;return result
    product=_catalog_identity(db,client_id,raw_barcode,warehouse_id=warehouse_id,batch_id=box["batch_id"],box_id=box["id"],actor=actor,allow_transit_local=True)
    _upsert_item(db,box["id"],product,1)
    _event(db,event_id,box["batch_id"],box["id"],"SCAN",{k:v for k,v in product.items() if k!="item_key"},actor)
    db.execute(text("UPDATE fc_transit_boxes SET updated_at=now() WHERE id=:id"),{"id":box["id"]})
    db.execute(text("UPDATE fc_transit_batches SET updated_at=now() WHERE id=:id"),{"id":box["batch_id"]})
    result=box_public(db,_box(db,box_id=box["id"],client_id=client_id,warehouse_id=warehouse_id));result["scanned"]=product
    return result


def manual_product(db, box_id, client_id, warehouse_id, article, size, name, event_id, actor):
    box=_lock_box_scope(db,box_id,client_id,warehouse_id)
    if box["batch_status"]!="OPEN" or box["status"]!="OPEN":
        raise HTTPException(409, detail="Товар можно добавлять только в открытый транзитный короб")
    if db.execute(text("SELECT 1 FROM fc_transit_events WHERE id=:id"), {"id":_uuid(event_id,"ID операции")}).first():
        result=box_public(db,box);result["replayed"]=True;return result
    article=_clean(article,120,"артикул",True);size=_clean(size,80,"размер",True);name=_clean(name,200,"название")
    # Enrich from a unique local SKU by article + size without requiring a barcode.
    sku_rows=db.execute(select(SKU).where(SKU.client_id==uuid.UUID(str(client_id)),SKU.is_active.is_(True),
        func.lower(func.trim(SKU.article))==article.casefold(),func.lower(func.trim(SKU.size))==size.casefold())).scalars().all()
    sku_id=str(sku_rows[0].id) if len(sku_rows)==1 else None
    if len(sku_rows)==1 and not name:name=sku_rows[0].name or ""
    digest=hashlib.sha256((article.casefold()+"\x1f"+size.casefold()).encode()).hexdigest()[:48]
    product={"item_key":f"MANUAL:{digest}","sku_id":sku_id,"barcode":None,"article":article,"size":size,"name":name,"source":"MANUAL"}
    _upsert_item(db,box["id"],product,1)
    _event(db,event_id,box["batch_id"],box["id"],"MANUAL_ADD",{k:v for k,v in product.items() if k!="item_key"},actor)
    db.execute(text("UPDATE fc_transit_boxes SET updated_at=now() WHERE id=:id"),{"id":box["id"]})
    db.execute(text("UPDATE fc_transit_batches SET updated_at=now() WHERE id=:id"),{"id":box["batch_id"]})
    result=box_public(db,_box(db,box_id=box["id"],client_id=client_id,warehouse_id=warehouse_id));result["scanned"]=product
    return result


def set_item_quantity(db, box_id, item_id, client_id, warehouse_id, quantity, event_id, actor):
    box=_lock_box_scope(db,box_id,client_id,warehouse_id)
    if box["batch_status"]!="OPEN" or box["status"]!="OPEN":
        raise HTTPException(409,detail="Редактирование доступно только в открытом транзитном коробе")
    quantity=int(quantity)
    if quantity<0 or quantity>1000000:raise HTTPException(422,detail="Некорректное количество")
    row=db.execute(text("SELECT * FROM fc_transit_box_items WHERE id=:id AND box_id=:b FOR UPDATE"),{"id":_uuid(item_id,"ID позиции"),"b":box["id"]}).mappings().first()
    if not row:raise HTTPException(404,detail="Позиция в транзитном коробе не найдена")
    old=int(row["quantity"])
    if quantity==0:db.execute(text("DELETE FROM fc_transit_box_items WHERE id=:id"),{"id":row["id"]})
    else:db.execute(text("UPDATE fc_transit_box_items SET quantity=:q,updated_at=now() WHERE id=:id"),{"q":quantity,"id":row["id"]})
    _event(db,event_id,box["batch_id"],box["id"],"SET_QTY",{"item_id":str(row["id"]),"old":old,"quantity":quantity},actor)
    db.execute(text("UPDATE fc_transit_boxes SET updated_at=now() WHERE id=:id"),{"id":box["id"]})
    db.execute(text("UPDATE fc_transit_batches SET updated_at=now() WHERE id=:id"),{"id":box["batch_id"]})
    return box_public(db,_box(db,box_id=box["id"],client_id=client_id,warehouse_id=warehouse_id))


def set_box_status(db, box_id, client_id, warehouse_id, close, event_id, actor):
    box=_lock_box_scope(db,box_id,client_id,warehouse_id)
    if box["batch_status"]!="OPEN":raise HTTPException(409,detail="Партия уже отправлена или закрыта")
    target="CLOSED" if close else "OPEN"
    if close:
        units=db.execute(text("SELECT COALESCE(SUM(quantity),0) FROM fc_transit_box_items WHERE box_id=:b"),{"b":box["id"]}).scalar_one()
        if int(units)<=0:raise HTTPException(409,detail="Пустой короб нельзя закрыть")
    db.execute(text("UPDATE fc_transit_boxes SET status=:s,updated_at=now() WHERE id=:id"),{"s":target,"id":box["id"]})
    _event(db,event_id,box["batch_id"],box["id"],"BOX_CLOSE" if close else "BOX_REOPEN",{},actor)
    db.execute(text("UPDATE fc_transit_batches SET updated_at=now() WHERE id=:id"),{"id":box["batch_id"]})
    return box_public(db,_box(db,box_id=box["id"],client_id=client_id,warehouse_id=warehouse_id))



def replace_box(db, box_id, client_id, warehouse_id, event_id, actor):
    """Give one transit box a fresh QR code without losing its contents.\n\n    Sequence, items and open/closed state remain unchanged. The old QR becomes\n    invalid immediately. SENT batches are immutable history.\n    """
    box=_lock_box_scope(db,box_id,client_id,warehouse_id)
    if box["batch_status"]!="OPEN":
        raise HTTPException(409,detail="Короб можно заменить только до отправки партии партнёру")
    eid=_uuid(event_id,"ID операции")
    if db.execute(text("SELECT 1 FROM fc_transit_events WHERE id=:id"),{"id":eid}).first():
        result=box_public(db,_box(db,box_id=box["id"],client_id=client_id,warehouse_id=warehouse_id));result["replayed"]=True;return result
    old=str(box["barcode"]);new=_new_box_code(db)
    db.execute(text("UPDATE fc_transit_boxes SET barcode=:new,updated_at=now() WHERE id=:id"),{"new":new,"id":box["id"]})
    _event(db,eid,box["batch_id"],box["id"],"BOX_REPLACE",{"old_barcode":old,"new_barcode":new,"sequence":int(box["sequence"])},actor)
    db.execute(text("UPDATE fc_transit_batches SET updated_at=now() WHERE id=:id"),{"id":box["batch_id"]})
    result=box_public(db,_box(db,box_id=box["id"],client_id=client_id,warehouse_id=warehouse_id));result["old_barcode"]=old
    return result


def _writeoff_rows(db,batch_id):
    return [dict(r) for r in db.execute(text('''SELECT i.sku_id,
        COALESCE(MAX(NULLIF(i.barcode,'')),'') barcode,
        COALESCE(MAX(NULLIF(i.article,'')),'') article,
        COALESCE(MAX(NULLIF(i.size,'')),'') size,
        COALESCE(MAX(NULLIF(i.name,'')),'') name,
        COALESCE(SUM(i.quantity),0)::int quantity
      FROM fc_transit_box_items i JOIN fc_transit_boxes b ON b.id=i.box_id
      WHERE b.batch_id=:a AND i.quantity>0
      GROUP BY i.sku_id ORDER BY article,size,barcode'''),{"a":str(batch_id)}).mappings().all()]



def _event_exists(db,event_id):
    return bool(db.execute(text("SELECT 1 FROM fc_transit_events WHERE id=:id"),{"id":_uuid(event_id,"ID операции")}).first())


def _ensure_inventory_sku(db, client_id, item):
    sid=item.get("sku_id")
    if sid:
        try:
            sku=db.get(SKU,uuid.UUID(str(sid)))
        except Exception:
            sku=None
        if sku and str(sku.client_id)==str(client_id):
            if not bool(getattr(sku,"is_active",True)):
                sku.is_active=True
            return sku
    barcode=str(item.get("barcode") or "").strip()
    if barcode:
        rows=db.execute(select(SKU).where(SKU.client_id==client_id,func.lower(SKU.barcode)==barcode.casefold())).scalars().all()
        if len(rows)>1:
            raise HTTPException(409,detail=f"ШК {barcode} привязан к нескольким SKU. Сначала исправьте карточку товара")
        if rows:
            sku=rows[0];sku.is_active=True;return sku
    article=str(item.get("article") or "").strip();size=str(item.get("size") or "").strip()
    if article:
        stmt=select(SKU).where(SKU.client_id==client_id,func.lower(func.coalesce(SKU.article,''))==article.casefold(),func.lower(func.coalesce(SKU.size,''))==size.casefold())
        rows=db.execute(stmt).scalars().all()
        if len(rows)==1:
            sku=rows[0];sku.is_active=True;return sku
    identity="|".join([str(client_id),barcode,article,size,str(item.get("name") or "")])
    digest=hashlib.sha256(identity.encode("utf-8")).hexdigest()
    safe_barcode=barcode or ("TRLOCAL-"+digest[:18].upper())
    sku=SKU(id=uuid.uuid4(),client_id=client_id,sku_code="LOCAL-TRANSIT-"+digest[:16].upper(),barcode=safe_barcode,
            article=article or safe_barcode,size=size or "Без размера",name=str(item.get("name") or article or safe_barcode)[:255],unit="PCS",is_active=True)
    db.add(sku);db.flush();return sku


def _decrement_transit_item(db,item_id,quantity):
    row=db.execute(text("SELECT * FROM fc_transit_box_items WHERE id=:id FOR UPDATE"),{"id":str(item_id)}).mappings().first()
    if not row:raise HTTPException(404,detail="Позиция транзита не найдена")
    old=int(row.get("quantity") or 0);quantity=int(quantity)
    if quantity<1 or quantity>old:raise HTTPException(409,detail=f"В транзите доступно {old} шт.")
    left=old-quantity
    if left:
        db.execute(text("UPDATE fc_transit_box_items SET quantity=:q,updated_at=now() WHERE id=:id"),{"q":left,"id":str(item_id)})
    else:
        db.execute(text("DELETE FROM fc_transit_box_items WHERE id=:id"),{"id":str(item_id)})
    return dict(row),left


def transfer_transit_to_inventory(db,box_id,item_id,client_id,warehouse_id,target,quantity,event_id,actor):
    target=str(target or "").upper()
    if target not in {"MAIN","FBS"}:raise HTTPException(422,detail="Назначение должно быть MAIN или FBS")
    first=_box(db,box_id=box_id,client_id=client_id,warehouse_id=warehouse_id,lock=False)
    batch=_batch(db,first["batch_id"],client_id,warehouse_id,lock=True)
    box=_box(db,box_id=box_id,client_id=client_id,warehouse_id=warehouse_id,lock=True)
    if batch["status"]!="OPEN":raise HTTPException(409,detail="Переводить можно только активный транзитный товар")
    if str(batch.get("accounting_mode") or "LEGACY_MAIN_FBS")!="TRANSIT_ONLY":
        raise HTTPException(409,detail="Эта старая партия учитывается по прежней схеме MAIN/FBS. Перевод из отдельного транзитного остатка для неё недоступен")
    if _event_exists(db,event_id):return {"replayed":True,"batch":batch_public(db,batch,include_items=True)}
    row=db.execute(text("SELECT * FROM fc_transit_box_items WHERE id=:id AND box_id=:b FOR UPDATE"),{"id":_uuid(item_id,"ID позиции"),"b":str(box_id)}).mappings().first()
    if not row:raise HTTPException(404,detail="Позиция транзита не найдена")
    q=int(quantity)
    if q<1 or q>int(row["quantity"]):raise HTTPException(409,detail=f"В транзите доступно {int(row['quantity'])} шт.")
    sku=_ensure_inventory_sku(db,uuid.UUID(str(client_id)),dict(row))
    from app.routers.receiving import _ensure_location
    from app.schemas.inventory import OperationCreate,StockLine
    from app.services.inventory import execute_operation
    from app.services.stock_core_v2 import reconcile_stock_bucket
    loc=_ensure_location(db,uuid.UUID(str(warehouse_id)),"ПРИЕМКА")
    op=execute_operation(db,OperationCreate(client_id=uuid.UUID(str(client_id)),warehouse_id=uuid.UUID(str(warehouse_id)),operation_type="RECEIVING",
        idempotency_key=f"TRANSIT_TO_{target}:{_uuid(event_id,'ID операции')}",performed_by=actor,
        reference=(f"Транзит → {target}: {batch.get('lot_number') or batch['number']}")[:255],
        comment=(f"Поставщик: {(_supplier(db,batch.get('supplier_id'),client_id,warehouse_id) or {}).get('name','')}; из короба {box['barcode']}")[:1000],
        lines=[StockLine(sku_id=uuid.UUID(str(sku.id)),quantity=Decimal(q),destination_location_id=loc.id,reason=f"TRANSIT_TO_{target}")]),commit=False)
    bucket=reconcile_stock_bucket(db,client_id=uuid.UUID(str(client_id)),warehouse_id=uuid.UUID(str(warehouse_id)),sku_id=uuid.UUID(str(sku.id)),
        location_id=loc.id,container_id=None,preferred_bucket="FBS_AVAILABLE" if target=="FBS" else "STORAGE_ONLY",lock=True)
    db.execute(text("UPDATE fc_transit_box_items SET sku_id=:s,barcode=COALESCE(NULLIF(barcode,''),:bc) WHERE id=:id"),{"s":str(sku.id),"bc":str(sku.barcode or ""),"id":str(row["id"])})
    _decrement_transit_item(db,row["id"],q)
    publication=None
    if target=="FBS":
        from app.services.stock_recovery import queue_publication
        publication=queue_publication(db,"MANUAL",{"client_id":str(client_id),"warehouse_id":str(warehouse_id),"sku_ids":[str(sku.id)]})
    payload={"item_id":str(row["id"]),"quantity":q,"target":target,"operation_id":str(op.id),"location":"ПРИЕМКА","publication":publication,
             "barcode":row.get("barcode") or str(sku.barcode or ""),"article":row.get("article") or str(sku.article or ""),
             "size":row.get("size") or str(sku.size or ""),"name":row.get("name") or str(sku.name or "")}
    _event(db,event_id,batch["id"],box["id"],f"TRANSFER_TO_{target}",payload,actor)
    db.execute(text("UPDATE fc_transit_boxes SET updated_at=now() WHERE id=:id"),{"id":box["id"]})
    db.execute(text("UPDATE fc_transit_batches SET updated_at=now() WHERE id=:id"),{"id":batch["id"]})
    return {"ok":True,"target":target,"quantity":q,"operation_id":str(op.id),"location":"ПРИЕМКА","bucket":bucket,"publication":publication,
            "batch":batch_public(db,_batch(db,batch["id"],client_id,warehouse_id),include_items=True)}


def _remove_inventory_for_transit(db,client_id,warehouse_id,sku_id,quantity,source,actor,event_id,batch,box):
    source=str(source or "").upper()
    if source not in {"MAIN","FBS"}:raise HTTPException(422,detail="Источник должен быть MAIN или FBS")
    from app.services import stock_actions as actions
    from app.services.stock_core_v2 import canonical_stock_identity_sku_ids,effective_fbs_free_for_sku
    ids=sorted(canonical_stock_identity_sku_ids(db,client_id=client_id,sku_id=str(sku_id)))
    actions.locks(db,str(client_id),str(warehouse_id),ids)
    stmt=text('''SELECT CAST(sku_id AS text) sku_id,CAST(location_id AS text) location_id,CAST(container_id AS text) container_id
        FROM inventory_balances WHERE client_id=:c AND warehouse_id=:w AND sku_id IN :ids AND quantity>0
        ORDER BY location_id,container_id NULLS FIRST''').bindparams(bindparam("ids",expanding=True))
    candidates=[dict(r) for r in db.execute(stmt,{"c":str(client_id),"w":str(warehouse_id),"ids":ids}).mappings().all()]
    need=Decimal(int(quantity));states={};budget=Decimal("999999999")
    if source=="FBS":budget=max(Decimal(0),Decimal(effective_fbs_free_for_sku(db,client_id=client_id,warehouse_id=warehouse_id,sku_id=str(sku_id))))
    def state_for(row):
        key=(row["sku_id"],row["location_id"],row.get("container_id") or "")
        if key not in states:
            snap=actions.bucket(db,str(client_id),str(warehouse_id),row["sku_id"],row["location_id"],row.get("container_id"))
            if snap["storage_raw"]!=snap["main"]:
                raise HTTPException(409,detail="Учёт Основного склада не совпадает с физическим остатком. Сначала выполните сверку этого адреса")
            states[key]={"snapshot":snap,"remove":Decimal(0)}
        return states[key]
    for row in candidates:
        if need<=0:break
        st=state_for(row);b=st["snapshot"]
        if source=="MAIN":
            if b["storage_raw"]!=b["main"]:raise HTTPException(409,detail="Учёт Основного склада не совпадает с физическим остатком. Сначала выполните сверку")
            free=max(Decimal(0),b["main"]-b["main_held"]-st["remove"])
        else:
            free=min(max(Decimal(0),b["fbs"]-b["fbs_held"]-st["remove"]),budget)
        take=min(need,free)
        if take>0:
            st["remove"]+=take;need-=take
            if source=="FBS":budget-=take
    if need>0:
        available=int(Decimal(int(quantity))-need)
        raise HTTPException(409,detail=f"Недостаточно свободного {source}: нужно {int(quantity)} шт., свободно {available} шт. Резерв не затрагивается")
    from app.schemas.inventory import OperationCreate,StockLine
    from app.services.inventory import execute_operation
    lines=[];container_ids=[]
    for st in states.values():
        b=st["snapshot"];q=st["remove"]
        if q<=0:continue
        lines.append(StockLine(sku_id=uuid.UUID(str(b["sku_id"])),quantity=q,source_location_id=uuid.UUID(str(b["location_id"])),
            source_container_id=uuid.UUID(str(b["container_id"])) if b.get("container_id") else None,reason=f"{source}_TO_TRANSIT"))
        if b.get("container_id"):container_ids.append(str(b["container_id"]))
    op=execute_operation(db,OperationCreate(client_id=uuid.UUID(str(client_id)),warehouse_id=uuid.UUID(str(warehouse_id)),operation_type="SHIPPING",
        idempotency_key=f"{source}_TO_TRANSIT:{_uuid(event_id,'ID операции')}",performed_by=actor,
        reference=(f"{source} → Транзит: {batch.get('lot_number') or batch['number']}")[:255],
        comment=(f"Перемещено в транзитный короб {box['barcode']}")[:1000],lines=lines),commit=False)
    for st in states.values():
        b=st["snapshot"];q=st["remove"]
        if q<=0:continue
        physical=b["physical"]-q;main=b["main"]-(q if source=="MAIN" else Decimal(0))
        actual=actions.bucket(db,str(client_id),str(warehouse_id),b["sku_id"],b["location_id"],b.get("container_id"))
        if actual["physical"]!=physical:raise HTTPException(409,detail="Физическое перемещение в транзит не подтвердилось. Операция отменена")
        actions.save_mirror(db,str(client_id),str(warehouse_id),b,physical,main)
    archived=0
    if container_ids:
        from app.services.container_lifecycle import archive_consumed_empty_containers
        archived=int(archive_consumed_empty_containers(db,client_id,warehouse_id,container_ids,actor=actor,reason=f"{source}_TO_TRANSIT").get("count",0))
    publication=None
    if source=="FBS":
        from app.services.stock_recovery import queue_publication
        publication=queue_publication(db,"MANUAL",{"client_id":str(client_id),"warehouse_id":str(warehouse_id),"sku_ids":ids})
    return {"operation_id":str(op.id),"quantity":int(quantity),"source":source,"archived_empty_boxes":archived,"publication":publication}


def transfer_inventory_to_transit(db,box_id,client_id,warehouse_id,source,barcode,quantity,event_id,actor):
    box=_lock_box_scope(db,box_id,client_id,warehouse_id)
    batch=_batch(db,box["batch_id"],client_id,warehouse_id,lock=False)
    if batch["status"]!="OPEN" or box["status"]!="OPEN":raise HTTPException(409,detail="Переводить товар можно только в открытый транзитный короб")
    if str(batch.get("accounting_mode") or "LEGACY_MAIN_FBS")!="TRANSIT_ONLY":raise HTTPException(409,detail="Перевод MAIN/FBS → транзит доступен только для новых партий r145f6")
    if _event_exists(db,event_id):return {"replayed":True,"batch":batch_public(db,batch,include_items=True)}
    q=int(quantity)
    if q<1:raise HTTPException(422,detail="Количество должно быть больше нуля")
    product=_catalog_identity(db,client_id,barcode,warehouse_id=warehouse_id,allow_transit_local=False)
    if not product.get("sku_id"):raise HTTPException(409,detail="Этот ШК не привязан к SKU локальной базы. Для перемещения из MAIN/FBS сначала добавьте карточку товара")
    summary=_remove_inventory_for_transit(db,client_id,warehouse_id,product["sku_id"],q,source,actor,event_id,batch,box)
    _upsert_item(db,box["id"],product,q)
    _event(db,event_id,batch["id"],box["id"],f"TRANSFER_FROM_{str(source).upper()}",{**summary,"barcode":product.get("barcode"),"article":product.get("article"),"size":product.get("size"),"name":product.get("name")},actor)
    db.execute(text("UPDATE fc_transit_boxes SET updated_at=now() WHERE id=:id"),{"id":box["id"]})
    db.execute(text("UPDATE fc_transit_batches SET updated_at=now() WHERE id=:id"),{"id":batch["id"]})
    return {"ok":True,**summary,"batch":batch_public(db,_batch(db,batch["id"],client_id,warehouse_id),include_items=True)}

def _writeoff_plan(db,batch,client_id,warehouse_id):
    """Build one reservation-safe physical source plan for partner dispatch.\n\n    Transit has no source address at scan time. Dispatch therefore consumes current\n    free warehouse truth: MAIN first, then free FBS. Active reservations/picking holds\n    are never consumed.\n    """
    from app.services import stock_actions as actions
    from app.services.stock_core_v2 import (
        canonical_stock_identity_lock_key, canonical_stock_identity_sku_ids,
        effective_fbs_free_for_sku,
    )

    rows=_writeoff_rows(db,batch["id"])
    if not rows:
        raise HTTPException(409,detail="В транзитной партии нет товара для списания")
    missing=[r for r in rows if not r.get("sku_id")]
    if missing:
        labels=[]
        for r in missing[:8]:
            label=(str(r.get("article") or "")+(" / "+str(r.get("size") or "") if r.get("size") else "")).strip(" /") or str(r.get("barcode") or "неизвестный товар")
            labels.append(f"{label}: {int(r.get('quantity') or 0)} шт.")
        raise HTTPException(409,detail="Нельзя списать отправку: часть товара не привязана к SKU локальной базы. Исправьте карточку/артикул: "+"; ".join(labels))

    groups={}
    for row in rows:
        sid=str(row["sku_id"])
        key=canonical_stock_identity_lock_key(db,client_id=client_id,sku_id=sid)
        g=groups.setdefault(key,{"representative":sid,"sku_ids":set(),"quantity":Decimal(0),"labels":[]})
        g["sku_ids"].update(canonical_stock_identity_sku_ids(db,client_id=client_id,sku_id=sid))
        g["quantity"]+=Decimal(int(row["quantity"]))
        g["labels"].append({k:row.get(k) for k in ("article","size","barcode","name")})

    all_ids=sorted({sid for g in groups.values() for sid in g["sku_ids"]})
    actions.locks(db,str(client_id),str(warehouse_id),all_ids)
    states={}
    total_main=Decimal(0);total_fbs=Decimal(0);fbs_skus=set()

    def state_for(sid,loc,box):
        key=(str(sid),str(loc),str(box or ""))
        if key not in states:
            snap=actions.bucket(db,str(client_id),str(warehouse_id),sid,loc,box)
            if snap["storage_raw"]!=snap["main"]:
                raise HTTPException(409,detail="Учёт Основного склада не совпадает с физическим остатком на одном из адресов. Сначала выполните сверку остатка, затем отправляйте транзит.")
            states[key]={"snapshot":snap,"main_remove":Decimal(0),"fbs_remove":Decimal(0)}
        return states[key]

    for g in groups.values():
        ids=sorted(g["sku_ids"]);need=Decimal(g["quantity"])
        stmt=text('''SELECT CAST(sku_id AS text) sku_id,CAST(location_id AS text) location_id,
              CAST(container_id AS text) container_id
            FROM inventory_balances
            WHERE client_id=:c AND warehouse_id=:w AND sku_id IN :ids AND quantity>0
            ORDER BY location_id,container_id NULLS FIRST''').bindparams(bindparam("ids",expanding=True))
        candidates=[dict(r) for r in db.execute(stmt,{"c":str(client_id),"w":str(warehouse_id),"ids":ids}).mappings().all()]

        for row in candidates:
            if need<=0:break
            st=state_for(row["sku_id"],row["location_id"],row.get("container_id"));b=st["snapshot"]
            free=max(Decimal(0),b["main"]-b["main_held"]-st["main_remove"])
            take=min(need,free)
            if take>0:st["main_remove"]+=take;need-=take;total_main+=take

        if need>0:
            budget=max(Decimal(0),Decimal(effective_fbs_free_for_sku(db,client_id=client_id,warehouse_id=warehouse_id,sku_id=g["representative"])))
            for row in candidates:
                if need<=0 or budget<=0:break
                st=state_for(row["sku_id"],row["location_id"],row.get("container_id"));b=st["snapshot"]
                free=max(Decimal(0),b["fbs"]-b["fbs_held"]-st["fbs_remove"])
                take=min(need,budget,free)
                if take>0:
                    st["fbs_remove"]+=take;need-=take;budget-=take;total_fbs+=take;fbs_skus.update(ids)

        if need>0:
            label=g["labels"][0];name=(str(label.get("article") or "")+(" / "+str(label.get("size") or "") if label.get("size") else "")).strip(" /") or str(label.get("barcode") or label.get("name") or "товар")
            available=g["quantity"]-need
            raise HTTPException(409,detail=f"Недостаточно свободного остатка для отправки партнёру: {name}. Нужно {int(g['quantity'])} шт., свободно {int(available)} шт. Резервированный товар не списывается.")

    return {"states":states,"main":total_main,"fbs":total_fbs,"fbs_skus":sorted(fbs_skus),"units":total_main+total_fbs}


def _writeoff_batch_stock(db,batch,client_id,warehouse_id,actor):
    if batch.get("stock_writeoff_at"):
        return batch.get("stock_writeoff_summary") or {"replayed":True}
    key=f"transit-partner-send-{batch['id']}"
    existing=db.execute(text("SELECT CAST(id AS text) id,status FROM inventory_operations WHERE client_id=:c AND idempotency_key=:k"),{"c":str(client_id),"k":key}).mappings().first()
    if existing:
        raise HTTPException(409,detail="Для этой транзитной партии уже существует операция списания, но партия не помечена как списанная. Повторное списание остановлено; проверьте историю склада.")

    plan=_writeoff_plan(db,batch,client_id,warehouse_id)
    from app.schemas.inventory import OperationCreate,StockLine
    from app.services.inventory import execute_operation
    from app.services import stock_actions as actions

    lines=[];container_ids=[]
    for st in plan["states"].values():
        b=st["snapshot"];qty=st["main_remove"]+st["fbs_remove"]
        if qty<=0:continue
        lines.append(StockLine(sku_id=uuid.UUID(str(b["sku_id"])),quantity=qty,
            source_location_id=uuid.UUID(str(b["location_id"])),
            source_container_id=uuid.UUID(str(b["container_id"])) if b.get("container_id") else None,
            reason="TRANSIT_PARTNER_SEND"))
        if b.get("container_id"):container_ids.append(str(b["container_id"]))
    if not lines:raise HTTPException(409,detail="Не найдено свободного остатка для списания транзитной партии")

    op=execute_operation(db,OperationCreate(client_id=uuid.UUID(str(client_id)),warehouse_id=uuid.UUID(str(warehouse_id)),
        operation_type="SHIPPING",idempotency_key=key,performed_by=actor,
        reference=("Транзит "+str(batch["number"]))[:255],
        comment=("Отправлено на склад партнёра: "+str(batch["partner"]))[:1000],lines=lines),commit=False)

    for st in plan["states"].values():
        b=st["snapshot"];removed=st["main_remove"]+st["fbs_remove"]
        if removed<=0:continue
        physical=b["physical"]-removed;main=b["main"]-st["main_remove"]
        actual=actions.bucket(db,str(client_id),str(warehouse_id),b["sku_id"],b["location_id"],b.get("container_id"))
        if actual["physical"]!=physical:
            raise HTTPException(409,detail="Физическое списание транзита не подтвердилось. Операция отменена без изменения партии.")
        actions.save_mirror(db,str(client_id),str(warehouse_id),b,physical,main)

    archived=0
    if container_ids:
        from app.services.container_lifecycle import archive_consumed_empty_containers
        archived=int(archive_consumed_empty_containers(db,client_id,warehouse_id,container_ids,actor=actor,reason="TRANSIT_PARTNER_SEND").get("count",0))
    wb_message=None
    if plan["fbs_skus"]:
        from app.services.stock_recovery import queue_publication
        wb_message=queue_publication(db,"MANUAL",{"client_id":str(client_id),"warehouse_id":str(warehouse_id),"sku_ids":plan["fbs_skus"]})

    summary={"operation_id":str(op.id),"units":int(plan["units"]),"main_units":int(plan["main"]),"fbs_units":int(plan["fbs"]),
             "archived_empty_boxes":archived,"wb":wb_message,"mode":"MAIN_FIRST_THEN_FREE_FBS","version":"r137f4"}
    db.execute(text("UPDATE fc_transit_batches SET stock_writeoff_at=now(),stock_writeoff_summary=CAST(:s AS jsonb),updated_at=now() WHERE id=:id"),
               {"s":json.dumps(summary,ensure_ascii=False),"id":str(batch["id"])})
    _event(db,uuid.uuid5(uuid.NAMESPACE_URL,"fulfillcore:transit-writeoff:"+str(batch["id"])),batch["id"],None,"STOCK_WRITE_OFF",summary,actor)
    return summary



def writeoff_sent_batch(db,batch_id,client_id,warehouse_id,event_id,actor):
    batch=_batch(db,batch_id,client_id,warehouse_id,lock=True)
    if str(batch.get("accounting_mode") or "LEGACY_MAIN_FBS")=="TRANSIT_ONLY":
        raise HTTPException(409,detail="Новые транзитные партии имеют отдельный остаток и не списываются из MAIN/FBS при отправке")
    if batch["status"]!="SENT":raise HTTPException(409,detail="Отдельное списание доступно только для партии, уже отмеченной как отправленная")
    if batch.get("stock_writeoff_at"):
        return batch_public(db,batch,include_items=True)
    summary=_writeoff_batch_stock(db,batch,client_id,warehouse_id,actor)
    _event(db,event_id,batch["id"],None,"LEGACY_SENT_WRITE_OFF",summary,actor)
    return batch_public(db,_batch(db,batch["id"],client_id,warehouse_id),include_items=True)



def send_batch(db, batch_id, client_id, warehouse_id, event_id, actor, destination="", note=""):
    batch=_batch(db,batch_id,client_id,warehouse_id,lock=True)
    mode=str(batch.get("accounting_mode") or "LEGACY_MAIN_FBS")
    if batch["status"]=="SENT":
        if mode=="LEGACY_MAIN_FBS" and not batch.get("stock_writeoff_at"):
            _writeoff_batch_stock(db,batch,client_id,warehouse_id,actor)
        return batch_public(db,_batch(db,batch["id"],client_id,warehouse_id),include_items=True)
    if batch["status"]!="OPEN":raise HTTPException(409,detail="Эту транзитную партию нельзя отправить")
    boxes=[dict(r) for r in db.execute(text("SELECT * FROM fc_transit_boxes WHERE batch_id=:a ORDER BY sequence FOR UPDATE"),{"a":batch["id"]}).mappings().all()]
    if not boxes:raise HTTPException(409,detail="В партии нет коробов")
    for b in boxes:
        units=int(db.execute(text("SELECT COALESCE(SUM(quantity),0) FROM fc_transit_box_items WHERE box_id=:b"),{"b":b["id"]}).scalar_one())
        if units<=0:raise HTTPException(409,detail=f"Короб {b['barcode']} пустой")
        if b["status"]!="CLOSED":raise HTTPException(409,detail=f"Сначала закройте короб {b['barcode']}")
    if mode=="TRANSIT_ONLY":
        destination=_clean(destination,200,"куда отправляем",True);note=_clean(note,500,"комментарий отправки")
        db.execute(text("UPDATE fc_transit_batches SET status='SENT',partner=:dest,destination=:dest,departure_note=:note,sent_at=now(),updated_at=now() WHERE id=:id"),
                   {"dest":destination,"note":note,"id":batch["id"]})
        db.execute(text("UPDATE fc_transit_boxes SET status='SENT',updated_at=now() WHERE batch_id=:id"),{"id":batch["id"]})
        units=sum(int(db.execute(text("SELECT COALESCE(SUM(quantity),0) FROM fc_transit_box_items WHERE box_id=:b"),{"b":b["id"]}).scalar_one()) for b in boxes)
        _event(db,event_id,batch["id"],None,"BATCH_DEPART",{"destination":destination,"note":note,"boxes":len(boxes),"units":units,"stock_changed":False},actor)
    else:
        summary=_writeoff_batch_stock(db,batch,client_id,warehouse_id,actor)
        db.execute(text("UPDATE fc_transit_batches SET status='SENT',destination=COALESCE(NULLIF(destination,''),partner),sent_at=now(),updated_at=now() WHERE id=:id"),{"id":batch["id"]})
        db.execute(text("UPDATE fc_transit_boxes SET status='SENT',updated_at=now() WHERE batch_id=:id"),{"id":batch["id"]})
        _event(db,event_id,batch["id"],None,"BATCH_SEND",{"partner":batch["partner"],"boxes":len(boxes),"writeoff":summary},actor)
    return batch_public(db,_batch(db,batch["id"],client_id,warehouse_id),include_items=True)



def cancel_batch(db,batch_id,client_id,warehouse_id,event_id,actor):
    batch=_batch(db,batch_id,client_id,warehouse_id,lock=True)
    if batch["status"]=="SENT":raise HTTPException(409,detail="Отправленную транзитную партию отменить нельзя")
    if batch["status"]=="CANCELLED":return batch_public(db,batch,include_items=True)
    if str(batch.get("accounting_mode") or "LEGACY_MAIN_FBS")=="TRANSIT_ONLY":
        units=int(db.execute(text('''SELECT COALESCE(SUM(i.quantity),0) FROM fc_transit_box_items i JOIN fc_transit_boxes b ON b.id=i.box_id WHERE b.batch_id=:a'''),{"a":batch["id"]}).scalar() or 0)
        if units>0:raise HTTPException(409,detail="В партии есть товар. Сначала переведите его в Основной/FBS или удалите ошибочные позиции, затем отменяйте пустую партию")
    db.execute(text("UPDATE fc_transit_batches SET status='CANCELLED',updated_at=now() WHERE id=:id"),{"id":batch["id"]})
    _event(db,event_id,batch["id"],None,"BATCH_CANCEL",{},actor)
    return batch_public(db,_batch(db,batch["id"],client_id,warehouse_id),include_items=True)



def _render_qr_labels(batch, boxes, supplier_name=""):
    from reportlab.graphics import renderPDF
    from reportlab.graphics.barcode import qr
    from reportlab.graphics.shapes import Drawing
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas
    from app.services.pdf_labels import _ensure_label_fonts
    regular,bold=_ensure_label_fonts();out=io.BytesIO();w,h=58*mm,40*mm;c=canvas.Canvas(out,pagesize=(w,h),pageCompression=1)
    mode=str(batch.get("accounting_mode") or "LEGACY_MAIN_FBS")
    for b in boxes:
        c.setFont(bold,9.5);c.drawString(2.2*mm,36.0*mm,"ТРАНЗИТ")
        c.setFont(bold,7);c.drawRightString(55.8*mm,36.2*mm,f"КОРОБ {int(b['sequence'])}")
        widget=qr.QrCodeWidget(str(b["barcode"]))
        x1,y1,x2,y2=widget.getBounds();qw=max(x2-x1,1);qh=max(y2-y1,1);size=22*mm
        scale=min(size/qw,size/qh);drawing=Drawing(size,size,transform=[scale,0,0,scale,-x1*scale,-y1*scale]);drawing.add(widget)
        renderPDF.draw(drawing,c,(w-size)/2,11.5*mm)
        c.setFont(bold,7.2);c.drawCentredString(w/2,8.7*mm,str(b["barcode"]))
        c.setFont(regular,5.3)
        c.drawString(2.2*mm,5.3*mm,("Партия: "+str(batch.get("lot_number") or batch["number"]))[:54])
        if mode=="TRANSIT_ONLY":
            c.drawString(2.2*mm,2.3*mm,("Поставщик: "+str(supplier_name or "—"))[:57])
        else:
            c.drawString(2.2*mm,2.3*mm,("Куда: "+str(batch.get("partner") or ""))[:57])
        c.showPage()
    c.save();return out.getvalue()



def box_labels_pdf(db,batch_id,client_id,warehouse_id):
    batch=_batch(db,batch_id,client_id,warehouse_id)
    boxes=[dict(r) for r in db.execute(text("SELECT * FROM fc_transit_boxes WHERE batch_id=:a ORDER BY sequence"),{"a":batch["id"]}).mappings().all()]
    if not boxes:raise HTTPException(409,detail="Сначала создайте транзитные короба")
    supplier=_supplier(db,batch.get("supplier_id"),client_id,warehouse_id) if batch.get("supplier_id") else None
    return _render_qr_labels(batch,boxes,(supplier or {}).get("name", ""))



def box_label_pdf(db,box_id,client_id,warehouse_id):
    box=_box(db,box_id=box_id,client_id=client_id,warehouse_id=warehouse_id)
    batch=_batch(db,box["batch_id"],client_id,warehouse_id)
    supplier=_supplier(db,batch.get("supplier_id"),client_id,warehouse_id) if batch.get("supplier_id") else None
    return _render_qr_labels(batch,[box],(supplier or {}).get("name", ""))



def export_xlsx(db,batch_id,client_id,warehouse_id):
    batch=_batch(db,batch_id,client_id,warehouse_id)
    data=batch_public(db,batch,include_items=True)
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb=Workbook();ws=wb.active;ws.title="Состав коробов"
    ws.append(["Поставщик","Партия поставщика","Статус","Куда отправлено","Дата отправки","QR короба","Статус короба","ШК товара","Артикул","Размер","Товар","Количество","Принято / добавлено","Источник"])
    source_names={"CATALOG":"Локальная карточка","UNKNOWN_BARCODE":"Товар транзита","MANUAL":"Вручную без ШК"}
    for box in data["boxes"]:
        for item in box.get("items") or []:
            append_safe(ws,[data.get("supplier_name") or "",data.get("lot_number") or data["number"],data["status"],data.get("destination") or "",data.get("sent_at") or "",
                box["barcode"],box["status"],item.get("barcode") or "",item.get("article") or "",item.get("size") or "",item.get("name") or "",
                int(item.get("quantity") or 0),str(item.get("created_at") or ""),source_names.get(item.get("source"),item.get("source"))])
    summary=wb.create_sheet("Короба");summary.append(["Поставщик","Партия","QR короба","Статус","Позиций","Штук"])
    for b in data["boxes"]:append_safe(summary,[data.get("supplier_name") or "",data.get("lot_number") or data["number"],b["barcode"],b["status"],int(b.get("lines") or 0),int(b.get("units") or 0)])
    events=wb.create_sheet("История приёмки");events.append(["Дата / время","Сотрудник","Действие","QR короба","ШК","Артикул","Размер","Товар","Количество"])
    erows=db.execute(text('''SELECT e.created_at,e.actor,e.action,e.payload,b.barcode box_barcode FROM fc_transit_events e LEFT JOIN fc_transit_boxes b ON b.id=e.box_id
        WHERE e.batch_id=:a AND e.action IN ('SCAN','MANUAL_ADD','TRANSFER_FROM_MAIN','TRANSFER_FROM_FBS') ORDER BY e.created_at,e.id'''),{"a":str(batch["id"])}).mappings().all()
    for e in erows:
        payload=e.get("payload") or {};qty=int(payload.get("quantity") or 1)
        append_safe(events,[str(e.get("created_at") or ""),e.get("actor") or "",e.get("action") or "",e.get("box_barcode") or "",payload.get("barcode") or "",payload.get("article") or "",payload.get("size") or "",payload.get("name") or "",qty])
    posting=wb.create_sheet("Учёт");posting.append(["Параметр","Значение"])
    append_safe(posting,["Режим","Отдельный транзит" if data.get("transit_only") else "Старая схема MAIN/FBS"] )
    append_safe(posting,["Поставщик",data.get("supplier_name") or ""]);append_safe(posting,["Партия",data.get("lot_number") or data["number"]])
    append_safe(posting,["Куда отправлено",data.get("destination") or ""]);append_safe(posting,["Дата отправки",data.get("sent_at") or ""])
    writeoff=data.get("stock_writeoff_summary") or {}
    append_safe(posting,["Списание MAIN/FBS", "Не требуется" if data.get("transit_only") else ("Да" if data.get("stock_writeoff_done") else "Нет")])
    append_safe(posting,["Всего списано, шт.",int(writeoff.get("units") or 0)]);append_safe(posting,["Из MAIN, шт.",int(writeoff.get("main_units") or 0)]);append_safe(posting,["Из FBS, шт.",int(writeoff.get("fbs_units") or 0)])
    fill=PatternFill("solid",fgColor="D9EAD3")
    for sheet in (ws,summary,events,posting):
        sheet.freeze_panes="A2";sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.font=Font(bold=True);cell.fill=fill;cell.alignment=Alignment(vertical="center",wrap_text=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:cell.alignment=Alignment(vertical="top",wrap_text=True)
        for i in range(1,sheet.max_column+1):
            values=[str(sheet.cell(r,i).value or "") for r in range(1,min(sheet.max_row,500)+1)]
            sheet.column_dimensions[get_column_letter(i)].width=min(45,max(12,max((len(x) for x in values),default=10)+2))
    return workbook_bytes(wb)


def supplier_export_xlsx(db,supplier_id,client_id,warehouse_id):
    data=supplier_detail(db,supplier_id,client_id,warehouse_id);supplier=data["supplier"]
    from openpyxl import Workbook
    from openpyxl.styles import Alignment,Font,PatternFill
    from openpyxl.utils import get_column_letter
    wb=Workbook();ws=wb.active;ws.title="Товар поставщика"
    ws.append(["Поставщик","Партия","Статус","Куда отправлено","QR короба","ШК","Артикул","Размер","Товар","Количество","Принято","Обновлено"])
    rows=db.execute(text('''SELECT a.lot_number,a.number,a.status,COALESCE(NULLIF(a.destination,''),a.partner,'') destination,b.barcode box_barcode,
        i.barcode,i.article,i.size,i.name,i.quantity,i.created_at,i.updated_at
        FROM fc_transit_batches a JOIN fc_transit_boxes b ON b.batch_id=a.id JOIN fc_transit_box_items i ON i.box_id=b.id
        WHERE a.client_id=:c AND a.warehouse_id=:w AND a.supplier_id=:s AND i.quantity>0 ORDER BY i.created_at'''),
        {"c":str(client_id),"w":str(warehouse_id),"s":str(supplier["id"])}).mappings().all()
    for r in rows:append_safe(ws,[supplier["name"],r.get("lot_number") or r.get("number") or "",r.get("status") or "",r.get("destination") or "",r.get("box_barcode") or "",r.get("barcode") or "",r.get("article") or "",r.get("size") or "",r.get("name") or "",int(r.get("quantity") or 0),str(r.get("created_at") or ""),str(r.get("updated_at") or "")])
    history=wb.create_sheet("Когда приняли");history.append(["Дата / время","Сотрудник","Партия","QR короба","Действие","ШК","Артикул","Размер","Товар","Количество"])
    erows=db.execute(text('''SELECT e.created_at,e.actor,e.action,e.payload,a.lot_number,a.number,b.barcode box_barcode
        FROM fc_transit_events e JOIN fc_transit_batches a ON a.id=e.batch_id LEFT JOIN fc_transit_boxes b ON b.id=e.box_id
        WHERE a.client_id=:c AND a.warehouse_id=:w AND a.supplier_id=:s AND e.action IN ('SCAN','MANUAL_ADD')
        ORDER BY e.created_at,e.id'''),{"c":str(client_id),"w":str(warehouse_id),"s":str(supplier["id"])}).mappings().all()
    for e in erows:
        payload=e.get("payload") or {};append_safe(history,[str(e.get("created_at") or ""),e.get("actor") or "",e.get("lot_number") or e.get("number") or "",e.get("box_barcode") or "",e.get("action") or "",payload.get("barcode") or "",payload.get("article") or "",payload.get("size") or "",payload.get("name") or "",int(payload.get("quantity") or 1)])
    lots=wb.create_sheet("Партии");lots.append(["Партия","Статус","Штук","Коробов","Создана","Куда отправлено","Покинула склад"])
    for b in data["batches"]:append_safe(lots,[b.get("lot_number") or b.get("number") or "",b.get("status") or "",int(b.get("units") or 0),int(b.get("box_count") or 0),b.get("created_at") or "",b.get("destination") or "",b.get("sent_at") or ""])
    product_summary=wb.create_sheet("Сводка товаров");product_summary.append(["ШК","Артикул","Размер","Товар","Принято от поставщика","Из MAIN в транзит","Из FBS в транзит","Сейчас в транзите","Покинуло склад","Из транзита в MAIN","Из транзита в FBS","Удалено как ошибка","Первый приём","Последний приём"])
    for x in data.get("products") or []:
        append_safe(product_summary,[x.get("barcode") or "",x.get("article") or "",x.get("size") or "",x.get("name") or "",int(x.get("received_units") or 0),int(x.get("from_main_units") or 0),int(x.get("from_fbs_units") or 0),int(x.get("active_units") or 0),int(x.get("departed_units") or 0),int(x.get("to_main_units") or 0),int(x.get("to_fbs_units") or 0),int(x.get("deleted_correction_units") or 0),x.get("first_received") or "",x.get("last_received") or ""])
    movements=wb.create_sheet("Движения");movements.append(["Дата / время","Сотрудник","Партия","QR короба","Действие","ШК","Артикул","Размер","Товар","Количество","Куда / откуда"])
    mrows=db.execute(text("""SELECT e.created_at,e.actor,e.action,e.payload,a.lot_number,a.number,b.barcode box_barcode
        FROM fc_transit_events e JOIN fc_transit_batches a ON a.id=e.batch_id LEFT JOIN fc_transit_boxes b ON b.id=e.box_id
        WHERE a.client_id=:c AND a.warehouse_id=:w AND a.supplier_id=:s
        ORDER BY e.created_at,e.id"""),{"c":str(client_id),"w":str(warehouse_id),"s":str(supplier["id"])}).mappings().all()
    for e in mrows:
        payload=e.get("payload") or {};action=e.get("action") or "";direction=payload.get("target") or payload.get("source") or payload.get("destination") or ""
        append_safe(movements,[str(e.get("created_at") or ""),e.get("actor") or "",e.get("lot_number") or e.get("number") or "",e.get("box_barcode") or "",action,payload.get("barcode") or "",payload.get("article") or "",payload.get("size") or "",payload.get("name") or "",int(payload.get("quantity") or (1 if action in ('SCAN','MANUAL_ADD') else 0)),direction])
    fill=PatternFill("solid",fgColor="D9EAD3")
    for sheet in (ws,history,lots,product_summary,movements):
        sheet.freeze_panes="A2";sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.font=Font(bold=True);cell.fill=fill;cell.alignment=Alignment(vertical="center",wrap_text=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:cell.alignment=Alignment(vertical="top",wrap_text=True)
        for i in range(1,sheet.max_column+1):
            vals=[str(sheet.cell(r,i).value or "") for r in range(1,min(sheet.max_row,1000)+1)]
            sheet.column_dimensions[get_column_letter(i)].width=min(45,max(12,max((len(x) for x in vals),default=10)+2))
    return workbook_bytes(wb)


# FC_R142_TRANSIT_ADMIN_REPAIR: repair legacy/erroneous partner dispatch records without
# stealing reserved stock.  These actions change transit bookkeeping only and are audited.
def admin_remove_item(db, box_id, item_id, client_id, warehouse_id, event_id, actor):
    first = _box(db, box_id=box_id, client_id=client_id, warehouse_id=warehouse_id, lock=False)
    batch = _batch(db, first["batch_id"], client_id, warehouse_id, lock=True)
    box = _box(db, box_id=box_id, client_id=client_id, warehouse_id=warehouse_id, lock=True)
    mode=str(batch.get("accounting_mode") or "LEGACY_MAIN_FBS")
    if batch.get("stock_writeoff_at"):
        raise HTTPException(409, detail="После проведённого списания состав отправки менять нельзя")
    if mode=="TRANSIT_ONLY" and batch["status"]!="OPEN":
        raise HTTPException(409, detail="Товар из вкладки «Покинули склад» является неизменяемой историей")
    if batch["status"] not in {"OPEN", "SENT"}:
        raise HTTPException(409, detail="Эту партию уже нельзя редактировать")
    eid = _uuid(event_id, "ID операции")
    if db.execute(text("SELECT 1 FROM fc_transit_events WHERE id=:id"), {"id": eid}).first():
        return batch_public(db, _batch(db, batch["id"], client_id, warehouse_id), include_items=True)
    row = db.execute(text("SELECT * FROM fc_transit_box_items WHERE id=:id AND box_id=:b FOR UPDATE"),
                     {"id": _uuid(item_id, "ID позиции"), "b": str(box_id)}).mappings().first()
    if not row:
        raise HTTPException(404, detail="Позиция уже удалена из транзитной партии")
    if mode=="TRANSIT_ONLY":
        moved=db.execute(text("""SELECT 1 FROM fc_transit_events
            WHERE box_id=:b AND action IN ('TRANSFER_FROM_MAIN','TRANSFER_FROM_FBS')
              AND lower(COALESCE(payload->>'barcode',''))=lower(:barcode) LIMIT 1"""),
            {"b":str(box["id"]),"barcode":str(row.get("barcode") or "")}).first()
        if moved:
            raise HTTPException(409,detail="Эта позиция поступила переводом из Основного/FBS. Не удаляйте её: используйте «→ Основной» или «→ FBS», чтобы остаток не потерялся")
    payload = {"item_id": str(row["id"]), "barcode": row.get("barcode"), "article": row.get("article"),
               "size": row.get("size"), "quantity": int(row.get("quantity") or 0), "batch_status": batch["status"]}
    db.execute(text("DELETE FROM fc_transit_box_items WHERE id=:id"), {"id": str(row["id"])})
    _event(db, eid, batch["id"], box["id"], "ADMIN_ITEM_REMOVE", payload, actor)
    db.execute(text("UPDATE fc_transit_boxes SET updated_at=now() WHERE id=:id"), {"id": box["id"]})
    db.execute(text("UPDATE fc_transit_batches SET updated_at=now() WHERE id=:id"), {"id": batch["id"]})
    return batch_public(db, _batch(db, batch["id"], client_id, warehouse_id), include_items=True)


def acknowledge_already_absent(db, batch_id, client_id, warehouse_id, event_id, actor):
    """Close a legacy SENT batch when goods were physically removed earlier.

    No inventory balance is changed.  This is deliberately separate from write-off and
    is only available before any r137+ stock write-off has been recorded.
    """
    batch = _batch(db, batch_id, client_id, warehouse_id, lock=True)
    if str(batch.get("accounting_mode") or "LEGACY_MAIN_FBS") == "TRANSIT_ONLY":
        raise HTTPException(409, detail="Для отдельного транзитного остатка подтверждение старого списания не используется")
    if batch["status"] != "SENT":
        raise HTTPException(409, detail="Подтверждение доступно только для уже отправленной старой партии")
    if batch.get("stock_writeoff_at"):
        return batch_public(db, batch, include_items=True)
    eid = _uuid(event_id, "ID операции")
    if db.execute(text("SELECT 1 FROM fc_transit_events WHERE id=:id"), {"id": eid}).first():
        return batch_public(db, batch, include_items=True)
    units = int(db.execute(text('''SELECT COALESCE(SUM(i.quantity),0) FROM fc_transit_box_items i
        JOIN fc_transit_boxes b ON b.id=i.box_id WHERE b.batch_id=:a'''), {"a": str(batch["id"])}).scalar() or 0)
    summary = {"mode": "ADMIN_CONFIRMED_ALREADY_ABSENT", "units": units,
               "note": "Товар уже отсутствовал на складе; повторное физическое списание не выполнялось", "version": "r142"}
    db.execute(text("UPDATE fc_transit_batches SET stock_writeoff_at=now(),stock_writeoff_summary=CAST(:s AS jsonb),updated_at=now() WHERE id=:id"),
               {"s": json.dumps(summary, ensure_ascii=False), "id": str(batch["id"])})
    _event(db, eid, batch["id"], None, "ADMIN_ALREADY_ABSENT", summary, actor)
    return batch_public(db, _batch(db, batch["id"], client_id, warehouse_id), include_items=True)
