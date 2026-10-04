"""Read-only presentation. Timestamps in the database are never rewritten."""
from __future__ import annotations
import io, re
from datetime import datetime, date, timezone
from decimal import Decimal, InvalidOperation
from zoneinfo import ZoneInfo
MSK=ZoneInfo('Europe/Moscow'); UTC=timezone.utc
ILLEGAL=re.compile(r'[\x00-\x08\x0b-\x0c\x0e-\x1f]')
LABELS={
 'RECEIVING':'\u041f\u0440\u0438\u0451\u043c\u043a\u0430',
 'RECEIVING_IMPORT':'\u041f\u0440\u0438\u0451\u043c\u043a\u0430 Excel',
 'PVZ_RECEIVING':'\u041f\u0440\u0438\u0451\u043c\u043a\u0430 \u041f\u0412\u0417',
 'MOVE':'\u041f\u0435\u0440\u0435\u043c\u0435\u0449\u0435\u043d\u0438\u0435',
 'MOVE_CONTAINER':'\u041f\u0435\u0440\u0435\u043c\u0435\u0449\u0435\u043d\u0438\u0435 \u043a\u043e\u0440\u043e\u0431\u0430',
 'SHIPPING':'\u0421\u043f\u0438\u0441\u0430\u043d\u0438\u0435 / \u043e\u0442\u0433\u0440\u0443\u0437\u043a\u0430',
 'ADJUSTMENT':'\u041a\u043e\u0440\u0440\u0435\u043a\u0442\u0438\u0440\u043e\u0432\u043a\u0430',
 'RECONCILIATION':'\u0422\u0435\u0445\u043d\u0438\u0447\u0435\u0441\u043a\u0430\u044f \u0441\u0432\u0435\u0440\u043a\u0430',
 'EXTERNAL_PICK_WRITE_OFF':'\u0421\u043f\u0438\u0441\u0430\u043d\u0438\u0435 \u043f\u043e \u0441\u0431\u043e\u0440\u043e\u0447\u043d\u043e\u043c\u0443',
 'FBS_RESET':'\u041e\u0431\u043d\u0443\u043b\u0435\u043d\u0438\u0435 FBS',
 'RECEIPT_CANCEL':'\u041e\u0442\u043c\u0435\u043d\u0430 \u043f\u0440\u0438\u0451\u043c\u043a\u0438',
 'RESERVATION_STATE':'\u0421\u043e\u0441\u0442\u043e\u044f\u043d\u0438\u0435 \u0440\u0435\u0437\u0435\u0440\u0432\u0430 (\u043d\u0435 \u0434\u0432\u0438\u0436\u0435\u043d\u0438\u0435)',
}
STATUS={'COMPLETED':'\u041f\u0440\u043e\u0432\u0435\u0434\u0435\u043d\u0430','CANCELLED':'\u041e\u0442\u043c\u0435\u043d\u0435\u043d\u0430',
 'ACTIVE':'\u0410\u043a\u0442\u0438\u0432\u043d\u0430','SALE_RESERVED':'\u0420\u0435\u0437\u0435\u0440\u0432 \u0437\u0430\u043a\u0430\u0437\u0430',
 'CONSUMED':'\u0417\u0430\u043a\u0440\u044b\u0442\u0430','RELEASED':'\u041e\u0441\u0432\u043e\u0431\u043e\u0436\u0434\u0435\u043d\u0430'}

def clean(v):return ILLEGAL.sub('',str(v if v is not None else ''))[:32760]
def parse_utc(value):
    if value in (None,''):return None
    if isinstance(value,datetime):dt=value
    elif isinstance(value,date):dt=datetime(value.year,value.month,value.day,tzinfo=UTC)
    else:
        try:dt=datetime.fromisoformat(str(value).strip().replace('Z','+00:00'))
        except (TypeError,ValueError):return None
    try:return (dt.replace(tzinfo=UTC) if dt.tzinfo is None else dt).astimezone(UTC)
    except (OverflowError,ValueError):return None

def local_time(value):
    dt=parse_utc(value)
    if dt is None:return None
    try:return dt.astimezone(MSK)
    except (OverflowError,ValueError):return None

def excel_time(value):
    dt=local_time(value);return dt.replace(tzinfo=None) if dt else None

def stamp(value):
    dt=local_time(value);return dt.strftime('%d.%m.%Y %H:%M:%S')+' \u041c\u0421\u041a' if dt else ''

def iso(value):
    dt=parse_utc(value);return dt.isoformat() if dt else None

def qty(value):
    if value is None or isinstance(value,bool):raise ValueError('MISSING_QUANTITY')
    try:n=Decimal(str(value))
    except (InvalidOperation,ValueError,TypeError) as exc:raise ValueError('INVALID_QUANTITY') from exc
    if not n.is_finite() or n<0:raise ValueError('INVALID_QUANTITY')
    return n

def number(value):
    n=Decimal(str(value));return int(n) if n==n.to_integral_value() else float(n)

def stock_backed(physical,new,picking,other=0,unreserved=0):
    """Reporting split only. Holds first, then PICKING. No reservation is modified.
    The canonical safe sellable amount must not be increased by this split.
    """
    p,n,k,h,u=map(qty,(physical,new,picking,other,unreserved))
    bh=min(p,h);left=p-bh;bk=min(left,k);left-=bk;bn=min(left,n)
    return {'new':bn,'picking':bk,'reserved':bn+bk,'other':bh,
      'unbacked_orders':n-bn+k-bk+u,'unbacked_other':h-bh,'logical_new':n,'logical_picking':k}

def event_label(r):
    code=str(r.get('event_type') or '').upper();return LABELS.get(code,code or '\u041e\u043f\u0435\u0440\u0430\u0446\u0438\u044f')

def public_row(source):
    r=dict(source)
    for k in ('created_at','server_created_at','first_scan_at','last_scan_at','reservation_created_at','consumed_at','released_at'):
        val=r.get(k);r[k]=iso(val);r[k+'_msk']=stamp(val)
    r['event_label']=event_label(r);r['status_label']=STATUS.get(str(r.get('status') or '').upper(),clean(r.get('status')))
    return r

def physical_row(source):
    s=dict(source);q=qty(s['quantity']);src=bool(s.get('source_location_id'));dst=bool(s.get('destination_location_id'))
    code=str(s.get('operation_type') or 'STOCK').upper();reason=clean(s.get('reason'));idem=clean(s.get('idempotency_key'));comment=clean(s.get('comment'))
    if reason.startswith('FBS_CLEAN_RESET') or idem.startswith('FBS-CLEAN-RESET'):code='FBS_RESET'
    elif reason.startswith('CANCEL_RECEIVING:') or idem.startswith('cancel-receiving-'):code='RECEIPT_CANCEL'
    elif reason.startswith('EXTERNAL_PICK_TASK_WRITE_OFF:'):code='EXTERNAL_PICK_WRITE_OFF'
    elif code=='RECEIVING' and (idem.upper().startswith('PVZ:') or reason.upper().startswith('PVZ_RECEIVING:') or '\u041f\u0420\u0418\u0401\u041c\u041a\u0410 \u041f\u0412\u0417' in comment.upper()):code='PVZ_RECEIVING'
    technical=(not src and not dst) or (code=='RECONCILIATION' and src and dst and s.get('source_location_id')==s.get('destination_location_id') and (s.get('source_container_id') or '')==(s.get('destination_container_id') or ''))
    r={**s,'kind':'STOCK','event_type':code,'technical':technical,
      'server_created_at':s.get('created_at'),'physical_delta':number(q if dst and not src else -q if src and not dst else 0),
      'reserve_delta':0,'quantity':number(q),'actor':s.get('performed_by'),'name':s.get('sku_name') or s.get('name') or '',
      'counts_in_physical':str(s.get('status') or 'COMPLETED')!='CANCELLED'}
    return public_row(r)

def reservation_row(source):
    r=dict(source);q=qty(r['quantity'])
    r.update(kind='RESERVE',event_type='RESERVATION_STATE',technical=True,created_at=r.get('updated_at') or r.get('reservation_created_at'),
      quantity=number(q),physical_delta=0,reserve_delta=0,counts_in_physical=False,actor='SYSTEM',server_created_at=None,
      first_scan_at=None,last_scan_at=None,name=r.get('sku_name') or '',comment=r.get('note') or '',
      time_note='\u041f\u043e\u0441\u043b\u0435\u0434\u043d\u0435\u0435 \u043e\u0431\u043d\u043e\u0432\u043b\u0435\u043d\u0438\u0435 \u0437\u0430\u043f\u0438\u0441\u0438; \u043d\u0435 \u0434\u0430\u0442\u0430 \u0441\u0431\u043e\u0440\u043a\u0438')
    return public_row(r)

def history_workbook(rows,*,client='',warehouse='',mode='MOVEMENTS',generated_at=None):
    """Streaming XLSX, bounded memory. No hidden 3000-row truncation or formulas from input."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment,Font,PatternFill
    from openpyxl.utils import get_column_letter
    wb=Workbook(write_only=True)
    title='\u0414\u0432\u0438\u0436\u0435\u043d\u0438\u044f \u0442\u043e\u0432\u0430\u0440\u0430' if mode=='MOVEMENTS' else '\u0422\u0435\u0445\u043d\u0438\u0447\u0435\u0441\u043a\u0438\u0439 \u0436\u0443\u0440\u043d\u0430\u043b'
    ws=wb.create_sheet(title);ws.freeze_panes='E2';ws.sheet_view.showGridLines=False
    headers=['\u0414\u0430\u0442\u0430/\u0432\u0440\u0435\u043c\u044f (\u041c\u0421\u041a)','\u0414\u0430\u0442\u0430 (\u041c\u0421\u041a)','\u0422\u0438\u043f','\u0421\u0442\u0430\u0442\u0443\u0441','\u0428\u041a','\u0410\u0440\u0442\u0438\u043a\u0443\u043b','\u0420\u0430\u0437\u043c\u0435\u0440','\u0422\u043e\u0432\u0430\u0440','\u041a\u043e\u043b-\u0432\u043e','\u0424\u0438\u0437. \u0438\u0437\u043c\u0435\u043d\u0435\u043d\u0438\u0435','\u041e\u0442\u043a\u0443\u0434\u0430','\u041a\u0443\u0434\u0430','\u0417\u0430\u043a\u0430\u0437/\u043e\u0441\u043d\u043e\u0432\u0430\u043d\u0438\u0435','\u0421\u043e\u0442\u0440\u0443\u0434\u043d\u0438\u043a','\u041a\u043e\u043c\u043c\u0435\u043d\u0442\u0430\u0440\u0438\u0439','\u041f\u0435\u0440\u0432\u044b\u0439 \u0441\u043a\u0430\u043d (\u041c\u0421\u041a)','\u041f\u043e\u0441\u043b\u0435\u0434\u043d\u0438\u0439 \u0441\u043a\u0430\u043d (\u041c\u0421\u041a)','\u0418\u0441\u0442\u043e\u0447\u043d\u0438\u043a / ID','\u041f\u043e\u044f\u0441\u043d\u0435\u043d\u0438\u0435 \u0432\u0440\u0435\u043c\u0435\u043d\u0438','\u041d\u0430\u0437\u0432\u0430\u043d\u0438\u0435 \u043f\u0440\u0438\u0451\u043c\u043a\u0438','\u0423\u0447\u0438\u0442\u044b\u0432\u0430\u0435\u0442\u0441\u044f \u0432 \u0444\u0438\u0437\u0438\u0447\u0435\u0441\u043a\u043e\u043c']
    widths=[25,15,29,19,23,23,13,42,12,18,29,29,29,22,57,25,25,42,48,35,25]
    if mode=='TECHNICAL':
        headers+=['\u0420\u0435\u0437\u0435\u0440\u0432 \u0441\u043e\u0437\u0434\u0430\u043d (\u041c\u0421\u041a)','\u0420\u0435\u0437\u0435\u0440\u0432 \u0438\u0441\u043f\u043e\u043b\u044c\u0437\u043e\u0432\u0430\u043d (\u041c\u0421\u041a)','\u0420\u0435\u0437\u0435\u0440\u0432 \u043e\u0441\u0432\u043e\u0431\u043e\u0436\u0434\u0451\u043d (\u041c\u0421\u041a)'];widths += [25,25,25]
    for i,w in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=w
    ws.auto_filter.ref=f'A1:{get_column_letter(len(headers))}{len(rows)+1}'
    ws.row_dimensions[1].height=42
    body_font=Font(name='Calibri',size=11,color='18344A')
    head_font=Font(name='Calibri',bold=True,color='FFFFFF')
    wrap=Alignment(wrap_text=True,vertical='top')
    def cells(values,heading=False):
        out=[]
        for i,v in enumerate(values,1):
            c=WriteOnlyCell(ws,value=clean(v) if isinstance(v,str) else v)
            if isinstance(v,str):c.data_type='s'
            c.font=head_font if heading else body_font;c.alignment=wrap
            if heading:c.fill=PatternFill('solid',fgColor='18344A')
            elif i in (1,16,17,22,23,24):c.number_format='dd.mm.yyyy hh:mm:ss'
            elif i==2:c.number_format='dd.mm.yyyy'
            elif i in (5,6,7,13,18):c.number_format='@'
            elif i in (9,10):c.number_format='#,##0.###;[Red]-#,##0.###;0'
            out.append(c)
        return out
    ws.append(cells(headers,True))
    for row_index,r in enumerate(rows,2):
        ws.row_dimensions[row_index].height=max(32,min(90,15*max((len(clean(r.get('comment')))+55)//56,(len(clean(r.get('name')))+39)//40,2)))
        dt=excel_time(r.get('created_at'))
        src=' \u00b7 '.join(clean(r.get(x)) for x in ('source_location','source_container') if r.get(x));dst=' \u00b7 '.join(clean(r.get(x)) for x in ('destination_location','destination_container') if r.get(x))
        ref=clean(r.get('reference'));receipt=ref if r.get('event_type') in ('RECEIVING','RECEIVING_IMPORT','PVZ_RECEIVING') else ''
        vals=[dt,dt.date() if dt else None,event_label(r),r.get('status_label') or STATUS.get(r.get('status'),r.get('status')),r.get('barcode'),r.get('article'),r.get('size'),r.get('name'),r.get('quantity'),r.get('physical_delta'),src,dst,ref,r.get('actor'),r.get('comment') or r.get('reason'),excel_time(r.get('first_scan_at')),excel_time(r.get('last_scan_at')),r.get('line_id') or r.get('reservation_id') or r.get('operation_id'),r.get('time_note') or '\u0412\u0440\u0435\u043c\u044f \u0441\u043a\u043b\u0430\u0434\u0441\u043a\u043e\u0439 \u043e\u043f\u0435\u0440\u0430\u0446\u0438\u0438',receipt,'\u0414\u0430' if r.get('counts_in_physical') else '\u041d\u0435\u0442']
        if mode=='TECHNICAL':vals += [excel_time(r.get(k)) for k in ('reservation_created_at','consumed_at','released_at')]
        ws.append(cells(vals))
    info=wb.create_sheet('\u041e\u0431 \u043e\u0442\u0447\u0451\u0442\u0435');info.sheet_view.showGridLines=False
    info.column_dimensions['A'].width=26;info.column_dimensions['B'].width=100
    notes=[('\u041a\u043b\u0438\u0435\u043d\u0442',client),('\u0421\u043a\u043b\u0430\u0434',warehouse),('\u0421\u0444\u043e\u0440\u043c\u0438\u0440\u043e\u0432\u0430\u043d',stamp(generated_at or datetime.now(UTC))),('\u0420\u0435\u0436\u0438\u043c',mode),('\u0421\u0442\u0440\u043e\u043a',len(rows)),('\u0412\u0440\u0435\u043c\u044f','Europe/Moscow. UTC \u0432 \u0431\u0430\u0437\u0435 \u043d\u0435 \u0438\u0437\u043c\u0435\u043d\u0451\u043d.'),('\u0418\u0441\u0442\u043e\u0447\u043d\u0438\u043a','FulfillCore: inventory_operations, inventory_operation_lines; fc_stock_reservations (TECHNICAL).'),('\u0411\u0435\u0437\u043e\u043f\u0430\u0441\u043d\u043e\u0441\u0442\u044c','\u0422\u043e\u043b\u044c\u043a\u043e \u0447\u0442\u0435\u043d\u0438\u0435. MAIN, FBS, \u0440\u0435\u0437\u0435\u0440\u0432\u044b \u0438 \u043f\u0440\u0438\u0451\u043c\u043a\u0438 \u043d\u0435 \u0438\u0437\u043c\u0435\u043d\u044f\u044e\u0442\u0441\u044f.'),('\u0422\u0435\u0445\u0436\u0443\u0440\u043d\u0430\u043b','\u0421\u043e\u0441\u0442\u043e\u044f\u043d\u0438\u0435 \u0440\u0435\u0437\u0435\u0440\u0432\u0430 \u043d\u0435 \u043d\u043e\u0432\u0430\u044f \u0441\u0431\u043e\u0440\u043a\u0430 \u0438 \u043d\u0435 \u0441\u043f\u0438\u0441\u0430\u043d\u0438\u0435. \u041f\u0443\u0441\u0442\u043e\u0435 \u0432\u0440\u0435\u043c\u044f \u0441\u043a\u0430\u043d\u0430 = \u043d\u0435\u0442 \u0441\u0432\u044f\u0437\u0430\u043d\u043d\u043e\u0439 \u0437\u0430\u043f\u0438\u0441\u0438.')]
    for row_index,pair in enumerate(notes,1):
        info.row_dimensions[row_index].height=40 if row_index>=6 else 26
        line=[]
        for v in pair:
            c=WriteOnlyCell(info,value=v)
            if isinstance(v,str):c.data_type='s'
            c.alignment=wrap;line.append(c)
        info.append(line)
    stream=io.BytesIO();wb.save(stream);return stream.getvalue()
