"""Read-only FBS export presentation. Does not allocate, reserve, publish or write stock."""
from __future__ import annotations
import io
import re
import math
from collections import defaultdict
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from app.services.r145f24_history import stock_backed, stamp as msk_stamp
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET
ILLEGAL_CHARACTERS_RE = re.compile(r"[\x00-\x08\x0b-\x0c\x0e-\x1f]")

D = Decimal
STATUS = {
    'ACTIVE':'\u0412 \u043f\u0440\u043e\u0434\u0430\u0436\u0435 WB',
    'PAUSED':'\u0421\u0442\u043e\u043f \u043f\u0440\u043e\u0434\u0430\u0436',
    'PENDING':'\u041e\u0436\u0438\u0434\u0430\u0435\u0442 \u043e\u0431\u043d\u043e\u0432\u043b\u0435\u043d\u0438\u044f WB',
    'STALE':'\u0421\u0432\u0435\u0440\u043a\u0430 WB \u0443\u0441\u0442\u0430\u0440\u0435\u043b\u0430',
    'UNVERIFIED':'WB \u043d\u0435 \u0441\u0432\u0435\u0440\u0435\u043d',
    'MISMATCH':'\u0420\u0430\u0441\u0445\u043e\u0436\u0434\u0435\u043d\u0438\u0435 \u0441 WB',
    'ERROR':'\u041e\u0448\u0438\u0431\u043a\u0430 \u0441\u0432\u0435\u0440\u043a\u0438 WB',
    'EMPTY':'\u041d\u0430 WB \u043d\u043e\u043b\u044c',
}
SUMMARY_HEADERS = [
 '\u0410\u0440\u0442\u0438\u043a\u0443\u043b','\u0420\u0430\u0437\u043c\u0435\u0440','\u0428\u041a',
 '\u0424\u0438\u0437\u0438\u0447\u0435\u0441\u043a\u0438 FBS, \u0448\u0442.',
 '\u0421\u0432\u043e\u0431\u043e\u0434\u043d\u043e \u0434\u043b\u044f \u043f\u0440\u043e\u0434\u0430\u0436\u0438, \u0448\u0442.',
 '\u0420\u0435\u0437\u0435\u0440\u0432 \u0432 \u043f\u0440\u0435\u0434\u0435\u043b\u0430\u0445 FBS, \u0448\u0442.',
 '\u0420\u0435\u0437\u0435\u0440\u0432: \u043d\u043e\u0432\u044b\u0435, \u0448\u0442.',
 '\u0420\u0435\u0437\u0435\u0440\u0432: \u043d\u0430 \u0441\u0431\u043e\u0440\u043a\u0435, \u0448\u0442.',
 '\u041f\u0440\u043e\u0447\u0438\u0435 \u0443\u0434\u0435\u0440\u0436\u0430\u043d\u0438\u044f, \u0448\u0442.',
 '\u0417\u0430\u043a\u0430\u0437\u044b \u0431\u0435\u0437 \u0442\u043e\u0432\u0430\u0440\u0430 / \u0440\u0435\u0437\u0435\u0440\u0432\u0430, \u0448\u0442.',
 '\u0421\u0432\u043e\u0431\u043e\u0434\u043d\u043e, \u043d\u043e \u0441\u0442\u043e\u043f \u043f\u0440\u043e\u0434\u0430\u0436, \u0448\u0442.',
 '\u0421\u0442\u0430\u0442\u0443\u0441 \u043f\u0440\u043e\u0434\u0430\u0436\u0438','\u041f\u0440\u0438\u043c\u0435\u0447\u0430\u043d\u0438\u0435']

def number(value):
    if value is None or isinstance(value, bool):
        raise ValueError('FBS_EXPORT_QUANTITY_MISSING')
    try:
        n = D(str(value))
    except (InvalidOperation, ValueError, TypeError) as exc:
        raise ValueError('FBS_EXPORT_QUANTITY_INVALID') from exc
    if not n.is_finite() or n < 0:
        raise ValueError('FBS_EXPORT_QUANTITY_INVALID')
    return n

def numeric(n):
    n = D(str(n))
    return int(n) if n == n.to_integral_value() else float(n)

def text(value):
    return ILLEGAL_CHARACTERS_RE.sub('', str(value if value is not None else ''))[:32760]

def stamp(value):
    return msk_stamp(value)

def prepare(physical_rows, truth_map, sales, *, client='', warehouse='', search='', sale_tab='ALL', generated_at=None):
    """Collapse canonical SKU identities BEFORE adding reservations, not per address."""
    if sale_tab not in ('ALL','ACTIVE','PAUSED','ATTENTION'):raise ValueError('BAD_SALE_FILTER')
    if len(physical_rows)>20000:raise ValueError('FBS_EXPORT_TOO_MANY_PRODUCTS')
    groups={}
    by_sale={str(r['barcode']):r for r in sales.get('rows',[])}
    critical=('physical','storage_only','ledger_sale_reserved','ledger_picking_reserved',
      'sale_reserved','picking_reserved','accounting_pending','external_pick_reserved','available_to_sell')
    for source in physical_rows:
        ids=[str(x) for x in source.get('sku_ids',[])]
        if not ids:raise ValueError('FBS_EXPORT_SKU_MISSING')
        for sid in ids:
            if sid not in truth_map:raise ValueError('FBS_EXPORT_TRUTH_MISSING')
            t=truth_map[sid]
            for name in critical:number(t.get(name))
            identity=tuple(sorted(str(x) for x in t.get('identity_sku_ids',[])))
            if not identity or sid not in identity:raise ValueError('FBS_EXPORT_BAD_IDENTITY')
            g=groups.setdefault(identity,{'truth':t,'articles':set(),'sizes':set(),'codes':set(),'places':{}})
            if any(number(g['truth'][k])!=number(t[k]) for k in critical):raise ValueError('FBS_EXPORT_INCONSISTENT_IDENTITY')
            g['articles'].add(text(source.get('article')));g['sizes'].add(text(source.get('size')))
            if source.get('barcode'):g['codes'].add(text(source['barcode']))
            for p in source.get('locations',[]):
                if str(p.get('sku_id'))!=sid:continue
                key=(sid,str(p.get('location_id') or ''),str(p.get('container_id') or ''))
                amount=number(p.get('quantity'))
                if key in g['places'] and number(g['places'][key]['quantity'])!=amount:
                    raise ValueError('FBS_EXPORT_DUPLICATE_PLACE_MISMATCH')
                g['places'][key]=dict(p)
    result=[];placement=[]
    for identity,g in groups.items():
        t=g['truth'];sale_rows=[by_sale[c] for c in sorted(g['codes']) if c in by_sale]
        # A missing policy result must never be interpreted as 'sales enabled'.
        if g['codes'] and len(sale_rows)!=len(g['codes']):raise ValueError('FBS_EXPORT_SALE_POLICY_MISSING')
        paused=any(bool(x.get('paused')) for x in sale_rows) or any(s in set(sales.get('paused_sku_ids',[])) for s in identity)
        states={x.get('status','UNVERIFIED') for x in sale_rows}
        status='PAUSED' if paused else ('ACTIVE' if states=={'ACTIVE'} else next((x for x in ('ERROR','MISMATCH','PENDING','STALE','UNVERIFIED','EMPTY') if x in states),'UNVERIFIED'))
        if sale_tab=='ACTIVE' and status!='ACTIVE':continue
        if sale_tab=='PAUSED' and not paused:continue
        if sale_tab=='ATTENTION' and (paused or status=='ACTIVE'):continue
        article=', '.join(sorted(g['articles']));size=', '.join(sorted(g['sizes']));code=', '.join(sorted(g['codes']))
        places=list(g['places'].values())
        hay=' '.join([article,size,code]+[text(p.get(k)) for p in places for k in ('location','container','route_section','route_row','route_position')]).casefold()
        if search.strip() and search.strip().casefold() not in hay:continue
        physical=max(D(0),number(t['physical'])-number(t['storage_only']))
        new=number(t['ledger_sale_reserved']);picking=number(t['ledger_picking_reserved'])
        unreserved=max(D(0),number(t['sale_reserved'])-new)+max(D(0),number(t['picking_reserved'])-picking)
        other=number(t['accounting_pending'])+number(t['external_pick_reserved'])
        available=number(t['available_to_sell'])
        warnings=[]
        expected=max(D(0),physical-new-picking-unreserved-other)
        if available!=expected:raise ValueError('FBS_EXPORT_CANONICAL_FORMULA_CHANGED')
        if unreserved:warnings.append('\u0415\u0441\u0442\u044c \u0437\u0430\u043a\u0430\u0437\u044b \u0431\u0435\u0437 \u043e\u0444\u043e\u0440\u043c\u043b\u0435\u043d\u043d\u043e\u0433\u043e \u0440\u0435\u0437\u0435\u0440\u0432\u0430')
        if new+picking+other>physical:warnings.append('\u0420\u0435\u0437\u0435\u0440\u0432/\u0443\u0434\u0435\u0440\u0436\u0430\u043d\u0438\u044f \u0431\u043e\u043b\u044c\u0448\u0435 \u0444\u0438\u0437\u0438\u0447\u0435\u0441\u043a\u043e\u0433\u043e \u043e\u0441\u0442\u0430\u0442\u043a\u0430; \u043d\u0443\u0436\u043d\u0430 \u0441\u0432\u0435\u0440\u043a\u0430')
        if sum((number(p['quantity']) for p in places),D(0))!=physical:
            warnings.append('\u0420\u0430\u0437\u043c\u0435\u0449\u0435\u043d\u0438\u0435 FBS \u0438 \u0443\u0447\u0451\u0442\u043d\u044b\u0439 \u0438\u0442\u043e\u0433 \u0440\u0430\u0441\u0445\u043e\u0434\u044f\u0442\u0441\u044f; \u043d\u0443\u0436\u043d\u0430 \u0441\u0432\u0435\u0440\u043a\u0430')
        backed=stock_backed(physical,new,picking,other,unreserved)
        if backed['unbacked_other']:warnings.append('UNBACKED_OTHER='+str(backed['unbacked_other']))
        result.append({'identity':identity,'article':article,'size':size,'barcode':code,'codes':sorted(g['codes']),
          'physical':physical,'free':D(0) if paused else available,'new':backed['new'],'picking':backed['picking'],'reserved':backed['reserved'],
          'other':backed['other'],'unreserved':backed['unbacked_orders'],'stopped_free':available if paused else D(0),
          'status':status,'warning':'; '.join(warnings)})
        for p in places:placement.append([article,size,code,numeric(number(p['quantity'])),text(p.get('location')),text(p.get('container')),text(p.get('route_section')),text(p.get('route_row')),text(p.get('route_position'))])
    result.sort(key=lambda r:(r['article'],r['size'],r['barcode']))
    placement.sort(key=lambda r:(r[0],r[1],r[2],r[4],r[5]))
    if len(placement)>200000:raise ValueError('FBS_EXPORT_TOO_MANY_PLACES')
    selected={c for r in result for c in r['codes']}
    wb={}
    for s in sales.get('rows',[]):
        if s.get('barcode') not in selected or s.get('chrt_id') is None:continue
        key=str(s['chrt_id'])
        if key not in wb:wb[key]={**s,'codes':set()}
        elif any(wb[key].get(k)!=s.get(k) for k in ('target','confirmed_amount','checked_at','status')):
            raise ValueError('FBS_EXPORT_WB_SIZE_MISMATCH')
        wb[key]['codes'].update(s.get('siblings') or [s['barcode']])
    return {'rows':result,'placements':placement,'wb':list(wb.values()),'client':client,'warehouse':warehouse,
      'generated_at':generated_at or datetime.now(timezone.utc),'wb_warehouse_id':sales.get('wb_warehouse_id'),
      'search':search,'sale_tab':sale_tab,'configuration_error':sales.get('configuration_error','')}

def _formula_cache(data,cached):
    """Store cached values for the explicit SUM/subtraction formulas generated below."""
    ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    dest=io.BytesIO()
    with ZipFile(io.BytesIO(data)) as src, ZipFile(dest,'w',ZIP_DEFLATED,compresslevel=3) as out:
        for item in src.infolist():
            raw=src.read(item.filename)
            vals=cached.get(item.filename)
            if vals:
                root=ET.fromstring(raw)
                for c in root.findall('.//m:c',ns):
                    if c.get('r') not in vals:continue
                    assert c.find('m:f',ns) is not None
                    v=c.find('m:v',ns)
                    if v is None:v=ET.SubElement(c,'{'+ns['m']+'}v')
                    v.text=str(vals[c.get('r')]);c.attrib.pop('t',None)
                raw=ET.tostring(root,encoding='utf-8',xml_declaration=True)
            out.writestr(item,raw)
    return dest.getvalue()

def workbook_bytes(report):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    wb=Workbook();wb.remove(wb.active);cache={}
    dark='18344A';teal='E6F3F1';muted='667687';green='186350';purple='6747A0'
    def make(title,headers,widths):
        ws=wb.create_sheet(title);ws.sheet_view.showGridLines=False
        n=len(headers);ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=n)
        ws.cell(1,1,text(title));ws.cell(1,1).font=Font(name='Calibri',size=19,bold=True,color=dark)
        ws.row_dimensions[1].height=32
        ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=n)
        ws.cell(2,1,text(report['client']+' / '+report['warehouse']+' | '+stamp(report['generated_at'])))
        ws.cell(2,1).font=Font(size=11,color=muted)
        ws.merge_cells(start_row=3,start_column=1,end_row=3,end_column=n)
        ws.cell(3,1,text('FBS | '+report.get('search','')+' | '+report.get('sale_tab','ALL')))
        ws.cell(3,1).font=Font(size=10,color=muted)
        for i,h in enumerate(headers,1):
            c=ws.cell(5,i,h);c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor=dark);c.alignment=Alignment(wrap_text=True,vertical='center')
            ws.column_dimensions[get_column_letter(i)].width=widths[i-1]
        ws.row_dimensions[5].height=56;ws.freeze_panes='D6'
        ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        ws.print_title_rows='1:5'
        cache['xl/worksheets/sheet'+str(len(wb.worksheets))+'.xml']={}
        return ws
    def put(ws,row,values):
        for col,v in enumerate(values,1):
            c=ws.cell(row,col)
            if isinstance(v,(D,int,float)) and not isinstance(v,bool):
                c.value=numeric(v);c.number_format='#,##0.###';c.font=Font(color=green,size=11)
            else:c.value=text(v);c.data_type='s';c.font=Font(color=green,size=11);c.number_format='@'
            c.alignment=Alignment(vertical='top',wrap_text=True)
        ws.row_dimensions[row].height=34
    def formula(ws,cell,expr,val):
        c=ws[cell];c.value=expr;c.number_format='#,##0.###';c.font=Font(color='000000',size=11,bold=True)
        idx=wb.worksheets.index(ws)+1;cache[f'xl/worksheets/sheet{idx}.xml'][cell]=numeric(val)
    def finish(ws,count,numeric_cols):
        end=5+count;ws.auto_filter.ref=f'A5:{get_column_letter(ws.max_column)}{max(5,end)}'
        total=end+2;put(ws,total,['\u0418\u0422\u041e\u0413\u041e'])
        for col in numeric_cols:
            letter=get_column_letter(col)
            vals=[ws.cell(r,col) for r in range(6,end+1)]
            sheetcache=cache[f'xl/worksheets/sheet{wb.worksheets.index(ws)+1}.xml']
            val=sum((D(str(sheetcache[c.coordinate] if c.data_type=='f' else c.value or 0)) for c in vals),D(0))
            formula(ws,f'{letter}{total}',f'=SUM({letter}6:{letter}{end})' if count else '=0',val)
        for c in ws[total]:c.fill=PatternFill('solid',fgColor=teal);c.border=Border(top=Side(style='thin',color='B6CBCE'))
        ws.row_dimensions[total].height=25
    ws=make('\u041e\u0441\u0442\u0430\u0442\u043a\u0438 FBS',SUMMARY_HEADERS,[28,12,23,16,19,16,16,18,17,17,20,25,52])
    for n,r in enumerate(report['rows'],6):
        put(ws,n,[r['article'],r['size'],r['barcode'],r['physical'],r['free'],None,r['new'],r['picking'],r['other'],r['unreserved'],r['stopped_free'],STATUS.get(r['status'],r['status']),r['warning']])
        formula(ws,f'F{n}',f'=SUM(G{n}:H{n})',r['reserved'])
        for c in (5,6):ws.cell(n,c).fill=PatternFill('solid',fgColor=teal)
        if r['warning']:ws.cell(n,13).fill=PatternFill('solid',fgColor='FFF0D6')
    finish(ws,len(report['rows']),list(range(4,12)))
    ws=make('\u0420\u0430\u0437\u043c\u0435\u0449\u0435\u043d\u0438\u0435',[
        SUMMARY_HEADERS[0],SUMMARY_HEADERS[1],SUMMARY_HEADERS[2],SUMMARY_HEADERS[3],
        '\u042f\u0447\u0435\u0439\u043a\u0430','\u041a\u043e\u0440\u043e\u0431','\u0421\u0435\u043a\u0446\u0438\u044f','\u0420\u044f\u0434','\u041c\u0435\u0441\u0442\u043e \u0432 \u0440\u044f\u0434\u0443'],[28,12,23,19,26,36,18,12,18])
    for n,r in enumerate(report['placements'],6):put(ws,n,r)
    finish(ws,len(report['placements']),[4])
    ws=make('\u0421\u0432\u0435\u0440\u043a\u0430 WB',[
      'chrtId','\u0428\u041a \u0440\u0430\u0437\u043c\u0435\u0440\u0430','\u0426\u0435\u043b\u044c WMS, \u0448\u0442.',
      '\u041d\u0430 WB \u043f\u043e \u0441\u0432\u0435\u0440\u043a\u0435, \u0448\u0442.','\u0420\u0430\u0437\u043d\u0438\u0446\u0430: \u0446\u0435\u043b\u044c - WB',
      '\u0414\u0430\u0442\u0430 \u0441\u0432\u0435\u0440\u043a\u0438','\u0421\u0442\u0430\u0442\u0443\u0441',
      '\u0421\u043a\u043b\u0430\u0434 WB','\u041f\u0440\u0438\u043c\u0435\u0447\u0430\u043d\u0438\u0435'],[18,42,18,19,20,28,28,18,50])
    for n,r in enumerate(report['wb'],6):
        put(ws,n,[str(r['chrt_id']),', '.join(sorted(r['codes'])),r.get('target'),r.get('confirmed_amount'),None,
           stamp(r.get('checked_at')),STATUS.get(r.get('status'),text(r.get('status'))),str(report.get('wb_warehouse_id') or ''),text(r.get('error'))])
        if r.get('target') is not None and r.get('confirmed_amount') is not None:
            formula(ws,f'E{n}',f'=C{n}-D{n}',D(str(r['target']))-D(str(r['confirmed_amount'])))
    # Missing observations are blank, not zero. Do not add a misleading total of a partial WB observation set.
    finish(ws,len(report['wb']),[])
    ws=make('\u041f\u043e\u044f\u0441\u043d\u0435\u043d\u0438\u044f',['\u041f\u043e\u043b\u0435','\u041a\u0430\u043a \u0447\u0438\u0442\u0430\u0442\u044c'],[36,120])
    notes=[
      ('\u0421\u0432\u043e\u0431\u043e\u0434\u043d\u043e \u0434\u043b\u044f \u043f\u0440\u043e\u0434\u0430\u0436\u0438','\u0421\u0432\u043e\u0431\u043e\u0434\u043d\u044b\u0439 \u043e\u0441\u0442\u0430\u0442\u043e\u043a WMS \u043f\u043e\u0441\u043b\u0435 \u0440\u0435\u0437\u0435\u0440\u0432\u043e\u0432 \u0438 \u0443\u0434\u0435\u0440\u0436\u0430\u043d\u0438\u0439. \u041f\u0440\u0438 \u0441\u0442\u043e\u043f\u0435 \u043f\u0440\u043e\u0434\u0430\u0436 \u2014 0. \u042d\u0442\u043e \u043d\u0435 \u043f\u043e\u0434\u0442\u0432\u0435\u0440\u0436\u0434\u0435\u043d\u0438\u0435 \u043f\u0443\u0431\u043b\u0438\u043a\u0430\u0446\u0438\u0438 \u043d\u0430 WB.'),
      ('\u0420\u0435\u0437\u0435\u0440\u0432','\u0420\u0435\u0437\u0435\u0440\u0432 \u0442\u043e\u0432\u0430\u0440\u0430 \u043d\u0435 \u0431\u043e\u043b\u044c\u0448\u0435 \u0444\u0438\u0437\u0438\u0447\u0435\u0441\u043a\u043e\u0433\u043e FBS. \u041f\u0440\u0438 FBS = 0 \u0440\u0435\u0437\u0435\u0440\u0432 \u0442\u043e\u0432\u0430\u0440\u0430 = 0. \u041f\u0440\u0438 \u0434\u0435\u0444\u0438\u0446\u0438\u0442\u0435 \u0434\u043b\u044f \u0440\u0430\u0437\u0431\u0438\u0432\u043a\u0438 \u043e\u0442\u0447\u0451\u0442\u0430 \u0441\u043d\u0430\u0447\u0430\u043b\u0430 \u0443\u0447\u0442\u0435\u043d\u044b \u043f\u0440\u043e\u0447\u0438\u0435 \u0443\u0434\u0435\u0440\u0436\u0430\u043d\u0438\u044f, \u0437\u0430\u0442\u0435\u043c \u0441\u0431\u043e\u0440\u043a\u0430, \u0437\u0430\u0442\u0435\u043c \u043d\u043e\u0432\u044b\u0435. \u042d\u0442\u043e \u043e\u0442\u0447\u0451\u0442\u043d\u0430\u044f \u0440\u0430\u0437\u0431\u0438\u0432\u043a\u0430, \u0430 \u043d\u0435 \u043f\u0435\u0440\u0435\u043d\u0430\u0437\u043d\u0430\u0447\u0435\u043d\u0438\u0435 \u0440\u0435\u0437\u0435\u0440\u0432\u043e\u0432.'),
      ('\u0417\u0430\u043a\u0430\u0437\u044b \u0431\u0435\u0437 \u0442\u043e\u0432\u0430\u0440\u0430 / \u0440\u0435\u0437\u0435\u0440\u0432\u0430','\u041b\u043e\u0433\u0438\u0447\u0435\u0441\u043a\u0438\u0435 \u043e\u0431\u044f\u0437\u0430\u0442\u0435\u043b\u044c\u0441\u0442\u0432\u0430, \u043d\u0435 \u043f\u043e\u043a\u0440\u044b\u0442\u044b\u0435 \u0444\u0438\u0437\u0438\u0447\u0435\u0441\u043a\u0438\u043c FBS, \u0438\u043b\u0438 \u043f\u043e\u0442\u0440\u0435\u0431\u043d\u043e\u0441\u0442\u044c \u0431\u0435\u0437 \u043e\u0444\u043e\u0440\u043c\u043b\u0435\u043d\u043d\u043e\u0433\u043e \u0440\u0435\u0437\u0435\u0440\u0432\u0430. \u041d\u0435 \u0443\u0432\u0435\u043b\u0438\u0447\u0438\u0432\u0430\u044e\u0442 \u043e\u0441\u0442\u0430\u0442\u043e\u043a. \u0421\u0430\u043c\u0438 \u0437\u0430\u043a\u0430\u0437\u044b \u0438 \u0431\u0435\u0437\u043e\u043f\u0430\u0441\u043d\u044b\u0439 \u0441\u0432\u043e\u0431\u043e\u0434\u043d\u044b\u0439 \u043e\u0441\u0442\u0430\u0442\u043e\u043a \u043d\u0435 \u0438\u0437\u043c\u0435\u043d\u044f\u044e\u0442\u0441\u044f.'),
      ('\u041f\u0440\u043e\u0447\u0438\u0435 \u0443\u0434\u0435\u0440\u0436\u0430\u043d\u0438\u044f','\u041e\u0436\u0438\u0434\u0430\u043d\u0438\u0435 \u043f\u043e\u0434\u0442\u0432\u0435\u0440\u0436\u0434\u0435\u043d\u0438\u044f \u0441\u043f\u0438\u0441\u0430\u043d\u0438\u044f \u0438 \u0434\u0440\u0443\u0433\u0438\u0435 \u0434\u0435\u0439\u0441\u0442\u0432\u0443\u044e\u0449\u0438\u0435 \u0443\u0434\u0435\u0440\u0436\u0430\u043d\u0438\u044f. \u0423\u0436\u0435 \u0432\u044b\u0447\u0442\u0435\u043d\u044b \u0438\u0437 \u0441\u0432\u043e\u0431\u043e\u0434\u043d\u043e\u0433\u043e.'),
      ('\u0424\u0438\u0437\u0438\u0447\u0435\u0441\u043a\u0438 FBS','\u041e\u0431\u0449\u0438\u0439 \u0444\u0438\u0437\u0438\u0447\u0435\u0441\u043a\u0438\u0439 \u043e\u0441\u0442\u0430\u0442\u043e\u043a \u043c\u0438\u043d\u0443\u0441 STORAGE_ONLY \u043f\u043e \u0435\u0434\u0438\u043d\u043e\u043c\u0443 \u0443\u0447\u0451\u0442\u0443 WMS. \u0420\u0430\u0441\u0445\u043e\u0436\u0434\u0435\u043d\u0438\u044f \u0441 \u043b\u0438\u0441\u0442\u043e\u043c \u0440\u0430\u0437\u043c\u0435\u0449\u0435\u043d\u0438\u044f \u043f\u043e\u043c\u0435\u0447\u0430\u044e\u0442\u0441\u044f, \u043d\u043e \u043d\u0435 \u0438\u0441\u043f\u0440\u0430\u0432\u043b\u044f\u044e\u0442 \u0443\u0447\u0451\u0442.'),
      ('\u0421\u0432\u0435\u0440\u043a\u0430 WB','\u041e\u0434\u043d\u0430 \u0441\u0442\u0440\u043e\u043a\u0430 = \u043e\u0434\u0438\u043d chrtId \u043d\u0430 \u0441\u043a\u043b\u0430\u0434\u0435 WB. \u041a\u043e\u043b\u0438\u0447\u0435\u0441\u0442\u0432\u043e \u0440\u0430\u0437\u043c\u0435\u0440\u0430 \u0446\u0435\u043b\u0438\u043a\u043e\u043c, \u0434\u0430\u0436\u0435 \u0435\u0441\u043b\u0438 \u0447\u0430\u0441\u0442\u044c \u0428\u041a \u043d\u0435 \u0432\u043e\u0448\u043b\u0430 \u0432 \u0444\u0438\u043b\u044c\u0442\u0440. \u041d\u0435\u0438\u0437\u0432\u0435\u0441\u0442\u043d\u043e\u0435 \u043a\u043e\u043b\u0438\u0447\u0435\u0441\u0442\u0432\u043e \u2014 \u043f\u0443\u0441\u0442\u0430\u044f \u044f\u0447\u0435\u0439\u043a\u0430, \u043d\u0435 0. \u042d\u043a\u0441\u043f\u043e\u0440\u0442 \u043d\u0435 \u0437\u0430\u043f\u0443\u0441\u043a\u0430\u0435\u0442 \u0441\u0432\u0435\u0440\u043a\u0443 \u0447\u0435\u0440\u0435\u0437 API WB.'),
      ('\u0418\u0441\u0442\u043e\u0447\u043d\u0438\u043a','FulfillCore: inventory_balances; fc_stocks_v2; fc_stock_reservations; marketplace_orders; fc_external_pick_allocations; fc_wb_verified_stock. Read-only snapshot.'),
      ('\u041d\u0430\u0441\u0442\u0440\u043e\u0439\u043a\u0430 WB',report.get('configuration_error') or '\u041e\u0448\u0438\u0431\u043e\u043a \u043d\u0435\u0442'),
    ]
    for n,r in enumerate(notes,6):put(ws,n,r);ws.row_dimensions[n].height=58
    ws.freeze_panes='A6'
    buf=io.BytesIO();wb.save(buf)
    return _formula_cache(buf.getvalue(),cache)
