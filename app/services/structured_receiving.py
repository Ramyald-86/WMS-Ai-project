from __future__ import annotations

import io
import re
import uuid
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook


STRUCTURED_SHEET_NAME = "Адресное хранение"
STRUCTURED_HEADERS = (
    "Секция",
    "Ряд",
    "Паллета",
    "Лимит",
    "Короб",
    "Партия",
    "Источник",
    "Принят на хранение",
    "Дней",
    "Объём м3",
    "Начислено руб",
    "Артикул",
    "Размер",
    "ШК товара",
    "Остаток",
)
_REQUIRED_HEADERS = {"Секция", "Паллета", "Короб", "ШК товара", "Остаток"}


def _norm_header(value: Any) -> str:
    text = str(value or "").strip().lower().replace("ё", "е")
    return re.sub(r"[^0-9a-zа-я]+", "", text)


_HEADER_ALIASES = {
    "Ряд": {_norm_header(x) for x in ["Ряд", "Ряд на палете", "Row", "Сторона палеты"]},
    "Секция": {_norm_header(x) for x in ["Секция", "Секция хранения", "Section"]},
    "Паллета": {_norm_header(x) for x in ["Паллета", "Палета", "Паллет", "Pallet"]},
    "Лимит": {_norm_header(x) for x in ["Лимит", "Лимит коробов", "Вместимость"]},
    "Короб": {_norm_header(x) for x in ["Короб", "Коробка", "Box", "ШК короба"]},
    "Партия": {_norm_header(x) for x in ["Партия", "Batch"]},
    "Источник": {_norm_header(x) for x in ["Источник", "Source"]},
    "Принят на хранение": {_norm_header(x) for x in ["Принят на хранение", "Дата приемки", "Дата приёмки", "Принят"]},
    "Дней": {_norm_header(x) for x in ["Дней", "Дни"]},
    "Объём м3": {_norm_header(x) for x in ["Объём м3", "Объем м3", "Объём", "Объем"]},
    "Начислено руб": {_norm_header(x) for x in ["Начислено руб", "Начислено", "Сумма"]},
    "Артикул": {_norm_header(x) for x in ["Артикул", "Артикул продавца", "Vendor code"]},
    "Размер": {_norm_header(x) for x in ["Размер", "Size"]},
    "ШК товара": {_norm_header(x) for x in ["ШК товара", "Штрихкод товара", "Штрихкод", "ШК", "Barcode"]},
    "Остаток": {_norm_header(x) for x in ["Остаток", "Количество", "Кол-во", "Qty"]},
}
_ALIAS_TO_HEADER = {alias: canonical for canonical, aliases in _HEADER_ALIASES.items() for alias in aliases}


def _cell_text(cell: Any) -> str:
    value = getattr(cell, "value", cell)
    if value in (None, ""):
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        number_format = str(getattr(cell, "number_format", "") or "")
        if re.fullmatch(r"0+", number_format):
            return str(value).zfill(len(number_format))
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            number_format = str(getattr(cell, "number_format", "") or "")
            if re.fullmatch(r"0+", number_format):
                return str(int(value)).zfill(len(number_format))
            return str(int(value))
        return format(value, ".15g")
    return str(value).strip()


def _decimal_or_none(value: Any) -> Decimal | None:
    if value in (None, ""):
        return None
    text = str(value).strip().replace(" ", "").replace(",", ".")
    if not text:
        return None
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return None


def _positive_int_or_none(value: Any) -> int | None:
    number = _decimal_or_none(value)
    if number is None:
        return None
    if number != number.to_integral_value() or number <= 0:
        return None
    return int(number)


def _find_header_row(ws) -> tuple[int, dict[str, int]]:
    best: tuple[int, dict[str, int]] | None = None
    for row_number in range(1, min(ws.max_row, 20) + 1):
        mapping: dict[str, int] = {}
        for cell in ws[row_number]:
            normalized = _norm_header(cell.value)
            canonical = _ALIAS_TO_HEADER.get(normalized)
            if canonical and canonical not in mapping:
                mapping[canonical] = int(cell.column)
        if _REQUIRED_HEADERS.issubset(mapping):
            return row_number, mapping
        if best is None or len(mapping) > len(best[1]):
            best = (row_number, mapping)
    found = sorted((best or (0, {}))[1])
    missing = sorted(_REQUIRED_HEADERS.difference(found))
    raise ValueError(
        "Не найден заголовок структурной приёмки. "
        f"Обязательные колонки: {', '.join(sorted(_REQUIRED_HEADERS))}. "
        f"Не найдены: {', '.join(missing) if missing else 'не удалось распознать строку заголовков'}."
    )


def _row_value(ws, row_number: int, mapping: dict[str, int], header: str) -> Any:
    column = mapping.get(header)
    return ws.cell(row=row_number, column=column) if column else None


def parse_structured_workbook(content: bytes) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    if not content:
        raise ValueError("Excel-файл пуст")
    wb = load_workbook(io.BytesIO(content), read_only=False, data_only=True)
    ws = wb[STRUCTURED_SHEET_NAME] if STRUCTURED_SHEET_NAME in wb.sheetnames else wb.active
    header_row, mapping = _find_header_row(ws)
    rows: list[dict[str, Any]] = []
    if ws.max_row - header_row > 100_000:
        raise ValueError("В одной структурной приёмке допускается не более 100 000 строк")

    for row_number in range(header_row + 1, ws.max_row + 1):
        raw = {header: _row_value(ws, row_number, mapping, header) for header in STRUCTURED_HEADERS}
        values = {header: _cell_text(cell) if cell is not None else "" for header, cell in raw.items()}
        if not any(values.values()):
            continue

        section = values["Секция"].strip()
        side_text = values["Ряд"].strip()
        side = _positive_int_or_none(side_text)
        pallet = values["Паллета"].strip()
        box = values["Короб"].strip()
        barcode = values["ШК товара"].strip()
        article = values["Артикул"].strip()
        size = values["Размер"].strip()
        quantity = _decimal_or_none(getattr(raw["Остаток"], "value", None) if raw["Остаток"] is not None else None)
        pallet_limit = _positive_int_or_none(getattr(raw["Лимит"], "value", None) if raw["Лимит"] is not None else None)
        limit_text = values["Лимит"].strip()
        errors: list[str] = []
        warnings: list[str] = []

        if side_text and (side is None or side>999):
            errors.append("ряд должен быть целым числом от 1 до 999")
        if not section:
            errors.append("не заполнена секция")
        if not pallet:
            errors.append("не заполнена палета")
        if not box:
            errors.append("не заполнен короб")
        if limit_text and pallet_limit is None:
            errors.append("лимит палеты должен быть положительным целым числом")

        product_fields_present = bool(barcode or article or size or values["Остаток"].strip())
        is_empty_box = not product_fields_present
        if not is_empty_box:
            if not barcode:
                errors.append("есть товарные данные, но нет ШК товара")
            if quantity is None:
                errors.append("не указано корректное количество")
            elif quantity < 0:
                errors.append("количество не может быть отрицательным")
            elif quantity == 0:
                warnings.append("остаток 0 — товарная строка будет пропущена, структура короба сохранится")
            elif quantity != quantity.to_integral_value():
                errors.append("количество товара должно быть целым числом")
        else:
            quantity = None

        rows.append({
            "row_number": row_number,
            "section": section,
            "row_on_pallet": side,
            "pallet": pallet,
            "pallet_limit": pallet_limit,
            "box": box,
            "batch": values["Партия"].strip(),
            "source": values["Источник"].strip(),
            "accepted_at": values["Принят на хранение"].strip(),
            "days": values["Дней"].strip(),
            "volume_m3": values["Объём м3"].strip(),
            "accrued_rub": values["Начислено руб"].strip(),
            "article": article,
            "size": size,
            "barcode": barcode,
            "quantity": quantity,
            "is_empty_box": is_empty_box,
            "parse_errors": errors,
            "errors": list(errors),
            "warnings": warnings,
        })

    if not rows:
        raise ValueError("После строки заголовков в Excel нет данных")

    box_sides = defaultdict(set)
    side_boxes = defaultdict(set)
    for row in rows:
        if row.get('row_on_pallet'):
            box_sides[row['box']].add(row['row_on_pallet'])
            side_boxes[(row['section'],row['pallet'],row['row_on_pallet'])].add(row['box'])
    for row in rows:
        if len(box_sides[row['box']])>1:
            row['parse_errors'].append('у одного короба указаны разные ряды')
        if len(side_boxes[(row['section'],row['pallet'],row.get('row_on_pallet'))])>8:
            row['parse_errors'].append('в одном ряду палеты не больше 8 коробов; укажите следующий ряд')
    box_places: dict[str, set[tuple[str, str]]] = defaultdict(set)
    pallet_limits: dict[tuple[str, str], set[int]] = defaultdict(set)
    pallet_boxes: dict[tuple[str, str], set[str]] = defaultdict(set)
    for row in rows:
        if row["box"]:
            box_places[row["box"]].add((row["section"], row["pallet"]))
        key = (row["section"], row["pallet"])
        if row["pallet_limit"] is not None:
            pallet_limits[key].add(int(row["pallet_limit"]))
        if row["box"]:
            pallet_boxes[key].add(row["box"])

    box_errors = {
        box: f"короб {box} указан сразу в нескольких местах: " + ", ".join(f"секция {s} / палета {p}" for s, p in sorted(places))
        for box, places in box_places.items() if len(places) > 1
    }
    pallet_errors: dict[tuple[str, str], str] = {}
    for key, limits in pallet_limits.items():
        if len(limits) > 1:
            pallet_errors[key] = f"для секции {key[0]} / палеты {key[1]} указаны разные лимиты: {', '.join(map(str, sorted(limits)))}"
            continue
        if limits:
            limit = next(iter(limits))
            box_count = len(pallet_boxes.get(key, set()))
            if box_count > limit:
                pallet_errors[key] = f"в секции {key[0]} / палете {key[1]} коробов {box_count}, а лимит {limit}"

    for row in rows:
        if row["box"] in box_errors:
            row["parse_errors"].append(box_errors[row["box"]])
        pallet_error = pallet_errors.get((row["section"], row["pallet"]))
        if pallet_error:
            row["parse_errors"].append(pallet_error)
        row["errors"] = list(dict.fromkeys(row["parse_errors"]))

    summary = structured_rows_summary(rows)
    summary.update({
        "sheet_name": ws.title,
        "header_row": header_row,
        "recognized_headers": [header for header in STRUCTURED_HEADERS if header in mapping],
        "missing_optional_headers": [header for header in STRUCTURED_HEADERS if header not in mapping],
    })
    return rows, summary


def structured_rows_summary(rows: list[dict[str, Any]]) -> dict[str, Any]:
    sections = {row["section"] for row in rows if row.get("section")}
    pallets = {(row["section"], row["pallet"]) for row in rows if row.get("section") and row.get("pallet")}
    boxes = {row["box"] for row in rows if row.get("box")}
    product_rows = [row for row in rows if not row.get("is_empty_box") and (row.get("quantity") or Decimal("0")) > 0]
    zero_product_rows = [row for row in rows if not row.get("is_empty_box") and row.get("quantity") == 0]
    empty_boxes = {
        row["box"] for row in rows
        if row.get("box") and row.get("is_empty_box")
    }
    total = Decimal("0")
    for row in product_rows:
        qty = row.get("quantity")
        if qty is not None and qty > 0:
            total += Decimal(qty)
    error_rows = sum(1 for row in rows if row.get("errors"))
    warning_rows = sum(1 for row in rows if row.get("warnings"))
    return {
        "row_count": len(rows),
        "product_row_count": len(product_rows),
        "zero_product_row_count": len(zero_product_rows),
        "section_count": len(sections),
        "pallet_count": len(pallets),
        "box_count": len(boxes),
        "empty_box_count": len(empty_boxes),
        "quantity_total": float(total),
        "error_count": error_rows,
        "warning_count": warning_rows,
    }


def _code_token(value: Any) -> str:
    text = re.sub(r"\s+", "-", str(value or "").strip().upper())
    text = re.sub(r"[^0-9A-ZА-ЯЁ_-]", "", text)
    text = re.sub(r"-{2,}", "-", text).strip("-_")
    return text or "0"


def structured_section_code(section: Any) -> str:
    return f"SEC-{_code_token(section)}"


def structured_section_name(section: Any) -> str:
    return f"Секция {str(section).strip()}"


def structured_location_code(section: Any, pallet: Any) -> str:
    return f"С{_code_token(section)}-П{_code_token(pallet)}"


def structured_location_name(section: Any, pallet: Any) -> str:
    return f"Секция {str(section).strip()} · Палета {str(pallet).strip()}"


def structured_location_barcode(warehouse_id: uuid.UUID | str, code: str) -> str:
    digest = uuid.uuid5(uuid.NAMESPACE_URL, f"fulfillcore:{warehouse_id}:{code}").hex[:16].upper()
    return f"FCLOC-{digest}"


def natural_key(value: Any) -> list[tuple[int, Any]]:
    return [
        (1, int(part)) if part.isdigit() else (0, part.casefold())
        for part in re.split(r"(\d+)", str(value or ""))
    ]


def install_row_schema(db):
    from sqlalchemy import text
    db.execute(text("""CREATE TABLE IF NOT EXISTS fc_structured_row_routes (
        receipt_row_id varchar(36) PRIMARY KEY, row_on_pallet integer NOT NULL CHECK(row_on_pallet BETWEEN 1 AND 999))"""))

    db.execute(text("""CREATE TABLE IF NOT EXISTS fc_structured_receipt_policy (
        receipt_id varchar(36) PRIMARY KEY, stock_class varchar(8) NOT NULL CHECK(stock_class IN ('MAIN','FBS')))"""))

    db.execute(text("""CREATE TABLE IF NOT EXISTS fc_structured_receipt_archives (
        receipt_id varchar(36) PRIMARY KEY,
        archived_at timestamptz NOT NULL DEFAULT now())"""))
