"""Pure, bounded input validation and deterministic address planning for r64."""
from __future__ import annotations
import hashlib
import io
import json
import re
import zipfile
from decimal import Decimal, InvalidOperation

MAX_ROWS = 10000
MAX_BYTES = 8 * 1024 * 1024


def natural(value):
    return tuple((1, int(x)) if x.isdigit() else (0, x) for x in re.split(r'(\d+)', str(value or '').upper()))


def quantity(value, label='Количество'):
    try:
        q = Decimal(str(value).replace(',', '.'))
    except (InvalidOperation, ValueError):
        raise ValueError(f'{label}: укажите целое положительное число')
    if not q.is_finite() or q != q.to_integral_value() or not 0 < q <= 10000000:
        raise ValueError(f'{label}: допустимо целое число от 1 до 10 000 000')
    return int(q)


def barcode(value):
    if isinstance(value, bool):
        raise ValueError('ШК товара не распознан')
    if isinstance(value, (int, float)):
        if value != int(value):
            raise ValueError('ШК товара должен быть текстом без дробной части')
        value = str(int(value))
    value = re.sub(r'^\](?:E0|C0|C1)', '', str(value or '').strip(), flags=re.I)
    if not re.fullmatch(r'(?:\d{8}|\d{12}|\d{13}|\d{14})', value):
        raise ValueError('Нужен обычный ШК товара (8, 12, 13 или 14 цифр), не ЧЗ. В Excel используйте текстовый формат')
    return value[1:] if len(value) == 14 and value.startswith('0') else value


def text_value(value, max_length=255):
    s = str(value or '').strip()
    if len(s) > max_length or any(ord(c) < 32 for c in s):
        raise ValueError('Слишком длинное значение или управляющие символы в тексте')
    return s


def fingerprint(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode()).hexdigest()


def normalize_package(raw):
    if not isinstance(raw, dict) or raw.get('schema') != 'fulfillcore-address-receiving-package-v1':
        raise ValueError('Это не файл приёмки терминала v1.6')
    package_id = text_value(raw.get('packageId'), 120)
    if not re.fullmatch(r'[A-Za-z0-9_.:-]{8,120}', package_id):
        raise ValueError('Нет корректного номера приёмки')
    boxes = raw.get('boxes')
    if not isinstance(boxes, list) or not 1 <= len(boxes) <= 2000:
        raise ValueError('В приёмке должно быть от 1 до 2000 коробов/мест')
    out = []; seen = set(); item_count = 0
    for b in boxes:
        if not isinstance(b,dict): raise ValueError('Некорректная запись короба')
        loose = b.get('loose') is True
        sscc = '' if loose else text_value(b.get('sscc'), 120).upper()
        if not loose and (not sscc or re.search(r'\s', sscc)):
            raise ValueError('Пустой или некорректный ШК короба')
        if sscc and sscc in seen:
            raise ValueError(f'Короб {sscc} повторяется в одной приёмке')
        seen.add(sscc)
        loc = text_value(b.get('locationCode'), 120)
        if not loc:
            section = quantity(b.get('section'), 'Секция')
            pallet = quantity(b.get('pallet'), 'Паллетоместо')
            if section > 9999 or pallet > 9999:
                raise ValueError('Для большого/буквенного адреса выберите адрес из WMS')
            loc = f'SEC-{section}-PP-{pallet:03d}'
        else:
            section = None; pallet = None
        row = quantity(b.get('row', 1), 'Ряд на палете')
        if row > 2:
            raise ValueError('На палете допустимы ряды 1 и 2')
        items = b.get('items')
        if not isinstance(items, list) or not items:
            raise ValueError(f'Пустой короб {sscc or loc}: заполните или удалите его из черновика')
        merged = {}
        for i in items:
            if not isinstance(i,dict): raise ValueError('Некорректная запись товара')
            item_count += 1
            if item_count > MAX_ROWS:
                raise ValueError('В одной приёмке не более 10 000 строк')
            code = barcode(i.get('barcode'))
            article = text_value(i.get('article')); size = text_value(i.get('size'))
            purpose = i.get('stockType', 'FBS')
            if purpose not in ('FBS', 'FBO'):
                raise ValueError(f'{code}: неизвестное назначение FBS/FBO')
            key = (code, article, size, purpose)
            merged[key] = merged.get(key, 0) + quantity(i.get('qty'), f'{code}: количество')
        out.append({'sscc': sscc, 'loose': loose, 'locationCode': loc,
                    'autoSection': section, 'autoPallet': pallet, 'row': row,
                    'batch': text_value(b.get('batch')), 'source': text_value(b.get('source')),
                    'items': [{'barcode': k[0], 'article': k[1], 'size': k[2], 'stockType': k[3], 'qty': v}
                              for k, v in sorted(merged.items())]})
    out.sort(key=lambda b: (b['locationCode'], b['sscc'], fingerprint(b)))
    return {'packageId': package_id, 'boxes': out}


def workbook(content):
    if not content or len(content) > MAX_BYTES:
        raise ValueError('Размер Excel должен быть от 1 байта до 8 МБ')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            if len(z.infolist()) > 500 or sum(x.file_size for x in z.infolist()) > 50 * 1024 * 1024:
                raise ValueError('Excel слишком большой после распаковки')
        from openpyxl import load_workbook
        return load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except ValueError:
        raise
    except Exception:
        raise ValueError('Не удалось открыть файл .xlsx. Пересохраните его в Excel')


def sheet_rows(ws):
    if ws.max_column and ws.max_column > 100:
        raise ValueError('В листе слишком много столбцов (максимум 100)')
    out = []
    for row in ws.iter_rows():
        if len(out) >= MAX_ROWS + 20:
            raise ValueError('В листе не более 10 000 товарных строк')
        values = []
        for c in row:
            if c.data_type == 'f':
                raise ValueError(f'{ws.title}!{c.coordinate}: замените формулу её значением')
            value = c.value
            if isinstance(value, (int, float)) and not isinstance(value, bool) and value == int(value) and re.fullmatch(r'0{8,14}', c.number_format or ''):
                value = str(int(value)).zfill(len(c.number_format))
            values.append(value)
        out.append(values)
    return out


def read_transfer_workbook(content):
    aliases = {'barcode': {'ШК', 'ШКТОВАРА', 'ШТРИХКОД', 'ШТРИХКОДТОВАРА', 'BARCODE', 'EAN', 'EAN13', 'GTIN'},
               'qty': {'КОЛИЧЕСТВО', 'КОЛВО', 'КОЛИЧЕСТВОШТ', 'ОСТАТОК', 'ОСТАТОКШТ', 'QTY', 'QUANTITY'},
               'location': {'ЯЧЕЙКА', 'АДРЕС', 'КОДЯЧЕЙКИ', 'ШКЯЧЕЙКИ', 'LOCATION', 'LOCATIONCODE'},
               'container': {'КОРОБ', 'ШККОРОБА', 'SSCC', 'BOX', 'CONTAINER'}}
    wb = workbook(content)
    try:
        rows = sheet_rows(wb['Перевод в FBS'] if 'Перевод в FBS' in wb.sheetnames else wb.active)
        mapping = {}; start = None
        for ix, row in enumerate(rows[:20]):
            m = {}
            for j, val in enumerate(row):
                h = re.sub(r'[^A-ZА-Я0-9]', '', str(val or '').upper().replace('Ё', 'Е'))
                for name, names in aliases.items():
                    if h in names:
                        if name in m:
                            raise ValueError(f'Столбец {val} указан дважды')
                        m[name] = j
            if 'barcode' in m and 'qty' in m:
                mapping = m; start = ix + 1; break
        if start is None:
            raise ValueError('Нужны столбцы «ШК товара» и «Количество» (или «Остаток»); «Ячейка» и «Короб» необязательны')
        out = []
        for ix, row in enumerate(rows[start:], start=start + 1):
            if not any(x not in (None, '') for x in row):
                continue
            def cell(k):
                j = mapping.get(k)
                return row[j] if j is not None and j < len(row) else None
            try:
                out.append({'row': ix, 'barcode': barcode(cell('barcode')), 'qty': quantity(cell('qty')),
                            'location': text_value(cell('location'), 120), 'container': text_value(cell('container'), 120)})
            except ValueError as exc:
                raise ValueError(f'Строка {ix}: {exc}')
        if not out:
            raise ValueError('В Excel нет товаров')
        return out
    finally:
        wb.close()


def read_terminal_workbook(content):
    wb = workbook(content)
    try:
        if 'Технический пакет' not in wb.sheetnames:
            raise ValueError('Нужен Excel, сохранённый кнопкой экспорта приёмки терминала')
        rows = sheet_rows(wb['Технический пакет'])
        if not rows or rows[0][0] != 'JSON_CHUNK':
            raise ValueError('Повреждён служебный лист приёмки')
        try:
            raw = json.loads(''.join(str(r[0] or '') for r in rows[1:] if r))
        except Exception:
            raise ValueError('Повреждён пакет приёмки')
        normalized = normalize_package(raw)
        # Reject edits to visible sheets instead of silently importing stale hidden content.
        expected_items = sorted((str(b.get('sscc') or ''), str(i.get('stockType') or 'FBS'), str(i.get('article') or ''),
                                 str(i.get('size') or ''), barcode(i.get('barcode')), quantity(i.get('qty')))
                                for b in raw['boxes'] for i in b['items'])
        if 'Товары' not in wb.sheetnames or 'Короба' not in wb.sheetnames:
            raise ValueError('Отсутствуют листы «Товары» или «Короба»')
        actual_items = sorted((str(r[0] or ''), str(r[1] or ''), str(r[2] or ''), str(r[3] or ''), barcode(r[4]), quantity(r[5]))
                              for r in sheet_rows(wb['Товары'])[1:] if any(x not in (None, '') for x in r))
        expected_boxes = sorted((str(b.get('sscc') or ''), int(b['section']), int(b['row']), int(b['pallet'])) for b in raw['boxes'])
        actual_boxes = sorted((str(r[0] or ''), int(r[1]), int(r[2]), int(r[3]))
                              for r in sheet_rows(wb['Короба'])[1:] if any(x not in (None, '') for x in r))
        if actual_items != expected_items or actual_boxes != expected_boxes:
            raise ValueError('Листы Excel изменены после экспорта. Исправьте приёмку в терминале и выгрузите заново')
        return raw
    except (IndexError, TypeError, KeyError):
        raise ValueError('Повреждена структура Excel приёмки')
    finally:
        wb.close()


def allocate(requests, buckets, *, key=None):
    """Exact named addresses first; each unit allocated at most once across rows."""
    left = {b['key']: max(0, int(b['free'])) for b in buckets}
    rows = []; shortages = []; by_barcode = {}
    for b in buckets: by_barcode.setdefault(b['barcode'], []).append(b)
    ordered = sorted(enumerate(requests), key=lambda x: (-bool(x[1].get('container')), -bool(x[1].get('location')), x[0]))
    for index, req in ordered:
        need = req['qty']
        choices = [b for b in by_barcode.get(req['barcode'], []) if (not req.get('location') or req['location'].casefold() in (b.get('location_code', '').casefold(), b.get('location_barcode', '').casefold()))
                   and (not req.get('container') or req['container'].casefold() in (b.get('container_name', '').casefold(), b.get('container_barcode', '').casefold()))]
        choices.sort(key=key or (lambda b: (natural(b.get('location_code')), natural(b.get('container_name') or b.get('container_barcode')), b['key'])))
        for b in choices:
            take = min(need, left[b['key']])
            if take <= 0:
                continue
            rows.append({**{k: v for k, v in b.items() if k not in ('free', 'held', 'main')},
                         'qty': take, 'excel_row': req.get('row', index + 2)})
            left[b['key']] -= take; need -= take
            if need == 0:
                break
        if need:
            shortages.append({'row': req.get('row', index + 2), 'barcode': req['barcode'], 'missing': need,
                              'location': req.get('location', ''), 'container': req.get('container', '')})
    rows.sort(key=(lambda r:(key(r),r['excel_row'],r['key'])) if key else (lambda r: (natural(r.get('location_code')), natural(r.get('container_name')), r['excel_row'], r['key'])))
    return {'rows': rows, 'shortages': shortages, 'requested': sum(r['qty'] for r in requests),
            'allocated': sum(r['qty'] for r in rows)}
