"""Readable Excel backup of the native receiving draft. Restoring never posts inventory."""
import io
import json
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from app.services.workflow_rules import workbook,barcode,quantity,text_value

HEADERS=['ШК','Артикул','Размер','Название','Количество','Ячейка','Короб']

def normalize(data,c,w):
    if not isinstance(data,dict):raise ValueError('Повреждён резерв черновика')
    if data.get('schema')!='fulfillcore-native-draft-v1':raise ValueError('Это не резерв черновика FulfillCore')
    if data.get('client_id')!=str(c) or data.get('warehouse_id')!=str(w):raise ValueError('Черновик относится к другому клиенту или складу')
    draft=data.get('draft') or {}
    if not isinstance(draft,dict):raise ValueError('Повреждён черновик')
    rows=draft.get('rows')
    if not isinstance(rows,list) or not 1<=len(rows)<=10000:raise ValueError('В резерве должно быть 1–10 000 строк')
    draft_id=text_value(draft.get('id'),120)
    if not draft_id:raise ValueError('Нет номера черновика')
    cleaned=[]
    for r in rows:
        if not isinstance(r,dict):raise ValueError('Повреждена строка черновика')
        cleaned.append({k:text_value(r.get(k),255) for k in ('id','article','size','name','location_code','container_code','container_name','section','aisle','pallet','storage_display')} |
            {'barcode':barcode(r.get('barcode')),'quantity':quantity(r.get('quantity')),'product_found':r.get('product_found') is not False,'wb_found':bool(r.get('wb_found'))})
    if any(not r['location_code'] for r in cleaned):raise ValueError('В строке не указан адрес приёмки')
    kind=data.get('stock_class','MAIN')
    if kind not in ('MAIN','FBS'):raise ValueError('Некорректный раздел склада')
    return {'schema':'fulfillcore-native-draft-v1','client_id':str(c),'warehouse_id':str(w),
        'stock_class':kind,'reference':text_value(data.get('reference'),240),'performed_by':text_value(data.get('performed_by'),120),
        'draft':{'id':draft_id,'rows':cleaned}}

def visible(data):
    return [[r['barcode'],r['article'],r['size'],r['name'],r['quantity'],r['location_code'],r['container_code']] for r in data['draft']['rows']]

def dump(data):
    wb=Workbook();ws=wb.active;ws.title='Черновик';ws.append(HEADERS)
    for r in visible(data):
        ws.append(r)
        for i in (1,2,3,4,6,7):ws.cell(ws.max_row,i).data_type='s'
    for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='322B85')
    for col,width in zip('ABCDEFG',[24,24,18,44,16,26,30]):ws.column_dimensions[col].width=width
    for row in ws:
        for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    meta=wb.create_sheet('Резерв WMS');blob=json.dumps(data,ensure_ascii=False)
    for i in range(0,len(blob),30000):meta.append([blob[i:i+30000]]);meta.cell(meta.max_row,1).data_type='s'
    meta.sheet_state='hidden';out=io.BytesIO();wb.save(out);return out.getvalue()

def restore(content,c,w):
    wb=workbook(content)
    try:
        if 'Резерв WMS' not in wb or 'Черновик' not in wb:raise ValueError('В Excel нет резерва черновика WMS')
        raw=''.join(str(r[0] or '') for r in wb['Резерв WMS'].iter_rows(values_only=True))
        try:data=normalize(json.loads(raw),c,w)
        except (TypeError,KeyError,json.JSONDecodeError):raise ValueError('Повреждён резерв черновика')
        values=list(wb['Черновик'].iter_rows(values_only=True))
        if not values or list(values[0])!=HEADERS:raise ValueError('Изменены колонки резервной копии')
        actual=[[v if v is not None else '' for v in r] for r in values[1:]]
        if actual!=visible(data):raise ValueError('Видимые строки изменены. Для изменённой раскладки используйте импорт Excel в приёмке')
        return data
    finally:wb.close()
