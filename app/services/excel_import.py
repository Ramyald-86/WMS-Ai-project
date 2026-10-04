import hashlib
import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from fastapi import HTTPException
from openpyxl import load_workbook


@dataclass
class ParsedWorkbook:
    sheet_name: str
    header_row: int
    headers: list[str]
    mapping: dict[str, int]
    rows: list[dict[str, Any]]
    checksum: str


def _norm(value: Any) -> str:
    s = str(value or "").strip().lower().replace("ё", "е")
    s = re.sub(r"[\s\n\r\t]+", " ", s)
    s = s.replace("№", "номер")
    s = re.sub(r"[^a-zа-я0-9_ ]+", "", s)
    return s.strip()


SYNONYMS = {
    "box": [
        "короб", "короба", "номер короба", "название короба", "шк короба", "штрихкод короба",
        "box", "container", "контейнер",
    ],
    "location": ["ячейка", "ячейка хранения", "адрес", "адрес хранения", "location", "место хранения"],
    "barcode": ["шк", "штрихкод", "шк товара", "штрихкод товара", "barcode", "ean", "ean13"],
    "article": ["артикул", "арт", "article", "offer id", "offer_id", "артикул продавца"],
    "size": ["размер", "size"],
    "qty": ["количество", "колво", "кол во", "остаток", "qty", "quantity", "шт", "собрать шт"],
    "name": ["наименование", "название товара", "товар", "name", "номенклатура"],
    "sku_code": ["код sku", "sku", "sku_code", "код товара", "код номенклатуры"],
}


def _score_header(norm_header: str, field: str) -> int:
    if not norm_header:
        return 0
    # Сначала исключаем неоднозначность: "ШК короба" — это короб, не товарный barcode.
    if field == "barcode" and any(x in norm_header for x in ("короб", "контейнер")):
        return 0
    if field == "box" and any(x in norm_header for x in ("короб", "контейнер")):
        return 100
    best = 0
    for synonym in SYNONYMS[field]:
        ns = _norm(synonym)
        if norm_header == ns:
            best = max(best, 100)
        elif len(ns) >= 3 and ns in norm_header:
            best = max(best, 60 + min(len(ns), 20))
    return best


def detect_mapping(headers: list[Any]) -> dict[str, int]:
    normalized = [_norm(h) for h in headers]
    mapping: dict[str, int] = {}
    used: set[int] = set()
    # Сначала самые специфичные поля.
    for field in ("box", "location", "barcode", "article", "size", "qty", "name", "sku_code"):
        candidates = sorted(
            (( _score_header(h, field), idx) for idx, h in enumerate(normalized) if idx not in used),
            reverse=True,
        )
        if candidates and candidates[0][0] > 0:
            mapping[field] = candidates[0][1]
            used.add(candidates[0][1])
    return mapping


def _find_header_row(ws) -> tuple[int, list[Any], dict[str, int]]:
    best: tuple[int, list[Any], dict[str, int], int] | None = None
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
        values = list(row)
        mapping = detect_mapping(values)
        score = len(mapping) + (3 if "box" in mapping else 0) + (3 if "qty" in mapping else 0)
        if any(k in mapping for k in ("barcode", "article", "sku_code")):
            score += 3
        if best is None or score > best[3]:
            best = (row_idx, values, mapping, score)
    if not best or best[3] < 5:
        raise HTTPException(
            status_code=400,
            detail="Не удалось определить заголовки Excel. Нужны как минимум: короб, количество и ШК/артикул/SKU.",
        )
    return best[0], best[1], best[2]


def _text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    s = str(value).strip()
    return s or None


def _qty(value: Any) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        q = Decimal(str(value).replace(" ", "").replace(",", "."))
    except InvalidOperation:
        return None
    return q if q > 0 else None


def parse_workbook(raw: bytes, filename: str | None = None, max_rows: int = 100_000) -> ParsedWorkbook:
    if filename and filename.lower().endswith(".xls"):
        raise HTTPException(status_code=400, detail="Старый формат .xls пока не поддерживается. Сохраните файл как .xlsx.")
    try:
        wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Не удалось открыть Excel: {exc}") from exc
    ws = wb.active
    header_row, header_values, mapping = _find_header_row(ws)
    headers = [_text(v) or f"Колонка {i+1}" for i, v in enumerate(header_values)]

    if "box" not in mapping:
        raise HTTPException(status_code=400, detail="Не найдена колонка с названием/ШК короба")
    if "qty" not in mapping:
        raise HTTPException(status_code=400, detail="Не найдена колонка с количеством")
    if not any(k in mapping for k in ("barcode", "article", "sku_code")):
        raise HTTPException(status_code=400, detail="Не найдена колонка ШК товара, артикул или код SKU")

    rows: list[dict[str, Any]] = []
    for idx, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        if len(rows) >= max_rows:
            raise HTTPException(status_code=400, detail=f"В одном импорте поддерживается до {max_rows} строк")
        vals = list(row)
        def get(field: str):
            col = mapping.get(field)
            return vals[col] if col is not None and col < len(vals) else None

        box = _text(get("box"))
        qty = _qty(get("qty"))
        barcode = _text(get("barcode"))
        article = _text(get("article"))
        sku_code = _text(get("sku_code"))
        size = _text(get("size"))
        name = _text(get("name"))
        location = _text(get("location"))

        if not any([box, qty, barcode, article, sku_code, size, name, location]):
            continue
        rows.append({
            "row": idx,
            "box": box,
            "location": location,
            "barcode": barcode,
            "article": article,
            "sku_code": sku_code,
            "size": size,
            "name": name,
            "quantity": qty,
        })
    return ParsedWorkbook(
        sheet_name=ws.title,
        header_row=header_row,
        headers=headers,
        mapping=mapping,
        rows=rows,
        checksum=hashlib.sha256(raw).hexdigest(),
    )
