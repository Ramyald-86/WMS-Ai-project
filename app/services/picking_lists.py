from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from io import BytesIO
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


EPSILON = 1e-9


def _number(value: Any) -> float:
    try:
        return max(0.0, float(value or 0))
    except (TypeError, ValueError):
        return 0.0


def _row_key(order: dict[str, Any], stock: dict[str, Any], status: str) -> tuple:
    return (
        str(stock.get("location_id") or stock.get("location") or ""),
        str(stock.get("container_id") or stock.get("container") or ""),
        str(order.get("sku_id") or ""),
        str(order.get("article") or ""),
        str(order.get("size") or ""),
        str(order.get("barcode") or ""),
        str(order.get("name") or ""),
        status,
    )


def allocate_picking_rows(source_orders: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Allocate one WB order at a time against real warehouse stock buckets.

    The returned shape remains compatible with the existing order reservation logic:
    ``location``, ``container``, ``order_ids`` and ``status`` are preserved. Extra
    fields describe pallet/cell placement and whether the whole physical box becomes
    empty after every row in this picking list is collected.
    """

    bucket_state: dict[tuple, dict[str, Any]] = {}
    aggregated: dict[tuple, dict[str, Any]] = {}

    for order in source_orders or []:
        order_id = str(order.get("order_id") or "")
        stocks = [dict(x) for x in (order.get("stock") or []) if _number(x.get("quantity")) > EPSILON]

        for stock in stocks:
            # 0.7.9.16: one physical balance bucket is identified only by
            # SKU + location + container. The previous list-position component could
            # represent the same real stock twice when row ordering differed between
            # orders, causing the Excel list to claim stock in the wrong cell.
            bucket_id = (
                str(order.get("sku_id") or order.get("barcode") or order.get("article") or ""),
                str(stock.get("location_id") or stock.get("location") or ""),
                str(stock.get("container_id") or stock.get("container") or ""),
            )
            if bucket_id not in bucket_state:
                bucket_state[bucket_id] = {
                    "initial": _number(stock.get("quantity")),
                    "remaining": _number(stock.get("quantity")),
                    "stock": stock,
                }

        need = 1.0
        if not order.get("sku_id"):
            status = "SKU НЕ СОПОСТАВЛЕН"
            stock = {}
            key = _row_key(order, stock, status)
            row = aggregated.setdefault(key, {
                "location": "",
                "location_id": None,
                "location_name": "",
                "location_type": "",
                "container": "",
                "container_id": None,
                "container_barcode": "",
                "container_total_quantity": 0.0,
                "article": order.get("article") or "",
                "size": order.get("size") or "",
                "barcode": order.get("barcode") or "",
                "name": order.get("name") or "",
                "required_quantity": 0.0,
                "available_quantity": 0.0,
                "order_ids": [],
                "order_age_hours": [],
                "max_age_hours": 0.0,
                "sla_critical": False,
                "status": status,
                "under_zero": False,
            })
            row["required_quantity"] += need
            if order_id:
                row["order_ids"].append(order_id)
            age = _number(order.get("age_hours"))
            row["order_age_hours"].append(age)
            row["max_age_hours"] = max(_number(row.get("max_age_hours")), age)
            row["sla_critical"] = bool(row.get("sla_critical") or order.get("sla_critical"))
            continue

        for stock in stocks:
            if need <= EPSILON:
                break
            bucket_id = (
                str(order.get("sku_id") or order.get("barcode") or order.get("article") or ""),
                str(stock.get("location_id") or stock.get("location") or ""),
                str(stock.get("container_id") or stock.get("container") or ""),
            )
            state = bucket_state[bucket_id]
            if state["remaining"] <= EPSILON:
                continue
            take = min(need, state["remaining"])
            state["remaining"] -= take
            need -= take
            status = "OK"
            key = _row_key(order, stock, status)
            row = aggregated.setdefault(key, {
                "location": stock.get("location") or "",
                "location_id": stock.get("location_id"),
                "location_name": stock.get("location_name") or "",
                "location_type": stock.get("location_type") or "",
                "container": stock.get("container") or "",
                "container_id": stock.get("container_id"),
                "container_barcode": stock.get("container_barcode") or "",
                "container_total_quantity": _number(stock.get("container_total_quantity")),
                "article": order.get("article") or "",
                "size": order.get("size") or "",
                "barcode": order.get("barcode") or "",
                "name": order.get("name") or "",
                "required_quantity": 0.0,
                "available_quantity": state["initial"],
                "order_ids": [],
                "order_age_hours": [],
                "max_age_hours": 0.0,
                "sla_critical": False,
                "status": status,
                "under_zero": False,
            })
            row["required_quantity"] += take
            if order_id:
                row["order_ids"].append(order_id)
            age = _number(order.get("age_hours"))
            row["order_age_hours"].append(age)
            row["max_age_hours"] = max(_number(row.get("max_age_hours")), age)
            row["sla_critical"] = bool(row.get("sla_critical") or order.get("sla_critical"))

        if need > EPSILON:
            status = "НЕ ХВАТАЕТ"
            stock = {}
            key = _row_key(order, stock, status)
            row = aggregated.setdefault(key, {
                "location": "",
                "location_id": None,
                "location_name": "",
                "location_type": "",
                "container": "",
                "container_id": None,
                "container_barcode": "",
                "container_total_quantity": 0.0,
                "article": order.get("article") or "",
                "size": order.get("size") or "",
                "barcode": order.get("barcode") or "",
                "name": order.get("name") or "",
                "required_quantity": 0.0,
                "available_quantity": 0.0,
                "order_ids": [],
                "order_age_hours": [],
                "max_age_hours": 0.0,
                "sla_critical": False,
                "status": status,
                "under_zero": False,
            })
            row["required_quantity"] += need
            if order_id:
                row["order_ids"].append(order_id)
            age = _number(order.get("age_hours"))
            row["order_age_hours"].append(age)
            row["max_age_hours"] = max(_number(row.get("max_age_hours")), age)
            row["sla_critical"] = bool(row.get("sla_critical") or order.get("sla_critical"))

    rows = list(aggregated.values())

    # A box is marked "ПОД 0" only when this complete picking list removes every
    # stock unit currently stored in that physical box, including other SKUs.
    planned_by_container: dict[str, float] = defaultdict(float)
    total_by_container: dict[str, float] = {}
    for row in rows:
        container_id = str(row.get("container_id") or "")
        if not container_id or row.get("status") != "OK":
            continue
        planned_by_container[container_id] += _number(row.get("required_quantity"))
        total = _number(row.get("container_total_quantity"))
        if total > EPSILON:
            total_by_container[container_id] = max(total_by_container.get(container_id, 0.0), total)
    for row in rows:
        container_id = str(row.get("container_id") or "")
        total = total_by_container.get(container_id, 0.0)
        row["under_zero"] = bool(
            container_id
            and total > EPSILON
            and planned_by_container.get(container_id, 0.0) + EPSILON >= total
        )
        # Backward-compatible aliases used by older exports and diagnostics.
        row["quantity"] = row.get("required_quantity", 0.0)
        row["required"] = row.get("required_quantity", 0.0)
        row["available"] = row.get("available_quantity", 0.0)

    return sorted(rows, key=lambda row: (
        0 if row.get("sla_critical") else 1,
        -_number(row.get("max_age_hours")),
        1 if row.get("status") != "OK" else 0,
        str(row.get("location") or ""),
        str(row.get("container") or ""),
    ))



def _display_number(value: Any) -> int | float:
    number = _number(value)
    return int(number) if abs(number - round(number)) < EPSILON else round(number, 3)


def _named_location_part(name: Any, label: str) -> str:
    match = re.search(rf"{re.escape(label)}\s+([^·]+)", str(name or ""), flags=re.IGNORECASE)
    return str(match.group(1) if match else "").strip()


def _clean_location_part(value: Any) -> str:
    text = str(value or "").strip()
    if re.fullmatch(r"\d+", text):
        try:
            return str(int(text))
        except ValueError:
            pass
    return text


def _picking_location_dimensions(item: dict[str, Any]) -> dict[str, str]:
    """Turn the stored FulfillCore address into human warehouse dimensions.

    Location names are intentionally the source of truth for the printable hierarchy:
    structured receiving writes ``Секция N · Палета M`` while rack/cell builders write
    ``Проход · Стеллаж · Этаж · Ячейка``. New pallet-place addresses use the same
    convention. Unknown/legacy addresses fall back to a single ``Адрес`` column.
    """
    location = str(item.get("location") or "").strip()
    name = str(item.get("location_name") or "").strip()
    location_type = str(item.get("location_type") or "").strip().upper()
    dims = {
        "section": _clean_location_part(_named_location_part(name, "Секция")),
        "pallet_place": _clean_location_part(_named_location_part(name, "Палеточное место")),
        "pallet": _clean_location_part(_named_location_part(name, "Палета")),
        "aisle": _clean_location_part(_named_location_part(name, "Проход")),
        "rack": _clean_location_part(_named_location_part(name, "Стеллаж")),
        "floor": _clean_location_part(_named_location_part(name, "Этаж")),
        "cell": _clean_location_part(_named_location_part(name, "Ячейка")),
        "address": "",
    }

    # Safe fallbacks for older addresses whose visible name did not yet contain the
    # hierarchy. Do not invent section/rack data: only identify the concrete address.
    if location_type == "PALLET_PLACE" and not dims["pallet_place"]:
        dims["pallet_place"] = location
    elif location_type == "PALLET" and not dims["pallet"]:
        dims["pallet"] = location
    elif location_type == "CELL" and not dims["cell"]:
        dims["cell"] = location
    elif location and not any(dims[key] for key in ("section", "pallet_place", "pallet", "aisle", "rack", "floor", "cell")):
        dims["address"] = location
    return dims


def picking_list_xlsx(
    *,
    title: str,
    created_at: datetime,
    rows: list[dict[str, Any]],
    supply_ids: list[str] | None = None,
    order_count: int = 0,
    sheet_number: str | None = None,
) -> bytes:
    """Build a compact WB picking list with only warehouse columns that are used.

    A warehouse may use cells, racks, sections, pallets, pallet places, boxes, or a
    mixture of them. The export therefore derives the hierarchy from each real stock
    address and removes columns that are completely unused in this particular list.
    The check column is intentionally blank because it is the collector's input field.
    """

    export_rows: list[dict[str, Any]] = []
    for source in rows or []:
        item = dict(source)
        item["_location_dims"] = _picking_location_dimensions(item)
        export_rows.append(item)

    dimension_specs = [
        ("section", "Секция", 13),
        ("pallet_place", "Палеточное место", 18),
        ("pallet", "Палета", 13),
        ("aisle", "Проход", 12),
        ("rack", "Стеллаж", 13),
        ("floor", "Этаж / полка", 13),
        ("cell", "Ячейка", 16),
        ("address", "Адрес", 20),
    ]
    columns: list[dict[str, Any]] = []
    for key, header, width in dimension_specs:
        if any(str(item.get("_location_dims", {}).get(key) or "").strip() for item in export_rows):
            columns.append({"key": key, "header": header, "width": width, "kind": "location"})

    optional_specs = [
        ("container", "Короб", 18, "text"),
        ("article", "Артикул", 16, "text"),
        ("size", "Размер", 11, "center"),
        ("barcode", "ШК товара", 20, "text"),
        ("name", "Наименование", 35, "text"),
    ]
    for key, header, width, kind in optional_specs:
        if any(str(item.get(key) or "").strip() for item in export_rows):
            columns.append({"key": key, "header": header, "width": width, "kind": kind})

    if any(_number(item.get("max_age_hours")) > 0 for item in export_rows):
        columns.append({"key": "max_age_hours", "header": "Возраст, ч", "width": 12, "kind": "number"})
    if any(bool(item.get("sla_critical")) for item in export_rows):
        columns.append({"key": "sla_critical", "header": "SLA 300+", "width": 12, "kind": "center"})
    columns.extend([
        {"key": "required_quantity", "header": "Собрать, шт.", "width": 12, "kind": "number"},
        {"key": "available_quantity", "header": "Доступно", "width": 11, "kind": "number"},
    ])
    if any(bool(item.get("under_zero")) for item in export_rows):
        columns.append({"key": "under_zero", "header": "Под 0", "width": 16, "kind": "center"})
    if any(str(item.get("status") or "OK").upper() != "OK" for item in export_rows):
        columns.append({"key": "status", "header": "Статус", "width": 20, "kind": "status"})
    columns.append({"key": "check", "header": "✓", "width": 6, "kind": "check"})

    wb = Workbook()
    ws = wb.active
    ws.title = "Сборочный лист"
    ws.sheet_view.showGridLines = False

    navy = "26364A"
    blue = "2F75B5"
    pale_red = "F4CCCC"
    pale_yellow = "FFF2CC"
    white = "FFFFFF"
    thin_gray = Side(style="thin", color="B7C4D3")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    max_col = max(1, len(columns))
    last_col_letter = get_column_letter(max_col)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    ws["A1"] = "FULFILLCORE · СБОРОЧНЫЙ ЛИСТ WB"
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color=white)
    ws["A1"].fill = PatternFill("solid", fgColor=navy)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 29

    number = (sheet_number or "").strip() or created_at.strftime("%Y%m%d-%H%M%S")
    metadata = [
        ("№ сборочного листа", number),
        ("Название", title),
        ("Создан", created_at.astimezone().strftime("%d.%m.%Y %H:%M:%S")),
        ("Заказов WB", order_count),
        ("Поставки WB", ", ".join(supply_ids or []) or "—"),
    ]
    for row_index, (label, value) in enumerate(metadata, start=2):
        ws.cell(row_index, 1, label).font = Font(bold=True, color=navy)
        if max_col > 1:
            ws.cell(row_index, 2, value)
            if max_col > 2:
                ws.merge_cells(start_row=row_index, start_column=2, end_row=row_index, end_column=max_col)
            ws.cell(row_index, 2).alignment = Alignment(horizontal="left", wrap_text=True, vertical="top")
        else:
            ws.cell(row_index, 1, f"{label}: {value}")

    header_row = 8
    ws.freeze_panes = f"A{header_row + 1}"
    for col_index, spec in enumerate(columns, start=1):
        cell = ws.cell(header_row, col_index, spec["header"])
        cell.font = Font(bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(col_index)].width = spec["width"]
    ws.row_dimensions[header_row].height = 30

    check_col = next(index for index, spec in enumerate(columns, start=1) if spec["key"] == "check")
    under_zero_col = next((index for index, spec in enumerate(columns, start=1) if spec["key"] == "under_zero"), None)
    check_validation = DataValidation(type="list", formula1='"✓,+"', allow_blank=True)
    check_validation.promptTitle = "Собрано"
    check_validation.prompt = "Поставьте ✓ или + после сборки позиции"
    ws.add_data_validation(check_validation)

    for row_index, item in enumerate(export_rows, start=header_row + 1):
        status = str(item.get("status") or "OK")
        under_zero = bool(item.get("under_zero"))
        dims = item.get("_location_dims") or {}
        for col_index, spec in enumerate(columns, start=1):
            key = spec["key"]
            if spec["kind"] == "location":
                value = dims.get(key) or ""
            elif key == "required_quantity":
                value = _display_number(item.get("required_quantity"))
            elif key == "available_quantity":
                value = _display_number(item.get("available_quantity"))
            elif key == "under_zero":
                value = "ПОД 0\nЗАБРАТЬ КОРОБ" if under_zero else ""
            elif key == "sla_critical":
                value = "300+ ЧАСОВ" if item.get("sla_critical") else ""
            elif key == "status":
                value = status if status.upper() != "OK" else ""
            elif key == "check":
                value = ""
            else:
                value = item.get(key) or ""

            cell = ws.cell(row_index, col_index, value)
            cell.border = border
            cell.alignment = Alignment(
                horizontal="center" if spec["kind"] in {"center", "number", "check", "status"} else "left",
                vertical="center",
                wrap_text=True,
            )
            if spec["kind"] == "number":
                cell.font = Font(bold=True)
            if spec["kind"] in {"location", "text", "center"}:
                cell.number_format = "@"
                cell.quotePrefix = True

        if status.upper() != "OK":
            for col_index in range(1, max_col + 1):
                ws.cell(row_index, col_index).fill = PatternFill("solid", fgColor=pale_red)
        elif under_zero and under_zero_col:
            ws.cell(row_index, under_zero_col).fill = PatternFill("solid", fgColor=pale_yellow)
            ws.cell(row_index, under_zero_col).font = Font(bold=True, color="9C6500")
        elif row_index % 2 == 0:
            for col_index in range(1, max_col + 1):
                ws.cell(row_index, col_index).fill = PatternFill("solid", fgColor="F7FAFC")
        ws.row_dimensions[row_index].height = 44 if under_zero or status.upper() != "OK" else 25
        check_validation.add(ws.cell(row_index, check_col))

    end_row = max(header_row + 1, header_row + len(export_rows))
    ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{end_row}"
    check_letter = get_column_letter(check_col)
    ws.conditional_formatting.add(
        f"{check_letter}{header_row + 1}:{check_letter}{end_row}",
        FormulaRule(
            formula=[f'OR({check_letter}{header_row + 1}="✓",{check_letter}{header_row + 1}="+")'],
            fill=PatternFill("solid", fgColor="C6EFCE"),
        ),
    )

    ws.print_title_rows = f"1:{header_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "Страница &P из &N"
    ws.oddFooter.center.size = 9
    ws.oddFooter.center.color = "667085"
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35
    ws.page_margins.header = 0.15
    ws.page_margins.footer = 0.15

    output = BytesIO()
    wb.save(output)
    return output.getvalue()

