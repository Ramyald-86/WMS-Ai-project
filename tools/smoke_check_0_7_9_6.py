#!/usr/bin/env python3
from __future__ import annotations

import compileall
import io
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
app = root / "app"


def require(path: Path, needle: str) -> None:
    value = path.read_text(encoding="utf-8-sig")
    if needle not in value:
        raise AssertionError(f"{path.relative_to(root)}: missing {needle!r}")


if not compileall.compile_dir(app, quiet=1, force=True):
    raise SystemExit("FAILED: Python compilation")

main = app / "main.py"
layout = app / "routers" / "layout.py"
receiving = app / "routers" / "receiving.py"
market = app / "routers" / "marketplaces.py"
picking = app / "services" / "picking_lists.py"
structured = app / "services" / "structured_receiving.py"
reservations = app / "services" / "stock_reservations.py"
js = app / "static" / "app.js"

for path, needle in [
    (main, 'version="0.7.9.6"'),
    (layout, 'Literal["RACK", "CELL", "PALLET", "PALLET_PLACE", "SECTION"]'),
    (layout, '"location_type": "PALLET_PLACE" if is_place else "PALLET"'),
    (layout, 'location_type="SECTION"'),
    (layout, 'and str(row.location_type or "").strip().upper() != "SECTION"'),
    (structured, 'def structured_section_code'),
    (receiving, 'loc.parent_id = section_map[section].id'),
    (receiving, '"created_sections": created_sections'),
    (picking, '"pallet_place", "Палеточное место"'),
    (picking, 'if any(bool(item.get("under_zero")) for item in export_rows)'),
    (market, '"container_total_quantity": container_totals.get(container.id, 0.0)'),
    (js, "PALLET_PLACE"),
    (js, "SECTION"),
    (js, "assignableLocations"),
    # Previous critical guarantees must remain cumulative.
    (market, "PACKED_STAYS_WRITTEN_OFF"),
    (market, "CHZ_AUTO_RELEASED_CANCELLED_BEFORE_LABEL"),
    (js, "Math.min(delay,85)"),
    (js, "Math.min(delay,220)"),
]:
    require(path, needle)

import sys
sys.path.insert(0, str(root))
from app.services.picking_lists import picking_list_xlsx


def headers_for(rows):
    content = picking_list_xlsx(
        title="Smoke",
        created_at=datetime.now(timezone.utc),
        rows=rows,
        order_count=1,
    )
    ws = load_workbook(io.BytesIO(content), data_only=False).active
    headers = [ws.cell(8, col).value for col in range(1, ws.max_column + 1)]
    # Every exported column except the intentional collector check column must have
    # at least one value in the data row(s).
    for col, header in enumerate(headers, start=1):
        if header == "✓":
            continue
        values = [ws.cell(row, col).value for row in range(9, ws.max_row + 1)]
        assert any(value not in (None, "") for value in values), (header, values)
    return headers


structured_headers = headers_for([{
    "location": "С1-П2", "location_name": "Секция 1 · Палета 2", "location_type": "PALLET",
    "container": "BOX-7", "article": "A-1", "size": "42", "barcode": "460000000001",
    "name": "Товар", "required_quantity": 2, "available_quantity": 5, "status": "OK", "under_zero": False,
}])
assert structured_headers == [
    "Секция", "Палета", "Короб", "Артикул", "Размер", "ШК товара", "Наименование",
    "Собрать, шт.", "Доступно", "✓",
]
assert "Ячейка" not in structured_headers and "Стеллаж" not in structured_headers and "Под 0" not in structured_headers

cell_headers = headers_for([{
    "location": "SEC-B-A-14-03-16",
    "location_name": "Секция B · Проход A · Стеллаж 14 · Этаж 3 · Ячейка 16",
    "location_type": "CELL", "container": "", "article": "A-2", "size": "", "barcode": "460000000002",
    "name": "", "required_quantity": 1, "available_quantity": 2, "status": "OK", "under_zero": False,
}])
assert cell_headers == [
    "Секция", "Проход", "Стеллаж", "Этаж / полка", "Ячейка", "Артикул", "ШК товара",
    "Собрать, шт.", "Доступно", "✓",
]
assert "Короб" not in cell_headers and "Размер" not in cell_headers and "Наименование" not in cell_headers

place_headers = headers_for([{
    "location": "SEC-C-PP-007", "location_name": "Секция C · Палеточное место 7", "location_type": "PALLET_PLACE",
    "container": "BOX-8", "article": "A-3", "size": "M", "barcode": "460000000003", "name": "Товар 3",
    "required_quantity": 1, "available_quantity": 1, "status": "OK", "under_zero": True,
}])
assert "Палеточное место" in place_headers and "Палета" not in place_headers and "Под 0" in place_headers

print("OK: FulfillCore 0.7.9.6 dynamic picking/storage-layout smoke-check passed")
